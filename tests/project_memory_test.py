"""Project 层测试：跨任务复用的已确认项目记忆（蓝图 §3.2）。

覆盖三件事：

1. **边界**：什么能进入项目记忆，什么不能（候选术语、未确认风格、非人类记录
   一律排除）；
2. **注入**：项目记忆真的会进入同项目的新任务，并且任务自带术语优先；
3. **界面**：项目列表/详情可读，提升动作在真实 `app.py` 里可用。

运行：`python -m pytest tests/project_memory_test.py -q`
"""

from __future__ import annotations

import json
import shutil
import sys
import tempfile
from contextlib import contextmanager
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "tests"))

import core  # noqa: E402
from make_scenario_fixtures import build_terminology_scenario  # noqa: E402
from offline_provider import OfflineProvider  # noqa: E402
from transpraxis import project as project_module  # noqa: E402


@contextmanager
def project_env(provider: OfflineProvider | None = None):
    """隔离的输出目录（项目与任务都落在里面）。"""
    tmp = Path(tempfile.mkdtemp(prefix="project-test-"))
    old_output, old_call = core.OUTPUT_DIR, core.call_llm
    core.OUTPUT_DIR = tmp
    if provider is not None:
        core.call_llm = provider
    try:
        yield tmp
    finally:
        core.OUTPUT_DIR, core.call_llm = old_output, old_call
        shutil.rmtree(tmp, ignore_errors=True)


def _project_detail(project_id, *, tab="overview", **state):
    """打开 `/projects/:projectId` 的某个一级 tab。

    项目详情是四个 tab（概览 / 任务 / 项目知识 / 设置）的一级表面，因此断言
    知识内容必须先落到对应 tab——这正是信息架构的一部分。
    """
    from streamlit.testing.v1 import AppTest

    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=60)
    at.session_state["app_view"] = "projects"
    at.session_state["active_project_id"] = project_id
    at.session_state["active_project_tab"] = tab
    for key, value in state.items():
        at.session_state[key] = value
    at.run()
    return at


LOCKED = {"source": "canopy closure", "target": "林冠郁闭", "preferred": "林冠郁闭",
          "status": "locked", "behavior": "translate"}
CANDIDATE = {"source": "soil respiration", "target": "土壤呼吸",
             "preferred": "土壤呼吸", "status": "candidate",
             "behavior": "translate"}


# ================= 标识与容器 =================


def test_project_ids_are_uuid_and_never_derived_from_the_name():
    """项目 ID 是 UUID：对任何名称唯一，也不由名称派生。

    这里守住的是"display name 不是主键"：旧实现用名称派生 slug，于是中文名
    会被哈希成一个 p- 前缀的 id、ASCII 名会直接得到可读 slug。ID 与名称绑定之后，
    改名就等于换项目，命名冲突会静默覆盖另一个项目。
    """
    names = ["无人机论文", "生态学 / 2026", "论文 / 2026", "Ecology 2026",
             "Field Notes", "默认项目"]
    with project_env():
        ids = []
        for name in names:
            ids.append(core.create_project(name)["project_id"])
        assert len(set(ids)) == len(ids), f"项目 ID 冲突：{dict(zip(names, ids))}"
        for value in ids:
            assert project_module.is_uuid(value), value
            assert not value.startswith("p-"), value


def test_creating_chinese_named_project_does_not_touch_default():
    """新建中文名项目绝不能覆盖默认项目（曾经的缺陷）。"""
    with project_env():
        core.ensure_default_project()  # 默认项目由变更路径创建，读取不再创建它
        before = core.load_project(core.DEFAULT_PROJECT_ID)
        created = core.create_project("无人机论文")
        after = core.load_project(core.DEFAULT_PROJECT_ID)
        assert created["project_id"] != core.DEFAULT_PROJECT_ID
        assert created["name"] == "无人机论文"
        assert after["name"] == before["name"] == core.DEFAULT_PROJECT_NAME
        assert created["project_id"] in [
            p["project_id"] for p in core.list_projects()]


def test_legacy_job_without_project_falls_back_to_default():
    """旧任务没有 project_id 时必须归入默认项目，而不是丢任务。"""
    with project_env():
        core.save_job_state("legacy", {"filename": "old.docx",
                                       "paras": [], "pairs": []})
        state = core.load_job_state("legacy")
        assert core.resolved_project_id(state) == core.DEFAULT_PROJECT_ID
        assert "legacy" in [job["job_id"]
                            for job in core.list_project_jobs(core.DEFAULT_PROJECT_ID)]


def test_job_assignment_resolves_by_name_or_id():
    """按名称和按 ID 归档必须命中同一个项目。"""
    with project_env():
        project = core.create_project("无人机论文")
        core.save_job_state("j1", core.new_job_state("a.docx"))
        by_name = core.assign_job_to_project("j1", "无人机论文")
        assert by_name["project_id"] == project["project_id"]
        core.save_job_state("j2", core.new_job_state("b.docx"))
        by_id = core.assign_job_to_project("j2", project["project_id"])
        assert by_id["project_id"] == project["project_id"]
        members = [job["job_id"]
                   for job in core.list_project_jobs(project["project_id"])]
        assert sorted(members) == ["j1", "j2"]


# ================= 边界：只有"已确认"能进入 =================


def test_only_confirmed_knowledge_enters_project_memory():
    """候选术语、未确认风格、非人类记录都不得进入项目记忆。"""
    with project_env():
        project = core.create_project("boundary")
        core.save_job_state("j", core.new_job_state("a.docx"))
        core.assign_job_to_project("j", project["project_id"])
        state = core.load_job_state("j")
        state["glossary"] = [LOCKED, CANDIDATE]
        state["confirmed_style_rules"] = [
            {"rule": "保持学术书面语", "status": "confirmed"},
            {"rule": "待定的风格", "status": "proposed"},
        ]
        state["human_actions"] = [
            {"record_type": "human_decision", "actor": "xueyang",
             "actor_type": "human", "decision": "dismiss"},
            {"record_type": "human_decision", "actor": "reviewer-model",
             "actor_type": "model", "decision": "accept_resolution"},
        ]
        core.save_job_state("j", state)

        promoted = core.promote_job_to_project("j", actor="xueyang")
        assert [e["source"] for e in promoted["glossary"]] == ["canopy closure"], \
            "候选术语不得进入项目记忆"
        assert [r["rule"] for r in promoted["style_rules"]] == ["保持学术书面语"], \
            "未确认风格不得进入项目记忆"
        assert len(promoted["human_decisions"]) == 1, "模型记录不得进入项目审计"
        assert promoted["human_decisions"][0]["actor"] == "xueyang"


def test_promotion_is_idempotent_and_versioned():
    """重复提升不得制造重复记录或版本膨胀。"""
    with project_env():
        project = core.create_project("idem")
        core.save_job_state("j", core.new_job_state("a.docx"))
        core.assign_job_to_project("j", project["project_id"])
        state = core.load_job_state("j")
        state["glossary"] = [LOCKED]
        core.save_job_state("j", state)

        first = core.promote_job_to_project("j", actor="xueyang")
        second = core.promote_job_to_project("j", actor="xueyang")
        assert len(first["glossary"]) == len(second["glossary"]) == 1
        assert len(second["glossary_versions"]) == 1, "内容未变不应新增版本"
        assert len(second["promotion_log"]) == 1, "内容未变不应新增提升记录"

        # 人工修正首选译名：项目记忆必须跟上，并产生新版本
        state = core.load_job_state("j")
        state["glossary"] = [{**LOCKED, "target": "冠层郁闭", "preferred": "冠层郁闭"}]
        core.save_job_state("j", state)
        third = core.promote_job_to_project("j", actor="xueyang")
        assert third["glossary"][0]["preferred"] == "冠层郁闭"
        assert len(third["glossary_versions"]) == 2
        # 不变式：一个 source 至多一条锁定条目（人工修正必须就地替换）
        sources = [e["source"].casefold() for e in third["glossary"]]
        assert sources.count("canopy closure") == 1, third["glossary"]
        # 注入时同样不得出现重复要求
        injected = project_module.injection_for(third)["glossary"]
        assert [e["source"].casefold() for e in injected].count("canopy closure") == 1


# ================= 注入：跨任务复用真的发生 =================


def test_project_memory_is_injected_into_a_new_job(fixtures_dir):
    """同项目的新任务必须：注入锁定术语与风格，并让译文真的用上项目术语。"""
    docx, glossary_path = fixtures_dir
    with project_env():
        project = core.create_project("reuse")
        seeded = project_module.merge_confirmed_knowledge(
            core.load_project(project["project_id"]),
            glossary=[{**LOCKED, "source": "canopy closure"}],
            style_rules=[{"rule": "保持学术书面语", "status": "confirmed"}],
            actor="xueyang", source_job_id="earlier")
        core.save_project(seeded)

        provider = OfflineProvider(glossary=seeded["glossary"])
        core.call_llm = provider
        state = core.new_job_state(docx.name)
        state["project_id"] = project["project_id"]
        core.save_job_state("reuse-job", state)

        result = core.run_job_pipeline(
            "reuse-job", docx.name, docx.read_bytes(),
            provider="DeepSeek", api_key="k", model="m", target_lang="简体中文",
            auto_term=False, enable_report=False, translation_theory="",
            user_glossary=[], style_rules="", enable_review=False,
            enable_annotate=False, use_tm=False,
            delivery_config={"deliver_report": False})

        # 项目记忆真的进入了本次任务，并被审计记录
        injection = result.get("project_memory") or {}
        assert injection.get("project_id") == project["project_id"]
        assert injection.get("injected_entry_ids"), "必须记录被注入的项目术语"
        assert injection.get("glossary_version") == 1
        assert injection.get("glossary_hash"), "注入必须带项目术语版本 hash"

        # 术语与风格都已生效，且 state 与运行时配置一致
        assert "canopy closure" in [e["source"] for e in result["glossary"]]
        assert result["style_rules"] == "保持学术书面语"
        assert (result.get("pipeline_config") or {}).get("style_rules") == \
            "保持学术书面语", "存储的风格必须与运行时使用的一致"

        joined = " ".join(pair["target"] for pair in result["pairs"])
        assert "林冠郁闭" in joined, "项目锁定术语必须被强制使用"


def test_task_owned_terms_win_over_project_memory(fixtures_dir):
    """同一个 source 同时来自任务与项目时，任务自带术语优先，且不产生重复条目。"""
    docx, _ = fixtures_dir
    with project_env():
        project = core.create_project("precedence")
        seeded = project_module.merge_confirmed_knowledge(
            core.load_project(project["project_id"]),
            glossary=[{**LOCKED, "target": "项目译名", "preferred": "项目译名"}],
            actor="xueyang", source_job_id="earlier")
        core.save_project(seeded)

        provider = OfflineProvider(
            glossary=[{**LOCKED, "target": "任务译名", "preferred": "任务译名"}])
        core.call_llm = provider
        state = core.new_job_state(docx.name)
        state["project_id"] = project["project_id"]
        core.save_job_state("precedence-job", state)

        result = core.run_job_pipeline(
            "precedence-job", docx.name, docx.read_bytes(),
            provider="DeepSeek", api_key="k", model="m", target_lang="简体中文",
            auto_term=False, enable_report=False, translation_theory="",
            user_glossary=[{**LOCKED, "target": "任务译名", "preferred": "任务译名"}],
            style_rules="", enable_review=False, enable_annotate=False,
            use_tm=False, delivery_config={"deliver_report": False})

        sources = [e["source"].casefold() for e in result["glossary"]]
        assert sources.count("canopy closure") == 1, \
            f"同一 source 不得出现两条锁定条目：{sources}"
        entry = next(e for e in result["glossary"]
                     if e["source"].casefold() == "canopy closure")
        assert entry["preferred"] == "任务译名", "任务自带术语必须优先于项目记忆"
        assert not (result.get("project_memory") or {}).get("injected_entry_ids"), \
            "已被任务术语覆盖的项目条目不应被记为注入"


def test_started_job_is_not_retro_injected(fixtures_dir):
    """进行中的任务不得因为项目记忆变化而改术语。"""
    docx, _ = fixtures_dir
    with project_env():
        project = core.create_project("inflight")
        state = core.new_job_state(docx.name)
        state["project_id"] = project["project_id"]
        state["glossary"] = [{"id": "own", "source": "own term", "target": "自有术语",
                              "preferred": "自有术语", "status": "locked",
                              "behavior": "translate"}]
        state["p2_done"] = True
        core.save_job_state("inflight-job", state)

        seeded = project_module.merge_confirmed_knowledge(
            core.load_project(project["project_id"]),
            glossary=[LOCKED], actor="xueyang", source_job_id="other")
        core.save_project(seeded)

        provider = OfflineProvider()
        core.call_llm = provider
        result = core.run_job_pipeline(
            "inflight-job", docx.name, docx.read_bytes(),
            provider="DeepSeek", api_key="k", model="m", target_lang="简体中文",
            auto_term=False, enable_report=False, translation_theory="",
            user_glossary=[], style_rules="", enable_review=False,
            enable_annotate=False, use_tm=False,
            delivery_config={"deliver_report": False})

        assert not (result.get("project_memory") or {}).get("injected_entry_ids"), \
            "已开始的任务不应被回溯注入项目记忆"
        assert [e["source"] for e in result["glossary"]] == ["own term"]


# ================= 界面 =================


@pytest.fixture(scope="module")
def fixtures_dir(tmp_path_factory):
    target = tmp_path_factory.mktemp("project-fixtures")
    docx = target / "project-scenario.docx"
    glossary = target / "project-scenario.glossary.json"
    build_terminology_scenario(docx, glossary)
    return docx, glossary


def test_ui_projects_surface_lists_and_opens_a_project(fixtures_dir):
    """项目列表必须给出记忆摘要，并且能进入项目详情看到共享记忆。"""
    from streamlit.testing.v1 import AppTest

    with project_env():
        project = core.create_project("无人机论文")
        seeded = project_module.merge_confirmed_knowledge(
            core.load_project(project["project_id"]),
            glossary=[LOCKED],
            style_rules=[{"rule": "保持学术书面语", "status": "confirmed"}],
            actor="xueyang", source_job_id="earlier")
        core.save_project(seeded)
        state = core.new_job_state("in-project.docx")
        state["project_id"] = project["project_id"]
        core.save_job_state("member", state)

        at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=60)
        at.session_state["app_view"] = "projects"
        at.run()
        assert not at.exception, [e.value for e in at.exception]

        page = "\n".join(str(m.value) for m in at.markdown)
        assert "无人机论文" in page, "项目列表必须显示项目名"
        assert "术语 1 · 规则 1" in page, \
            "项目卡片必须显示已积累的记忆摘要"

        at.button(key=f"project_open_{project['project_id']}").click()
        at.run()
        assert not at.exception, [e.value for e in at.exception]
        assert at.session_state["active_project_id"] == project["project_id"]
        assert "无人机论文" in "\n".join(str(m.value) for m in at.markdown)

        # 概览是工作中心：给出真实的知识统计；「任务」tab 只列属于本项目的任务。
        overview = "\n".join(str(m.value) for m in at.markdown)
        assert "工作概览" in overview and "项目知识" in overview, overview[:300]
        # 每张 summary card = label / 主值 / 短注三元组，说明保持短句。
        for label, value, note in (("术语", 1, "已确认"), ("规则", 1, "已确认"),
                                   ("记忆", 0, "尚无")):
            assert (f'<div class="tp-stat-label">{label}</div>'
                    f'<div class="tp-stat-value">{value}</div>'
                    f'<div class="tp-stat-note">{note}</div>') in overview, label
        at.button(key="project_tab_tasks").click()
        at.run()
        assert at.button(key="project_task_open_member") is not None, \
            [b.key for b in at.button]


def test_ui_promotion_moves_confirmed_terms_into_project_memory():
    """Memory gate 在界面上可用：提升后项目记忆才拥有可注入的术语。"""
    from streamlit.testing.v1 import AppTest

    with project_env(OfflineProvider()):
        project = core.create_project("promote-ui")
        state = core.new_job_state("promote.docx")
        state.update({"project_id": project["project_id"], "p1_done": True,
                      "p2_done": False, "quality_mode": True,
                      "stage": "GLOSSARY_PENDING",
                      "paras": ["The canopy closure index was recomputed."],
                      "glossary": [{**LOCKED, "id": "t1", "occurrences": [0]}]})
        core.save_job_state("promote-job", state)
        core.save_source("promote-job", b"source")

        at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=60)
        at.session_state["active_job_id"] = "promote-job"
        at.session_state["app_view"] = "workspace"
        at.session_state["workspace_mode"] = True
        at.session_state["workspace_section"] = "terms"
        at.run()
        assert not at.exception, [e.value for e in at.exception]

        promote = next((b for b in at.button if b.label == "提升到项目记忆"), None)
        assert promote is not None, \
            f"术语页必须提供提升到项目记忆的动作：{[b.label for b in at.button]}"
        assert promote.disabled is False, "存在锁定术语时必须允许提升"
        # 界面必须说明"会提升什么"，而不是给一个语义模糊的按钮
        page = "\n".join(str(c.value) for c in at.caption)
        assert "可提升" in page and "锁定术语 1 条" in page, page[:200]
        assert "候选术语与未审校译文不会进入项目记忆" in page

        promote.click()
        at.run()
        assert not at.exception, [e.value for e in at.exception]

        stored = core.load_project(project["project_id"])
        assert [e["source"] for e in stored["glossary"]] == ["canopy closure"]
        assert core.project_injection(project["project_id"])["glossary"], \
            "提升后必须能注入到同项目的新任务"


# ================= 翻译记忆按项目隔离 =================

SIGNAL = "Signal sentence about canopy structure and light competition."


def _probe_doc(tmp: Path, name: str = "probe.docx") -> Path:
    from docx import Document as Docx

    document = Docx()
    document.add_paragraph(SIGNAL)
    path = tmp / name
    document.save(str(path))
    return path


def _run_probe(job_id: str, docx: Path, project_id: str) -> dict:
    state = core.new_job_state(docx.name)
    state["project_id"] = project_id
    core.save_job_state(job_id, state)
    return core.run_job_pipeline(
        job_id, docx.name, docx.read_bytes(),
        provider="DeepSeek", api_key="k", model="m", target_lang="简体中文",
        auto_term=False, enable_report=False, translation_theory="",
        user_glossary=[], style_rules="", enable_review=True,
        enable_annotate=False, use_tm=True,
        delivery_config={"deliver_report": False})


def test_tm_paths_are_isolated_and_default_keeps_legacy_path():
    """每个项目一份记忆文件；默认项目沿用历史路径（无需迁移）。"""
    with project_env():
        project = core.create_project("project-a")
        assert core.tm_path(core.DEFAULT_PROJECT_ID) == \
            core.OUTPUT_DIR / "translation_memory.json"
        assert core.tm_path(project["project_id"]) == \
            core.OUTPUT_DIR / "projects" / project["project_id"] / \
            "translation_memory.json"
        assert core.tm_path(None) == core.tm_path(core.DEFAULT_PROJECT_ID)
        # 无 project_id 的旧调用点读写的仍是历史全局记忆
        core.save_tm({"legacy source": {"target": "旧译文", "reviewed": True}})
        assert core.load_tm() == core.load_tm(core.DEFAULT_PROJECT_ID)
        assert core.load_tm(project["project_id"]) == {}


def test_project_memories_do_not_leak_between_projects():
    """检索必须按项目过滤：A 的记忆不能被 B 命中。"""
    with project_env() as tmp:
        core.call_llm = OfflineProvider()
        project_a = core.create_project("project-a")
        project_b = core.create_project("project-b")
        core.save_tm({SIGNAL: {"target": "【A项目记忆】林冠结构", "reviewed": True}},
                     project_a["project_id"])
        docx = _probe_doc(tmp)

        hit = _run_probe("probe-a", docx, project_a["project_id"])
        assert hit["pairs"][0].get("from_tm") is True, "A 必须命中自己项目的记忆"
        assert hit["pairs"][0]["target"] == "【A项目记忆】林冠结构"

        miss = _run_probe("probe-b", docx, project_b["project_id"])
        assert miss["pairs"][0].get("from_tm") is not True, \
            "B 不得命中 A 项目的翻译记忆"
        assert "A项目记忆" not in miss["pairs"][0]["target"]

        # 各写各的文件
        assert core.tm_path(project_a["project_id"]).is_file()
        assert core.tm_path(project_b["project_id"]).is_file()
        assert core.load_tm(project_a["project_id"]) != \
            core.load_tm(project_b["project_id"])


def test_same_source_can_differ_between_projects():
    """同一个 source 在两个项目里可以有不同的已审校译文。

    这正是"按项目分文件"而不是"单文件加 project_id 字段"的原因：记录以 source
    为键，单文件无法同时表示两种译法。
    """
    with project_env() as tmp:
        core.call_llm = OfflineProvider()
        project_a = core.create_project("project-a")
        project_b = core.create_project("project-b")
        core.save_tm({SIGNAL: {"target": "A 的译法", "reviewed": True}},
                     project_a["project_id"])
        core.save_tm({SIGNAL: {"target": "B 的译法", "reviewed": True}},
                     project_b["project_id"])
        docx = _probe_doc(tmp)

        assert _run_probe("a", docx, project_a["project_id"])["pairs"][0]["target"] \
            == "A 的译法"
        assert _run_probe("b", docx, project_b["project_id"])["pairs"][0]["target"] \
            == "B 的译法"
        # 两份记忆各自保留自己的译法，互不覆盖
        assert core.load_tm(project_a["project_id"])[SIGNAL]["target"] == "A 的译法"
        assert core.load_tm(project_b["project_id"])[SIGNAL]["target"] == "B 的译法"


def test_adopting_default_memory_is_explicit_and_merge_only():
    """把默认项目的记忆带入新项目必须是显式动作，且只增不改。"""
    with project_env():
        project = core.create_project("project-a")
        core.save_tm({"s1": {"target": "t1", "reviewed": True},
                      "s2": {"target": "t2", "reviewed": True}},
                     core.DEFAULT_PROJECT_ID)
        core.save_tm({"s2": {"target": "本项目自定义", "reviewed": True}},
                     project["project_id"])

        assert core.copy_default_tm_to_project(core.DEFAULT_PROJECT_ID) == 0, \
            "默认项目不需要导入自己"
        added = core.copy_default_tm_to_project(project["project_id"])
        assert added == 1, "只应并入缺失的条目"
        merged = core.load_tm(project["project_id"])
        assert merged["s1"]["target"] == "t1"
        assert merged["s2"]["target"] == "本项目自定义", "不得覆盖本项目已有译法"
        # 幂等
        assert core.copy_default_tm_to_project(project["project_id"]) == 0


def test_tm_writeback_lands_in_the_job_project():
    """任务写回的已审校译文必须落在该项目自己的记忆里。"""
    with project_env() as tmp:
        core.call_llm = OfflineProvider()
        project = core.create_project("project-a")
        docx = _probe_doc(tmp)
        _run_probe("writeback", docx, project["project_id"])

        project_tm = core.load_tm(project["project_id"])
        assert project_tm, "已审校译文必须写入项目记忆"
        assert core.load_tm(core.DEFAULT_PROJECT_ID) == {}, \
            "不得写进默认项目的（历史全局）记忆"


def test_ui_library_shows_per_project_memory():
    """翻译记忆必须按项目展示，而不是只给一个全局数字。"""
    from streamlit.testing.v1 import AppTest

    with project_env():
        project = core.create_project("project-a")
        core.save_tm({"s1": {"target": "t1", "reviewed": True}},
                     project["project_id"])
        core.save_tm({"legacy": {"target": "旧译文", "reviewed": True}},
                     core.DEFAULT_PROJECT_ID)

        at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=60)
        at.session_state["app_view"] = "library"
        at.run()
        assert not at.exception, [e.value for e in at.exception]

        # 按项目统计现在属于「翻译记忆」Tab（一级 Tab 之间职责分离）。
        at.segmented_control[0].set_value("tm").run()
        assert not at.exception, [e.value for e in at.exception]

        metrics = {m.label: m.value for m in at.metric}
        assert metrics.get("project-a") == "1", metrics
        assert metrics.get(core.DEFAULT_PROJECT_NAME) == "1", metrics
        assert any(s.label == "查看哪个项目的记忆" for s in at.selectbox), \
            "必须可以选择查看哪个项目的记忆"


def test_ui_project_can_adopt_system_memory():
    """项目记忆 tab 必须提供显式的"并入系统工作区记忆"动作，并说明隔离事实。"""
    with project_env():
        project = core.create_project("project-a")
        core.save_tm({"legacy": {"target": "旧译文", "reviewed": True}},
                     core.DEFAULT_PROJECT_ID)

        at = _project_detail(project["project_id"], tab="knowledge")
        assert not at.exception, [e.value for e in at.exception]

        # 并入动作住在「已审核记忆」模块里，而模块详情默认收起——先展开，
        # 否则断言的是"没点开所以没有"，不是"能力缺失"。
        at.button(key="pd_knowledge_toggle_memory").click()
        at.run()
        assert not at.exception, [e.value for e in at.exception]

        page = "\n".join(str(c.value) for c in at.caption)
        assert "不会自动共享" in page, f"必须说明记忆按项目隔离：{page[:200]}"
        adopt = next((b for b in at.button if "并入本项目" in b.label), None)
        assert adopt is not None, f"缺少并入动作：{[b.label for b in at.button]}"

        adopt.click()
        at.run()
        assert not at.exception, [e.value for e in at.exception]
        assert "legacy" in core.load_tm(project["project_id"])


# ================= 缺口一：批量移动任务 =================


def test_batch_move_jobs_between_projects():
    """一次调用可以移动多个任务，并逐个报告被跳过的原因。"""
    with project_env():
        project_a = core.create_project("project-a")
        project_b = core.create_project("project-b")
        for job_id in ("m1", "m2"):
            state = core.new_job_state(f"{job_id}.docx")
            state["project_id"] = project_a["project_id"]
            core.save_job_state(job_id, state)
        already = core.new_job_state("m3.docx")
        already["project_id"] = project_b["project_id"]
        core.save_job_state("m3", already)

        result = core.assign_jobs_to_project(
            ["m1", "m2", "m3", "missing"], project_b["project_id"])
        assert sorted(result["moved"]) == ["m1", "m2"]
        reasons = {item["job_id"]: item["reason"] for item in result["skipped"]}
        assert reasons["m3"] == "已在该项目中"
        assert reasons["missing"] == "任务不存在"
        assert sorted(job["job_id"] for job in
                      core.list_project_jobs(project_b["project_id"])) == \
            ["m1", "m2", "m3"]
        assert core.list_project_jobs(project_a["project_id"]) == []


def test_batch_move_does_not_move_terminology_or_overwrite_memory():
    """移动只改归属：不迁移术语，并入译文时只增不改。"""
    with project_env():
        project_a = core.create_project("project-a")
        project_b = core.create_project("project-b")
        state = core.new_job_state("t.docx")
        state["project_id"] = project_a["project_id"]
        state["glossary"] = [{**LOCKED, "id": "t1"}]
        state["paras"] = ["src"]
        state["pairs"] = [{"source": "s1", "target": "本项目译法", "reviewed": True},
                          {"source": "s2", "target": "t2", "reviewed": True}]
        core.save_job_state("t", state)
        # 目标项目已有 s1 的不同译法
        core.save_tm({"s1": {"target": "目标项目译法", "reviewed": True}},
                     project_b["project_id"])

        result = core.assign_jobs_to_project(
            ["t"], project_b["project_id"], move_translations=True)
        assert result["tm_added"] == 1, "只应并入目标项目缺失的条目"

        after = core.load_job_state("t")
        assert core.resolved_project_id(after) == project_b["project_id"]
        assert [e["source"] for e in after["glossary"]] == ["canopy closure"], \
            "移动不得丢弃任务自己的术语"
        assert core.load_project(project_b["project_id"])["glossary"] == [], \
            "移动不得自动把术语写进项目记忆（那要走 Memory gate）"
        tm = core.load_tm(project_b["project_id"])
        assert tm["s1"]["target"] == "目标项目译法", "不得覆盖目标项目已有译法"
        assert tm["s2"]["target"] == "t2"


def test_running_job_cannot_be_moved():
    """正在运行的任务不能改归属：它已经读取了某个项目的记忆。"""
    with project_env():
        project_a = core.create_project("project-a")
        project_b = core.create_project("project-b")
        state = core.new_job_state("run.docx")
        state["project_id"] = project_a["project_id"]
        core.save_job_state("run", state)
        old_status = core.build_job_runtime_view
        core.build_job_runtime_view = lambda job_id, s=None: {"runtime_status": "running"}
        try:
            result = core.assign_jobs_to_project(["run"], project_b["project_id"])
        finally:
            core.build_job_runtime_view = old_status
        assert result["moved"] == []
        assert result["skipped"][0]["reason"] == "任务正在运行，无法改归属"
        assert core.resolved_project_id(core.load_job_state("run")) == \
            project_a["project_id"], "被拒绝时归属不得改变"


def test_ui_can_move_jobs_into_a_project():
    """项目详情必须能批量移入任务，并显示任务当前所在项目。"""
    from streamlit.testing.v1 import AppTest

    with project_env():
        project_a = core.create_project("project-a")
        project_b = core.create_project("project-b")
        for job_id in ("mv1", "mv2"):
            state = core.new_job_state(f"{job_id}.docx")
            state["project_id"] = project_b["project_id"]
            core.save_job_state(job_id, state)

        at = _project_detail(project_a["project_id"], tab="tasks")
        assert not at.exception, [e.value for e in at.exception]

        mover = next((m for m in at.multiselect if m.label == "选择要移入的任务"), None)
        assert mover is not None, f"缺少批量移动控件：{[m.label for m in at.multiselect]}"
        assert any("project-b" in str(option) for option in mover.options), \
            "候选项必须标出任务当前所在项目"

        mover.select(mover.options[0])
        at.run()
        go = next((b for b in at.button if b.label == "移入本项目"), None)
        assert go is not None and not go.disabled
        go.click()
        at.run()
        assert not at.exception, [e.value for e in at.exception]
        assert len(core.list_project_jobs(project_a["project_id"])) == 1


# ================= 缺口二：项目记忆导出 / 导入 =================


def _seed_portable_project() -> dict:
    project = core.create_project("生态学专著")
    seeded = project_module.merge_confirmed_knowledge(
        core.load_project(project["project_id"]),
        glossary=[{**LOCKED, "id": "t1"}],
        style_rules=[{"rule": "保持学术书面语", "status": "confirmed"}],
        human_decisions=[{"record_type": "human_decision", "actor": "xueyang",
                          "actor_type": "human", "decision": "dismiss"}],
        actor="xueyang", source_job_id="j1")
    core.save_project(seeded)
    core.save_tm({"The canopy closure index was recomputed.":
                  {"target": "林冠郁闭指数被重新计算。", "reviewed": True}},
                 project["project_id"])
    return project


def test_project_memory_round_trips_between_machines():
    """导出再导入必须完整还原术语、风格、决定与本项目已审校记忆。"""
    with project_env() as tmp:
        project = _seed_portable_project()
        blob = core.export_project_memory(project["project_id"])
        payload = json.loads(blob)
        assert payload["format"] == project_module.MEMORY_FORMAT
        assert payload["format_version"] == project_module.MEMORY_FORMAT_VERSION
        assert "content_sha256" in payload
        # 导出物不得包含任务状态、源文档或任何凭据
        assert "state" not in payload and "api_key" not in blob

        # 模拟另一台机器：清空全部输出
        shutil.rmtree(tmp / "projects")
        (tmp / "translation_memory.json").unlink(missing_ok=True)

        restored, report = core.import_project_memory(blob)
        assert restored["name"] == "生态学专著"
        assert [e["source"] for e in restored["glossary"]] == ["canopy closure"]
        assert [r["rule"] for r in restored["style_rules"]] == ["保持学术书面语"]
        assert len(restored["human_decisions"]) == 1
        assert len(core.load_tm(restored["project_id"])) == 1
        assert report["glossary_added"] == 1 and report["tm_added"] == 1

        # 导入后必须能注入到新任务——这才是"可移植资产"的意义
        injected = core.project_injection(restored["project_id"])
        assert [e["source"] for e in injected["glossary"]] == ["canopy closure"]


def test_import_is_merge_only_and_idempotent():
    """重复导入不产生重复；冲突保留本地，不静默覆盖。"""
    with project_env():
        project = _seed_portable_project()
        blob = core.export_project_memory(project["project_id"])

        core.import_project_memory(blob)
        _, again = core.import_project_memory(blob)
        assert (again["glossary_added"], again["style_added"], again["tm_added"],
                again["decisions_added"]) == (0, 0, 0, 0), "重复导入必须幂等"

        # 本地改成不同译法后再次导入：本地保留，冲突被报告
        core.save_tm({"The canopy closure index was recomputed.":
                      {"target": "本地译法", "reviewed": True}},
                     project["project_id"])
        _, conflict = core.import_project_memory(blob)
        assert conflict["tm_conflicts_count"] == 1
        assert core.load_tm(project["project_id"])[
            "The canopy closure index was recomputed."]["target"] == "本地译法"


def test_import_rejects_malformed_or_tampered_payloads():
    """导入是外部输入：格式、版本、校验和都必须被校验。"""
    with project_env():
        project = _seed_portable_project()
        blob = core.export_project_memory(project["project_id"])

        tampered = json.loads(blob)
        tampered["glossary"][0]["target"] = "被篡改"
        cases = {
            "校验和": json.dumps(tampered),
            "格式": '{"format":"something-else"}',
            "版本": json.dumps({"format": project_module.MEMORY_FORMAT,
                                "format_version": 99}),
            "非 JSON": "definitely not json",
            "非对象": "[1, 2, 3]",
        }
        for label, payload in cases.items():
            with pytest.raises(ValueError):
                core.import_project_memory(payload)
        # 超大输入必须被拒绝，而不是把界面拖死
        with pytest.raises(ValueError):
            core.import_project_memory(b"x" * (project_module.MAX_IMPORT_BYTES + 1))


def test_import_can_rename_to_avoid_touching_an_existing_project():
    """`name` 可以把导入内容放到另一个项目，而不动同名项目。"""
    with project_env():
        project = _seed_portable_project()
        blob = core.export_project_memory(project["project_id"])
        other, _ = core.import_project_memory(blob, name="另一个项目")
        assert other["name"] == "另一个项目"
        assert other["project_id"] != project["project_id"]
        original = core.load_project(project["project_id"])
        assert original["name"] == "生态学专著"
        assert original["glossary"], "原项目不得被改名或清空"


def test_ui_project_exports_and_imports_memory():
    """项目页必须提供导出下载与导入入口（导入是 modal，不是首屏表单）。"""
    from streamlit.testing.v1 import AppTest

    with project_env():
        project = _seed_portable_project()

        detail = _project_detail(project["project_id"], tab="knowledge")
        assert not detail.exception, [e.value for e in detail.exception]
        labels = [d.label for d in getattr(detail, "download_button", [])]
        assert any("导出项目记忆" in label for label in labels), labels

        listing = AppTest.from_file(str(ROOT / "app.py"), default_timeout=60)
        listing.session_state["app_view"] = "projects"
        listing.run()
        assert not listing.exception, [e.value for e in listing.exception]
        assert not getattr(listing, "file_uploader", []), \
            "导入表单不得占据项目首屏"
        next(b for b in listing.button
             if b.key == "project_new_import").click()
        listing.run()
        assert not listing.exception, [e.value for e in listing.exception]
        assert any("项目文件" in u.label
                   for u in getattr(listing, "file_uploader", [])), \
            "「新建项目」菜单必须提供导入入口"


# ================= 缺口三：翻译记忆的排版等价匹配 =================


def test_tm_normalization_is_equivalence_not_similarity():
    """归一化只吸收排版差异，不做语义近似。"""
    same = [
        ("The canopy closure index was recomputed.",
         "The  canopy   closure index was recomputed."),
        ("The canopy closure index was recomputed.",
         "The canopy closure index was recomputed. "),
        ('He said "hello".', "He said \u201chello\u201d."),
        ("A B", "A\u00a0B"),
        ("A B", "A\u3000B"),
        ("caf\u00e9", "cafe\u0301"),
    ]
    for left, right in same:
        assert core.tm_normalize(left) == core.tm_normalize(right), (left, right)

    different = [
        ("The canopy closure index was recomputed.",
         "The canopy closure index was recomputed twice."),
        ("Apple", "apple"),
        ("The index was recomputed.", "The index was recomputed?"),
    ]
    for left, right in different:
        assert core.tm_normalize(left) != core.tm_normalize(right), (left, right)


def test_tm_lookup_prefers_exact_and_reports_normalized_hits():
    """精确命中优先；归一化命中必须被标注，便于审计。"""
    tm = {"The canopy closure index was recomputed.":
          {"target": "林冠郁闭指数被重新计算。", "reviewed": True}}
    index = core.tm_index(tm)

    hit, kind = core.tm_lookup(tm, "The canopy closure index was recomputed.", index)
    assert kind == "exact" and hit["target"] == "林冠郁闭指数被重新计算。"

    hit, kind = core.tm_lookup(
        tm, "The  canopy   closure index was recomputed.", index)
    assert kind == "normalized" and hit["target"] == "林冠郁闭指数被重新计算。"

    # 语义不同不得命中
    hit, kind = core.tm_lookup(
        tm, "The canopy closure index was recomputed twice.", index)
    assert hit is None and kind == ""


def test_pipeline_reuses_typography_variant_and_records_match_kind():
    """排版不同但等价的段落必须复用已审校译文，并记录命中方式。"""
    from docx import Document as Docx

    with project_env() as tmp:
        core.call_llm = OfflineProvider()
        stored = "The canopy closure index was recomputed."
        core.save_tm({stored: {"target": "林冠郁闭指数被重新计算。", "reviewed": True}},
                     core.DEFAULT_PROJECT_ID)
        document = Docx()
        document.add_paragraph("The  canopy   closure index was recomputed.")
        path = tmp / "variant.docx"
        document.save(str(path))
        state = core.new_job_state(path.name)
        core.save_job_state("variant", state)

        result = core.run_job_pipeline(
            "variant", path.name, path.read_bytes(),
            provider="DeepSeek", api_key="k", model="m", target_lang="简体中文",
            auto_term=False, enable_report=False, translation_theory="",
            user_glossary=[], style_rules="", enable_review=True,
            enable_annotate=False, use_tm=True,
            delivery_config={"deliver_report": False})

        pair = result["pairs"][0]
        assert pair.get("from_tm") is True, "排版等价的段落应复用记忆"
        assert pair.get("tm_match") == "normalized"
        assert pair["target"] == "林冠郁闭指数被重新计算。"


# ================= 项目归档与删除 =================


def test_archive_hides_project_but_keeps_everything():
    """归档是可恢复的：从活动列表消失，但记忆与任务归属完整保留。"""
    with project_env():
        project = core.create_project("生态学专著")
        core.save_tm({"s": {"target": "t", "reviewed": True}}, project["project_id"])
        state = core.new_job_state("in-project.docx")
        state["project_id"] = project["project_id"]
        core.save_job_state("member", state)

        core.archive_project(project["project_id"], True)
        assert project["project_id"] not in [
            p["project_id"] for p in core.list_active_projects()]
        assert project["project_id"] in [
            p["project_id"] for p in core.list_projects()], "归档项目仍应可枚举"
        # 记忆与任务归属不受影响
        assert core.load_tm(project["project_id"]) == {"s": {"target": "t",
                                                            "reviewed": True}}
        assert core.resolved_project_id(core.load_job_state("member")) == \
            project["project_id"]
        assert [job["job_id"] for job in
                core.list_project_jobs(project["project_id"])] == ["member"]

        core.archive_project(project["project_id"], False)
        assert project["project_id"] in [
            p["project_id"] for p in core.list_active_projects()]


def test_default_project_cannot_be_archived_or_deleted():
    """默认项目承载历史归属与既有全局记忆，必须不可归档、不可删除。"""
    with project_env():
        # 默认项目由变更路径创建；读取不再顺手创建它。
        core.ensure_default_project()
        before = core.load_project(core.DEFAULT_PROJECT_ID)
        with pytest.raises(ValueError):
            core.archive_project(core.DEFAULT_PROJECT_ID, True)
        with pytest.raises(ValueError):
            core.delete_project(core.DEFAULT_PROJECT_ID,
                                confirm_name=core.DEFAULT_PROJECT_NAME)
        after = core.load_project(core.DEFAULT_PROJECT_ID)
        assert after is not None, "被拒绝后默认项目必须仍然存在"
        assert after["name"] == before["name"]
        assert not after.get("archived_at"), "默认项目不得被归档"

    # 保护不能依赖磁盘状态：记录还不存在时也必须拒绝
    with project_env():
        assert core.load_project(core.DEFAULT_PROJECT_ID) is None
        with pytest.raises(ValueError):
            core.archive_project(core.DEFAULT_PROJECT_ID, True)
        with pytest.raises(ValueError):
            core.delete_project(core.DEFAULT_PROJECT_ID,
                                confirm_name=core.DEFAULT_PROJECT_NAME)


def test_delete_requires_exact_name_confirmation():
    """删除必须逐字输入项目名称。

    首尾空白被忽略是有意为之：复制粘贴带来的一个空格不应该否定一次明确的确认。
    安全性来自"必须知道并输入项目名称"，而不是逐字节相等。
    """
    with project_env():
        project = core.create_project("生态学专著")
        for wrong in ("", "生态学", "生态学专著V2", "另一个项目"):
            with pytest.raises(ValueError):
                core.delete_project(project["project_id"], confirm_name=wrong)
        assert core.load_project(project["project_id"]) is not None, \
            "确认失败时项目不得被删除"

        # 首尾空白不算错误输入（已文档化的行为）
        core.delete_project(project["project_id"], confirm_name="  生态学专著  ")
        assert core.load_project(project["project_id"]) is None


def test_delete_refuses_while_jobs_are_assigned():
    """项目下仍有任务时拒绝删除，除非显式指定迁移目标。"""
    with project_env():
        project = core.create_project("生态学专著")
        state = core.new_job_state("j.docx")
        state["project_id"] = project["project_id"]
        core.save_job_state("j", state)

        with pytest.raises(ValueError) as excinfo:
            core.delete_project(project["project_id"], confirm_name="生态学专著")
        assert "任务" in str(excinfo.value)
        assert core.load_project(project["project_id"]) is not None

        result = core.delete_project(
            project["project_id"], confirm_name="生态学专著",
            move_jobs_to=core.DEFAULT_PROJECT_ID)
        assert result["reassigned_jobs"] == ["j"]
        assert core.resolved_project_id(core.load_job_state("j")) == \
            core.DEFAULT_PROJECT_ID, "任务必须被移到目标项目，而不是悬空"
        assert core.load_project(project["project_id"]) is None


def test_delete_writes_a_recoverable_backup():
    """删除前必须留备份，且备份能被导回来（删除在实践中可恢复）。"""
    with project_env():
        project = _seed_portable_project()
        blob_before = core.export_project_memory(project["project_id"])

        result = core.delete_project(project["project_id"],
                                     confirm_name="生态学专著")
        backup = result["backup"]
        assert backup.is_file(), "删除必须留下备份文件"
        assert core.load_project(project["project_id"]) is None

        # 备份内容与删除前的导出等价
        backup_payload = json.loads(backup.read_text(encoding="utf-8"))
        before_payload = json.loads(blob_before)
        assert backup_payload["content_sha256"] == before_payload["content_sha256"]

        # 用备份导回来 → 完整还原
        restored, report = core.import_project_memory(
            backup.read_bytes(), name="生态学专著")
        assert [e["source"] for e in restored["glossary"]] == ["canopy closure"]
        assert [r["rule"] for r in restored["style_rules"]] == ["保持学术书面语"]
        assert len(core.load_tm(restored["project_id"])) == 1
        assert report["glossary_added"] == 1


def test_ui_project_archive_restore_and_guarded_delete():
    """界面：设置 tab 给出归档 / 恢复 / 删除；删除必须二次确认后才可用。"""
    with project_env():
        project = core.create_project("生态学专著")

        detail = _project_detail(project["project_id"], tab="settings")
        assert not detail.exception, [e.value for e in detail.exception]

        assert detail.button(key=f"settings_archive_{project['project_id']}") is not None, \
            f"缺少归档入口：{[b.label for b in detail.button]}"
        assert detail.button(key=f"settings_delete_{project['project_id']}") is not None
        # 危险操作走 modal：二次确认（逐字输入项目名称）在弹窗里完成。
        detail.button(key=f"settings_delete_{project['project_id']}").click()
        detail.run()
        assert not detail.exception, [e.value for e in detail.exception]
        confirm = detail.button(key="project_form_confirm")
        assert confirm.disabled is True, "未输入项目名称前删除必须禁用"
        detail.text_input(key="project_form_confirm_name").set_value("生态学专著")
        detail.run()
        assert detail.button(key="project_form_confirm").disabled is False
        detail.button(key="project_form_cancel").click()
        detail.run()
        assert core.load_project(project["project_id"]) is not None, "取消不得删除项目"

        # 归档：可恢复，且不删除任何东西。
        detail.button(key=f"settings_archive_{project['project_id']}").click()
        detail.run()
        detail.button(key="project_form_confirm").click()
        detail.run()
        assert core.load_project(project["project_id"]).get("archived_at")
        assert detail.session_state["active_project_tab"] == "settings", \
            "归档后停留在该项目的设置 tab（它仍然可以打开）"

        at = _project_detail(project["project_id"], tab="settings")
        assert at.button(key=f"settings_restore_{project['project_id']}") is not None, \
            f"归档后必须能恢复：{[b.label for b in at.button]}"
        at.button(key=f"settings_restore_{project['project_id']}").click()
        at.run()
        at.button(key="project_form_confirm").click()
        at.run()
        assert not core.load_project(project["project_id"]).get("archived_at")
        assert project["project_id"] in [
            p["project_id"] for p in core.list_active_projects()]

        # 系统工作区不是项目：它进入 Task Inbox，不渲染项目设置 tab，
        # 因此 archive / delete / rename 入口一概不存在。
        system = _project_detail(core.SYSTEM_PROJECT_ID, tab="settings")
        assert not system.exception, [e.value for e in system.exception]
        keys = [b.key for b in system.button]
        assert f"settings_archive_{core.SYSTEM_PROJECT_ID}" not in keys
        assert f"settings_delete_{core.SYSTEM_PROJECT_ID}" not in keys
        assert f"settings_rename_{core.SYSTEM_PROJECT_ID}" not in keys
        assert all(b.key != "project_tab_settings" for b in system.button), \
            "系统工作区不得渲染项目「设置」tab"


def test_new_task_selector_only_offers_active_projects():
    """新建任务不应把已归档项目作为可选项。"""
    from docx import Document as Docx
    from streamlit.testing.v1 import AppTest

    with project_env() as tmp:
        active = core.create_project("活动项目")
        archived = core.create_project("已归档项目")
        core.archive_project(archived["project_id"], True)
        document = Docx()
        document.add_paragraph("A sentence for the project selector.")
        path = tmp / "pick.docx"
        document.save(str(path))

        at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=60)
        at.session_state["app_view"] = "new"
        at.session_state["task_step"] = 1
        at.session_state["task_files"] = [{"name": path.name,
                                          "bytes": path.read_bytes()}]
        at.run()
        assert not at.exception, [e.value for e in at.exception]
        # 项目上下文只有一个来源：侧栏 switcher。它的列表里不出现归档项目。
        at.button(key="current_project_selector").click()
        at.run()
        rows = [b.label for b in at.button
                if b.key and b.key.startswith("switcher_pick_")]
        assert any("活动项目" in label for label in rows), rows
        assert not any("已归档项目" in label for label in rows), \
            f"已归档项目不应作为上下文：{rows}"
        assert active["project_id"]


# ================= 导入冲突的持久化与一键采纳 =================


def _peer_export(*, glossary_target: str, tm_target: str) -> str:
    """构造一个"协作者导出物"：同名项目、但术语与译文都不同。"""
    peer = core.create_project(f"协作者导出-{glossary_target}")
    seeded = project_module.merge_confirmed_knowledge(
        core.load_project(peer["project_id"]),
        glossary=[{**LOCKED, "target": glossary_target,
                   "preferred": glossary_target}],
        actor="peer", source_job_id="peer-job")
    core.save_project(seeded)
    core.save_tm({"Sentence A.": {"target": tm_target, "reviewed": True}},
                 peer["project_id"])
    return core.export_project_memory(peer["project_id"])


def _local_project(*, glossary_target: str = "本地译名",
                   tm_target: str = "本地 A 译文") -> dict:
    project = core.create_project("生态学专著")
    seeded = project_module.merge_confirmed_knowledge(
        core.load_project(project["project_id"]),
        glossary=[{**LOCKED, "target": glossary_target,
                   "preferred": glossary_target}],
        actor="xueyang", source_job_id="local-job")
    core.save_project(seeded)
    core.save_tm({"Sentence A.": {"target": tm_target, "reviewed": True}},
                 project["project_id"])
    return project


def test_import_records_addressable_conflicts_without_overwriting():
    """冲突必须成为持久、可寻址的待决项，且本地不被静默覆盖。"""
    with project_env():
        local = _local_project()
        blob = _peer_export(glossary_target="协作者译名", tm_target="协作者 A 译文")
        _, report = core.import_project_memory(blob, name="生态学专著")

        assert report["conflicts_recorded"] == 2, \
            f"术语与译文冲突都该被登记：{report}"
        conflicts = core.list_project_conflicts(local["project_id"])
        kinds = sorted(item["kind"] for item in conflicts)
        assert kinds == ["glossary", "translation_memory"], conflicts
        for item in conflicts:
            assert item["conflict_id"], "冲突必须可寻址"
            assert item["local"] and item["incoming"]
            assert item["imported_from"] == "协作者导出-协作者译名"

        # 本地未被覆盖
        assert [e["preferred"] for e in
                core.load_project(local["project_id"])["glossary"]] == ["本地译名"]
        assert core.load_tm(local["project_id"])["Sentence A."]["target"] == "本地 A 译文"


def test_conflicts_survive_reload_and_do_not_duplicate():
    """待决冲突是持久状态：重新加载仍在；重复导入不堆叠。"""
    with project_env():
        local = _local_project()
        blob = _peer_export(glossary_target="协作者译名", tm_target="协作者 A 译文")
        core.import_project_memory(blob, name="生态学专著")
        first = core.list_project_conflicts(local["project_id"])
        assert len(first) == 2

        core.import_project_memory(blob, name="生态学专著")
        assert len(core.list_project_conflicts(local["project_id"])) == 2, \
            "重复导入同一批冲突不得堆叠"
        # 重新从磁盘读取仍是同样两条
        reloaded = core.list_project_conflicts(local["project_id"])
        assert {item["conflict_id"] for item in reloaded} == \
            {item["conflict_id"] for item in first}


def test_adopting_incoming_replaces_glossary_and_translation_memory():
    """采纳导入版本：术语就地替换、译文写入 TM，并留下审计记录。"""
    with project_env():
        local = _local_project()
        blob = _peer_export(glossary_target="协作者译名", tm_target="协作者 A 译文")
        core.import_project_memory(blob, name="生态学专著")

        for conflict in list(core.list_project_conflicts(local["project_id"])):
            core.resolve_project_conflict(local["project_id"],
                                          conflict["conflict_id"],
                                          adopt_incoming=True, actor="xueyang")

        assert [e["preferred"] for e in
                core.load_project(local["project_id"])["glossary"]] == ["协作者译名"]
        assert core.load_tm(local["project_id"])["Sentence A."]["target"] == \
            "协作者 A 译文"
        assert core.list_project_conflicts(local["project_id"]) == []

        decisions = [entry.get("conflict_resolution")
                     for entry in core.load_project(local["project_id"])["promotion_log"]
                     if entry.get("conflict_resolution")]
        assert len(decisions) == 2
        assert {item["adopted"] for item in decisions} == {"incoming"}
        assert all(item["source"] for item in decisions)
        # 采纳术语必须产生新的项目术语版本
        assert len(core.load_project(local["project_id"])["glossary_versions"]) >= 1


def test_keeping_local_removes_conflict_without_changing_anything():
    """保留本地：只移除待决项，本地译名不变，但决定仍被审计。"""
    with project_env():
        local = _local_project()
        blob = _peer_export(glossary_target="协作者译名", tm_target="协作者 A 译文")
        core.import_project_memory(blob, name="生态学专著")

        for conflict in list(core.list_project_conflicts(local["project_id"])):
            core.resolve_project_conflict(local["project_id"],
                                          conflict["conflict_id"],
                                          adopt_incoming=False, actor="xueyang")

        assert core.list_project_conflicts(local["project_id"]) == []
        assert [e["preferred"] for e in
                core.load_project(local["project_id"])["glossary"]] == ["本地译名"]
        assert core.load_tm(local["project_id"])["Sentence A."]["target"] == "本地 A 译文"
        decisions = [entry["conflict_resolution"]
                     for entry in core.load_project(local["project_id"])["promotion_log"]
                     if entry.get("conflict_resolution")]
        assert {item["adopted"] for item in decisions} == {"local"}


def test_bulk_resolution_defaults_to_keeping_local():
    """批量处理默认保留本地——批量最容易误伤，默认值必须保守。"""
    with project_env():
        local = _local_project()
        blob = _peer_export(glossary_target="协作者译名", tm_target="协作者 A 译文")
        core.import_project_memory(blob, name="生态学专著")
        pending = core.list_project_conflicts(local["project_id"])
        assert len(pending) == 2

        handled = core.resolve_all_project_conflicts(local["project_id"])  # 默认 False
        assert handled == 2
        assert core.list_project_conflicts(local["project_id"]) == []
        assert [e["preferred"] for e in
                core.load_project(local["project_id"])["glossary"]] == ["本地译名"], \
            "默认批量处理不得改变本地译名"

        # 显式采纳全部
        core.import_project_memory(blob, name="生态学专著")
        assert core.list_project_conflicts(local["project_id"]), "应重新产生冲突"
        core.resolve_all_project_conflicts(local["project_id"], adopt_incoming=True)
        assert [e["preferred"] for e in
                core.load_project(local["project_id"])["glossary"]] == ["协作者译名"]


def test_resolving_an_unknown_conflict_is_rejected():
    """处理不存在的冲突必须报错，而不是静默通过。"""
    with project_env():
        local = _local_project()
        with pytest.raises(ValueError):
            core.resolve_project_conflict(local["project_id"], "cf-does-not-exist")


def test_ui_offers_per_conflict_and_bulk_decisions():
    """界面：每条冲突都能单独决定，并提供批量动作。"""
    from streamlit.testing.v1 import AppTest

    with project_env():
        local = _local_project()
        blob = _peer_export(glossary_target="协作者译名", tm_target="协作者 A 译文")
        core.import_project_memory(blob, name="生态学专著")

        at = _project_detail(local["project_id"], tab="knowledge")
        assert not at.exception, [e.value for e in at.exception]

        labels = [b.label for b in at.button]
        assert "采纳导入版本" in labels and "保留本地版本" in labels, labels
        assert "全部保留本地" in labels and "全部采纳导入" in labels, labels
        page = "\n".join(str(m.value) for m in at.markdown)
        assert "待处理的导入冲突" in page
        assert "本地：" in page and "导入：" in page, "必须并列显示两个版本"

        take = next(b for b in at.button if b.label == "采纳导入版本")
        take.click()
        at.run()
        assert not at.exception, [e.value for e in at.exception]
        assert len(core.list_project_conflicts(local["project_id"])) == 1, \
            "只应处理被点击的那一条"


# ================= 读取不得写盘 =================


def test_reading_projects_never_writes_to_disk():
    """列出项目是纯读取；连"打开新建任务页"也不得在 outputs/ 下创建任何东西。

    曾经的实现：`list_projects()` 会顺手创建默认项目，于是任何只读访问
    （界面渲染、脚本、eval、测试）都会污染存储。这里把它固定成不变量。
    """
    with project_env() as tmp:
        assert core.list_projects() == []
        assert core.system_project_view()["project_id"] == core.system_project_id()
        assert any(item["project_id"] == core.system_project_id()
                   for item in core.list_active_project_options())
        assert core.project_memory_view()["project_id"] == core.system_project_id()
        assert not (tmp / "projects").exists(), "只读调用不得创建目录"

        from streamlit.testing.v1 import AppTest

        # 渲染"项目"页
        projects_page = AppTest.from_file(str(ROOT / "app.py"), default_timeout=60)
        projects_page.session_state["app_view"] = "projects"
        projects_page.run()
        assert not projects_page.exception, [e.value for e in projects_page.exception]
        assert not (tmp / "projects").exists(), "渲染项目页不得创建目录"
        # 空存储时也必须把系统工作区「未分类」呈现出来（它是真实容器）。
        assert core.SYSTEM_PROJECT_NAME in "\n".join(
            str(m.value) for m in projects_page.markdown)

        # 渲染"新建任务"第 1 步（这里是原先的泄漏点）
        job = core.new_job_state("picker.docx")
        core.save_job_state("picker", job)
        new_task = AppTest.from_file(str(ROOT / "app.py"), default_timeout=60)
        new_task.session_state["app_view"] = "new"
        new_task.session_state["task_step"] = 1
        new_task.session_state["task_files"] = [{"name": "picker.docx",
                                                "bytes": b"content"}]
        new_task.run()
        assert not new_task.exception, [e.value for e in new_task.exception]
        assert not (tmp / "projects").exists(), \
            "渲染新建任务页不得创建项目目录"
        # 新建任务页只读地显示项目上下文（未选择 → 未分类任务），不自己造
        # 第二个选择器，因此也没有"读一次就落盘"的副作用。
        page = "\n".join(str(m.value) for m in new_task.markdown)
        assert "项目上下文" in page
        assert "未分类任务" in page
        assert not any(s.label == "所属项目" for s in new_task.selectbox)

        # 真正发生变更时才会落盘
        core.ensure_default_project()
        assert (tmp / "projects").is_dir()
