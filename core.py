"""Folith 核心逻辑层（与 Streamlit UI 解耦，便于测试）。

职责：大模型路由、文档清洗、术语抽取、双语翻译、报告生成、任务进度持久化。
"""
import hashlib
import io
import json
import csv
import inspect
from concurrent.futures import ThreadPoolExecutor
import os
import queue
import re
import shutil
import subprocess
import threading
import tempfile
import unicodedata
import time
import traceback
import uuid
import zipfile
from datetime import datetime, timedelta, timezone
from html import escape as html_escape
from pathlib import Path
from urllib.parse import urlparse

import fitz  # PyMuPDF
import httpx
import pandas as pd
from defusedxml import ElementTree as ET
from docx import Document
from google import genai
from openai import OpenAI

from transpraxis import models as _models
from transpraxis import state_migration as _state_migration
from transpraxis import context as _context
from transpraxis import knowledge as _knowledge
from transpraxis import repair as _repair
from transpraxis import translation_evidence as _translation_evidence
from transpraxis import checkpoint as _checkpoint
from transpraxis import snapshots as _snapshots
from transpraxis import entity_registry as _entity_registry
from transpraxis import model_roles as _model_roles
from transpraxis import pdf_ingestion as _pdf_ingestion
from transpraxis import source_cleanup as _source_cleanup
from transpraxis import segmentation as _segmentation
from transpraxis import source_quality as _source_quality
from transpraxis import translation_protocol as _translation_protocol
from transpraxis.storage import job_repository as _job_repo
from transpraxis import translation_target as _translation_target
from transpraxis import finalization as _finalization
from transpraxis import rendered_qa as _rendered_qa
from transpraxis import project as _project
from transpraxis import translation_memory as _tm_scope
from transpraxis import usage as _usage
from transpraxis.performance import PipelineProfiler
from transpraxis.textual import has_textual_content, normalize_language

# ================= 常量 =================
# 任务进度与过程文件的本地存储目录（已加入 .gitignore）
OUTPUT_DIR = Path(os.environ.get("FOLITH_OUTPUT_DIR") or os.environ.get("FOLIOTHREAD_OUTPUT_DIR", "outputs"))

DELIVERY_CONFIG_DEFAULTS = {
    "enable_annotate": False,
    "enable_report": False,
    "deliver_plain_docx": True,
    "deliver_bilingual_docx": True,
    "deliver_pdf": False,
    "deliver_terms_xlsx": True,
    "deliver_tbx": False,
    "deliver_tmx": False,
    "deliver_jsonl": False,
    "deliver_evidence": True,
    "deliver_cases": False,
    "deliver_academic_workspace": False,
    "deliver_review_report": False,
}


def default_delivery_config():
    return dict(DELIVERY_CONFIG_DEFAULTS)


def normalize_delivery_config(config, *, enable_report=None, enable_annotate=None):
    normalized = default_delivery_config()
    if isinstance(config, dict):
        normalized.update({key: bool(config[key]) for key in normalized if key in config})
    if enable_report is not None:
        normalized["enable_report"] = bool(enable_report)
    if enable_annotate is not None:
        normalized["enable_annotate"] = bool(enable_annotate)
    return normalized

# 提供商注册表：kind=openai_compat 走 OpenAI SDK 兼容接口（官方与中转站通用）。
# base_url 为 None 表示使用官方 SDK 默认路由；custom_base_url 表示地址由用户在
# UI 中填写（通用中转站）。OpenCode Go 只列出官方标记为
# /chat/completions 的模型；/messages 与 /responses 模型不走此路由。
PROVIDERS = {
    "OpenCode Go": {
        "kind": "openai_compat",
        "base_url": "https://opencode.ai/zen/go/v1",
        "proxy_bypass": True,
        "default_model": "glm-5.2",
        "models": [
            "glm-5.2", "glm-5.1",
            "deepseek-v4-pro", "deepseek-v4-flash",
            "kimi-k3", "kimi-k2.7-code", "kimi-k2.6",
            "mimo-v2.5-pro", "mimo-v2.5", "hy3", "grok-4.5",
        ],
    },
    "DeepSeek": {
        "kind": "openai_compat",
        "base_url": "https://api.deepseek.com",
        "default_model": "deepseek-v4-flash",
        "models": ["deepseek-v4-flash", "deepseek-v4-pro"],
    },
    "OpenAI": {
        "kind": "openai",
        # Keep the stable general models and expose current reasoning families
        # so the conditional reasoning-effort control has a real target.
        "models": [
            "gpt-4o-mini", "gpt-4o", "gpt-4.1-mini", "gpt-4.1",
            "gpt-5.2", "gpt-5.1", "gpt-5", "gpt-5-mini", "o3", "o4-mini",
        ],
        "capabilities": {
            "supports_json_schema": True,
            "supports_json_object": True,
            "supports_response_format": True,
        },
    },
    "Gemini": {
        "kind": "gemini",
        "default_model": "gemini-3.6-flash",
        "models": ["gemini-3.6-flash", "gemini-3.5-flash",
                   "gemini-3.5-flash-lite", "gemini-2.5-flash", "gemini-2.5-pro"],
    },
    "OpenRouter": {
        "kind": "openai_compat",
        "base_url": "https://openrouter.ai/api/v1",
        "models": [],
        "model_hint": "如 anthropic/claude-sonnet-4、openai/gpt-5、google/gemini-3-flash",
    },
    "SiliconFlow": {
        "kind": "openai_compat",
        "base_url": "https://api.siliconflow.cn/v1",
        "models": [],
        "model_hint": "如 Qwen/Qwen3-235B-A22B、deepseek-ai/DeepSeek-V3.2",
    },
    "Moonshot (Kimi)": {
        "kind": "openai_compat",
        "base_url": "https://api.moonshot.cn/v1",
        "models": [],
        "model_hint": "如 kimi-k2.5、moonshot-v1-8k",
    },
    "Zhipu (GLM)": {
        "kind": "openai_compat",
        "base_url": "https://open.bigmodel.cn/api/paas/v4",
        "models": [],
        "model_hint": "如 glm-4.5、glm-5",
    },
    "Qwen (DashScope)": {
        "kind": "openai_compat",
        "base_url": "https://dashscope.aliyuncs.com/compatible-mode/v1",
        "models": [],
        "model_hint": "如 qwen-max、qwen3-235b-a22b",
    },
    "自定义中转站": {
        "kind": "openai_compat",
        "base_url": None,
        "custom_base_url": True,
        "models": [],
        "model_hint": "填写中转站提供的模型名（OpenAI /chat/completions 兼容）",
    },
}

# 兼容旧引用：仅保留 模型名 -> 列表 的视图
MODELS = {name: cfg["models"] for name, cfg in PROVIDERS.items()}

# 会话线程级中转地址：自定义中转站的 base_url 由 UI 写入，call_llm 自动优先使用。
# 每个 Streamlit 会话线程独立，多设备同时使用不会互相串扰。
_LLM_CTX = threading.local()

# Runtime status is deliberately file-backed: the UI can poll it while the
# worker is inside a blocking provider request, and a browser refresh does not
# erase the last known activity.
_RUNTIME_CTX = threading.local()
_RUNTIME_LOCK = threading.RLock()
_RUNTIME_WORKERS = {}
_RUNTIME_WORKERS_LOCK = threading.RLock()


def normalize_openai_base_url(base_url):
    """接受基址或误填的 OpenAI 兼容具体端点，并统一回基址。"""
    base = (base_url or "").strip().rstrip("/")
    for suffix in ("/chat/completions", "/completions", "/responses", "/models"):
        if base.endswith(suffix):
            base = base[:-len(suffix)]
            break
    if not base:
        return ""
    parsed = urlparse(base)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        return ""
    return base


def resolve_openai_base_url(provider, base_url=None):
    """Resolve an OpenAI-compatible base URL without silently falling back.

    A custom relay has no registry URL.  Falling through to the OpenAI SDK's
    default in that case sends a relay credential to api.openai.com and turns
    a configuration mistake into a misleading ``invalid_api_key`` response.
    """
    cfg = PROVIDERS.get(provider) or {}
    explicit = base_url is not None
    resolved = normalize_openai_base_url(
        base_url if explicit else cfg.get("base_url"))
    if cfg.get("custom_base_url") and explicit and not resolved:
        raise ValueError("自定义中转站的 API 地址无效（需要 http(s)://…/v1）")
    if cfg.get("custom_base_url") and not explicit and not resolved:
        resolved = normalize_openai_base_url(getattr(_LLM_CTX, "base_url", None))
    if cfg.get("custom_base_url") and not resolved:
        raise ValueError("自定义中转站未配置有效的 API 地址（需要 http(s)://…/v1）")
    return resolved


def set_llm_base_url(base_url):
    """为当前线程设置 OpenAI 兼容中转站地址（空值清除）。"""
    _LLM_CTX.base_url = normalize_openai_base_url(base_url) or None
    return _LLM_CTX.base_url


_REASONING_EFFORT_VALUES = ("none", "minimal", "low", "medium", "high", "xhigh", "max")


def reasoning_effort_options(provider, model):
    """Return safe reasoning-effort values for a known reasoning model.

    OpenAI-compatible relays do not expose a portable capability schema.  We
    therefore only offer the control for recognizable reasoning model names;
    unknown custom models remain in automatic mode so an unsupported request
    parameter cannot break translation.
    """
    cfg = PROVIDERS.get(provider) or {}
    if cfg.get("kind") not in {"openai", "openai_compat"}:
        return ()
    name = str(model or "").strip().lower()
    if not name:
        return ()
    # OpenAI reasoning families and common compatible relay names.  The UI
    # deliberately uses the portable low/medium/high subset.
    prefixes = ("o1", "o3", "o4-mini", "gpt-5", "gpt-6", "deepseek-r1", "qwq")
    if not name.startswith(prefixes):
        return ()
    return ("low", "medium", "high")


def set_llm_reasoning_effort(value):
    """Set or clear the per-thread Chat Completions reasoning effort."""
    value = str(value or "").strip().lower()
    _LLM_CTX.reasoning_effort = value if value in _REASONING_EFFORT_VALUES else None
    return _LLM_CTX.reasoning_effort


def _runtime_job_id():
    return getattr(_RUNTIME_CTX, "job_id", None)


def _runtime_cancel_requested(job_id):
    return bool((load_runtime_state(job_id) or {}).get("cancel_requested"))


# ================= 基础工具函数 =================
def clean_xml_chars(text):
    if not isinstance(text, str):
        return str(text)
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]', '', text)


def parse_json_array(text):
    """从 LLM 输出中稳健地解析 JSON 数组；解析失败返回 None。

    依次尝试：整体解析 -> 去掉 Markdown 代码块后解析 -> 从每个 '[' 位置做
    raw_decode（可容忍输出前后夹带解释文字）。
    """
    if not isinstance(text, str) or not text.strip():
        return None
    candidate = text.strip()
    candidate = re.sub(r'^```(?:json)?\s*', '', candidate, flags=re.DOTALL)
    candidate = re.sub(r'\s*```$', '', candidate, flags=re.DOTALL).strip()

    try:
        obj = json.loads(candidate)
        if isinstance(obj, list):
            return obj
    except Exception:
        pass

    decoder = json.JSONDecoder()
    for m in re.finditer(r'\[', candidate):
        try:
            obj, _ = decoder.raw_decode(candidate[m.start():])
        except Exception:
            continue
        if isinstance(obj, list):
            return obj
    return None


def parse_translation_array(res, expected):
    """Compatibility API backed by the canonical translation parser."""
    return _translation_protocol.parse_translation_array(res, expected)

def is_rate_limited(err):
    s = str(err)
    return '429' in s or 'RESOURCE_EXHAUSTED' in s or 'rate limit' in s.lower()


# ================= PDF 确定性段落提取 =================
# 经验（来自 localize-anything 与全书实测）：分段/清洗必须确定性，不能交给 LLM。
# 旧流程把 ~2500 字符的任意文本块交给模型"清洗"，导致两类系统性缺陷：
#   1. 块边界落在句中 -> 句子被拦腰截断（"…pecking at" / "crumbs and bones…"）；
#   2. 模型自由裁量 -> 对白、引语被随意拆分或合并，分段结果不可复现。
# 新流程直接读 PDF 版面：块(block)->行(line)->首行缩进判定段落，连字符修复、
# 跨页段落合并、页眉页脚/页码剔除全部确定完成。

# 句末终结符（用于判断段落是否未完结、需要与下一段合并）
_SENTENCE_TERMINAL = set('.!?"”’…:;)')

# 纯装饰符号行（章节分隔花饰等），无字母/数字，不是正文
_ORNAMENT_RE = re.compile(r"^[\s*•·▪◦‣❦❧—–\-]{1,12}$")

# 常见缩写（句点不计入句界）
_ABBREV_RE = re.compile(
    r"\b(?:Lt|Col|Gen|Maj|Capt|Sgt|Brig|Mr|Mrs|Ms|Dr|St|No|Vol|pp|"
    r"e\.g|i\.e|vs|etc|a\.m|p\.m|U\.S|A\.F|B\.C|A\.D)\.", re.IGNORECASE)

OCR_WORKERS_DEFAULT = 2
OCR_QUEUE_SIZE_DEFAULT = 4
OCR_CONFIDENCE_THRESHOLD_DEFAULT = 88.0
OCR_CHECKPOINT_VERSION = 1
_OCR_PAGE_NUMBER_RE = re.compile(
    r"^\s*(?:page\s+)?\d{1,4}(?:\s*(?:of|/)\s*\d{1,4})?\s*$",
    re.IGNORECASE,
)


def extract_pdf_paragraphs(file_bytes):
    """Compatibility API backed by the layout-aware PDF ingestion module."""
    return _pdf_ingestion.extract_pdf_paragraphs(file_bytes)


def _find_tesseract():
    """Find Tesseract even when the app was launched from Finder/Explorer.

    GUI launchers often receive a smaller PATH than an interactive shell.  The
    explicit environment variable is useful for portable installs and takes
    precedence over the usual PATH/Homebrew/Windows locations.
    """
    configured = os.environ.get("FOLIOTHREAD_TESSERACT_CMD", "").strip()
    candidates = [configured, shutil.which("tesseract")]
    if os.name == "nt":
        candidates.extend([
            r"C:\Program Files\Tesseract-OCR\tesseract.exe",
            r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
        ])
    else:
        candidates.extend([
            "/opt/homebrew/bin/tesseract",
            "/usr/local/bin/tesseract",
            "/usr/bin/tesseract",
        ])
    for candidate in candidates:
        if not candidate:
            continue
        path = Path(candidate)
        if path.is_file() and os.access(path, os.X_OK):
            return str(path)
    return ""


def _tesseract_languages(command):
    """Return installed OCR languages, ignoring Tesseract control models."""
    try:
        proc = subprocess.run(
            [command, "--list-langs"], capture_output=True, text=True,
            encoding="utf-8", errors="replace", timeout=15,
        )
    except Exception:  # noqa: BLE001 - diagnostics must not block OCR
        return []
    if proc.returncode != 0:
        return []
    languages = []
    for line in (proc.stdout or "").splitlines():
        value = line.strip()
        if value and re.fullmatch(r"[A-Za-z0-9_]+", value):
            languages.append(value)
    return languages


def _select_tesseract_language(installed):
    """Choose a useful installed language set without assuming chi_sim exists."""
    installed = set(installed or [])
    # The product handles Chinese and English documents most often.  Keep both
    # when available, but do not make a missing chi_sim pack break English OCR.
    for preferred in (
        ("chi_sim", "eng"),
        ("eng",),
        ("chi_sim",),
        ("jpn", "eng"),
        ("kor", "eng"),
        ("ara", "eng"),
        ("rus", "eng"),
        ("deu", "eng"),
        ("fra", "eng"),
        ("spa", "eng"),
    ):
        if set(preferred).issubset(installed):
            return "+".join(preferred)
    fallback = sorted(installed - {"osd", "snum"})
    return fallback[0] if fallback else ""


def _ocr_workers_default():
    try:
        configured = int(os.environ.get("FOLIOTHREAD_OCR_WORKERS", ""))
    except ValueError:
        configured = OCR_WORKERS_DEFAULT
    return max(1, min(configured or OCR_WORKERS_DEFAULT, 8))


def _ocr_queue_size_default(workers):
    try:
        configured = int(os.environ.get("FOLIOTHREAD_OCR_QUEUE_SIZE", ""))
    except ValueError:
        configured = OCR_QUEUE_SIZE_DEFAULT
    return max(1, min(configured or workers * 2, 32))


def _runtime_int_option(explicit, saved, env_name, default, *, minimum=1, maximum=32):
    value = explicit
    if value is None:
        value = saved
    if value is None:
        value = os.environ.get(env_name, "")
    try:
        value = int(value)
    except (TypeError, ValueError):
        value = int(default)
    return max(minimum, min(value, maximum))


def _runtime_float_option(explicit, saved, env_name, default, *, minimum=0.0, maximum=300.0):
    value = explicit
    if value is None:
        value = saved
    if value is None:
        value = os.environ.get(env_name, "")
    try:
        value = float(value)
    except (TypeError, ValueError):
        value = float(default)
    return max(minimum, min(value, maximum))


def _ocr_runtime_metadata(language, workers):
    """Describe the provider actually executing OCR, not only a config value."""
    configured_device = str(os.environ.get("FOLIOTHREAD_OCR_DEVICE", "auto") or "auto").lower()
    metadata = {
        "backend": "tesseract",
        "model": "tesseract-lstm",
        "device": "cpu",
        "provider": "tesseract",
        "execution_provider": "native_cpu",
        "gpu_name": None,
        "language": language,
        "workers": workers,
    }
    warnings = []
    if configured_device not in {"", "auto", "cpu"}:
        warnings.append(
            f"OCR 配置请求 device={configured_device}，但当前 Tesseract runtime 实际使用 CPU；"
            "未发生 GPU 加速。"
        )
    return metadata, warnings


def _ocr_join_lines(left, right):
    left = re.sub(r"\s+", " ", str(left or "").strip())
    right = re.sub(r"\s+", " ", str(right or "").strip())
    if not left:
        return right
    if not right:
        return left
    if left.endswith(("-", "‐", "‑")) and right[:1].islower():
        return left[:-1] + right
    if (re.search(r"[\u3400-\u9fff]$", left)
            or re.match(r"[\u3400-\u9fff]", right)):
        return left + right
    return f"{left} {right}"


def _parse_tesseract_tsv(payload):
    """Convert word-level TSV into paragraph-like blocks with real confidence."""
    rows = {}
    try:
        reader = csv.DictReader(str(payload or "").splitlines(), delimiter="\t")
        for row in reader:
            text = str(row.get("text") or "").strip()
            if not text:
                continue
            try:
                key = (
                    int(row.get("block_num") or 0),
                    int(row.get("par_num") or 0),
                    int(row.get("line_num") or 0),
                )
                confidence = float(row.get("conf") or -1)
                left = float(row.get("left") or 0)
                top = float(row.get("top") or 0)
                width = float(row.get("width") or 0)
                height = float(row.get("height") or 0)
            except (TypeError, ValueError):
                continue
            rows.setdefault(key, []).append({
                "text": text, "confidence": confidence,
                "left": left, "top": top, "right": left + width,
                "bottom": top + height,
            })
    except (TypeError, ValueError):
        return []
    grouped = {}
    for (block_num, par_num, line_num), words in rows.items():
        words.sort(key=lambda word: (word["left"], word["top"]))
        valid = [word["confidence"] for word in words if word["confidence"] >= 0]
        grouped.setdefault((block_num, par_num), []).append({
            "line_num": line_num,
            "text": _ocr_join_lines("", " ".join(word["text"] for word in words)),
            "confidence": sum(valid) / len(valid) if valid else None,
            "top": min(word["top"] for word in words),
            "bottom": max(word["bottom"] for word in words),
        })
    blocks = []
    for key, lines in grouped.items():
        lines.sort(key=lambda line: line["line_num"])
        text = ""
        confidences = []
        for line in lines:
            text = _ocr_join_lines(text, line["text"])
            if line["confidence"] is not None:
                confidences.append(line["confidence"])
        if not text:
            continue
        blocks.append({
            "text": text,
            "confidence": sum(confidences) / len(confidences) if confidences else None,
            "low_confidence_ratio": (
                sum(1 for value in confidences if value < OCR_CONFIDENCE_THRESHOLD_DEFAULT)
                / len(confidences) if confidences else None
            ),
            "top": min(line["top"] for line in lines),
            "bottom": max(line["bottom"] for line in lines),
        })
    return blocks


def _ocr_page_with_tesseract(command, png_bytes, language):
    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
        tmp.write(png_bytes)
        tmp.flush()
        tmp_path = tmp.name
    try:
        proc = subprocess.run(
            [command, tmp_path, "stdout", "-l", language, "--psm", "3", "tsv"],
            capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120,
        )
        if proc.returncode != 0:
            detail = (proc.stderr or "").strip().splitlines()
            return None, detail[-1][:180] if detail else "无输出"
        blocks = _parse_tesseract_tsv(proc.stdout)
        if not blocks:
            return None, "未生成带置信度的 OCR 文本"
        return blocks, None
    except subprocess.TimeoutExpired:
        return None, "超过 120 秒"
    except Exception as exc:  # noqa: BLE001 - one page must not abort the book
        return None, str(exc)[:180]
    finally:
        Path(tmp_path).unlink(missing_ok=True)


def _ocr_checkpoint_path(checkpoint_dir):
    return Path(checkpoint_dir) / "ocr_checkpoint.json" if checkpoint_dir else None


def _load_ocr_checkpoint(path, source_hash, indices, language):
    if path is None or not path.is_file():
        return {}
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, TypeError, ValueError):
        return {}
    if not isinstance(value, dict) or value.get("version") != OCR_CHECKPOINT_VERSION:
        return {}
    if value.get("source_sha256") != source_hash or value.get("language") != language:
        return {}
    if value.get("indices") != list(indices):
        return {}
    pages = value.get("pages")
    return pages if isinstance(pages, dict) else {}


def _save_ocr_checkpoint(path, source_hash, indices, language, pages, status="running"):
    if path is None:
        return
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        temporary = path.with_name(path.name + ".tmp")
        temporary.write_text(json.dumps({
            "version": OCR_CHECKPOINT_VERSION,
            "source_sha256": source_hash,
            "indices": list(indices),
            "language": language,
            "pages": pages,
            "status": status,
            "updated_at": _utc_now_iso(),
        }, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(path)
    except OSError:
        pass


def _ocr_pdf_text_with_warnings(
    file_bytes, max_pages=None, on_progress=None, *, workers=None, queue_size=None,
    checkpoint_dir=None, cancel_check=None, profiler=None, return_details=False,
):
    """OCR a scanned PDF with Tesseract and return ``(text, warnings)``.

    ``max_pages=3`` is used by the quick-profile preview.  ``None`` means all
    pages and is used by the actual translation pipeline.  The function never
    sends page images or OCR text anywhere; it only invokes the local
    Tesseract executable.
    """
    command = _find_tesseract()
    if not command:
        result = ("", [
            "未找到本机 Tesseract OCR。请安装 Tesseract，或设置 "
            "FOLIOTHREAD_TESSERACT_CMD 指向 tesseract 可执行文件。"
        ], {})
        return result if return_details else result[:2]

    warnings = []
    installed = _tesseract_languages(command)
    language = _select_tesseract_language(installed)
    if not language:
        result = ("", [
            "Tesseract 未找到可用语言包（至少需要 eng 或 chi_sim），无法执行 OCR。"
        ], {})
        return result if return_details else result[:2]

    workers = max(1, min(int(workers or _ocr_workers_default()), 8))
    queue_size = max(1, min(int(queue_size or _ocr_queue_size_default(workers)), 32))
    runtime_metadata, runtime_warnings = _ocr_runtime_metadata(language, workers)
    warnings.extend(runtime_warnings)
    source_hash = hashlib.sha256(file_bytes or b"").hexdigest()
    checkpoint_path = _ocr_checkpoint_path(checkpoint_dir)

    try:
        with fitz.open(stream=file_bytes, filetype="pdf") as doc:
            count = doc.page_count
            if count == 0:
                result = ("", ["PDF 不包含任何页面，无法执行 OCR。"], {})
                return result if return_details else result[:2]
            if max_pages is None:
                indices = list(range(count))
            else:
                indices = sorted({0, count // 2, count - 1})[:max_pages]
            page_records = _load_ocr_checkpoint(
                checkpoint_path, source_hash, indices, language)
            pending_indices = [index for index in indices if str(index) not in page_records]
            failed_pages = []
            render_stage = profiler.start_stage(
                "rasterize", concurrency=1, page_count=len(indices),
                metadata={"dpi": 220, "queue_size": queue_size,
                          "cached_pages": len(indices) - len(pending_indices)},
            ) if profiler is not None else None
            ocr_stage = profiler.start_stage(
                "ocr", concurrency=workers, page_count=len(indices),
                metadata={**runtime_metadata, "queue_size": queue_size,
                          "cached_pages": len(indices) - len(pending_indices)},
            ) if profiler is not None else None
            page_lock = threading.RLock()
            queue_lock = threading.RLock()
            work_queue = queue.Queue(maxsize=queue_size)
            stop_event = threading.Event()
            worker_errors = []
            completed_pages = len(indices) - len(pending_indices)

            def check_cancel():
                if cancel_check and cancel_check():
                    raise RuntimeError("任务已请求取消")

            def emit_page_progress():
                if not on_progress:
                    return
                with queue_lock:
                    current = completed_pages
                try:
                    on_progress(
                        f"【阶段一】扫描 PDF OCR（{current}/{len(indices)} 页；"
                        f"workers={workers}；device={runtime_metadata['device']}）…"
                    )
                except Exception:  # noqa: BLE001 - UI progress is optional
                    pass

            def ocr_worker():
                nonlocal completed_pages
                while True:
                    task = work_queue.get()
                    try:
                        if task is None:
                            return
                        idx, png_bytes, enqueued_at, page_width, page_height = task
                        check_cancel()
                        started = time.perf_counter()
                        blocks, error = _ocr_page_with_tesseract(command, png_bytes, language)
                        finished = time.perf_counter()
                        if error:
                            with page_lock:
                                failed_pages.append((idx + 1, error))
                        else:
                            record = {
                                "page_number": idx + 1,
                                "page_width": page_width,
                                "page_height": page_height,
                                "blocks": blocks,
                            }
                            with page_lock:
                                page_records[str(idx)] = record
                                _save_ocr_checkpoint(
                                    checkpoint_path, source_hash, indices, language,
                                    page_records,
                                )
                        with queue_lock:
                            completed_pages += 1
                        if ocr_stage is not None:
                            ocr_stage.item(
                                f"page-{idx + 1}", started, finished_monotonic=finished,
                                queue_wait_s=max(0.0, started - enqueued_at),
                                status="completed" if not error else "failed",
                                metadata={"kind": "page", "page_number": idx + 1},
                            )
                        emit_page_progress()
                    except BaseException as exc:  # noqa: BLE001 - propagate cancellation
                        worker_errors.append(exc)
                        stop_event.set()
                    finally:
                        work_queue.task_done()

            threads = [threading.Thread(target=ocr_worker, name=f"ocr-worker-{i}", daemon=True)
                       for i in range(workers)]
            cancelled = False
            try:
                for thread in threads:
                    thread.start()
                emit_page_progress()
                for idx in pending_indices:
                    check_cancel()
                    if stop_event.is_set():
                        break
                    started = time.perf_counter()
                    pix = doc[idx].get_pixmap(dpi=220, alpha=False)
                    png_bytes = pix.tobytes("png")
                    rendered_at = time.perf_counter()
                    queued_at = rendered_at
                    while True:
                        try:
                            page_width, page_height = float(doc[idx].rect.width), float(doc[idx].rect.height)
                            work_queue.put(
                                (idx, png_bytes, queued_at, page_width, page_height),
                                timeout=0.2,
                            )
                            queued_at = time.perf_counter()
                            break
                        except queue.Full:
                            check_cancel()
                            if stop_event.is_set():
                                break
                    if render_stage is not None:
                        render_stage.item(
                            f"page-{idx + 1}", started,
                            finished_monotonic=queued_at,
                            queue_wait_s=max(0.0, queued_at - rendered_at),
                            metadata={"kind": "page", "page_number": idx + 1},
                        )
                    if stop_event.is_set():
                        break
                while not stop_event.is_set() and completed_pages < len(indices):
                    check_cancel()
                    time.sleep(0.1)
                if worker_errors:
                    raise worker_errors[0]
            except BaseException:
                cancelled = True
                stop_event.set()
                raise
            finally:
                stop_event.set() if worker_errors else None
                for _ in threads:
                    while True:
                        try:
                            work_queue.put(None, timeout=0.2)
                            break
                        except queue.Full:
                            if not any(thread.is_alive() for thread in threads):
                                break
                for thread in threads:
                    thread.join(timeout=10)
                if render_stage is not None:
                    render_stage.finish(status="cancelled" if cancelled else "completed")
                if ocr_stage is not None:
                    ocr_stage.finish(status="cancelled" if cancelled else "completed")
                _save_ocr_checkpoint(
                    checkpoint_path, source_hash, indices, language, page_records,
                    status="cancelled" if cancelled else "completed",
                )

            if failed_pages:
                sample = "；".join(f"第 {page} 页：{detail}" for page, detail in failed_pages[:3])
                more = f"（另有 {len(failed_pages) - 3} 页失败）" if len(failed_pages) > 3 else ""
                warnings.append(f"OCR 部分页面失败：{sample}{more}")
            paragraph_confidences = []
            chunks = []
            repeated = {}
            all_blocks = []
            for idx in indices:
                record = page_records.get(str(idx)) or {}
                height = float(record.get("page_height") or 0)
                for block in record.get("blocks") or []:
                    text = re.sub(r"\s+", " ", str(block.get("text") or "").strip())
                    if not text:
                        continue
                    normalized = re.sub(r"\d+", "#", text).casefold()
                    repeated[normalized] = repeated.get(normalized, 0) + 1
                    all_blocks.append((idx, height, block, text, normalized))
            repeat_threshold = max(2, int((len(indices) * 0.2) + 0.999)) if len(indices) > 1 else 99
            by_page = {}
            for idx, height, block, text, normalized in all_blocks:
                top = float(block.get("top") or 0)
                bottom = float(block.get("bottom") or 0)
                is_edge = top < height * 0.12 or bottom > height * 0.88
                if _OCR_PAGE_NUMBER_RE.fullmatch(text) and is_edge:
                    continue
                if repeated.get(normalized, 0) >= repeat_threshold and is_edge:
                    continue
                by_page.setdefault(idx, []).append((text, block.get("confidence")))
            for idx in indices:
                page_chunks = by_page.get(idx) or []
                if not page_chunks:
                    continue
                chunks.append("\n\n".join(text for text, _ in page_chunks))
                paragraph_confidences.extend(conf for _, conf in page_chunks)
            valid_confidences = [float(value) for value in paragraph_confidences
                                 if value is not None]
            details = {
                **runtime_metadata,
                "page_count": len(indices),
                "completed_pages": sum(1 for index in indices if str(index) in page_records),
                "page_order": [index + 1 for index in indices
                               if str(index) in page_records],
                "failed_pages": len(failed_pages),
                "checkpoint_file": str(checkpoint_path) if checkpoint_path else "",
                "paragraph_confidences": paragraph_confidences,
                "mean_confidence": (
                    sum(valid_confidences) / len(valid_confidences)
                    if valid_confidences else None
                ),
                "confidence_threshold": OCR_CONFIDENCE_THRESHOLD_DEFAULT,
            }
            text = "\n\n".join(chunks)
            if not text:
                warnings.append("OCR 未识别到有效文本。")
            result = (text, warnings, details)
            return result if return_details else result[:2]
    except Exception as exc:  # noqa: BLE001 - caller can show a recoverable warning
        if "任务已请求取消" in str(exc):
            raise
        result = ("", [f"OCR 处理 PDF 失败：{str(exc)[:240]}"], {})
        return result if return_details else result[:2]


def _ocr_pdf_text(file_bytes, max_pages=3):
    """Compatibility wrapper used by the quick-profile path."""
    text, _warnings = _ocr_pdf_text_with_warnings(file_bytes, max_pages=max_pages)
    return text


def _ocr_text_to_paragraphs(text):
    """Turn Tesseract page/block output into translation-ready paragraphs."""
    paragraphs = []
    for block in re.split(r"\n\s*\n", str(text or "")):
        lines = [clean_xml_chars(line.strip())
                 for line in block.splitlines() if line.strip()]
        if not lines:
            continue
        merged = lines[0]
        for line in lines[1:]:
            if merged.endswith("-") and line[:1].islower():
                merged = merged[:-1] + line
            elif re.search(r"[\u3400-\u9fff]$", merged) or \
                    re.match(r"[\u3400-\u9fff]", line):
                merged += line
            else:
                merged += " " + line
        if len(merged) > 1 and has_textual_content(merged):
            paragraphs.append(merged)
    return paragraphs


EXTRACTION_REPORT_VERSION = 1
_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_M_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"


def _docx_part_texts(document, needle):
    """读取 word/<needle> 部件里的全部文本（脚注/尾注/批注不在正文 XML 里）。"""
    try:
        package = getattr(document.part, "package", None)
        if package is None:
            return []
        blob = None
        for part in package.parts:
            if needle in str(getattr(part, "partname", "")):
                blob = part.blob
                break
        if not blob:
            return []
        from lxml import etree
        root = etree.fromstring(bytes(blob))
    except Exception:  # noqa: BLE001 - report is best-effort, never blocks import
        return []
    return [str(node.text or "").strip() for node in root.iter(f"{_W_NS}t")
            if str(node.text or "").strip()]


def _docx_unsupported_parts(document):
    """Inspect a DOCX for content a paragraph-only extractor cannot translate.

    The point is not to support these structures yet: it is to make sure the
    user is never told "文档已导入" while tables, running heads or footnotes were
    silently dropped.  Counting is deliberately structural (no text heuristics).
    """
    parts = []

    def add(kind, label, count, detail, sample=""):
        if count <= 0:
            return
        parts.append({"kind": kind, "label": label, "count": int(count),
                      "detail": detail, "sample": str(sample or "")[:160]})

    try:
        body = document.element.body
        tables = list(getattr(document, "tables", []) or [])
        table_cells = sum(len(row.cells) for table in tables for row in table.rows)
        table_sample = next(
            (str(cell.text).strip() for table in tables for row in table.rows
             for cell in row.cells if str(cell.text).strip()), "")
        add("table", "表格内容", len(tables),
            f"{len(tables)} 个表格、约 {table_cells} 个单元格；表格文字不进入翻译",
            table_sample)

        headers = footers = 0
        header_sample = footer_sample = ""
        for section in getattr(document, "sections", []) or []:
            for paragraph in getattr(section.header, "paragraphs", []) or []:
                text = str(paragraph.text or "").strip()
                if text:
                    headers += 1
                    header_sample = header_sample or text
            for paragraph in getattr(section.footer, "paragraphs", []) or []:
                text = str(paragraph.text or "").strip()
                if text:
                    footers += 1
                    footer_sample = footer_sample or text
        add("header", "页眉", headers,
            f"页眉有 {headers} 段文字（重复页眉是扫描件最常见的问题来源）",
            header_sample)
        add("footer", "页脚", footers,
            f"页脚有 {footers} 段文字（通常含页码）", footer_sample)

        boxes = body.findall(f".//{_W_NS}txbxContent")
        box_sample = next(
            (" ".join(str(node.text or "") for node in box.iter(f"{_W_NS}t")).strip()
             for box in boxes
             if " ".join(str(node.text or "") for node in box.iter(f"{_W_NS}t")).strip()),
            "")
        add("textbox", "文本框", len(boxes),
            f"{len(boxes)} 个文本框；文本框文字不在正文段落里", box_sample)

        images = body.findall(f".//{_A_NS}blip")
        add("image", "图片", len(images),
            f"{len(images)} 张内嵌图片按原样保留在源文件里，但图片内的文字不参与翻译")

        equations = body.findall(f".//{_M_NS}oMath")
        add("equation", "公式", len(equations),
            f"{len(equations)} 处公式不参与翻译")

        footnotes = _docx_part_texts(document, "footnotes.xml")
        add("footnote", "脚注", len(footnotes),
            f"{len(footnotes)} 条脚注不参与翻译", footnotes[0] if footnotes else "")
        endnotes = _docx_part_texts(document, "endnotes.xml")
        add("endnote", "尾注", len(endnotes),
            f"{len(endnotes)} 条尾注不参与翻译", endnotes[0] if endnotes else "")
        comments = _docx_part_texts(document, "comments.xml")
        add("comment", "批注", len(comments),
            f"{len(comments)} 条批注不参与翻译", comments[0] if comments else "")
    except Exception:  # noqa: BLE001
        return parts
    return parts


def _pdf_extraction_facts(file_bytes):
    """Page/image facts for the import report (never raises)."""
    facts = {"pages": None, "images": 0, "text_pages": 0, "empty_pages": 0}
    try:
        with fitz.open(stream=file_bytes, filetype="pdf") as document:
            facts["pages"] = document.page_count
            for page in document:
                if str(page.get_text() or "").strip():
                    facts["text_pages"] += 1
                else:
                    facts["empty_pages"] += 1
                try:
                    facts["images"] += len(page.get_images(full=True))
                except Exception:  # noqa: BLE001
                    continue
    except Exception:  # noqa: BLE001
        return facts
    return facts


def _build_extraction_report(filename, paragraphs, *, ocr_used=False,
                             unsupported=None, facts=None, warnings=None):
    values = [str(item or "") for item in paragraphs or []]
    unsupported = [item for item in unsupported or [] if item.get("count")]
    notes = []
    ext = Path(str(filename or "")).suffix.lower() or "未知格式"
    if ocr_used:
        notes.append("本文档没有文本层，正文由本地 OCR 生成；OCR 会引入错字与假断行。")
    if unsupported:
        notes.append("未纳入翻译的内容已在下方逐项列出；它们保留在源文件里，"
                     "不会出现在译文文档中。")
    notes.append("导出是**重建的译文文档**，不是原文件的格式保真输出："
                 "排版、表格、图片与样式不还原。")
    return {
        "version": EXTRACTION_REPORT_VERSION,
        "format": ext.lstrip("."),
        "filename": str(filename or ""),
        "status": "partial" if unsupported else "complete",
        "extracted": {
            "paragraphs": len(values),
            "characters": sum(len(item) for item in values),
            "pages": (facts or {}).get("pages"),
            "text_pages": (facts or {}).get("text_pages"),
            "empty_pages": (facts or {}).get("empty_pages"),
            "images": (facts or {}).get("images", 0),
            "ocr_used": bool(ocr_used),
        },
        "unsupported": unsupported,
        "unsupported_total": sum(int(item.get("count") or 0) for item in unsupported),
        "notes": notes,
        "warnings": [str(item) for item in warnings or []],
        "generated_at": datetime.now(timezone.utc).isoformat(),
    }


def extract_document_paragraphs_with_report(filename, file_bytes, *,
                                            ocr_max_pages=3, on_progress=None,
                                            profiler=None, checkpoint_dir=None,
                                            ocr_workers=None, ocr_queue_size=None,
                                            cancel_check=None):
    """Extract paragraphs together with an import-scope report.

    Returns ``(paragraphs, warnings, report)``.  The report is what makes the
    import honest: it states how much text was read, which parts of the
    document were *not* included in translation, and that the export is a
    regenerated translation document rather than a format-faithful round trip.
    """
    warnings = []
    paragraphs = []
    unsupported = []
    facts = {}
    ocr_used = False
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf"):
            classify_stage = profiler.start_stage(
                "pdf_classify", concurrency=1, metadata={"backend": "PyMuPDF"}
            ) if profiler is not None else None
            facts = _pdf_extraction_facts(file_bytes)
            if classify_stage is not None:
                classify_stage.record["item_count"] = int(facts.get("pages") or 0)
                classify_stage.record["page_count"] = int(facts.get("pages") or 0)
                classify_stage.finish(
                    status="completed", text_pages=int(facts.get("text_pages") or 0),
                    empty_pages=int(facts.get("empty_pages") or 0),
                )
            layout_stage = profiler.start_stage(
                "layout_recovery", concurrency=1,
                metadata={"backend": "PyMuPDF", "page_count": facts.get("pages")},
            ) if profiler is not None else None
            paragraphs = [clean_xml_chars(p) for p in extract_pdf_paragraphs(file_bytes)]
            if layout_stage is not None:
                layout_stage.record["item_count"] = len(paragraphs)
                layout_stage.record["page_count"] = int(facts.get("pages") or 0)
                layout_stage.finish(status="completed")
            ocr_details = {}
            if not paragraphs:
                scope = "全文" if ocr_max_pages is None else "前/中/后代表页"
                warnings.append(f"PDF 无文本层，正在自动 OCR {scope}…")
                ocr_fn = _ocr_pdf_text_with_warnings
                ocr_kwargs = {
                    "max_pages": ocr_max_pages,
                    "on_progress": on_progress,
                }
                supports_kwargs = False
                try:
                    parameters = inspect.signature(ocr_fn).parameters
                    supports_kwargs = any(
                        parameter.kind == inspect.Parameter.VAR_KEYWORD
                        for parameter in parameters.values()
                    )
                    optional_ocr_kwargs = {
                        "workers": ocr_workers,
                        "queue_size": ocr_queue_size,
                        "checkpoint_dir": checkpoint_dir,
                        "cancel_check": cancel_check,
                        "profiler": profiler,
                        "return_details": True,
                    }
                    for key, value in optional_ocr_kwargs.items():
                        if supports_kwargs or key in parameters:
                            ocr_kwargs[key] = value
                except (TypeError, ValueError):
                    optional_ocr_kwargs = {}
                ocr_result = ocr_fn(file_bytes, **ocr_kwargs)
                if isinstance(ocr_result, tuple) and len(ocr_result) >= 3:
                    ocr_text, ocr_warnings, ocr_details = ocr_result[:3]
                else:
                    ocr_text, ocr_warnings = ocr_result
                warnings.extend(ocr_warnings)
                paragraphs = _ocr_text_to_paragraphs(ocr_text)
                ocr_used = True
                if not paragraphs:
                    warnings.append("OCR 未生成可用段落；请安装语言包或手动选择风格")
            images = int(facts.get("images") or 0)
            if images:
                unsupported.append({
                    "kind": "image", "label": "图片", "count": images,
                    "detail": f"{images} 张图片里的文字不在文本层内，不参与翻译",
                    "sample": ""})
            empty_pages = int(facts.get("empty_pages") or 0)
            if empty_pages and not ocr_used:
                unsupported.append({
                    "kind": "page", "label": "无文本层页面",
                    "count": empty_pages,
                    "detail": f"{empty_pages} 页没有文本层（可能是扫描图），"
                              "其内容未进入翻译",
                    "sample": ""})
                warnings.append(
                    f"{empty_pages} 页没有文本层；如需翻译请重新导入并允许 OCR。")
        elif name.endswith(".docx"):
            doc_word = Document(io.BytesIO(file_bytes))
            for p in doc_word.paragraphs:
                for sub_p in re.split(r"\n+", clean_xml_chars(p.text)):
                    t = sub_p.strip()
                    if len(t) > 1 and not _ORNAMENT_RE.match(t):
                        paragraphs.append(t)
            unsupported = _docx_unsupported_parts(doc_word)
        else:
            warnings.append("不支持的文档格式，无法自动画像")
    except Exception as exc:  # noqa: BLE001
        if "任务已请求取消" in str(exc):
            raise
        warnings.append(f"预提取失败：{exc}")
    report = _build_extraction_report(
        filename, paragraphs, ocr_used=ocr_used, unsupported=unsupported,
        facts=facts, warnings=warnings)
    if ocr_used:
        report["ocr"] = dict(ocr_details or {})
        report["ocr"]["paragraph_count"] = len(paragraphs)
        report["ocr"]["paragraph_confidences"] = list(
            report["ocr"].get("paragraph_confidences") or []
        )
    return paragraphs, warnings, report


def extract_document_paragraphs(filename, file_bytes, *, ocr_max_pages=3,
                                 on_progress=None):
    """Backward-compatible wrapper: ``(paragraphs, warnings)``.

    New callers should use :func:`extract_document_paragraphs_with_report` so
    the import scope can be shown to the user.
    """
    paragraphs, warnings, _report = extract_document_paragraphs_with_report(
        filename, file_bytes, ocr_max_pages=ocr_max_pages, on_progress=on_progress)
    return paragraphs, warnings



def cleanup_source_paragraphs(paragraphs, provider, api_key, model, *,
                              call_llm_fn=None, on_progress=None,
                              max_batch_chars=None, max_batch_items=None,
                              parallelism=None, confidence_by_index=None,
                              confidence_threshold=None, checkpoint_dir=None,
                              cancel_check=None, max_retries=None,
                              retry_backoff_seconds=None,
                              request_interval_seconds=None, profiler=None):
    """Repair OCR characters/line breaks while preserving indexed source items.

    This compatibility wrapper keeps the cleanup contract in the core API so
    callers and tests do not need to know the internal module layout.
    """
    cleanup_call = call_llm_fn
    if cleanup_call is None:
        # ``source_cleanup`` executes provider calls in worker threads.  The
        # normal LLM router keeps base_url/reasoning/runtime job id in
        # thread-local context, so copy the caller's context explicitly; a
        # worker must not silently fall back to the provider's default URL.
        inherited_base_url = getattr(_LLM_CTX, "base_url", None)
        inherited_reasoning = getattr(_LLM_CTX, "reasoning_effort", None)
        inherited_job_id = getattr(_RUNTIME_CTX, "job_id", None)

        def cleanup_call(provider_name, api_key_value, model_name,
                         system_prompt, user_prompt, **kwargs):
            if inherited_job_id:
                _RUNTIME_CTX.job_id = inherited_job_id
            inherited_kwargs = {}
            if inherited_base_url:
                inherited_kwargs["base_url"] = inherited_base_url
            if inherited_reasoning:
                inherited_kwargs["reasoning_effort"] = inherited_reasoning
            return call_llm(
                provider_name, api_key_value, model_name, system_prompt,
                user_prompt, **inherited_kwargs, **kwargs)

    return _source_cleanup.cleanup_source_paragraphs(
        paragraphs, provider, api_key, model,
        call_llm=cleanup_call or call_llm,
        on_progress=on_progress,
        max_batch_chars=max_batch_chars or _source_cleanup.MAX_BATCH_CHARS,
        max_batch_items=max_batch_items or _source_cleanup.MAX_BATCH_ITEMS,
        parallelism=parallelism or _source_cleanup.DEFAULT_PARALLELISM,
        confidence_by_index=confidence_by_index,
        confidence_threshold=(confidence_threshold
                              if confidence_threshold is not None
                              else _source_cleanup.DEFAULT_CONFIDENCE_THRESHOLD),
        checkpoint_dir=checkpoint_dir,
        cancel_check=cancel_check,
        max_retries=(max_retries if max_retries is not None
                     else _source_cleanup.DEFAULT_MAX_RETRIES),
        retry_backoff_seconds=(retry_backoff_seconds
                               if retry_backoff_seconds is not None
                               else _source_cleanup.DEFAULT_RETRY_BACKOFF_SECONDS),
        request_interval_seconds=request_interval_seconds or 0.0,
        profiler=profiler,
    )


def _source_segmentation_mode(state=None, explicit=None):
    """Resolve sentence/paragraph mode without changing an existing task."""
    saved = {}
    if isinstance(state, dict):
        saved = state.get("segmentation") or {}
    filename = str((state or {}).get("filename") or "").lower()
    default_mode = ("sentence" if filename.endswith(".pdf") else "paragraph")
    return _segmentation.resolve_mode(
        explicit,
        saved.get("mode") if isinstance(saved, dict) else None,
        os.environ.get("FOLIOTHREAD_SEGMENTATION_MODE"),
        default=default_mode,
    )


def segment_source_paragraphs(paragraphs, *, state=None, mode=None,
                              source_lang=None):
    """Convert cleaned paragraphs into ordered CAT translation units.

    ``source_paragraphs`` is kept by the caller for audit/export.  The return
    value is the translation-facing sequence plus a compact paragraph mapping.
    """
    resolved_mode = _source_segmentation_mode(state, mode)
    language = source_lang
    if language is None and isinstance(state, dict):
        language = state.get("source_lang")
    return _segmentation.segment_paragraphs(
        paragraphs, mode=resolved_mode, source_lang=language)


def _apply_source_segmentation(state, cleaned_paragraphs, *, mode=None):
    """Persist paragraph structure and replace only the translation units."""
    source_paragraphs = [str(item or "").strip() for item in cleaned_paragraphs]
    segments, metadata = segment_source_paragraphs(
        source_paragraphs, state=state, mode=mode)
    state["source_paragraphs"] = list(source_paragraphs)
    state["segmentation"] = metadata
    return segments, metadata


def audit_source_quality(paragraphs):
    """Deterministically quarantine OCR fragments before translation."""
    return _source_quality.audit_source_segments(paragraphs)


def _source_quality_gate_enabled(state):
    """Apply the OCR admission gate to PDF source extraction only."""
    return str((state or {}).get("filename") or "").lower().endswith(".pdf")


def _audit_source_quality_for_state(state):
    paragraphs = (state or {}).get("paras") or []
    if not _source_quality_gate_enabled(state):
        return {
            "version": _source_quality.VERSION,
            "status": "disabled",
            "checked_count": len(paragraphs),
            "flagged_count": 0,
            "flags": [],
            "clusters": [],
        }
    return audit_source_quality(paragraphs)


def _source_quality_finding(flag):
    """Build the persisted blocking finding for one quarantined source item."""
    segment_index = int(flag.get("segment_index"))
    reasons = "；".join(str(reason) for reason in flag.get("reasons") or [])
    source_text = str(flag.get("text") or "")
    return {
        "segment_index": segment_index,
        "segment_id": segment_index,
        "type": "source_quality_gate",
        "severity": "blocking",
        "category": "source_quality",
        "summary": "原文片段未通过翻译准入门",
        "source_span": source_text,
        "target_span": source_text,
        "explanation": (
            "该段是短小、未闭合或含异常空格连字符的 OCR 片段，"
            "仅凭当前文本无法证明它是可翻译的正文。"
        ),
        "recommendation": (
            "回看 PDF 原页，确认应合并、删除还是保留；确认前系统只保留源文，"
            "不让模型猜译，也不会把它写入翻译记忆。"
        ),
        "confidence": None,
        "detector": "Source quality gate",
        "diagnostic_version": _source_quality.VERSION,
        "reason": f"源文质量门拦截：{reasons or '疑似 OCR 碎片'}",
        "detected_text": source_text,
        "requires_human_confirmation": True,
    }


def call_llm(provider, api_key, model, system_prompt, user_prompt,
             temperature=0.1, base_url=None, response_format=None,
             reasoning_effort=None):
    """底层大模型统一路由（超时 150 秒，模型可配置）。

    支持官方接口与 OpenAI /chat/completions 兼容中转站：
    base_url 显式传入 > 提供商默认 base_url > 会话线程级自定义中转地址。
    """
    cfg = PROVIDERS.get(provider)
    if not cfg:
        return ""
    runtime_job_id = _runtime_job_id()
    if runtime_job_id:
        if _runtime_cancel_requested(runtime_job_id):
            raise RuntimeError("任务已请求取消")
        update_runtime_state(
            runtime_job_id, status="waiting_external", phase="waiting_llm",
            phase_label="等待模型响应", event="已向模型发送请求",
            event_name="llm_request_started")
    if cfg["kind"] == "gemini":
        try:
            client = genai.Client(api_key=api_key,
                                  http_options=genai.types.HttpOptions(timeout=150_000))
        except (AttributeError, TypeError):
            client = genai.Client(api_key=api_key)
        config_kwargs = {"temperature": temperature}
        if response_format and cfg.get("capabilities", {}).get("supports_response_format"):
            config_kwargs["response_mime_type"] = "application/json"
        res = client.models.generate_content(
            model=model,
            contents=user_prompt,
            system_instruction=system_prompt,
            config=genai.types.GenerateContentConfig(**config_kwargs),
        )
        result = (res.text or "").strip()
        if runtime_job_id:
            update_runtime_state(runtime_job_id, status="running", phase="running",
                                 phase_label="处理模型结果", event="已收到模型响应",
                                 event_name="llm_response_received")
            if _runtime_cancel_requested(runtime_job_id):
                raise RuntimeError("任务已请求取消")
        return result

    # OpenAI 官方与所有 OpenAI 兼容接口共用 SDK 路由
    kwargs = {"api_key": api_key, "timeout": (15.0, 180.0), "max_retries": 1}
    resolved_base = resolve_openai_base_url(provider, base_url)
    resolved_reasoning = (reasoning_effort if reasoning_effort is not None
                          else getattr(_LLM_CTX, "reasoning_effort", None))
    if resolved_reasoning and resolved_reasoning not in reasoning_effort_options(
            provider, model):
        # The thread setting belongs to the translator role.  A separate
        # reviewer or an unknown relay model must not inherit it accidentally.
        resolved_reasoning = None
    if resolved_base:
        kwargs["base_url"] = resolved_base
    http_client = None
    if cfg.get("proxy_bypass"):
        # 该接口经用户本地代理时 TLS 失败：显式关闭环境代理。
        http_client = httpx.Client(trust_env=False, timeout=(15.0, 180.0))
        kwargs["http_client"] = http_client
    try:
        client = OpenAI(**kwargs)
        request = {
            "model": model,
            "messages": [{"role": "system", "content": system_prompt},
                         {"role": "user", "content": user_prompt}],
        }
        # Reasoning models may reject sampling controls such as temperature;
        # when the user explicitly selects reasoning, omit it for official
        # recognized reasoning families and let the model's policy take over.
        if resolved_reasoning:
            request["reasoning_effort"] = resolved_reasoning
            if not str(model or "").lower().startswith(
                    ("o1", "o3", "o4-mini", "gpt-5", "gpt-6",
                     "deepseek-r1", "qwq")):
                request["temperature"] = temperature
        else:
            request["temperature"] = temperature
        if response_format and cfg.get("capabilities", {}).get(
                "supports_response_format"):
            request["response_format"] = response_format
        res = client.chat.completions.create(**request)
        result = (res.choices[0].message.content or "").strip()
        if runtime_job_id:
            update_runtime_state(runtime_job_id, status="running", phase="running",
                                 phase_label="处理模型结果", event="已收到模型响应",
                                 event_name="llm_response_received")
            if _runtime_cancel_requested(runtime_job_id):
                raise RuntimeError("任务已请求取消")
        return result
    finally:
        if http_client:
            http_client.close()


def _provider_error_payload(error):
    """Extract safe, non-secret fields from SDK/http responses.

    OpenAI-compatible relays do not all use the same exception shape: the
    OpenAI SDK exposes ``status_code``/``body`` while httpx exposes a response
    object.  Keep this extraction deliberately small and never return the raw
    exception text to the UI (it may contain request details or credentials).
    """
    response = error if isinstance(error, httpx.Response) else getattr(
        error, "response", None)
    status_code = getattr(error, "status_code", None)
    if status_code is None and response is not None:
        status_code = getattr(response, "status_code", None)
    try:
        status_code = int(status_code) if status_code is not None else None
    except (TypeError, ValueError):
        status_code = None

    payloads = []
    body = getattr(error, "body", None)
    if body is not None:
        payloads.append(body)
    if response is not None:
        try:
            payloads.append(response.json())
        except (TypeError, ValueError, AttributeError):
            pass

    code = ""
    message = ""

    def visit(payload):
        nonlocal code, message
        if isinstance(payload, str):
            stripped = payload.strip()
            if stripped.startswith(("{", "[")):
                try:
                    visit(json.loads(stripped))
                except (TypeError, ValueError):
                    pass
            return
        if not isinstance(payload, dict):
            return
        if not code:
            for key in ("code", "error_code", "type"):
                value = payload.get(key)
                if isinstance(value, (str, int)) and str(value).strip():
                    code = str(value).strip()
                    break
        if not message:
            value = payload.get("message") or payload.get("detail")
            if isinstance(value, str) and value.strip():
                message = value.strip()
        for key in ("error", "detail"):
            nested = payload.get(key)
            if isinstance(nested, dict):
                visit(nested)

    for payload in payloads:
        visit(payload)
    error_text = str(error or "")
    if status_code is None:
        match = re.search(r"\b(401|403|408|409|429|5\d{2})\b", error_text)
        if match:
            status_code = int(match.group(1))
    return {
        "http_status": status_code,
        "code": code,
        "message": message,
        "text": error_text,
    }


def provider_error_status(error):
    """Return a stable, user-facing status for a provider failure.

    The UI can use the ``status`` value for state styling and ``message`` for
    the copy.  Matching provider codes/messages takes precedence over HTTP
    status because relays commonly use HTTP 403 for both insufficient balance
    and other permission failures.
    """
    info = _provider_error_payload(error)
    http_status = info["http_status"]
    code = info["code"].replace("-", "_").upper()
    searchable = " ".join((info["text"], info["code"], info["message"])).lower()
    status_suffix = f"（HTTP {http_status}）" if http_status else ""

    if (code in {"INSUFFICIENT_BALANCE", "BALANCE_INSUFFICIENT"}
            or "insufficient_balance" in searchable
            or "insufficient balance" in searchable
            or "余额不足" in searchable
            or "balance is too low" in searchable):
        return {
            "status": "insufficient_balance",
            "label": "余额不足",
            "message": "余额不足，请充值或更换 API Key。",
            "http_status": http_status,
            "provider_code": info["code"],
        }

    if (http_status == 401 or code in {"INVALID_API_KEY", "UNAUTHORIZED"}
            or "invalid_api_key" in searchable
            or "incorrect api key" in searchable
            or "unauthorized" in searchable):
        return {
            "status": "invalid_api_key",
            "label": "API Key 无效",
            "message": f"API Key 无效或已过期{status_suffix}，请检查密钥和服务商。",
            "http_status": http_status,
            "provider_code": info["code"],
        }

    if (http_status == 429 or code in {"RATE_LIMITED", "TOO_MANY_REQUESTS"}
            or "rate limit" in searchable or "too many requests" in searchable):
        return {
            "status": "rate_limited",
            "label": "请求受限",
            "message": f"请求过于频繁或已达到服务商限额{status_suffix}，请稍后重试。",
            "http_status": http_status,
            "provider_code": info["code"],
        }

    if http_status == 403:
        return {
            "status": "forbidden",
            "label": "访问被拒绝",
            "message": "服务商拒绝了请求（HTTP 403），请检查 API 权限、余额或接口地址。",
            "http_status": http_status,
            "provider_code": info["code"],
        }

    if isinstance(error, httpx.RequestError):
        return {
            "status": "network_error",
            "label": "无法连接服务商",
            "message": "无法连接服务商，请检查 API 地址和网络。",
            "http_status": http_status,
            "provider_code": info["code"],
        }

    if http_status is not None and http_status >= 500:
        return {
            "status": "provider_unavailable",
            "label": "服务商暂时不可用",
            "message": f"服务商暂时不可用{status_suffix}，请稍后重试。",
            "http_status": http_status,
            "provider_code": info["code"],
        }

    return {
        "status": "unknown",
        "label": "API 请求失败",
        "message": f"API 请求失败{status_suffix}，请检查服务商配置后重试。",
        "http_status": http_status,
        "provider_code": info["code"],
    }


def provider_error_message(error, action="请求失败"):
    """Format a provider failure without leaking raw SDK/relay details."""
    status = provider_error_status(error)
    message = status["message"]
    return f"{action}：{message}" if action else message


def test_provider(provider, api_key, model, base_url=None):
    """连通性测试：发送一个最小请求，返回 (ok, message)。"""
    t0 = time.time()
    try:
        out = call_llm(provider, api_key, model,
                       "你是连接测试助手。", "请只回复两个字：OK",
                       temperature=0.0, base_url=base_url)
    except Exception as exc:
        return False, provider_error_message(exc)
    elapsed = time.time() - t0
    if not (out or "").strip():
        return False, "返回内容为空（请检查 API Key / 模型名 / 余额）"
    return True, f"响应「{(out or '').strip()[:24]}」· 耗时 {elapsed:.1f}s"


def fetch_provider_models(provider, api_key, base_url=None):
    """获取 OpenAI 兼容服务商的模型目录，返回 (ok, models, message)。"""
    cfg = PROVIDERS.get(provider) or {}
    if cfg.get("kind") not in ("openai", "openai_compat"):
        return False, [], "当前服务商不支持 OpenAI 兼容模型目录"
    if not (api_key or "").strip():
        return False, [], "请先填写 API Key"
    try:
        resolved_base = resolve_openai_base_url(provider, base_url)
    except ValueError as exc:
        return False, [], str(exc)
    if not resolved_base:
        return False, [], "请先填写 API 地址"
    try:
        response = httpx.get(
            f"{resolved_base}/models",
            headers={"Authorization": f"Bearer {api_key.strip()}",
                     "Accept": "application/json"},
            timeout=20.0,
        )
        response.raise_for_status()
        payload = response.json()
    except httpx.HTTPStatusError as exc:
        return False, [], provider_error_message(exc, "获取模型失败")
    except httpx.HTTPError as exc:
        return False, [], provider_error_message(exc, "获取模型失败")
    except ValueError:
        return False, [], "获取模型失败：服务商返回了无法解析的响应。"

    rows = payload.get("data") if isinstance(payload, dict) else payload
    if not rows and isinstance(payload, dict):
        rows = payload.get("models")
    models = []
    for row in rows or []:
        model_id = row if isinstance(row, str) else (
            row.get("id") or row.get("model") or row.get("name")
            if isinstance(row, dict) else None)
        if model_id and str(model_id).strip():
            models.append(str(model_id).strip())
    models = sorted(set(models), key=str.casefold)
    if not models:
        return False, [], "模型目录为空或返回格式不受支持"
    return True, models, f"已获取 {len(models)} 个可用模型"


def parse_termbase(file_stream):
    """解析用户上传的术语库 Excel，返回概念化术语条目列表。

    必选列：Source / Target；可选列：Behavior（translate/preserve）、
    Status（locked/provisional）、Preferred（首选译名）、Forbidden（禁止译名，
    可用 ; 或 , 分隔）、Scope、Note。解析失败抛出 ValueError（不再静默吞错）。
    """
    try:
        df = pd.read_excel(file_stream)
    except Exception as e:
        raise ValueError(f"无法读取 Excel 文件：{e}") from e
    df.columns = [str(c).strip() for c in df.columns]
    return _termbase_df_to_entries(df)


def _termbase_df_to_entries(df):
    if "Source" not in df.columns or "Target" not in df.columns:
        raise ValueError("术语库缺少 Source / Target 列，请检查表头")
    df = df.dropna(subset=["Source", "Target"])
    entries = []
    for _, row in df.iterrows():
        entry = {"source": str(row["Source"]).strip(),
                 "target": str(row["Target"]).strip()}
        for col, key in (("Behavior", "behavior"), ("Status", "status"),
                         ("Preferred", "preferred"), ("Scope", "scope"),
                         ("Note", "note")):
            if col in df.columns and pd.notna(row[col]):
                entry[key] = str(row[col]).strip()
        if "Forbidden" in df.columns and pd.notna(row["Forbidden"]):
            entry["forbidden"] = [x.strip() for x in
                                  re.split(r"[;；,，]", str(row["Forbidden"])) if x.strip()]
        entries.append(entry)
    return entries


def parse_termbase_csv(file_stream):
    """CSV 术语库：列约定与 Excel 一致（Source/Target 必选）。"""
    try:
        df = pd.read_csv(file_stream)
    except Exception as e:
        raise ValueError(f"无法读取 CSV 文件：{e}") from e
    df.columns = [str(c).strip() for c in df.columns]
    return _termbase_df_to_entries(df)


def _local_name(tag):
    return str(tag).rsplit("}", 1)[-1]


def parse_termbase_tbx(file_stream):
    """解析 TBX 术语库（Trados MultiTerm 等标准格式）。

    每个 termEntry 取前两种语言的第一个 term，作为 source/target；
    导入条目默认 status=locked（导入即视为已固定译名）。
    """
    try:
        root = ET.fromstring(file_stream.read())
    except Exception as e:
        raise ValueError(f"无法解析 TBX 文件：{e}") from e
    entries = []
    for te in root.iter():
        if _local_name(te.tag) != "termEntry":
            continue
        langs = []
        for ls in te.iter():
            if _local_name(ls.tag) != "langSet":
                continue
            # ElementTree 会把 xml: 前缀自动展开为命名空间键
            lang = ls.get("{http://www.w3.org/XML/1998/namespace}lang") \
                or ls.get("xml:lang") or ls.get("lang") or ""
            term = None
            for node in ls.iter():
                if _local_name(node.tag) == "term" and (node.text or "").strip():
                    term = node.text.strip()
                    break
            if term and lang and lang not in [x[0] for x in langs]:
                langs.append((lang, term))
            if len(langs) >= 2:
                break
        if len(langs) >= 2:
            entries.append({"source": langs[0][1], "target": langs[1][1],
                            "behavior": "translate", "status": "locked"})
    if not entries:
        raise ValueError("TBX 中未找到可导入的术语（需要至少两种语言的 langSet）")
    return entries


def import_tmx(file_stream, project_id=None, target_lang=None):
    """导入 TMX 翻译记忆（Trados / memoQ 等导出的标准格式）。

    按 <tu> 的 <tuv><seg> 文本对入库：仅接受源文含字母/数字且译文非空的
    单元；与现有翻译记忆冲突的源文跳过（不覆盖项目内已审校条目）。

    条目必须带上**目标语言**，否则不入库——语言未知的记忆一旦落盘，就会成为
    下一次跨语言命中的来源。语言取调用方显式给出的 `target_lang`（调用方
    知道"这份记忆是给哪个目标语言的"）；调用方没给时退回目标 <tuv xml:lang>。
    刻意**不**把 BCP-47 代码（`zh-CN`）翻译成本应用的显示名（`简体中文`）：
    猜测式映射会让导入"看起来成功、实际永不命中"，比显式失败更糟。

    导入目标为指定项目的记忆（默认项目即历史的全局记忆）。
    返回 {"added": n, "skipped": m}。
    """
    try:
        root = ET.fromstring(file_stream.read())
    except Exception as e:
        raise ValueError(f"无法解析 TMX 文件：{e}") from e
    existing = load_tm(project_id)
    declared_lang = str(target_lang or "").strip()
    added = skipped = 0
    for tu in root.iter():
        if _local_name(tu.tag) != "tu":
            continue
        texts = []
        languages = []
        for tuv in tu:
            if _local_name(tuv.tag) != "tuv":
                continue
            seg = next((c for c in tuv if _local_name(c.tag) == "seg"), None)
            if seg is not None and (seg.text or "").strip():
                texts.append(seg.text.strip())
                languages.append(tuv.get("{http://www.w3.org/XML/1998/namespace}lang")
                                 or tuv.get("lang") or "")
            if len(texts) >= 2:
                break
        if len(texts) >= 2:
            src, tgt = texts[0], texts[-1]
            language = declared_lang or str(languages[-1] or "").strip()
            key = tm_scope_key(language, src)
            if _tm_eligible(src, tgt) and key:
                if key not in existing:
                    record = tm_record(tgt, language)
                    record["source"] = "tmx_import"
                    existing[key] = record
                    added += 1
                else:
                    skipped += 1
    if not added:
        raise ValueError(
            "TMX 中未找到可导入的新翻译单元（源文需含字母/数字、目标语言需显式指定"
            "或可从 xml:lang 确定，且不与现有记忆冲突）")
    save_tm(existing, project_id)
    return {"added": added, "skipped": skipped}


def extract_auto_terms(paragraphs, target_lang, provider, api_key, model):
    """自动抽取术语库（兼容旧接口：返回 {source: target}）。

    新实现（transpraxis.terminology.extract_auto_terms_v2）：分布式采样、
    全量 occurrences、candidate 状态与 model_knowledge 证据。
    """
    from transpraxis.terminology import extract_auto_terms_v2
    entries, _warnings = extract_auto_terms_v2(
        paragraphs, target_lang, provider, api_key, model)
    return {e["source"]: e["target"] for e in entries}


# ================= 概念化术语表（对齐 localize-anything 的 Glossary 模型）=================
def normalize_glossary(entries):
    """标准化术语条目（委托 transpraxis.models，兼容旧字段并新增 id/occurrences/evidence）。"""
    return _models.normalize_glossary(entries)


def glossary_block(glossary):
    """把术语表渲染成注入翻译/审校 prompt 的文本块。"""
    locked_translate = [e for e in glossary if e["behavior"] == "translate" and e["status"] == "locked"]
    preserve = [e for e in glossary if e["behavior"] == "preserve"]
    provisional = [e for e in glossary if e["behavior"] == "translate" and e["status"] != "locked"]
    lines = []
    if locked_translate:
        lines.append("【锁定术语（必须使用首选译名，不得使用禁止译名）】：")
        for e in locked_translate:
            seg = f"- {e['source']} -> {e['preferred']}"
            if e["forbidden"]:
                seg += f"（禁止：{'、'.join(e['forbidden'])}）"
            lines.append(seg)
    if preserve:
        lines.append("【必须保留原文的术语/名称】：" + "、".join(e["source"] for e in preserve))
    if provisional:
        lines.append("【建议术语（仅供参考，请优先采用）】：")
        for e in provisional:
            lines.append(f"- {e['source']} -> {e['target']}")
    return "\n".join(lines)


def glossary_to_terms(glossary):
    """翻译行为术语 -> 扁平 dict（供报告生成等场景使用）。"""
    return {e["source"]: (e["preferred"] or e["target"])
            for e in glossary if e["behavior"] == "translate" and e["target"]}


def check_glossary_compliance(src, tgt, glossary, segment_id=None,
                              section_profile=None):
    """锁定术语的确定性合规检查（委托 transpraxis.terminology：entry_id/segment_id 级）。"""
    from transpraxis.terminology import check_glossary_compliance as _qa
    return _qa(src, tgt, glossary, segment_id=segment_id,
               section_profile=section_profile)


# ================= 确定性检查（对齐 localize-anything 的机械检查）=================
PRESERVE_RE = re.compile(
    r'(?P<placeholder>%[sd]|%1\$[sd]|\{[A-Za-z_][A-Za-z0-9_]*\}|\{\{[A-Za-z_][A-Za-z0-9_]*\}\})'
    r'|(?P<url>https?://\S+|www\.\S+)'
    r'|(?P<email>[\w.+-]+@[\w-]+(?:\.[\w-]+)+)'
    r'|(?P<doi>10\.\d{4,9}/[^\s]+)'
    r'|(?P<citation>\[\d+(?:[-,]\s*\d+)*\])',
    re.IGNORECASE,
)

PRESERVE_SEVERITY = {
    "placeholder": "blocking",   # 占位符损坏 = 结构破坏，绝不可自动放行
    "url": "actionable",
    "email": "actionable",
    "doi": "actionable",
    "citation": "actionable",
}


def extract_preserved_tokens(text):
    """提取源文本中必须原样保留的 token（占位符/URL/邮箱/DOI/引用标注）。"""
    return {m.group(0): m.lastgroup for m in PRESERVE_RE.finditer(text or "")}


def find_residuals(src, tgt, target_lang):
    """检测目标语言中残留的源语言片段。

    返回 [(片段, severity)]：连续 ≥2 个源语单词/较长汉字串 -> actionable；
    单个词（可能是专有名词）-> informational。启发式，不替代审校。
    """
    tgt_clean = PRESERVE_RE.sub(" ", tgt or "")
    if target_lang == "English":
        source_runs = set(re.findall(r'[\u4e00-\u9fff]{2,}', src or ""))
        return [(c, "actionable" if len(c) >= 4 else "informational")
                for c in re.findall(r'[\u4e00-\u9fff]{2,}', tgt_clean)
                if any(c in run for run in source_runs)]
    src_words = set(w.lower() for w in re.findall(r'[A-Za-z]{5,}', src or ""))
    allowed = {"mti"}  # 产品名等明确保留词白名单（审校负责语义判断）
    words = re.findall(r'[A-Za-z]{5,}', tgt_clean)
    hits = [w for w in words if w.lower() in src_words and w.lower() not in allowed]
    result, run = [], []
    for w in words:
        if w in hits:
            run.append(w)
        else:
            if run:
                result.append((" ".join(run),
                               "actionable" if len(run) >= 2 else "informational"))
                run = []
    if run:
        result.append((" ".join(run), "actionable" if len(run) >= 2 else "informational"))
    return result


def _count_sentences(text):
    """粗粒度句数统计：按终结符切分（引号/括号闭合归并到前一句）。"""
    text = _ABBREV_RE.sub(" ", text)
    parts = re.split(r"[.!?…。！？]+[”\"'’)\]]*", text)
    return sum(1 for p in parts if p.strip())


def is_incomplete_translation(src, tgt):
    """疑似漏译/截断判定（双重规则，实测调优）：
    1. 字符级：长原文（≥120 字符）配极短译文（<15%）——只拦灾难性截断，
       英译中正常比例可低至 0.2-0.3，不能用高阈值；
    2. 句子级：原文 ≥2 句而译文不足一半句数，且字符占比 <35%——
       截断译文必然句数对不上，完整译文即使语言再凝练也很少掉一半句。
    """
    tgt = (tgt or "").strip()
    if not tgt:
        return True
    if len(src) >= 120 and len(tgt) < 0.15 * len(src):
        return True
    src_sents = _count_sentences(src)
    if src_sents >= 2:
        tgt_sents = _count_sentences(tgt)
        if tgt_sents < src_sents * 0.5 and len(tgt) < 0.35 * len(src):
            return True
    return False


def _deterministic_finding(segment_index, severity, category, summary,
                           explanation, recommendation, *, source_span=None,
                           target_span=None, reason=None, kind=None,
                           detected_text=None, **extra):
    """Build an evidence-free but actionable deterministic QA finding."""
    finding = {
        "segment_index": segment_index, "segment_id": segment_index,
        "type": "check", "severity": severity,
        "category": category, "summary": summary,
        "source_span": source_span, "target_span": target_span,
        "explanation": explanation, "recommendation": recommendation,
        "confidence": None, "detector": "Deterministic QA",
        "diagnostic_version": 1,
        "reason": reason or summary,
    }
    if kind:
        finding["kind"] = kind
    if detected_text:
        finding["detected_text"] = detected_text
    finding.update({key: value for key, value in extra.items() if value is not None})
    return finding


def check_translation_batch(sources, targets, glossary, target_lang,
                            section_profile=None):
    """确定性检查一批译文：空译、保留项丢失、源语残留、锁定术语合规（scope 感知）。"""
    findings = []
    for i, (src, tgt) in enumerate(zip(sources, targets)):
        raw_src, raw_tgt = src, tgt
        src = "" if src is None else str(src)
        tgt = "" if tgt is None else str(tgt)
        if not tgt.strip():
            findings.append(_deterministic_finding(
                i, "blocking", "completeness", "译文为空",
                "原文存在，但当前段落没有任何译文内容，无法确认信息是否被完整传达。",
                "补译本段后，检查是否覆盖原文的全部句子、限制条件和专有名词。",
                source_span=src, target_span=""))
            continue
        target_report = _translation_target.validate_translation_target(
            raw_src, raw_tgt, segment_index=i)
        for issue in target_report["issues"]:
            findings.append(_deterministic_finding(
                i, "blocking", "format_integrity",
                issue["message"],
                "目标文本仍带有模型 transport 或解释包装，不能作为普通正文交付。",
                "重新解析或重新翻译本段，只保留最终目标文本后再检查。",
                source_span=src, target_span=tgt,
                reason=issue["message"], kind="translation_target_invariant",
                invariant_code=issue["code"]))
        # 完整性检查：拦截截断译文。
        # 实测根因：审校/修复环节的整段替换把长段译文换成了一句修正。
        if is_incomplete_translation(src, tgt):
            reason = (f"疑似漏译/截断：原文 {len(src)} 字符/{_count_sentences(src)} 句，"
                      f"译文仅 {len(tgt.strip())} 字符/{_count_sentences(tgt)} 句")
            findings.append(_deterministic_finding(
                i, "blocking", "completeness", "译文疑似遗漏或被截断",
                "原文长度和句子数量与当前译文不匹配，译文可能没有覆盖完整内容。",
                "对照原文逐句补齐缺失内容，并确认修复后的译文通过完整性复验。",
                source_span=src, target_span=tgt, reason=reason))
        source_words = re.findall(r"[A-Za-z]+", src or "")
        whole_source_preserved = any(
            str(entry.get("behavior") or "") == "preserve"
            and str(entry.get("status") or "") == "locked"
            and re.sub(r"\s+", " ", str(entry.get("source") or "")).strip().casefold()
            == re.sub(r"\s+", " ", src or "").strip().casefold()
            for entry in glossary or [])
        if target_lang != "English" and len(source_words) >= 2 \
                and re.sub(r"\s+", " ", src or "").strip().casefold() \
                == re.sub(r"\s+", " ", tgt or "").strip().casefold() \
                and not whole_source_preserved:
            findings.append(_deterministic_finding(
                i, "actionable", "translation_completion", "译文仍保留整段源文",
                "当前译文与英文原文完全相同，除非这是明确的保留项，否则可能尚未完成翻译。",
                "确认该段是否属于项目规定的保留内容；若不是，请重新翻译并保留专有名词的必要形式。",
                source_span=src, target_span=tgt,
                reason="译文与英文源段完全相同，疑似整段未翻译"))
        for token, kind in extract_preserved_tokens(src).items():
            if token not in tgt:
                severity = PRESERVE_SEVERITY.get(kind, "actionable")
                findings.append(_deterministic_finding(
                    i, severity, "format_integrity",
                    f"译文遗漏必须保留的{kind}「{token}」",
                    f"原文包含必须保留的 {kind}「{token}」，但当前译文中找不到该内容。",
                    "补回该保留项后重新检查占位符、链接或引用格式，确认其余译文未被破坏。",
                    source_span=token, target_span="",
                    reason=f"保留项 {kind}「{token}」在译文中丢失", kind=kind))
        for residual, sev in find_residuals(src, tgt, target_lang):
            findings.append(_deterministic_finding(
                i, sev, "source_language_residue", f"译文残留源语片段「{residual}」",
                "当前译文仍包含原文语言片段，可能意味着该部分未完成翻译，或未经确认地保留了源语。",
                "确认该片段是否为有意保留的专名或术语；若不是，请翻译后检查术语和上下文一致性。",
                source_span=residual if residual in src else None,
                target_span=residual, reason=f"疑似残留源语片段「{residual}」",
                kind="source_residue", detected_text=residual))
        findings.extend(check_glossary_compliance(
            src, tgt, glossary, segment_id=i, section_profile=section_profile))
        for f in findings:
            if "segment_index" not in f:
                f["segment_index"] = i
            if "segment_id" not in f or f.get("segment_id") is None:
                f["segment_id"] = i
    return findings


def _globalize_batch_findings(findings, offset):
    """Persist repair findings in document-global, never batch-local, space."""
    return [
        {**finding,
         "segment_id": offset + finding["segment_id"],
         "segment_index": offset + finding["segment_index"]}
        for finding in findings
    ]


# ================= 语义批次（对齐 localize-anything 的上下文批次）=================
BATCH_SIZE = 4
MAX_BATCH_CHARS = 1600
TRANSLATION_MAX_BATCH_CHARS = 2400
DEFAULT_TRANSLATION_CONCURRENCY = 1


def make_batches(paragraphs, batch_size=BATCH_SIZE, max_chars=MAX_BATCH_CHARS,
                 semantic_units=None):
    """把段落聚成批次；提供 semantic_units 时不跨单元边界。"""
    if semantic_units:
        batches = []
        ranges = []
        for unit in semantic_units:
            if not isinstance(unit, dict):
                continue
            try:
                start = int(unit.get("start_segment"))
                end = int(unit.get("end_segment"))
            except (TypeError, ValueError):
                continue
            if 0 <= start <= end < len(paragraphs):
                ranges.append((start, end))
        for start, end in sorted(ranges):
            batches.extend(make_batches(paragraphs[start:end + 1], batch_size, max_chars))
        covered = {index for start, end in ranges for index in range(start, end + 1)}
        range_size = sum(end - start + 1 for start, end in ranges)
        if batches and len(covered) == len(paragraphs) and range_size == len(paragraphs):
            return batches
    batches, cur, n = [], [], 0
    for p in paragraphs:
        if cur and (len(cur) >= batch_size or n + len(p) > max_chars):
            batches.append(cur)
            cur, n = [], 0
        cur.append(p)
        n += len(p)
    if cur:
        batches.append(cur)
    return batches


def _translation_evidence_index(
    paras, pairs, batch_pairs, glossary, document_profile,
    document_synopsis, section_digests, findings_all, blind=False,
    candidate_targets=None,
):
    return _translation_evidence.TranslationEvidenceIndex(
        paras, pairs + batch_pairs, glossary, document_profile,
        document_synopsis, section_digests, findings_all,
        blind=blind, candidate_targets=candidate_targets)


def _review_finding_record(finding, review_event_id):
    """Project a Translation Core finding onto the existing persisted surface."""
    segment_id = finding.get("segment_id")
    severity = finding.get("severity")
    return {
        **finding,
        "segment_id": segment_id, "segment_index": segment_id,
        "severity": severity, "type": "review",
        "category": finding.get("category") or "semantic_accuracy",
        "summary": str(finding.get("summary") or finding.get("reason") or
                       "审校发现问题"),
        "source_span": finding.get("source_span"),
        "target_span": finding.get("target_span"),
        "explanation": finding.get("explanation"),
        "recommendation": finding.get("recommendation"),
        "confidence": finding.get("confidence"),
        "detector": finding.get("detector") or "Semantic QA",
        "diagnostic_version": finding.get("diagnostic_version"),
        "reason": str(finding.get("reason") or finding.get("summary") or
                      "审校发现问题"),
        "evidence_refs": list(finding.get("evidence_refs") or []),
        "review_event_id": review_event_id,
    }


def _runtime_review_context(
    state, offset, batch_len, glossary_text, style_rules, target_lang,
):
    paras = state.get("paras") or []
    pairs = state.get("pairs") or []
    profile = state.get("document_profile") or {}
    digests = state.get("section_digests") or []
    previous_target = _context.select_target_context(pairs, offset, limit=2)
    dependency_ids = _runtime_review_dependency_ids(
        state, offset, batch_len, previous_target)
    target_dependency_ids = set(range(
        max(0, offset), max(0, offset) + max(0, batch_len)))
    target_dependency_ids.update(
        int(item["segment_index"])
        for item in previous_target or []
        if isinstance(item, dict)
        and isinstance(item.get("segment_index"), int)
        and not isinstance(item.get("segment_index"), bool)
    )
    return {
        "document_profile": profile,
        "document_synopsis": state.get("document_synopsis") or {},
        "section_profile": _batch_section_profile(profile, offset, batch_len) or {},
        "section_digest": _context.digest_for_segment(digests, offset) or {},
        "previous_source_context": list(paras[max(0, offset - 2):offset]),
        "previous_target_context": [
            dict(item) for item in previous_target
            if item.get("level") in {"human_accepted", "reviewed", "tm_approved"}
        ],
        "next_source_context": list(
            paras[offset + batch_len:offset + batch_len + 2]),
        "target_language": target_lang,
        "style_constraints": style_rules or "",
        "advisory_terminology_context": glossary_text or "",
        "_dependency_segment_ids": dependency_ids,
        "_dependency_target_segment_ids": sorted(target_dependency_ids),
    }


def _runtime_review_dependency_ids(state, offset, batch_len, previous_target=None):
    """Return segment IDs whose source/target can appear in review context."""
    paragraphs = state.get("paras") or []
    ids = set(range(max(0, offset), max(0, offset) + max(0, batch_len)))
    ids.update(range(max(0, offset - 2), max(0, offset)))
    ids.update(range(
        max(0, offset) + max(0, batch_len),
        min(len(paragraphs), max(0, offset) + max(0, batch_len) + 2),
    ))
    ids.update(
        int(item["segment_index"])
        for item in previous_target or []
        if isinstance(item, dict)
        and isinstance(item.get("segment_index"), int)
        and not isinstance(item.get("segment_index"), bool)
    )
    return sorted(ids)


def _batch_section_profile(document_profile, offset, batch_len):
    """按全局段区间匹配 section profile（用于相关术语的 section:<id> scope）。"""
    if not document_profile:
        return None
    for sec in document_profile.get("sections") or []:
        if not isinstance(sec, dict):
            continue
        try:
            start = int(sec.get("start_segment"))
            end = int(sec.get("end_segment"))
        except (TypeError, ValueError):
            continue
        if start <= offset and offset + batch_len - 1 <= end:
            return sec
    return None


# ================= 翻译记忆（对齐 localize-anything 的 TM：仅收录审校通过段落）=================
DEFAULT_TARGET_LANG = "简体中文"


def tm_path(project_id=None):
    """翻译记忆的存储路径；**按项目隔离**。

    翻译记忆是"已审校译文的受控记忆"，蓝图 §3.2 把它列在 Project 之下。
    条目键是「目标语言 + 原文」的作用域键（见 `tm_scope_key`）——因此同一个
    项目里同一段原文可以有多个目标语言的译法，而不会互相串用。

    系统工作区「未分类」沿用历史上的全局路径 `outputs/translation_memory.json`
    ——「未分类」就是旧默认项目的后继实体，因此既有任务的**行为与数据**完全
    不变（无迁移、无丢失）；命名项目各自使用
    `outputs/projects/<project_uuid>/translation_memory.json`。
    """
    project_id = _project.canonical_project_id(project_id) if project_id \
        else _project.SYSTEM_PROJECT_ID
    if _project.is_system_project_id(project_id):
        return OUTPUT_DIR / "translation_memory.json"
    return _project.project_dir(OUTPUT_DIR, project_id) / "translation_memory.json"


def _tm_eligible(source, target):
    """翻译记忆资格：源文必须有字母/数字（纯符号装饰行不入库），译文非空。"""
    return has_textual_content(source) \
        and bool((target or "").strip()) \
        and not _translation_target.is_translation_transport_wrapper(target)


# ---------------- 翻译记忆的作用域键 ----------------
# 一条翻译记忆的身份是「目标语言 + 原文」，**不是**「原文」。
# 同一段原文在 Français 与简体中文任务里的正确译文不同；只用原文当键，一种
# 语言的译文就会串进另一种语言的任务，并以"已审校"的名义静默通过交付检查。
# 因此作用域键是结构性的：不给出目标语言就**无法**命中任何条目。
# 定义放在 `transpraxis.translation_memory`，核心层与展示层共用同一份格式。
TM_SCOPE_SEP = _tm_scope.TM_SCOPE_SEP
tm_scope_key = _tm_scope.tm_scope_key
tm_unscope_key = _tm_scope.tm_unscope_key
tm_record_language = _tm_scope.tm_record_language
tm_record = _tm_scope.tm_record
tm_put = _tm_scope.tm_put
tm_discard = _tm_scope.tm_discard
tm_legacy_keys = _tm_scope.tm_legacy_keys


# 归一化只做**可证明等价**的处理：同一段文字在不同来源下的排版差异。
# 刻意不做语义近似（编辑距离、词序、同义替换）：翻译记忆是错误放大器，
# 一次错配会复制到整篇文档，所以匹配必须是"同一句话"而不是"相似的话"。
_TM_CHAR_MAP = str.maketrans({
    "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
    "\u2013": "-", "\u2014": "-", "\u2212": "-", "\u00ad": "",
    "\u2026": "...", "\u00a0": " ", "\u3000": " ",
})
_TM_SPACE_RE = re.compile(r"\s+")


def tm_normalize(text):
    """翻译记忆的规范化键：排版等价 → 同一键。

    处理的是真实来源差异，而不是"相似文本"：

    - PDF 与 DOCX 抽取出的换行、制表与连续空格；
    - 不间断空格（U+00A0）、全角空格（U+3000）；
    - 弯引号/直引号、en/em dash 与连字符、省略号、软连字符；
    - Unicode 兼容等价（NFKC，如全角字母数字）。

    不做大小写折叠：首字母大小写不同的段落是不同的段落。
    """
    if text is None:
        return ""
    value = unicodedata.normalize("NFKC", str(text))
    value = value.translate(_TM_CHAR_MAP)
    return _TM_SPACE_RE.sub(" ", value).strip()


def tm_index(tm, target_lang=None):
    """规范化原文 -> 作用域键的索引，**只收录目标语言可证明相同的条目**。

    没有目标语言上下文（`target_lang` 为空）时返回空索引：作用域不可证明
    就不建立任何可命中路径。
    """
    language = normalize_language(target_lang)
    index = {}
    if not language:
        return index
    for key, record in (tm or {}).items():
        if tm_record_language(key, record) != language:
            continue
        _, source = tm_unscope_key(key)
        normalized = tm_normalize(source)
        if normalized and normalized not in index:
            index[normalized] = key
    return index


def _tm_hit(record, language, key):
    return isinstance(record, dict) and bool(record.get("reviewed")) \
        and bool(record.get("target")) \
        and tm_record_language(key, record) == language


def tm_lookup(tm, source, index=None, *, target_lang=None):
    """查找翻译记忆条目，返回 (记录, 命中方式)。

    **必须显式给出目标语言**：语言未知（旧条目没有语言标注）或目标语言不同
    -> 不命中。这是发布阻断级的正确性约束，不是可选过滤。

    先精确匹配，再退到归一化匹配。归一化命中是"同一段文字的排版差异"，
    不涉及语义猜测，因此可以安全复用译文；命中方式会被记录，便于审计。
    """
    language = normalize_language(target_lang)
    if not language:
        return None, ""
    cleaned = str(source or "").replace("\n", " ")
    for key in (tm_scope_key(target_lang, cleaned), cleaned):
        if not key:
            continue
        if _tm_hit((tm or {}).get(key), language, key):
            return (tm or {}).get(key), "exact"
    normalized = tm_normalize(cleaned)
    if not normalized:
        return None, ""
    scoped_index = index if index is not None else tm_index(tm, target_lang)
    raw_key = scoped_index.get(normalized)
    if raw_key is None or raw_key == cleaned:
        return None, ""
    record = (tm or {}).get(raw_key)
    if _tm_hit(record, language, raw_key):
        return record, "normalized"
    return None, ""


def load_tm(project_id=None):
    """加载指定项目的翻译记忆并自清洗：非法条目直接丢弃。

    翻译记忆是错误放大器（一次错译会复制到全书），因此加载即消毒，
    防止旧版本或异常写入留下的污染条目继续命中。

    语言无法证明的旧条目**保留在文件里**（不丢用户数据），但不会被
    `tm_lookup` 命中：证明不了目标语言的记忆不能自动复用。可用
    `tm_legacy_keys` 把它们单独列出来。

    `project_id=None` 表示默认项目（即历史的全局记忆）。
    """
    p = tm_path(project_id)
    if p.is_file():
        try:
            raw = json.loads(p.read_text(encoding="utf-8"))
        except Exception:
            return {}
        return {k: v for k, v in raw.items()
                if isinstance(v, dict) and v.get("reviewed")
                and _tm_eligible(tm_unscope_key(k)[1], v.get("target"))}
    return {}


def save_tm(tm, project_id=None):
    p = tm_path(project_id)
    p.parent.mkdir(parents=True, exist_ok=True)
    tmp = p.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(tm, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(p)


def tm_project_id(state):
    """任务使用的翻译记忆所属项目（与任务的项目归属一致）。"""
    return resolved_project_id(state or {})


def state_target_lang(state, fallback=None):
    """任务的目标语言（TM 作用域用），返回**原样写法**（用于落盘与展示）。

    只认 state 里真正记下来的目标语言；没有就退回调用方显式给出的 fallback，
    两者都没有则返回空串。这里**不补默认值**——补一个默认语言等于猜语言，
    而猜错的代价是把一条已审校的错译文复用到另一种语言的任务里。
    """
    raw = (state or {}).get("target_lang")
    if normalize_language(raw):
        return str(raw)
    return str(fallback or "")


def copy_system_tm_to_project(project_id):
    """把系统工作区「未分类」的已审校记忆并入指定项目（人工动作）。

    新建项目从空白记忆开始（隔离的代价）。如果用户希望把既有积累带过去，
    这是一个显式的、可解释的动作，而不是静默共享或静默复制。
    返回本次并入的条目数。
    """
    target_id = _project.canonical_project_id(project_id)
    if _project.is_system_project_id(target_id):
        return 0
    if load_project(target_id) is None:
        raise ValueError(f"项目不存在：{project_id}")
    source_tm = load_tm(_project.SYSTEM_PROJECT_ID)
    if not source_tm:
        return 0
    target_tm = load_tm(target_id)
    added = 0
    for source, record in source_tm.items():
        if source in target_tm:
            continue
        target_tm[source] = record
        added += 1
    if added:
        save_tm(target_tm, target_id)
    return added


# 向后兼容别名（旧名字里的"默认项目"现在叫系统工作区「未分类」）。
copy_default_tm_to_project = copy_system_tm_to_project


# ================= Project：跨任务复用的已确认项目记忆 =================
# Project 层本身是领域模块（transpraxis/project.py，显式接收 root）；
# 这里只是把它接到本地 OUTPUT_DIR，并负责"任务属于哪个项目"的读写。
#
# 身份模型（v0.5 定死）：
#   - 每个项目都有不可变 UUID；display name 只是标签，不是主键，也不进路由；
#   - 系统工作区「未分类」是**真实项目**（真实 UUID + is_system=true），承载
#     所有没有 project_id 的任务；它不能重命名 / 归档 / 删除；
#   - 字符串 "default" 只是历史别名，任何入口都先归一到系统项目 UUID。

SYSTEM_PROJECT_ID = _project.SYSTEM_PROJECT_ID
SYSTEM_PROJECT_NAME = _project.SYSTEM_PROJECT_NAME
# 向后兼容别名：值现在是系统项目的真实 UUID，不再是字符串 "default"。
DEFAULT_PROJECT_ID = _project.DEFAULT_PROJECT_ID
DEFAULT_PROJECT_NAME = _project.DEFAULT_PROJECT_NAME
# 任务归属「未分类」时写进 state["project_id"] 的值就是系统项目 UUID。
UNCLASSIFIED_PROJECT_ID = SYSTEM_PROJECT_ID


def system_project_id():
    """系统工作区「未分类」的 UUID。"""
    return _project.SYSTEM_PROJECT_ID


def is_system_project_id(project_id):
    return _project.is_system_project_id(project_id)


def is_system_project(project):
    return _project.is_system_project(project)


def resolved_project_id(state):
    """任务所属项目 ID；**永远返回一个真实项目 UUID，或空字符串**。

    三种输入被刻意区分开，因为它们语义不同：

    - 字段缺失：迁移前创建的旧任务 -> 归入系统工作区「未分类」；
    - 显式 `None` / 空字符串：用户选择了「未分类」-> 同样是系统工作区。
      这是产品模型允许的状态，且**它仍然是一个真实的容器**：未分类的任务
      与普通项目走同一套读写，因此"打开项目"不会撞上不存在的记录；
    - 其它值：可能是历史别名 `"default"`（旧版本的虚拟 project id）或旧版本
      按名称派生的 slug。两者都经 `canonical_project_id` 归一，`"default"`
      收敛到系统项目 UUID，旧 slug 原样保留以便仍能打开旧项目记录。

    重要：本函数**不做磁盘 I/O**，"未分类"的落盘由变更路径
    （`ensure_system_project`）负责，读取路径保持无副作用。
    """
    raw = state or {}
    if "project_id" not in raw:
        return _project.SYSTEM_PROJECT_ID
    value = raw.get("project_id")
    if value is None or not str(value).strip():
        return _project.SYSTEM_PROJECT_ID
    return _project.canonical_project_id(value)


def list_projects():
    """列出磁盘上的项目（含系统工作区）。

    **纯读取：不创建任何东西。** 系统工作区由变更路径（创建项目、归档任务、
    提升记忆）通过 `ensure_system_project()` 落盘；只读展示用
    `system_project_view()` 在内存中给出。
    """
    return _project.list_projects(OUTPUT_DIR)


def project_sections():
    """界面用的分区视图：{system, active, archived}（同一次扫盘的分区）。"""
    return _project.split_projects(list_projects())


def system_project_view():
    """系统工作区「未分类」的只读表示：磁盘上没有时也在内存中给出。

    它带真实 UUID 与 `is_system=True`，因此界面渲染、路由、任务归属都能直接
    使用，不需要先写文件，也不会出现"虚拟 project id"。
    """
    existing = load_project(_project.SYSTEM_PROJECT_ID)
    return existing if existing is not None else _project.empty_system_project()


# 向后兼容别名：旧名字里的「默认项目」现在叫系统工作区「未分类」。
default_project_view = system_project_view


def ensure_default_project():
    """兼容别名：把系统工作区落盘（旧名字）。"""
    return ensure_system_project()


def ensure_system_project():
    """把系统工作区落盘（幂等）。只有变更路径调用它。"""
    _project.migrate_legacy_layout(OUTPUT_DIR)
    _project.ensure_system_project(OUTPUT_DIR)
    _migrate_legacy_system_tm()
    return system_project_view()


def _migrate_legacy_system_tm():
    """把历史遗留的 `projects/default/translation_memory.json` 补到根路径。

    系统工作区的 TM 只有一个权威位置：`outputs/translation_memory.json`
    （旧默认项目的全局记忆路径）。旧版本里某些路径曾经把 TM 写进
    `projects/<id>/translation_memory.json`；迁移时**只在权威文件不存在时**
    搬过来，绝不合并两份——那会造成"同一实体两份 TM 真值"。
    """
    authoritative = OUTPUT_DIR / "translation_memory.json"
    if authoritative.is_file():
        return False
    legacy = _project.projects_root(OUTPUT_DIR) \
        / _project.SYSTEM_PROJECT_ID / "translation_memory.json"
    if not legacy.is_file():
        return False
    try:
        authoritative.parent.mkdir(parents=True, exist_ok=True)
        legacy.replace(authoritative)
    except OSError:
        return False
    return True


def list_active_projects():
    """活动项目 + 系统工作区（不含已归档）：新建任务与选择器用它。"""
    return [p for p in list_projects() if not p.get("archived_at")]


def list_active_project_options():
    """活动项目选项，**始终包含系统工作区**（只读，不落盘）。

    系统工作区是「未分类」这个真实选择的落点：磁盘上还没有它的记录时，
    也在内存中作为一项给出，界面因此不会少一个选项、也不会多一个假项目。
    """
    projects = list_active_projects()
    if not any(_project.is_system_project(p) for p in projects):
        projects = [_project.empty_system_project(), *projects]
    return projects


def find_project_by_name(name, *, exclude=""):
    """按显示名称查找项目（读取路径）。名称不是主键，仅用于人工输入入口。"""
    return _project.find_by_name(list_projects(), name, exclude=exclude)


# ---- CRUD：创建 / 读取 / 更新 / 归档 / 删除 ----

PROJECT_DESCRIPTION_LIMIT = 600
PROJECT_NAME_LIMIT = 80


def validate_project_name(name, *, exclude=""):
    """校验项目显示名称；返回清理后的名称，不合法时抛 ValueError。

    名称不是主键，但仍然要唯一：项目页与 picker 都用名称做人工入口，
    重名会让"选哪个项目"变成猜谜。
    """
    cleaned = str(name or "").strip()
    if not cleaned:
        raise ValueError("项目名称不能为空")
    if len(cleaned) > PROJECT_NAME_LIMIT:
        raise ValueError(f"项目名称最多 {PROJECT_NAME_LIMIT} 个字符")
    if cleaned.casefold() == _project.SYSTEM_PROJECT_NAME.casefold():
        raise ValueError(f"「{_project.SYSTEM_PROJECT_NAME}」是系统工作区名称，"
                         "不能用作项目名称")
    if existing := find_project_by_name(cleaned, exclude=exclude):
        raise ValueError(f"已存在同名项目「{existing['name']}」")
    return cleaned


def create_project(name, description=""):
    """创建项目；返回项目记录。

    ID 是现场生成的 UUIDv4，**不由名称派生**：改名不换 ID，同名不同项目也不会
    互相覆盖。名称仍要求唯一（人工入口用名称检索）。
    """
    cleaned = validate_project_name(name)
    project = _project.empty_project(
        _project.new_project_id(), cleaned,
        description=str(description or "").strip()[:PROJECT_DESCRIPTION_LIMIT])
    return _project.save_project(OUTPUT_DIR, project)


def load_project(project_id):
    """按 ID 读取项目；不存在时返回 None（读取路径不创建任何东西）。"""
    return _project.load_project(OUTPUT_DIR, project_id)


def save_project(project):
    return _project.save_project(OUTPUT_DIR, project)


def require_project(project_id):
    """按 ID 读取项目；不存在时抛带原因的 ValueError。

    系统工作区在磁盘上还没有记录时返回内存视图——它的 ID 是固定 UUID，
    "记录还没落盘"不等于"项目不存在"。
    """
    if not str(project_id or "").strip():
        raise ValueError("必须指定项目")
    if _project.is_system_project_id(project_id):
        record = load_project(_project.SYSTEM_PROJECT_ID)
        return record if record is not None else system_project_view()
    project = load_project(project_id)
    if project is None:
        raise ValueError(f"项目不存在：{project_id}")
    return project


def rename_project(project_id, name):
    """改名（**不改 ID**）。返回更新后的项目。"""
    project = require_project(project_id)
    cleaned = validate_project_name(name, exclude=project["project_id"])
    return save_project(_project.set_metadata(project, name=cleaned))


def update_project(project_id, *, name=None, description=None):
    """更新项目基本信息（名称 / 描述）。名称变化走同一套唯一性校验。"""
    project = require_project(project_id)
    cleaned = None
    if name is not None:
        cleaned = validate_project_name(name, exclude=project["project_id"])
    updated = _project.set_metadata(project, name=cleaned, description=description)
    return save_project(updated)


def archive_project(project_id, archived=True):
    """归档 / 恢复项目。归档可恢复，不删除任务，也不删除项目记忆。

    系统工作区一律拒绝归档——这项保护先于磁盘状态判断，不能因为"记录还不存在"
    而失效。
    """
    if _project.is_system_project_id(project_id):
        raise ValueError(f"「{_project.SYSTEM_PROJECT_NAME}」是系统工作区，不能归档")
    project = load_project(project_id)
    if project is None:
        raise ValueError(f"项目不存在：{project_id}")
    return save_project(_project.set_archived(project, bool(archived)))


def restore_project(project_id):
    """恢复已归档项目（`archive_project(id, False)` 的语义化入口）。"""
    return archive_project(project_id, False)


def project_for_job(job_id, state=None):
    """返回任务所属项目；任务不存在时返回 None。

    **每个任务都属于某个真实项目**：没有归属的任务属于系统工作区「未分类」，
    因此这里返回的是带真实 UUID 的记录（磁盘上还没落盘时给出内存只读视图），
    而不是 None。调用方因此可以直接用 `project["project_id"]` 打开项目详情，
    不会撞上"项目不存在"。

    **纯读取**：不创建文件。需要真正落盘的地方（提升记忆、归档任务）会自己调用
    `ensure_system_project()`。
    """
    state = state if state is not None else load_job_state(job_id)
    if state is None:
        return None
    project_id = resolved_project_id(state)
    if not str(project_id or "").strip():
        return None
    project = load_project(project_id)
    if project is not None:
        return project
    if _project.is_system_project_id(project_id):
        return system_project_view()
    return None


def task_project_id_for_new_job(chosen_project):
    """新建任务落盘前的归属校验：**项目记录已不存在时按未分类处理**。

    返回 `(project_id_or_None, stale)`；`stale=True` 表示上下文里的项目已经没了，
    调用方负责如实告知用户（绝不静默改写归属）。

    为什么必须校验：任务归属的唯一来源是界面的 Project Context，而那个值活在
    session 里。删除可能发生在**另一个会话/标签页**，`_reload_project_state` 只清理
    执行删除的那一个会话。少了这一步，任务会带着一个永远解析不出项目的
    project_id 落盘，变成历史页上一张"不在任何项目下"的孤儿卡片
    （实测：两次孤儿任务的目录创建时间分别在其项目被删之后 29 秒与 13 秒）。

    纯读取：只查项目记录是否存在，不创建、不写入。空值 / 系统工作区一律归一成
    `None`（显式 null = "用户没有选择长期归属"，见 `resolved_project_id`）。
    """
    value = str(chosen_project or "").strip()
    if not value or _project.is_system_project_id(value):
        return None, False
    if load_project(value) is None:
        return None, True
    return value, False


def resolve_project_ref(project_id_or_name, *, create=False, description=""):
    """把"项目 ID 或名称"解析成真实项目记录。

    - 命中 ID（含历史别名 `default` -> 系统工作区）→ 直接返回；
    - 命中名称 → 返回同名项目（名称只是人工入口，解析后一律改用 ID）；
    - 都没命中：`create=True` 时新建一个 UUID 项目，否则抛 ValueError。
    """
    label = str(project_id_or_name or "").strip()
    if not label:
        raise ValueError("必须指定项目")
    if _project.is_system_project_id(label):
        return ensure_system_project()
    direct = load_project(label)
    if direct is not None:
        return direct
    if by_name := find_project_by_name(label):
        return by_name
    if create:
        return create_project(label, description=description)
    raise ValueError(f"项目不存在：{label}")


def assign_job_to_project(job_id, project_id_or_name):
    """把任务归入项目（人工动作）。

    参数是项目 ID 或项目名称：名称经 `find_project_by_name` 解析成 ID 之后
    一律用 ID 落盘。**不会**因为"名称没命中"就静默新建项目——那会让一次误输入
    产生一个幽灵项目；确实要新建请用 `create_project`。
    """
    state = load_job_state(job_id)
    if state is None:
        return None
    project = resolve_project_ref(project_id_or_name)
    state["project_id"] = project["project_id"]
    save_job_state(job_id, state)
    return project


def list_project_jobs(project_id):
    """项目的任务列表：由各任务的 project_id 派生，不另存一份 job_ids。

    系统工作区的列表包含所有"没有归属"的任务（字段缺失、显式 null、
    历史别名 `default`），因为它就是这些任务的容器。
    """
    target = _project.canonical_project_id(project_id) \
        if str(project_id or "").strip() else _project.SYSTEM_PROJECT_ID
    return [job for job in list_jobs()
            if resolved_project_id(job["state"]) == target]


def list_unassigned_jobs():
    """未分类任务（系统工作区「未分类」里的任务）。"""
    return list_project_jobs(_project.SYSTEM_PROJECT_ID)


def project_name_for_id(project_id):
    """项目 ID 的可读名称；找不到记录时如实说「未知项目」。

    空值 / 历史别名 `default` 都指向系统工作区，因此返回「未分类」——这正是
    "所有没有 projectId 的任务归入系统项目"的直接体现。
    """
    if not str(project_id or "").strip():
        return _project.SYSTEM_PROJECT_NAME
    project = load_project(project_id)
    if project is None and _project.is_system_project_id(project_id):
        project = system_project_view()
    return str(project["name"]) if project else "未知项目"


# 任务正在运行的状态：此时它已经读取过一个项目的记忆，改归属会让同一个任务
# 在两套记忆之间漂移，因此这些状态一律拒绝移动。
_ACTIVE_RUNTIME_STATUSES = {
    "resume_requested", "queued", "starting", "running",
    "waiting_external", "cancelling",
}


def job_is_active(job_id, state=None):
    """任务是否正在运行（不可改归属）。"""
    try:
        view = build_job_runtime_view(job_id, state)
        status = view.get("runtime_status") or view.get("status")
    except Exception:  # 运行状态不可读时不阻止用户操作，由后续流程兜底
        return False
    return str(status or "") in _ACTIVE_RUNTIME_STATUSES


def assign_jobs_to_project(job_ids, project_id_or_name, *,
                           move_translations=False):
    """把一个或多个任务移入项目（人工动作，支持批量）。

    - 正在运行的任务会被跳过：它已经读过某个项目的记忆，中途改归属会让同一个
      任务在两套记忆之间漂移；
    - 只改变归属。`move_translations=True` 时额外把该任务**已审校**的译文并入
      目标项目的翻译记忆——只增不改，不覆盖目标项目已有的译法；
    - 不迁移术语或风格：那些需要通过 Memory gate 显式提升
      （`promote_job_to_project`）。

    返回 {"project", "moved", "skipped", "tm_added"}。
    """
    target = resolve_project_ref(project_id_or_name, create=True)

    moved, skipped, tm_added = [], [], 0
    target_tm = load_tm(target["project_id"])
    for job_id in job_ids or []:
        state = load_job_state(job_id)
        if state is None:
            skipped.append({"job_id": job_id, "reason": "任务不存在"})
            continue
        if resolved_project_id(state) == target["project_id"]:
            skipped.append({"job_id": job_id, "reason": "已在该项目中"})
            continue
        if job_is_active(job_id, state):
            skipped.append({"job_id": job_id, "reason": "任务正在运行，无法改归属"})
            continue
        if move_translations:
            job_lang = state_target_lang(state)
            for pair in state.get("pairs") or []:
                source = str(pair.get("source") or "")
                target_text = str(pair.get("target") or "")
                if not pair.get("reviewed") or not _tm_eligible(source, target_text):
                    continue
                key = tm_scope_key(job_lang, source)
                if not key:
                    continue  # 目标语言无法证明：不迁移记忆，也不猜
                if key in target_tm:
                    continue  # 目标项目已有该语言下的译法：不覆盖
                tm_put(target_tm, source, target_text, job_lang)
                tm_added += 1
        state["project_id"] = target["project_id"]
        save_job_state(job_id, state)
        moved.append(job_id)
    if tm_added:
        save_tm(target_tm, target["project_id"])
    return {"project": target, "moved": moved, "skipped": skipped,
            "tm_added": tm_added}


def promote_job_to_project(job_id, *, actor="user", state=None):
    """把任务中**已被人工确认**的知识提升进项目记忆（Memory gate）。

    只提升：locked 术语、confirmed 风格规则、human 决定与其授权记录。
    候选术语与未审校输出一律留在任务里。

    任务不属于任何项目时抛 `ValueError`：提升是把知识写进某个容器，静默落到
    默认项目会让用户在不知情的情况下污染另一个项目。
    """
    state = state if state is not None else load_job_state(job_id)
    if state is None:
        return None
    project = project_for_job(job_id, state)
    if project is None:
        raise ValueError("这个任务还没有归入任何项目；请先把任务移入一个项目。")

    frozen = state.get("glossary_frozen") or {}
    glossary_source = frozen.get("entries") or state.get("glossary") or []
    human_actions = [item for item in state.get("human_actions") or []
                     if isinstance(item, dict)
                     and str(item.get("record_type") or "") == "human_decision"]

    merged = _project.merge_confirmed_knowledge(
        project,
        glossary=glossary_source,
        style_rules=state.get("confirmed_style_rules") or [],
        human_decisions=human_actions,
        source_job_id=job_id,
        actor=actor,
    )
    return _project.save_project(OUTPUT_DIR, merged)


def project_memory_view(project=None):
    """项目记忆摘要；TM 计数取自既有受控存储（项目文件不持有 TM）。"""
    project = project if project is not None else system_project_view()
    return _project.memory_view(project,
                                translation_memory_count=len(
                                    load_tm(project["project_id"])))


def project_summary(project):
    """项目页卡片需要的派生数据：记忆摘要 + 任务数（不落任何文件）。"""
    view = project_memory_view(project)
    jobs = list_project_jobs(project["project_id"])
    view["job_count"] = len(jobs)
    view["jobs"] = jobs
    return view


def _project_backup_dir():
    return _project.projects_root(OUTPUT_DIR) / "_deleted"


def delete_project(project_id, *, confirm_name, move_jobs_to=None, cascade=False):
    """删除项目（永久，不可逆），删除前先写可恢复备份。

    四重保护，缺一不可：

    1. 系统工作区「未分类」永远不可删除：它是所有无归属任务的容器；
    2. 必须逐字输入项目名称确认（`confirm_name`），避免误点；
    3. 项目下仍有任务时默认**拒绝删除**，提示先移动或删除任务。若显式指定
       `cascade=True`，则确认承担风险同时删除该项目下所有任务；
       可选参数 `move_jobs_to` 保留给程序化调用（先批量移走再删）；
    4. 删除前把完整项目记忆（含已审校译对）写入
       `outputs/projects/_deleted/<id>-<时间>.json`。

    备份就是 `import_project_memory` 能直接吃回去的格式，因此误删在实践中
    仍然可恢复。
    """
    if _project.is_system_project_id(project_id):
        raise ValueError(
            f"「{_project.SYSTEM_PROJECT_NAME}」是系统工作区，不能删除")
    project = load_project(project_id)
    if project is None:
        raise ValueError(f"项目不存在：{project_id}")

    label = str(confirm_name or "").strip()
    if label != project["name"]:
        raise ValueError(f"确认名称不匹配：请输入项目名称「{project['name']}」")

    jobs = list_project_jobs(project["project_id"])
    reassigned: list[str] = []
    deleted_jobs: list[str] = []
    if jobs:
        if cascade:
            for job in jobs:
                jid = job["job_id"]
                try:
                    if is_job_worker_alive(jid):
                        request_job_cancel(jid, force=True)
                except Exception:
                    pass
                delete_job(jid, allow_active=True)
                deleted_jobs.append(jid)
        elif move_jobs_to is None:
            raise ValueError(
                f"项目下仍有 {len(jobs)} 个任务，请先移动或删除这些任务，再删除项目")
        else:
            target = resolve_project_ref(move_jobs_to, create=False)
            result = assign_jobs_to_project([job["job_id"] for job in jobs],
                                            target["project_id"])
            reassigned = result["moved"]
            if result["skipped"]:
                reasons = "；".join(f"{item['job_id']}：{item['reason']}"
                                   for item in result["skipped"])
                raise ValueError(f"有任务无法移出，已取消删除：{reasons}")

    backup_dir = _project_backup_dir()
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    backup = backup_dir / f"{project['project_id']}-{stamp}.json"
    backup.write_text(export_project_memory(project["project_id"]),
                      encoding="utf-8")

    shutil.rmtree(_project.project_dir(OUTPUT_DIR, project["project_id"]),
                  ignore_errors=True)
    return {"project_id": project["project_id"], "name": project["name"],
            "backup": backup, "reassigned_jobs": reassigned,
            "deleted_jobs": deleted_jobs}


def list_project_conflicts(project_id):
    """项目的待决冲突（导入时发现、等待人工决定）。"""
    project = load_project(project_id)
    if project is None:
        return []
    return _project.pending_conflicts(project)


def resolve_project_conflict(project_id, conflict_id, *, adopt_incoming=True,
                             actor="用户"):
    """处理一条待决冲突：采纳导入版本，或保留本地版本。

    两种选择都会被记入项目 `promotion_log`，因此"为什么是这个译名"可追溯。
    翻译记忆的改动在这里写入，因为 TM 存储不归项目文件所有。
    """
    project = require_project(project_id)
    conflict = next((item for item in _project.pending_conflicts(project)
                     if item.get("conflict_id") == conflict_id), None)
    if conflict is None:
        raise ValueError("冲突不存在或已处理")
    updated, tm_change = _project.resolve_conflict(
        project, conflict, adopt_incoming=adopt_incoming, actor=actor)
    saved = save_project(updated)
    if tm_change:
        target_tm = load_tm(saved["project_id"])
        for key, record in tm_change.items():
            # 冲突记录里的 source 就是记忆**键**，因此作用域（目标语言）随之保留；
            # 记录字段缺失时从键前缀补上，让"这条记忆属于哪种语言"不依赖键格式。
            language, _ = tm_unscope_key(key)
            if language and not record.get("target_lang"):
                record["target_lang"] = language
            target_tm[key] = record
        save_tm(target_tm, saved["project_id"])
    return saved


def resolve_all_project_conflicts(project_id, *, adopt_incoming=False, actor="用户"):
    """批量处理全部待决冲突；返回处理条数。

    默认**保留本地**：批量操作最容易误伤，默认值必须是保守的那一边。
    """
    handled = 0
    for conflict in list(list_project_conflicts(project_id)):
        resolve_project_conflict(project_id, conflict["conflict_id"],
                                 adopt_incoming=adopt_incoming, actor=actor)
        handled += 1
    return handled


def export_project_memory(project_id):
    """导出项目记忆为可移植 JSON 文本（含该项目自己的已审校记忆）。

    `project_id` 可以是历史别名 `default`——它归一到系统工作区，因此旧链接
    与旧脚本不会撞上"项目不存在：default"。
    """
    project = require_project(project_id)
    payload = _project.export_memory(
        project, translation_memory=load_tm(project["project_id"]))
    return json.dumps(payload, ensure_ascii=False, indent=2) + "\n"


def import_project_memory(raw, *, name=None, description=""):
    """导入项目记忆（只增不改），返回 (项目, 报告)。

    同名项目会被**并入**，而不是新建重复项目；`name` 可以显式改名。名称没命中
    时新建一个 UUID 项目（ID 现场生成，不再由名称派生）。
    翻译记忆由本层合并，因为 TM 存储不归项目文件所有。
    """
    if isinstance(raw, (bytes, bytearray)):
        if len(raw) > _project.MAX_IMPORT_BYTES:
            raise ValueError("项目记忆文件过大，已拒绝导入")
        try:
            payload = json.loads(raw.decode("utf-8"))
        except UnicodeDecodeError as exc:
            raise ValueError("项目记忆文件必须是 UTF-8 编码的 JSON") from exc
        except ValueError as exc:
            raise ValueError(f"项目记忆文件不是合法 JSON：{exc}") from exc
    elif isinstance(raw, str):
        if len(raw.encode("utf-8")) > _project.MAX_IMPORT_BYTES:
            raise ValueError("项目记忆文件过大，已拒绝导入")
        payload = json.loads(raw)
    elif isinstance(raw, dict):
        payload = raw
    else:
        raise ValueError("无法识别的项目记忆输入")

    payload = _project.validate_memory_payload(payload)
    source_name = str(name or (payload.get("project") or {}).get("name") or "").strip()
    existing = find_project_by_name(source_name) if source_name else None
    project, report, imported_tm = _project.import_memory(
        existing, payload, name=name, project_id=_project.new_project_id(),
        description=description)
    saved = save_project(project)

    if imported_tm:
        current = load_tm(saved["project_id"])
        recorded = _project.record_tm_conflicts(
            saved, current, imported_tm,
            imported_from=str((payload.get("project") or {}).get("name") or ""))
        if recorded:
            report["conflicts_recorded"] = report.get("conflicts_recorded", 0) + recorded
            saved = save_project(saved)
        added = conflicts = 0
        for key, record in imported_tm.items():
            if key in current:
                if current[key].get("target") != record["target"]:
                    conflicts += 1
                continue
            language = _tm_scope.tm_record_language(key, record)
            if language:
                # 语言可证明：经 tm_put 写入，保证键与记录字段都是规范形态
                # （只靠键前缀传语言，任何一次键重写都会静默丢掉语言身份）。
                tm_put(current, _tm_scope.tm_unscope_key(key)[1],
                       record["target"], language)
            else:
                # 语言不可证明：条目保留，但永不参与自动命中（fail closed）。
                current[key] = record
            added += 1
        if added:
            save_tm(current, saved["project_id"])
        report["tm_added"] = added
        report["tm_conflicts_count"] = conflicts
    return saved, report


def project_injection(project_id):
    """新任务启动时注入的项目记忆（锁定术语 + 确认风格规则）。

    系统工作区同样可以持有记忆（它就是"未分类任务"的容器），因此这里不特殊
    跳过；找不到记录时才返回空注入。
    """
    if not str(project_id or "").strip():
        project_id = _project.SYSTEM_PROJECT_ID
    project = load_project(project_id)
    if project is None and _project.is_system_project_id(project_id):
        project = system_project_view()
    if project is None:
        return {"glossary": [], "style_rules": "", "glossary_version": None,
                "glossary_hash": ""}
    return _project.injection_for(project)


def _project_now_iso():
    from datetime import datetime, timezone
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _project_memory_injection(state):
    """读取任务所属项目的记忆，供管线注入。

    返回 (glossary_entries, style_text, meta)；任务**不属于任何项目**、项目不存在
    或记忆为空时返回空值，因此对「无项目」的任务与没有项目的旧任务都是完全无副
    作用的。
    """
    try:
        project_id = resolved_project_id(state)
        if not project_id:
            return [], "", {}
        project = load_project(project_id)
        if project is None:
            return [], "", {}
        injection = _project.injection_for(project)
    except Exception:  # 项目记忆损坏不应阻断翻译
        return [], "", {}
    meta = {
        "project_id": project_id,
        "glossary_version": injection.get("glossary_version"),
        "glossary_hash": injection.get("glossary_hash") or "",
    }
    return injection["glossary"], injection["style_rules"], meta


# ================= 翻译 / 修复 / 审校（对齐 localize-anything 的三通道）=================
def _translator_system(glossary_text, style_rules, target_lang):
    return (f"你是一位学术翻译专家，请将用户提供的段落翻译成{target_lang}。\n"
            f"规则：只翻译可翻译的正文；作者姓名、机构名、品牌名、URL、邮箱、DOI、"
            f"引用标注（如 [12]）等保留原文；译文须与原文一一对应并保持顺序。\n"
            f"{glossary_text}\n"
            f"{style_rules}\n"
            "翻译前请在内部检查：指代与照应关系、术语和专名、理论/非字面用法、"
            "临时造词、长句逻辑关系与修辞功能；优先复用已提供的术语、实体和连续性选择，"
            "人工锁定实体/术语优先于审校、TM 和生成式建议；生成式实体提示不得覆盖锁定术语。"
            "避免明显的逐词直译和词典首义机械替换。\n"
            "不要输出分析过程，只输出最终译文。\n"
            "请严格输出合法的 JSON 字符串数组，不要包含任何解释文字。")


def _invoke_llm(call_fn, provider, api_key, model, system_prompt, user_prompt,
                temperature, response_format=None):
    """Call old and new provider/test doubles through one compatibility gate."""
    kwargs = {"temperature": temperature}
    if response_format is not None:
        kwargs["response_format"] = response_format
    try:
        return call_fn(provider, api_key, model, system_prompt, user_prompt, **kwargs)
    except TypeError:
        try:
            return call_fn(provider, api_key, model, system_prompt, user_prompt,
                           temperature=temperature)
        except TypeError:
            return call_fn(provider, api_key, model, system_prompt, user_prompt)


def _native_translation_response_format(provider, model, expected):
    capabilities = _model_roles.provider_capabilities(PROVIDERS, provider, model)
    if capabilities.get("supports_json_schema") and capabilities.get(
            "supports_response_format"):
        return _translation_protocol.json_schema_for_translations(expected)
    return None


def translate_batch(segments, ctx_prev, ctx_next, glossary_text, style_rules, target_lang,
                    provider, api_key, model, context_packet=None, call_llm_fn=None):
    """翻译一个语义批次，返回与 segments 等长的译文列表；失败抛出 RuntimeError。"""
    numbered = "\n".join(f"{i + 1}. {s}" for i, s in enumerate(segments))
    if context_packet is not None:
        # Keep the system prefix invariant across batches.  Glossary/style are
        # carried once in the ordered context packet below.
        sys_prompt = _translator_system("", "", target_lang)
        user_prompt = _context.render_context_packet(context_packet)
    else:
        sys_prompt = _translator_system(glossary_text, style_rules, target_lang)
        context = ""
        if ctx_prev:
            context += "【前文上下文】：\n" + "\n".join(f"- {s}" for s in ctx_prev) + "\n\n"
        if ctx_next:
            context += "【后文上下文】：\n" + "\n".join(f"- {s}" for s in ctx_next) + "\n\n"
        user_prompt = f"{context}待翻译段落（按序号返回等长译文数组）：\n{numbered}"
    call_fn = call_llm_fn or call_llm
    response_format = _native_translation_response_format(
        provider, model, len(segments))
    last_err = None
    for _attempt in range(3):
        try:
            res = _invoke_llm(
                call_fn, provider, api_key, model, sys_prompt, user_prompt,
                temperature=0.3, response_format=response_format)
            return _translation_protocol.parse_translation_response(
                res, len(segments))
        except Exception as e:
            last_err = e
            if is_rate_limited(e):
                time.sleep(15)

    raise RuntimeError(
        f"批次翻译失败（{len(segments)} 段）："
        f"{last_err or '模型返回格式异常或数量不匹配'}"
    )


def repair_batch(sources, targets, findings, glossary_text, style_rules, target_lang,
                 provider, api_key, model, call_llm_fn=None):
    """根据确定性检查发现的问题自动修复一批译文；返回与 sources 等长的译文列表。"""
    numbered = "\n".join(
        f"{i + 1}. 原文：{s}\n   译文：{t}" for i, (s, t) in enumerate(zip(sources, targets)))
    issues = "\n".join(f"- 段落 {f['segment_index'] + 1}: [{f['severity']}] {f['reason']}"
                       for f in findings)
    sys_prompt = _translator_system(glossary_text, style_rules, target_lang)
    user_prompt = ("以下译文未通过检查，请仅修正有问题的段落，其余段落保持原样，"
                   f"返回与段落数相同的 JSON 字符串数组：\n\n{numbered}\n\n问题清单：\n{issues}")
    call_fn = call_llm_fn or call_llm
    response_format = _native_translation_response_format(
        provider, model, len(sources))
    for _attempt in range(3):
        try:
            res = _invoke_llm(
                call_fn, provider, api_key, model, sys_prompt, user_prompt,
                temperature=0.2, response_format=response_format)
            arr = _translation_protocol.parse_translation_response(
                res, len(sources))
            return [clean_xml_chars(item).strip() for item in arr]
        except Exception as e:
            if is_rate_limited(e):
                time.sleep(15)
    raise RuntimeError("自动修复失败：模型返回格式异常")


def review_translation_batch(sources, targets, glossary_text, style_rules, target_lang,
                             provider, api_key, model):
    """独立审校一个批次（与翻译分离的 prompt/上下文），返回 (findings, failed)。"""
    numbered = "\n".join(
        f"{i + 1}. 原文：{s}\n   译文：{t}" for i, (s, t) in enumerate(zip(sources, targets)))
    sys_prompt = (f"你是一位独立的翻译审校专家，负责审查机器译文。请检查：语义准确性、术语一致性、"
                  f"漏译/增译、目标语言自然度与风格。只报告真实存在的问题，"
                  f"不要为低风险或主观偏好制造 finding。\n"
                  f"severity 只允许以下三种：blocking（结构/占位符/语义严重错误）、"
                  f"actionable（应修正的问题）、informational（建议）。\n"
                  "如果整批译文没有问题，请严格返回空数组 []，不要输出任何 informational 备注。\n"
                  f"{glossary_text}\n"
                  f"{style_rules}\n"
                  '请严格输出 JSON 数组，每项格式：{"segment_index": 0, "severity": "actionable", '
                  '"reason": "问题说明", "suggested_target": "可选：修正后的译文"}')
    user_prompt = f"待审校段落（目标语言：{target_lang}）：\n{numbered}"
    for _attempt in range(3):
        try:
            res = call_llm(provider, api_key, model, sys_prompt, user_prompt, temperature=0.2)
            arr = parse_json_array(res)
            if arr is None:
                return [], True
            return arr, False
        except Exception as e:
            if is_rate_limited(e):
                time.sleep(15)
            else:
                break
    return [], True


# ================= 自动标注（三色学习重点）=================
# 红=生僻词/难词；黄=专业名词（特殊译法）；青绿=翻译难点句（特别译法）。
ANNOT_BATCH_SIZE = 10
ANNOT_MAX_PER_SEG = {"rare": 3, "domain": 3, "hard": 2}

# 常用英语词表（en_50k 字幕语料前 14000 词），用于把 LLM 滥标的"生僻词"挡回去
_DATA_DIR = Path(_models.__file__).resolve().parent / "resources"
_COMMON_WORDS = None


def _common_words():
    global _COMMON_WORDS
    if _COMMON_WORDS is None:
        try:
            _COMMON_WORDS = set(
                _DATA_DIR.joinpath("en_common.txt").read_text(encoding="utf-8").splitlines())
        except OSError:
            _COMMON_WORDS = set()
    return _COMMON_WORDS


_INFLECTION_SUFFIXES = (("ily", "y"), ("ness", ""), ("ment", ""), ("tion", ""),
                        ("sion", ""), ("ing", ""), ("ed", ""), ("er", ""),
                        ("est", ""), ("es", ""), ("ly", ""), ("s", ""))


def _base_form(word):
    """词形还原（一次）：grinning->grin、speedily->speedy、boxes->box。"""
    w = word
    for suffix, repl in _INFLECTION_SUFFIXES:
        if w.endswith(suffix) and len(w) > len(suffix) + 2:
            w = w[:-len(suffix)] + repl
            break
    if len(w) > 3 and w[-1] == w[-2] and w[-1] not in "eio":
        w = w[:-1]  # grinning -> grinn -> grin
    return w


def _is_common_word(token):
    """token 或其词形还原是否在常用词表内。"""
    common = _common_words()
    if not common:
        return False
    w = token.lower().strip('"\'(),;:!?')
    if not w:
        return True
    if w in common:
        return True
    base = _base_form(w)
    return base != w and base in common


_KINSHIP_TITLE_RE = re.compile(
    r"^(?:grandma|grandpa|grandmother|grandfather|aunt|uncle|mr|mrs|ms|dr|sir|lady|lord|"
    r"mother|father|mom|dad|brother|sister|captain|colonel|major|general|rabbi|"
    r"professor|doctor)\b", re.IGNORECASE)


def _rare_ok(span_text, token_freq):
    """生僻词门槛：单 token、非常用词（含词形还原）、全书出现次数少。"""
    span = (span_text or "").strip()
    if not span or " " in span or "\u00a0" in span:
        return False  # 只接受单词（可带连字符），短语/名句一律不要
    w = span.lower().strip('"\'(),;:!?')
    if len(w) < 4:
        return False
    if _is_common_word(w):
        return False
    if token_freq.get(w, 0) >= 8:
        return False  # 全书反复出现的词不算生僻
    return True


def _domain_ok(span_text):
    """专业名词门槛：称谓+人名、全常用词短语不算专业名词。"""
    span = (span_text or "").strip()
    if not span:
        return False
    if _KINSHIP_TITLE_RE.match(span):
        return False  # Grandma Sarah / Mr. Smith 之类
    tokens = re.findall(r"[A-Za-z]+(?:['’\-][A-Za-z]+)*", span)
    if not tokens:
        return False
    if all(_is_common_word(t) for t in tokens):
        return False  # Translation from Hebrew 之类全是常用词
    return True


def _normalize_annotations(annotations):
    """标注字典键归一化：JSON 落盘后键变字符串，统一回 int；越界/非列表值丢弃。

    这是本项目反复踩过的坑（标注渲染、过滤、续跑三处各自处理过一次），
    统一入口避免再次各修各的。
    """
    out = {}
    for k, v in (annotations or {}).items():
        try:
            gi = int(k)
        except (TypeError, ValueError):
            continue
        if isinstance(v, list):
            out[gi] = v
    return out


def _clean_annotations(annotations, pairs):
    """过滤 + 去重 + 数量上限；术语表覆盖的 domain（note 以"术语："开头）不参与过滤。"""
    token_freq = {}
    for pr in pairs:
        for tok in re.findall(r"[A-Za-z]+", pr["source"].lower()):
            token_freq[tok] = token_freq.get(tok, 0) + 1
    cleaned = {}
    for gi, items in _normalize_annotations(annotations).items():
        if not (0 <= gi < len(pairs)):
            continue  # 段已被删除等场景：越界键丢弃
        src_text = pairs[gi]["source"]
        seen, kept, counts = set(), [], {"rare": 0, "domain": 0, "hard": 0}
        for it in items:
            atype = it["type"]
            s_span = it.get("src_span")
            span_text = src_text[s_span[0]:s_span[1]] if s_span else ""
            overlay = atype == "domain" and str(it.get("note", "")).startswith("术语：")
            if atype == "rare" and not _rare_ok(span_text, token_freq):
                continue
            if atype == "domain" and not overlay and not _domain_ok(span_text):
                continue
            key = (atype, tuple(s_span) if s_span else None,
                   tuple(it["tgt_span"]) if it.get("tgt_span") else None)
            if key in seen:
                continue
            seen.add(key)
            if counts[atype] >= ANNOT_MAX_PER_SEG[atype]:
                continue
            counts[atype] += 1
            kept.append(it)
        if kept:
            cleaned[gi] = kept
    return cleaned


def _annotator_system(target_lang):
    return (f"你是一位翻译教学专家。请从下列{target_lang}双语对照中标注三类学习重点：\n"
            "1. rare：真正的生僻词/难词——英语母语者也未必认识的低频书面词，"
            "如 chicory、muezzin、cacophony。只标注单个单词（最多一个带连字符的复合词）。\n"
            "   严禁标注日常常用词（如 production、grin、rooster、speedily、elementary），"
            "严禁标注短语、引用句或名句（如 'Elementary, my dear Watson'）。\n"
            "2. domain：专业领域名词，译文采用了专门/约定俗成的译法（如术语表、行话、专名译法）；\n"
            "   严禁标注亲属称谓（如 Grandma Sarah）、全常用词短语（如 Translation from Hebrew）、"
            "普通日常表达。\n"
            "3. hard：翻译难度高的句子，译文使用了特别翻译技巧（语序调整、词性转换、拆合句、"
            "文化负载词处理、比喻/双关处理等）；普通直译句不要标注。\n"
            "规则：\n"
            "- 每个段落最多 2 个 rare、2 个 domain、1 个 hard；标注真正有价值的，宁缺毋滥；\n"
            "- src 必须是原文中的原文字符串，tgt 必须是对应译文中的字符串（hard 可以是整句/整段）；\n"
            "- note 用一句话说明标注理由或所用译法；\n"
            '严格输出 JSON 数组，每项格式：{"seg": 1, "type": "rare", "src": "...", '
            '"tgt": "...", "note": "..."}。seg 为段落序号（从 1 开始）。不要输出任何解释文字。')


def annotate_batch(pairs_slice, target_lang, provider, api_key, model):
    """标注一个批次，返回 [{seg, type, src, tgt, note}]（seg 为 0 基）；解析失败返回 []。"""
    numbered = "\n\n".join(
        f"--- 段落 {i + 1} ---\n原文：{p['source']}\n译文：{p['target']}"
        for i, p in enumerate(pairs_slice))
    sys_prompt = _annotator_system(target_lang)
    user_prompt = f"待标注双语段落：\n{numbered}"
    for _attempt in range(3):
        try:
            res = call_llm(provider, api_key, model, sys_prompt, user_prompt, temperature=0.2)
            arr = parse_json_array(res)
            if arr is None:
                return []
            out = []
            for item in arr:
                if not isinstance(item, dict):
                    continue
                seg = item.get("seg")
                if not isinstance(seg, int) or not (1 <= seg <= len(pairs_slice)):
                    continue
                atype = item.get("type")
                if atype not in ANNOTATION_COLORS:
                    continue
                src = str(item.get("src") or "").strip()
                tgt = str(item.get("tgt") or "").strip()
                note = str(item.get("note") or "").strip()
                if not src:
                    continue
                out.append({"seg": seg - 1, "type": atype, "src": src,
                            "tgt": tgt, "note": note})
            return out
        except Exception as e:
            if is_rate_limited(e):
                time.sleep(10)  # 限流退避后重试，避免整批标注静默丢失
                continue
            return []
    return []


def _compose_spans(spans, text_len):
    """按边界切分 + 优先级（rare>domain>hard）覆盖，返回互不重叠的 (start,end,type) 列表。

    难点句常覆盖整段，词级标注嵌在其中：切到所有起点/终点后，每段取覆盖它的最高优先级。
    """
    priority = {"rare": 0, "domain": 1, "hard": 2}
    clipped = []
    for s, e, t in spans:
        s = max(0, min(int(s), text_len))
        e = max(s, min(int(e), text_len))
        if s < e:
            clipped.append((s, e, t))
    if not clipped:
        return []
    bounds = sorted({0, text_len} | {x for s, e, _ in clipped for x in (s, e)})
    composed = []
    for a, b in zip(bounds, bounds[1:]):
        if a >= b:
            continue
        mid = (a + b) / 2
        covering = [(priority[t], t) for s, e, t in clipped if s <= mid < e]
        if covering:
            composed.append((a, b, min(covering)[1]))
    return composed


def annotate_stage(state, job_id, glossary, provider, api_key, model, target_lang,
                   on_caption=None):
    """三色自动标注：LLM 识别 + 术语表确定性覆盖（专业名词=黄色必标）。"""
    if state.get("annotations_done"):
        return state
    pairs = state["pairs"]
    annotations = _normalize_annotations(state.get("annotations"))
    batches = make_batches([p["source"] for p in pairs], batch_size=ANNOT_BATCH_SIZE,
                           max_chars=2600)
    # 断点按“已标注段数”记录（而非批次号），避免批大小调整后续跑错位
    done_offset = state.get("annotations_done_offset") or 0
    start_bi = next((bi for bi, b in enumerate(batches)
                     if sum(len(x) for x in batches[:bi]) >= done_offset), len(batches))
    if done_offset:
        if on_caption:
            on_caption(f"↩️ 从第 {done_offset + 1} 段（批次 {start_bi + 1}/{len(batches)}）继续标注...")

    # 1) LLM 标注
    failed_batches = 0
    for bi in range(start_bi, len(batches)):
        batch_srcs = batches[bi]
        offset = sum(len(b) for b in batches[:bi])
        slice_pairs = pairs[offset:offset + len(batch_srcs)]
        if on_caption and bi % 10 == 0:
            on_caption(f"🎨 自动标注第 {offset + 1}-{offset + len(slice_pairs)} 段"
                       f"（共 {len(pairs)} 段，批次 {bi + 1}/{len(batches)}）...")
        items = annotate_batch(slice_pairs, target_lang, provider, api_key, model)
        if not items:
            failed_batches += 1
        for item in items:
            gi = offset + item["seg"]
            src, tgt, atype, note = item["src"], item["tgt"], item["type"], item["note"]
            src_span = _find_span(pairs[gi]["source"], src)
            tgt_span = _find_span(pairs[gi]["target"], tgt) if tgt else None
            if atype == "hard":
                # 难句兜底：找不到精确片段时标整段
                if src_span is None:
                    src_span = (0, len(pairs[gi]["source"]))
                if tgt_span is None:
                    tgt_span = (0, len(pairs[gi]["target"])) if pairs[gi]["target"] else None
            elif src_span is None:
                continue  # 词级标注必须在原文中定位
            annotations.setdefault(gi, []).append(
                {"type": atype, "src_span": list(src_span) if src_span else None,
                 "tgt_span": list(tgt_span) if tgt_span else None, "note": note})
        # 每批落盘：断点粒度 = 一个批次
        state["annotations"] = annotations
        state["annotations_done_offset"] = offset + len(slice_pairs)
        save_job_state(job_id, state)

    # 2) 术语表确定性覆盖：专业名词（特殊译法）-> 黄色
    for gi, pr in enumerate(pairs):
        for entry in glossary:
            term = (entry.get("source") or "").strip()
            if len(term) < 2 or term not in pr["source"]:
                continue
            span = _find_span(pr["source"], term)
            tgt_span = None
            tgt_term = (entry.get("target") or "").strip()
            if tgt_term:
                tgt_span = _find_span(pr["target"], tgt_term)
            annotations.setdefault(gi, []).append(
                {"type": "domain", "src_span": list(span) if span else None,
                 "tgt_span": list(tgt_span) if tgt_span else None,
                 "note": f"术语：{term} -> {tgt_term or '保留原文'}"})

    # 3) 确定性过滤（常用词/称谓/全常用词短语）+ 去重 + 数量上限
    cleaned = _clean_annotations(annotations, pairs)
    state["annotations"] = cleaned
    state["annotations_done"] = True
    state["annotations_failed_batches"] = failed_batches
    save_job_state(job_id, state)
    if on_caption:
        total = sum(len(v) for v in cleaned.values())
        on_caption(f"✅ 自动标注完成：{total} 处（失败批次 {failed_batches}/{len(batches)}）")
    return state


def translate_stage(state, job_id, glossary, provider, api_key, model, target_lang,
                    style_rules, enable_review, use_tm=True, document_profile=None,
                    on_status=None, on_caption=None, translator_config=None,
                    reviewer_config=None, batch_size=None, max_batch_chars=None,
                    auxiliary_config=None, knowledge_feedback_interval=None,
                    translation_concurrency=None):
    """阶段二：语义批次翻译 + 确定性检查/修复 + 独立审校 + 翻译记忆。

    对齐 localize-anything 经验：
    - 语义批次：≤4 段一组，携带前后文，保留每批落盘的断点粒度；
    - 概念化术语表：锁定术语强制首选译名/禁止译名，保留项强制原样；
    - 确定性检查：占位符/URL/引用等保留项、残留原文、锁定术语合规，问题自动修复一轮；
    - 独立审校：actionable 建议经确定性复验后应用，blocking 记录给用户确认；
    - 翻译记忆：仅审校通过的段落入库，精确命中直接复用。
    """
    state["translation_core_review_required"] = bool(enable_review)
    # 严格术语治理门禁：存在待审核候选术语且未冻结（且未显式跳过）时，
    # 任何入口都禁止开始翻译。导入的锁定术语视为已固定，不构成阻塞。
    pending = [e for e in (glossary or [])
               if (e.get("status") or "").lower() == "candidate"]
    if state.get("quality_mode") and pending \
            and not state.get("glossary_frozen") \
            and not state.get("quality_bypass"):
        raise RuntimeError(
            "严格术语治理：仍有候选术语未审核（术语表尚未冻结），"
            "禁止开始翻译（请在术语审核面板冻结后继续）")
    translator_config = _model_roles.normalize_role_config(
        translator_config, fallback_provider=provider, fallback_model=model,
        fallback_api_key=api_key)
    reviewer_config = _model_roles.normalize_role_config(
        reviewer_config, fallback_provider=provider, fallback_model=model,
        fallback_api_key=api_key)
    auxiliary_config = _model_roles.normalize_role_config(
        auxiliary_config, fallback_provider=provider, fallback_model=model,
        fallback_api_key=api_key)
    usage_ledger = state.setdefault("llm_usage", _usage.empty_usage())
    translator_call = _model_roles.make_role_call(
        _usage.tracked_call(call_llm, usage_ledger, role="translation"),
        translator_config)
    reviewer_call = _model_roles.make_role_call(
        _usage.tracked_call(call_llm, usage_ledger, role="review"),
        reviewer_config)
    auxiliary_call = _model_roles.make_role_call(
        _usage.tracked_call(call_llm, usage_ledger, role="auxiliary"),
        auxiliary_config)
    _tm_project = tm_project_id(state)
    # 目标语言是 TM 作用域的一部分：没有它就不建立任何命中路径（fail closed）。
    _tm_lang = str(target_lang or "").strip() or str(state.get("target_lang") or "").strip()
    tm = load_tm(_tm_project) if use_tm else {}
    tm_norm_index = tm_index(tm, _tm_lang) if use_tm else {}
    if use_tm:
        def _tm_scope(source, target):
            """把目标语言编进记忆键；语言无法证明时不恢复任何记忆。"""
            key = tm_scope_key(_tm_lang, source)
            if not key:
                return None
            return key, tm_record(target, _tm_lang)

        recovered, pending_events = _checkpoint.reconcile_translation_memory(
            tm, state, job_dir(job_id), scope=_tm_scope)
        if recovered:
            save_tm(tm, _tm_project)
            state["tm_recovered_count"] = state.get("tm_recovered_count", 0) + pending_events
            save_job_state(job_id, state)
    paras = state["paras"]
    pairs = state["pairs"]
    # Source cleanup repairs text, but a repair pass cannot prove that every
    # OCR item is a real paragraph.  Quarantine short/structurally suspicious
    # items before they can become model translation requests.
    source_quality_gate = _audit_source_quality_for_state(state)
    state["source_quality_gate"] = source_quality_gate
    blocked_source_indexes = {
        int(item["segment_index"]): item
        for item in source_quality_gate.get("flags") or []
        if isinstance(item, dict) and isinstance(item.get("segment_index"), int)
    }
    truncated_indexes = []

    def _commit_translation_batch(batch_pairs, offset):
        nonlocal truncated_indexes
        pairs.extend(batch_pairs)
        changed_indexes = list(range(offset, offset + len(batch_pairs)))
        if changed_indexes:
            _mark_translation_truth_changed(
                job_id, state, changed_indexes,
                "翻译流水线写入 CURRENT_TRANSLATION；记录该批次影响范围",
                actor="pipeline", action="translation_batch",
                stale_translation_reviews=False)
        truncated_indexes = []
        save_job_state(job_id, state)

    def _save_translation_failure():
        nonlocal truncated_indexes
        if truncated_indexes:
            _mark_translation_truth_changed(
                job_id, state, truncated_indexes,
                "断点恢复截断未完成批次；CURRENT_TRANSLATION 已改变",
                actor="pipeline", action="translation_checkpoint_truncate")
            truncated_indexes = []
        save_job_state(job_id, state)

    # 批次大小直接影响调用次数与耗时（本仓库 82 段文档实测：4/2400 → 32 批，
    # 6/4800 → 17 批）。默认沿用模块常量以保持既有行为，只有显式配置时才改变。
    effective_batch_size = int(batch_size or BATCH_SIZE)
    effective_max_chars = int(max_batch_chars or TRANSLATION_MAX_BATCH_CHARS)
    state["batch_plan"] = {
        "batch_size": effective_batch_size,
        "max_batch_chars": effective_max_chars,
    }
    batches = make_batches(
        paras, batch_size=effective_batch_size, max_chars=effective_max_chars,
        semantic_units=state.get("semantic_units")
        or state.get("section_digests") or None)
    state["batch_plan"]["batch_count"] = len(batches)
    # Knowledge extraction is a continuity aid, not a second translation pass.
    # In strict/review mode it remains per-batch; in ordinary mode it observes
    # every batch until continuity exists, then observes at a bounded interval.
    if knowledge_feedback_interval is None:
        feedback_interval = 1 if (enable_review or state.get("quality_mode")) else 2
    else:
        try:
            feedback_interval = max(1, int(knowledge_feedback_interval))
        except (TypeError, ValueError):
            feedback_interval = 1
    state["knowledge_feedback_policy"] = {
        "interval": feedback_interval,
        "force_every_batch": bool(enable_review or state.get("quality_mode")),
        "batch_count": len(batches),
        "observed_batches": [],
        "skipped_batches": [],
    }
    registry = _entity_registry.EntityRegistry(state.get("entity_registry") or [])

    # 断点：从第一个未完成批次继续；若中间批次不完整则截断重译
    cum_end, start_batch = 0, 0
    for bi, b in enumerate(batches):
        prev_end = cum_end
        cum_end += len(b)
        if cum_end > len(pairs):
            if len(pairs) > prev_end:
                truncated_indexes = list(range(prev_end, len(pairs)))
                del pairs[prev_end:]
            start_batch = bi
            break
    else:
        start_batch = len(batches)

    stats = state.setdefault("review_stats", {
        "reviewed_segments": 0, "batches_reviewed": 0,
        "blocking": 0, "actionable": 0, "informational": 0, "review_failed": 0,
    })
    findings_all = state.setdefault("findings", [])
    existing_source_gate_indexes = {
        int(item.get("segment_index"))
        for item in findings_all
        if isinstance(item, dict)
        and item.get("type") == "source_quality_gate"
        and isinstance(item.get("segment_index"), int)
    }
    for segment_index, flag in blocked_source_indexes.items():
        if segment_index in existing_source_gate_indexes:
            continue
        findings_all.append(_source_quality_finding(flag))
    from transpraxis.terminology import (
        detect_glossary_conflicts as _detect_conflicts,
        glossary_block as _glossary_block,
        select_glossary_for_segments as _select_glossary,
    )
    frozen_hash = (state.get("glossary_frozen") or {}).get("glossary_hash")
    sections = (document_profile or {}).get("sections") or []
    section_digests = state.get("section_digests") or []
    document_synopsis = state.get("document_synopsis") or {}

    offsets = []
    _cur_offset = 0
    for b in batches:
        offsets.append(_cur_offset)
        _cur_offset += len(b)

    effective_concurrency = _runtime_int_option(
        translation_concurrency,
        (state.get("pipeline_config") or {}).get("translation_concurrency"),
        "FOLIOTHREAD_TRANSLATION_CONCURRENCY",
        DEFAULT_TRANSLATION_CONCURRENCY,
        minimum=1,
        maximum=16,
    )
    effective_concurrency = max(1, min(int(effective_concurrency), max(1, len(batches) - start_batch)))
    state["batch_plan"]["concurrency"] = effective_concurrency

    checkpoint_lock = threading.RLock()
    commit_lock = threading.RLock()
    active_batches = set()

    def safe_append_event(event):
        with checkpoint_lock:
            _checkpoint.append_event(job_dir(job_id), event)

    inherited_base_url = getattr(_LLM_CTX, "base_url", None)
    inherited_reasoning = getattr(_LLM_CTX, "reasoning_effort", None)
    inherited_job_id = getattr(_RUNTIME_CTX, "job_id", job_id)

    def _process_batch_payload(bi):
        if inherited_base_url is not None:
            _LLM_CTX.base_url = inherited_base_url
        if inherited_reasoning is not None:
            _LLM_CTX.reasoning_effort = inherited_reasoning
        if inherited_job_id is not None:
            _RUNTIME_CTX.job_id = inherited_job_id

        if _runtime_cancel_requested(job_id):
            raise RuntimeError("任务已请求取消")
        batch = batches[bi]
        offset = offsets[bi]
        safe_append_event({
            "batch": bi, "offset": offset, "phase": "generation_started",
            "segment_count": len(batch),
        })

        with commit_lock:
            active_batches.add(bi)
        if on_caption:
            active_str = " ".join(f"#{i + 1}" for i in sorted(active_batches))
            on_caption(f"🌍 正在翻译第 {offset + 1}-{offset + len(batch)} 段（并发批次 {active_str}，共 {len(paras)} 段）...")

        try:
            ctx_prev = [
                para for index, para in enumerate(paras[max(0, offset - 2):offset],
                                                 start=max(0, offset - 2))
                if index not in blocked_source_indexes
            ]
            ctx_next = [
                para for index, para in enumerate(
                    paras[min(len(paras), offset + len(batch)):
                          min(len(paras), offset + len(batch) + 2)],
                    start=min(len(paras), offset + len(batch)))
                if index not in blocked_source_indexes
            ]
            with commit_lock:
                previous_target = [
                    item for item in _context.select_target_context(
                        pairs, offset, limit=max(2, offset))
                    if item.get("segment_index") not in blocked_source_indexes
                ][-2:]
            section_digest = _context.digest_for_segment(section_digests, offset)

            # 1) 翻译记忆精确命中直接复用
            batch_pairs = [None] * len(batch)
            to_translate = []  # (index, clean_source)
            blocked_local_indexes = set()
            for i, para in enumerate(batch):
                clean_src = para.replace('\n', ' ')
                source_index = offset + i
                source_flag = blocked_source_indexes.get(source_index)
                if source_flag is not None:
                    blocked_local_indexes.add(i)
                    batch_pairs[i] = {
                        "source": clean_src, "target": clean_src,
                        "initial_target": clean_src,
                        "accepted_target": clean_src,
                        "target_provenance": "source_review_required",
                        "reviewed": False,
                        "review_status": "source_review_required",
                        "from_tm": False,
                        "source_quality": dict(source_flag),
                    }
                    continue
                hit, tm_match = tm_lookup(tm, clean_src, tm_norm_index,
                                          target_lang=_tm_lang)
                if hit:
                    batch_pairs[i] = {"source": clean_src, "target": hit["target"],
                                      "initial_target": hit["target"],
                                      "accepted_target": hit["target"],
                                      "target_provenance": "tm_approved",
                                      "reviewed": True, "review_status": "tm_approved",
                                      "from_tm": True, "tm_match": tm_match}
                    with commit_lock:
                        state["tm_used_count"] = state.get("tm_used_count", 0) + 1
                elif not has_textual_content(clean_src):
                    batch_pairs[i] = {"source": clean_src, "target": clean_src,
                                      "initial_target": clean_src,
                                      "accepted_target": clean_src,
                                      "target_provenance": "reviewed",
                                      "reviewed": True, "review_status": "reviewed_clean",
                                      "from_tm": False}
                else:
                    to_translate.append((i, clean_src))

            texts = [t for _, t in to_translate]
            section_profile = _batch_section_profile(document_profile, offset, len(batch))
            with commit_lock:
                provisional_hints = _knowledge.provisional_hints(
                    state.get("knowledge_candidates") or [],
                    authoritative_entries=glossary)
                entity_hints = registry.hints_for(
                    texts, glossary_entries=glossary, limit=12)
            selected, injected_ids = _select_glossary(
                texts, glossary + provisional_hints,
                document_profile, section_profile)
            glossary_text = _glossary_block(selected)
            context_packet = _context.compile_context_packet(
                document_profile=document_profile,
                document_synopsis=document_synopsis,
                section_digest=section_digest,
                glossary_text=glossary_text,
                previous_source=ctx_prev,
                previous_target=previous_target,
                next_source=ctx_next,
                current_batch=texts,
                style_rules=style_rules,
                entity_hints=entity_hints,
                context_budget_chars=max(5200, min(12000, effective_max_chars * 3)),
            )
            review_context = _runtime_review_context(
                state, offset, len(batch), glossary_text, style_rules, target_lang)

            # 2) 未命中段落批次翻译
            if to_translate:
                try:
                    targets = translate_batch(texts, ctx_prev, ctx_next, glossary_text, style_rules,
                                              target_lang, translator_config["provider"],
                                              translator_config["api_key"],
                                              translator_config["model"],
                                              context_packet=context_packet,
                                              call_llm_fn=translator_call)
                except RuntimeError as batch_error:
                    if "任务已请求取消" in str(batch_error):
                        raise
                    failure = {
                        "batch": bi, "offset": offset, "segment_count": len(texts),
                        "reason": str(batch_error)[:500], "recovered": False,
                    }
                    with commit_lock:
                        state.setdefault("translation_failures", []).append(failure)
                    safe_append_event({
                        "batch": bi, "offset": offset,
                        "phase": "translation_protocol_failed",
                        "reason": str(batch_error)[:240],
                    })
                    if len(texts) == 1:
                        with commit_lock:
                            _save_translation_failure()
                        raise
                    if on_caption:
                        on_caption(f"⚠️ 批次 {bi + 1} 翻译返回格式异常，降级为逐段翻译...")
                    targets = []
                    try:
                        for t in texts:
                            if _runtime_cancel_requested(job_id):
                                raise RuntimeError("任务已请求取消")
                            single_packet = dict(context_packet, current_batch=[t])
                            targets.append(translate_batch(
                                [t], ctx_prev, ctx_next, glossary_text, style_rules, target_lang,
                                translator_config["provider"], translator_config["api_key"],
                                translator_config["model"], context_packet=single_packet,
                                call_llm_fn=translator_call)[0])
                    except Exception as single_error:
                        failure["reason"] = str(single_error)[:500]
                        with commit_lock:
                            _save_translation_failure()
                        raise
                    failure["recovered"] = True
                for (i, src), tgt in zip(to_translate, targets):
                    cleaned_tgt = clean_xml_chars(tgt).replace('\n', ' ')
                    batch_pairs[i] = {"source": src, "target": cleaned_tgt,
                                      "initial_target": cleaned_tgt,
                                      "target_provenance": "generated",
                                      "reviewed": False, "review_status": "not_reviewed",
                                      "from_tm": False}

            batch_sources = [p["source"] for p in batch_pairs]
            batch_targets = [p["target"] for p in batch_pairs]
            for p in batch_pairs:
                p["glossary_entry_ids"] = list(injected_ids)
                p["glossary_hash_used"] = frozen_hash
            safe_append_event({
                "batch": bi, "offset": offset, "phase": "generation_done",
            })
            findings = [
                finding for finding in check_translation_batch(
                    batch_sources, batch_targets, glossary, target_lang,
                    section_profile=section_profile)
                if finding.get("segment_index") not in blocked_local_indexes
            ]

            # 3) 确定性问题自动修复（一轮）
            fixable = [
                f for f in findings
                if f["severity"] in ("blocking", "actionable")
                and f.get("segment_index") not in blocked_local_indexes
            ]
            repair_overlays_to_record = []
            review_evidence_to_record = []
            if fixable and len(fixable) <= 8:
                if on_caption:
                    on_caption(f"🔧 批次 {bi + 1} 发现 {len(fixable)} 个确定性问题，正在自动修复...")
                try:
                    repaired = repair_batch(batch_sources, batch_targets, fixable, glossary_text,
                                            style_rules, target_lang,
                                            translator_config["provider"],
                                            translator_config["api_key"],
                                            translator_config["model"],
                                            call_llm_fn=translator_call)
                    if repaired and len(repaired) == len(batch_pairs):
                        formal_targets = list(batch_targets)
                        shadow_targets = list(formal_targets)
                        for j, p in enumerate(batch_pairs):
                            if j in blocked_local_indexes:
                                continue
                            if not p["from_tm"] and repaired[j] and repaired[j].strip():
                                candidate = clean_xml_chars(repaired[j]).replace('\n', ' ')
                                if is_incomplete_translation(batch_sources[j], candidate) \
                                        and not is_incomplete_translation(batch_sources[j], formal_targets[j]):
                                    continue
                                shadow_targets[j] = candidate
                        overlay = _repair.create_overlay(
                            formal_targets, shadow_targets, fixable, "deterministic",
                            sources=batch_sources,
                            finding_segment_ids=[offset + f["segment_index"] for f in fixable])
                        shadow_findings = check_translation_batch(
                            batch_sources, shadow_targets, glossary, target_lang,
                            section_profile=section_profile)
                        shadow_findings = _globalize_batch_findings(shadow_findings, offset)
                        blind_findings, blind_failed, blind_trace = [], False, None
                        if enable_review and not any(
                                f["severity"] in ("blocking", "actionable")
                                for f in shadow_findings):
                            with commit_lock:
                                current_pairs = list(pairs)
                                current_findings_all = list(findings_all)
                            shadow_index = _translation_evidence_index(
                                paras, current_pairs, batch_pairs, glossary, document_profile,
                                document_synopsis, section_digests, current_findings_all,
                                blind=True,
                                candidate_targets={offset + j: shadow_targets[j]
                                                   for j in range(len(shadow_targets))})
                            shadow_packet = _translation_evidence.build_runtime_review_packet(
                                state, batch_pairs,
                                list(range(offset, offset + len(batch_pairs))),
                                glossary,
                                deterministic_checks=shadow_findings,
                                review_context=review_context,
                                candidate_targets={offset + j: shadow_targets[j]
                                                   for j in range(len(shadow_targets))},
                                blind=True,
                            )
                            blind_findings, blind_failed, blind_trace = \
                                _translation_evidence.review_translation_batch_with_evidence(
                                    batch_sources, shadow_targets, glossary_text, style_rules,
                                    target_lang, reviewer_config["provider"],
                                    reviewer_config["api_key"], reviewer_config["model"],
                                    shadow_index, call_llm=reviewer_call, blind=True,
                                    segment_ids=list(range(offset, offset + len(batch_pairs))),
                                    review_identity={
                                        "input_hash": overlay["input_hash"],
                                        "candidate_hash": overlay["candidate_hash"],
                                    }, translation_core_packet=shadow_packet)
                            if blind_trace:
                                review_evidence_to_record.append({
                                    "batch": bi, "phase": "shadow_repair", **blind_trace})
                        overlay = _repair.evaluate_overlay(
                            overlay, shadow_findings, blind_findings, blind_failed,
                            review_identity=(blind_trace or {}).get("review_identity"))
                        repair_overlays_to_record.append({
                            "batch": bi, "offset": offset, **overlay,
                            "blind_trace": blind_trace,
                        })
                        promoted = _repair.promoted_targets(overlay)
                        for j, p in enumerate(batch_pairs):
                            if not p["from_tm"]:
                                p["target"] = promoted[j]
                    batch_targets = [p["target"] for p in batch_pairs]
                    findings = [
                        finding for finding in check_translation_batch(
                            batch_sources, batch_targets, glossary, target_lang,
                            section_profile=section_profile)
                        if finding.get("segment_index") not in blocked_local_indexes
                    ]
                except Exception as exc:
                    repair_overlays_to_record.append({
                        "batch": bi, "offset": offset, "source": "deterministic",
                        "status": "rejected", "rejection": "repair_error",
                        "error": str(exc)[:240],
                    })
            safe_append_event({
                "batch": bi, "offset": offset, "phase": "deterministic_qa_done",
            })

            # 4) 独立审校
            review_succeeded = False
            rfindings = []
            review_failed = False
            review_trace = None
            review_event_id = None
            formal_records = []
            promoted_review_events = []
            formal_segment_ids = list(range(offset, offset + len(batch_pairs)))
            if enable_review:
                with commit_lock:
                    current_pairs = list(pairs)
                    current_findings_all = list(findings_all)
                evidence_index = _translation_evidence_index(
                    paras, current_pairs, batch_pairs, glossary, document_profile,
                    document_synopsis, section_digests, current_findings_all)
                formal_packet = _translation_evidence.build_runtime_review_packet(
                    state, batch_pairs, formal_segment_ids, glossary,
                    deterministic_checks=_globalize_batch_findings(findings, offset),
                    review_context=review_context,
                )
                rfindings, failed, review_trace = \
                    _translation_evidence.review_translation_batch_with_evidence(
                        batch_sources, batch_targets, glossary_text, style_rules, target_lang,
                        reviewer_config["provider"], reviewer_config["api_key"],
                        reviewer_config["model"], evidence_index, call_llm=reviewer_call,
                        segment_ids=formal_segment_ids,
                        translation_core_packet=formal_packet)
                review_event_id = (
                    f"translation-review-{job_id}-{bi}-formal-"
                    f"{len(state.get('review_evidence') or []) + len(review_evidence_to_record)}")
                review_trace["review_event_id"] = review_event_id
                if failed:
                    review_failed = True
                    for pair in batch_pairs:
                        if not pair.get("from_tm"):
                            pair["review_status"] = "review_failed"
                else:
                    review_succeeded = True
                    global_to_local = {
                        offset + index: index for index in range(len(batch_pairs))
                    }
                    for rf in rfindings:
                        sev = rf.get("severity")
                        if sev not in ("blocking", "actionable", "informational"):
                            continue
                        segment_id = rf.get("segment_id")
                        idx = global_to_local.get(segment_id)
                        if idx is None:
                            continue
                        record = _review_finding_record(rf, review_event_id)
                        if sev == "actionable" and rf.get("suggested_target") \
                                and idx not in blocked_local_indexes \
                                and not batch_pairs[idx]["from_tm"]:
                            suggested = clean_xml_chars(rf["suggested_target"]).replace('\n', ' ').strip()
                            if suggested:
                                old_target = batch_pairs[idx]["target"]
                                overlay = _repair.create_overlay(
                                    [old_target], [suggested], [rf], "review_suggested",
                                    sources=[batch_sources[idx]],
                                    finding_segment_ids=[segment_id])
                                recheck = _globalize_batch_findings(check_translation_batch(
                                    [batch_sources[idx]], [suggested], glossary, target_lang,
                                    section_profile=section_profile), segment_id)
                                blind_trace = None
                                blind_findings = []
                                blind_failed = False
                                if not any(f["severity"] in ("blocking", "actionable")
                                           for f in recheck):
                                    with commit_lock:
                                        cur_pairs = list(pairs)
                                        cur_f_all = list(findings_all)
                                    blind_index = _translation_evidence_index(
                                        paras, cur_pairs, batch_pairs, glossary, document_profile,
                                        document_synopsis, section_digests, cur_f_all,
                                        blind=True, candidate_targets={segment_id: suggested})
                                    suggested_packet = \
                                        _translation_evidence.build_runtime_review_packet(
                                            state, [batch_pairs[idx]], [segment_id], glossary,
                                            deterministic_checks=recheck,
                                            review_context=review_context,
                                            candidate_targets={segment_id: suggested},
                                            blind=True,
                                        )
                                    blind_findings, blind_failed, blind_trace = \
                                        _translation_evidence.review_translation_batch_with_evidence(
                                            [batch_sources[idx]], [suggested], glossary_text,
                                            style_rules, target_lang,
                                            reviewer_config["provider"],
                                            reviewer_config["api_key"],
                                            reviewer_config["model"], blind_index,
                                            call_llm=reviewer_call, blind=True,
                                            segment_ids=[segment_id], review_identity={
                                                "input_hash": overlay["input_hash"],
                                                "candidate_hash": overlay["candidate_hash"],
                                            }, translation_core_packet=suggested_packet)
                                overlay = _repair.evaluate_overlay(
                                    overlay, recheck, blind_findings, blind_failed,
                                    review_identity=(blind_trace or {}).get("review_identity"))
                                repair_overlays_to_record.append({
                                    "batch": bi, "offset": offset,
                                    "batch_local_ordinal": idx,
                                    "segment_id": segment_id, **overlay,
                                    "blind_trace": blind_trace,
                                })
                                if overlay["status"] == "accepted":
                                    batch_pairs[idx]["target"] = suggested
                                    if blind_trace:
                                        promoted_review_events.append((
                                            blind_trace,
                                            f"translation-review-{job_id}-{bi}-suggested-"
                                            f"{segment_id}-{len(promoted_review_events)}",
                                            segment_id,
                                        ))
                                else:
                                    record["suggested_target"] = suggested
                                    formal_records.append(record)
                                    if blind_trace:
                                        review_evidence_to_record.append({
                                            "batch": bi, "phase": "suggested_shadow_review",
                                            **blind_trace,
                                        })
                                continue
                        formal_records.append(record)

            return {
                "bi": bi,
                "offset": offset,
                "batch": batch,
                "batch_pairs": batch_pairs,
                "batch_sources": batch_sources,
                "injected_ids": injected_ids,
                "entity_hints": entity_hints,
                "context_packet": context_packet,
                "section_digest": section_digest,
                "section_profile": section_profile,
                "blocked_local_indexes": blocked_local_indexes,
                "review_succeeded": review_succeeded,
                "rfindings": rfindings,
                "review_failed": review_failed,
                "review_trace": review_trace,
                "review_event_id": review_event_id,
                "formal_records": formal_records,
                "promoted_review_events": promoted_review_events,
                "formal_segment_ids": formal_segment_ids,
                "repair_overlays": repair_overlays_to_record,
                "review_evidence": review_evidence_to_record,
            }
        finally:
            with commit_lock:
                active_batches.discard(bi)

    def _commit_batch(bundle):
        bi = bundle["bi"]
        offset = bundle["offset"]
        batch = bundle["batch"]
        batch_pairs = bundle["batch_pairs"]
        batch_sources = bundle["batch_sources"]
        injected_ids = bundle["injected_ids"]
        context_packet = bundle["context_packet"]
        section_digest = bundle["section_digest"]
        section_profile = bundle["section_profile"]
        blocked_local_indexes = bundle["blocked_local_indexes"]
        review_succeeded = bundle["review_succeeded"]
        rfindings = bundle["rfindings"]
        review_failed = bundle["review_failed"]
        review_trace = bundle["review_trace"]
        review_event_id = bundle["review_event_id"]
        formal_records = bundle["formal_records"]
        promoted_review_events = bundle["promoted_review_events"]
        formal_segment_ids = bundle["formal_segment_ids"]

        state.setdefault("context_packet_log", []).append({
            "batch": bi,
            "offset": offset,
            **_context.context_metadata(context_packet),
        })
        state.setdefault("glossary_injection_log", []).append({
            "batch": bi,
            "offset": offset,
            "entry_ids": injected_ids,
            "glossary_version": (state.get("glossary_frozen") or {}).get("version"),
            "glossary_hash": (state.get("glossary_frozen") or {}).get("glossary_hash"),
            "entity_hint_count": len(bundle.get("entity_hints") or []),
            "document_synopsis_summary": (document_synopsis or {}).get("summary", ""),
            "section_digest_summary": (section_digest or {}).get("summary", ""),
        })
        if bundle.get("repair_overlays"):
            state.setdefault("repair_overlays", []).extend(bundle["repair_overlays"])
        if bundle.get("review_evidence"):
            state.setdefault("review_evidence", []).extend(bundle["review_evidence"])

        if enable_review:
            stats["batches_reviewed"] += 1
            if review_failed:
                stats["review_failed"] += 1
            else:
                for rec in formal_records:
                    findings_all.append(rec)
            if review_trace and review_event_id:
                _translation_evidence.register_runtime_review_event(
                    state, {"batch": bi, **review_trace}, formal_records,
                    review_event_id, formal_segment_ids)
            for promoted_trace, promoted_event_id, promoted_segment_id in \
                    promoted_review_events:
                _translation_evidence.register_runtime_review_event(
                    state, {"batch": bi, **promoted_trace}, [],
                    promoted_event_id, [promoted_segment_id],
                    phase="suggested_shadow_review")
            safe_append_event({
                "batch": bi, "offset": offset, "phase": "semantic_review_done",
            })
        else:
            safe_append_event({
                "batch": bi, "offset": offset, "phase": "semantic_review_skipped",
            })

        # 审校可能修改过译文：对最终译文整体复验一次确定性检查
        findings = [
            finding for finding in check_translation_batch(
                batch_sources, [p["target"] for p in batch_pairs], glossary,
                target_lang, section_profile=section_profile)
            if finding.get("segment_index") not in blocked_local_indexes
        ]

        # 批内冲突检测（跨段同术语多译法）——在 TM 入库前执行
        for cf in _detect_conflicts(batch_pairs, glossary, sections=sections):
            segment_id = offset + cf["segment_id"]
            cf["segment_index"] = segment_id
            cf["segment_id"] = segment_id
            findings_all.append(cf)

        # 记录仍未解决的确定性问题
        for f in findings:
            if f["severity"] in ("blocking", "actionable", "informational"):
                segment_id = offset + f["segment_index"]
                record = dict(f)
                record.update({"segment_id": segment_id,
                               "segment_index": segment_id,
                               "severity": f["severity"], "type": "check",
                               "detector": f.get("detector") or "Deterministic QA"})
                findings_all.append(record)

        # 5) 批后知识反馈：只进入 candidate queue，不改变 frozen glossary.
        accepted_for_knowledge = []
        review_bad_segments = {
            finding.get("segment_id")
            for finding in (rfindings if enable_review and review_succeeded else [])
            if finding.get("severity") in ("blocking", "actionable")
        }
        for j, pair in enumerate(batch_pairs):
            local_findings = [
                finding for finding in findings
                if finding.get("segment_index") == j
                and finding.get("severity") in ("blocking", "actionable")
            ]
            if (not pair.get("from_tm") and not local_findings
                    and offset + j not in review_bad_segments):
                accepted_for_knowledge.append(j)
        knowledge_due = _knowledge.feedback_due(
            bi, len(batches),
            interval=feedback_interval,
            existing_candidates=state.get("knowledge_candidates") or [],
            force_every_batch=bool(enable_review or state.get("quality_mode")),
        )
        if knowledge_due and accepted_for_knowledge:
            state["knowledge_feedback_policy"]["observed_batches"].append(bi)
            knowledge_segment_ids = [offset + j for j in accepted_for_knowledge]
            knowledge_candidates, knowledge_events, knowledge_warning = \
                _knowledge.observe_batch(
                    [batch_sources[j] for j in accepted_for_knowledge],
                    [batch_pairs[j]["target"] for j in accepted_for_knowledge],
                    paras, pairs,
                    glossary, offset, auxiliary_config["provider"],
                    auxiliary_config["api_key"], auxiliary_config["model"],
                    existing_candidates=state.get("knowledge_candidates") or [],
                    call_llm=auxiliary_call, segment_ids=knowledge_segment_ids,
                    observation_provenance=(
                        "reviewed" if review_succeeded else "generated_continuity"
                    ))
        else:
            if not knowledge_due:
                state["knowledge_feedback_policy"]["skipped_batches"].append(bi)
            knowledge_candidates, knowledge_events, knowledge_warning = \
                state.get("knowledge_candidates") or [], [], None
        for event in knowledge_events:
            event["batch"] = bi
            state.setdefault("knowledge_events", []).append(event)
            if event.get("type") == "entity_observation":
                registry.observe(
                    event.get("source"), event.get("observed_target"),
                    entity_type=_entity_registry.entity_type_from_kind(
                        event.get("kind")),
                    segment_id=event.get("segment_id"),
                    provenance=event.get("provenance") or "generated_observation",
                    confidence=0.7 if event.get("provenance") == "reviewed" else 0.35,
                )
            if event.get("type") == "target_conflict":
                segment_id = event.get("segment_id", offset)
                findings_all.append({
                    "segment_id": segment_id, "segment_index": segment_id,
                    "severity": "actionable",
                    "type": "knowledge_conflict",
                    "category": "terminology_consistency",
                    "summary": "观察到的译法与锁定术语不一致",
                    "source_span": event.get("source"),
                    "target_span": event.get("observed_target"),
                    "explanation": "翻译流中观察到的译法与项目锁定术语的首选译名不同，可能造成术语漂移。",
                    "recommendation": "核对当前语境是否构成合理例外；若不是，请统一为锁定术语的首选译名。",
                    "confidence": None, "detector": "Knowledge QA",
                    "diagnostic_version": 1,
                    "reason": event.get("reason") or "翻译流观察译法与锁定术语不一致",
                    "source": event.get("source"),
                    "observed_target": event.get("observed_target"),
                    "preferred_target": event.get("preferred_target"),
                })
        state["knowledge_candidates"] = knowledge_candidates
        state["translation_continuity"] = list(knowledge_candidates)
        state["entity_registry"] = registry.to_list()
        if knowledge_warning:
            state["knowledge_feedback_failures"] = \
                state.get("knowledge_feedback_failures", 0) + 1
            state.setdefault("knowledge_events", []).append({
                "type": "extract_failed", "batch": bi, "offset": offset,
                "reason": knowledge_warning,
            })
        safe_append_event({
            "batch": bi, "offset": offset, "phase": "knowledge_feedback_done",
            "candidate_count": len(knowledge_candidates),
        })

        # 6) 审校通过的段落 -> 翻译记忆
        if enable_review:
            for j, p in enumerate(batch_pairs):
                segment_id = offset + j
                seg_findings = [
                    f for f in findings_all
                    if f.get("segment_id", f.get("segment_index")) == segment_id
                ]
                if not p.get("from_tm"):
                    if not review_succeeded:
                        p["review_status"] = "review_failed"
                        p["reviewed"] = False
                    elif seg_findings:
                        p["review_status"] = "reviewed_with_findings"
                        p["reviewed"] = False
                    else:
                        p["review_status"] = "reviewed_clean"
                if review_succeeded and not seg_findings and not p["from_tm"] \
                        and _tm_eligible(p["source"], p["target"]):
                    p["reviewed"] = True
                    p["accepted_target"] = p["target"]
                    p["target_provenance"] = "tm_approved" if p.get("from_tm") else "reviewed"
                    if use_tm:
                        _key = tm_put(tm, p["source"], p["target"], _tm_lang)
                        if _key:
                            _norm = tm_normalize(p["source"])
                            if _norm and _norm not in tm_norm_index:
                                tm_norm_index[_norm] = _key
                    stats["reviewed_segments"] += 1

        _commit_translation_batch(batch_pairs, offset)
        safe_append_event({
            "batch": bi, "offset": offset, "phase": "state_commit_done",
            "pairs_count": len(pairs),
        })
        tm_entries = _checkpoint.batch_entries([
            pair for pair in batch_pairs
            if not pair.get("from_tm")
            and pair.get("review_status") == "reviewed_clean"
        ]) if enable_review and use_tm and review_succeeded else []
        if tm_entries:
            safe_append_event({
                "batch": bi, "offset": offset, "phase": "tm_promotion_pending",
                "entries": tm_entries,
            })
            save_tm(tm, _tm_project)
            safe_append_event({
                "batch": bi, "offset": offset, "phase": "tm_promotion_done",
                "entries": tm_entries,
            })

        if on_status:
            concurrency_tag = f"，并发 {effective_concurrency}" if effective_concurrency > 1 else ""
            on_status(f"【阶段二】双语翻译与术语严格注入...（批次 {bi + 1}/{len(batches)}{concurrency_tag}）")
        if on_caption:
            on_caption(f"🌍 正在翻译第 {offset + 1}-{offset + len(batch)} 段（共 {len(paras)} 段）...")

    if effective_concurrency <= 1 or (len(batches) - start_batch) <= 1:
        for bi in range(start_batch, len(batches)):
            if _runtime_cancel_requested(job_id):
                raise RuntimeError("任务已请求取消")
            bundle = _process_batch_payload(bi)
            _commit_batch(bundle)
    else:
        window = min(len(batches) - start_batch, effective_concurrency * 2)
        futures = {}
        with ThreadPoolExecutor(max_workers=effective_concurrency, thread_name_prefix="ft-trans") as executor:
            for bi in range(start_batch, start_batch + window):
                futures[bi] = executor.submit(_process_batch_payload, bi)

            for bi in range(start_batch, len(batches)):
                if _runtime_cancel_requested(job_id):
                    for f in futures.values():
                        f.cancel()
                    raise RuntimeError("任务已请求取消")
                next_bi = bi + window
                if next_bi < len(batches):
                    futures[next_bi] = executor.submit(_process_batch_payload, next_bi)

                bundle = futures[bi].result()
                _commit_batch(bundle)
                del futures[bi]

    # 全局冲突检测（跨批次），与批内结果去重
    batch_conflict_keys = {(f.get("type"), f.get("entry_id"),
                            f.get("segment_index"), True)
                           for f in findings_all if f.get("conflict")}
    for cf in _detect_conflicts(pairs, glossary, sections=sections):
        cf["segment_index"] = cf["segment_id"]
        key = (cf.get("type"), cf.get("entry_id"), cf.get("segment_index"), True)
        if key not in batch_conflict_keys:
            findings_all.append(cf)
    state["entity_registry"] = registry.to_list()
    existing_entity_conflicts = {
        (item.get("type"), item.get("segment_index"), item.get("source"))
        for item in findings_all
    }
    for conflict in registry.consistency_findings():
        key = (conflict.get("type"), conflict.get("segment_index"),
               conflict.get("source"))
        if key not in existing_entity_conflicts:
            findings_all.append(conflict)
            existing_entity_conflicts.add(key)

    stats["blocking"] = sum(1 for f in findings_all if f["severity"] == "blocking")
    stats["actionable"] = sum(1 for f in findings_all if f["severity"] == "actionable")
    stats["informational"] = sum(1 for f in findings_all if f["severity"] == "informational")
    state["has_blocking"] = stats["blocking"] > 0
    return state


def findings_report_md(state):
    """把审查结果渲染成 Markdown 报告（下载/展示用）。"""
    stats = state.get("review_stats") or {}
    lines = [
        "# 翻译审查报告", "",
        "## 概览",
        f"- 已审校段落：{stats.get('reviewed_segments', 0)}",
        f"- 审校批次：{stats.get('batches_reviewed', 0)}",
        f"- 审校失败批次：{stats.get('review_failed', 0)}",
        f"- 翻译记忆复用：{state.get('tm_used_count', 0)} 段",
        f"- blocking：{stats.get('blocking', 0)}",
        f"- actionable：{stats.get('actionable', 0)}",
        f"- informational：{stats.get('informational', 0)}",
        "", "## 待处理问题",
    ]
    findings = state.get("findings") or []
    if not findings:
        lines.append("无。")
    else:
        for f in findings:
            line = f"- 第 {f.get('segment_index', -1) + 1} 段 [{f.get('severity')}] {f.get('reason')}"
            if f.get("suggested_target"):
                line += f"（建议译文：{f['suggested_target']}）"
            lines.append(line)
    return "\n".join(lines)


# ================= 文档/表格生成 =================
EN_FONT = "Times New Roman"
CN_FONT = "宋体"

# 自动标注三色：生僻词=红、专业名词（特殊译法）=黄、翻译难点句=青绿
ANNOTATION_COLORS = {"rare": "C00000", "domain": "BF8F00", "hard": "008080"}
ANNOTATION_LABELS = {"rare": "生僻词/难词", "domain": "专业名词（特殊译法）",
                     "hard": "翻译难点句（特别译法）"}


def _apply_doc_fonts(doc):
    """默认字体与可编辑学术文档的基础版式。"""
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)
    style_sizes = {
        "Normal": 12, "Heading 1": 15, "Heading 2": 13.5,
        "Heading 3": 12.5, "Heading 4": 12, "Title": 18,
    }
    for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3",
                       "Heading 4", "Title", "Intense Quote", "List Bullet",
                       "List Number"):
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        style.font.name = EN_FONT
        if style_name in style_sizes:
            style.font.size = Pt(style_sizes[style_name])
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:ascii"), EN_FONT)
        rfonts.set(qn("w:hAnsi"), EN_FONT)
        rfonts.set(qn("w:eastAsia"), CN_FONT)
        if style_name == "Normal":
            style.paragraph_format.line_spacing = 1.5
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.first_line_indent = Cm(0.74)
        elif style_name not in {"Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"}:
            style.paragraph_format.line_spacing = 1.5
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.first_line_indent = None


def _apply_run_fonts(run):
    """单个 run 的字体（表格单元格里的 run 不受 Normal 样式继承影响时兜底）。"""
    from docx.oxml.ns import qn
    run.font.name = EN_FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), EN_FONT)
    rfonts.set(qn("w:hAnsi"), EN_FONT)
    rfonts.set(qn("w:eastAsia"), CN_FONT)


def dict_to_excel(term_dict):
    df = pd.DataFrame(list(term_dict.items()), columns=["Source", "Target"])
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False)
    output.seek(0)
    return output


def paragraphs_to_word(paragraphs):
    doc = Document()
    _apply_doc_fonts(doc)
    doc.add_heading('阶段一：清洗后原文提取', 0)
    for p in paragraphs:
        doc.add_paragraph(p)
    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out


def translations_to_word(pairs):
    doc = Document()
    _apply_doc_fonts(doc)
    doc.add_heading("译文", 0)
    for pair in pairs or []:
        doc.add_paragraph(str(pair.get("target") or ""))
    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out


def translations_to_pdf(pairs):
    """Render a paginated, Unicode-aware translation PDF with PyMuPDF Story."""
    paragraphs = "".join(
        '<table class="entry"><tr><td dir="auto">'
        f'{html_escape(str(pair.get("target") or ""))}</td></tr></table>'
        for pair in pairs or [])
    html = f'<article><h1>译文</h1>{paragraphs}</article>'
    css = (
        "@page { size: A4; } "
        "body { font-family: sans-serif; font-size: 11pt; line-height: 1.65; "
        "color: #1f2937; } h1 { font-size: 20pt; margin: 0 0 20pt; "
        "color: #111827; } table.entry { width: 100%; border-collapse: collapse; "
        "margin: 0 0 10pt; break-inside: avoid; page-break-inside: avoid; } "
        "td { padding: 0; text-align: start; unicode-bidi: plaintext; }"
    )
    story = fitz.Story(html=html, user_css=css, em=11)
    page_rect = fitz.paper_rect("a4")
    content_rect = fitz.Rect(54, 54, page_rect.width - 54, page_rect.height - 54)

    def rectfn(_rect_num, _filled):
        return page_rect, content_rect, None

    pdf = story.write_with_links(rectfn)
    for number, page in enumerate(pdf, start=1):
        page.insert_text(
            (page_rect.width / 2 - 4, page_rect.height - 25), str(number),
            fontsize=9, fontname="helv", color=(0.45, 0.48, 0.52))
    data = pdf.tobytes(garbage=4, deflate=True)
    pdf.close()
    return data


def glossary_to_excel(entries, fallback=None):
    normalized = normalize_glossary(entries or [])
    if normalized:
        rows = [{
            "Source": entry.get("source") or "",
            "Target": entry.get("preferred") or entry.get("target") or "",
            "Status": entry.get("status") or "",
            "Domain": entry.get("domain") or "",
            "Note": entry.get("note") or "",
        } for entry in normalized]
        frame = pd.DataFrame(rows)
    else:
        frame = pd.DataFrame(list((fallback or {}).items()),
                             columns=["Source", "Target"])
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False)
    output.seek(0)
    return output


def _find_span(text, needle):
    """宽容定位子串：统一引号/破折号/省略号、折叠空白后查找，返回 (start, end) 或 None。"""
    if not needle or not text:
        return None
    if needle in text:
        pos = text.find(needle)
        return pos, pos + len(needle)
    mapping = []
    norm_chars = []
    for orig_idx, ch in enumerate(text):
        ch2 = ch.replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
        ch2 = ch2.replace("–", "-").replace("—", "-").replace("…", "...")
        if ch2.isspace():
            if norm_chars and norm_chars[-1] != " ":
                norm_chars.append(" ")
                mapping.append(None)
            continue
        norm_chars.append(ch2)
        mapping.append(orig_idx)
    norm_text = "".join(norm_chars)
    needle2 = needle.replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
    needle2 = needle2.replace("–", "-").replace("—", "-").replace("…", "...")
    needle2 = re.sub(r"\s+", " ", needle2).strip()
    pos = norm_text.find(needle2)
    if pos < 0:
        return None
    start = None
    for i in range(pos, len(mapping)):
        if mapping[i] is not None:
            start = mapping[i]
            break
    end = None
    for i in range(min(pos + len(needle2), len(mapping)) - 1, -1, -1):
        if mapping[i] is not None:
            end = mapping[i] + 1
            break
    if start is None or end is None:
        return None
    return start, end


def _colored_cell(cell, text, spans, colors=None):
    """把一个单元格按 spans（(start,end,type) 已排序不重叠）拆成带色 run。"""
    from docx.shared import RGBColor
    palette = colors or ANNOTATION_COLORS
    cursor = 0
    first = True
    for start, end, atype in spans:
        if start > cursor:
            run = cell.paragraphs[0].add_run(text[cursor:start])
            _apply_run_fonts(run)
        run = cell.paragraphs[0].add_run(text[start:end])
        _apply_run_fonts(run)
        run.font.color.rgb = RGBColor.from_string(
            str(palette.get(atype, ANNOTATION_COLORS[atype])).lstrip("#"))
        if atype in ("rare", "domain"):
            run.bold = True
        cursor = end
        first = False
    if cursor < len(text):
        run = cell.paragraphs[0].add_run(text[cursor:])
        _apply_run_fonts(run)
    if first and not text:
        cell.paragraphs[0].add_run("")


def pairs_to_word(pairs, annotations=None, colors=None):
    """双语对照表 -> Word 表格。

    annotations: {seg: [{"type": "rare|domain|hard", "src_span": [s,e]|None,
                         "tgt_span": [s,e]|None, "note": str}]}
    colors: {"rare"|"domain"|"hard": "RRGGBB"}，可自定义三类标注颜色。
    """
    palette = dict(ANNOTATION_COLORS)
    if colors:
        palette.update({k: str(v).lstrip("#") for k, v in colors.items()})
    doc = Document()
    _apply_doc_fonts(doc)
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    table.rows[0].cells[0].text = "原文"
    table.rows[0].cells[1].text = "译文"
    annot = _normalize_annotations(annotations)
    for i, pair in enumerate(pairs):
        row = table.add_row().cells
        seg_annot = annot.get(i) or []
        src_spans = _compose_spans(
            [(it["src_span"][0], it["src_span"][1], it["type"])
             for it in seg_annot if it.get("src_span")], len(pair['source']))
        tgt_spans = _compose_spans(
            [(it["tgt_span"][0], it["tgt_span"][1], it["type"])
             for it in seg_annot if it.get("tgt_span")], len(pair['target']))
        _colored_cell(row[0], pair['source'], src_spans, palette)
        _colored_cell(row[1], pair['target'], tgt_spans, palette)
    # 图例（放表格后，避免挤占首行）
    p_legend = doc.add_paragraph()
    legend_parts = [
        f"{ANNOTATION_LABELS[k]}（#{palette.get(k, ANNOTATION_COLORS[k])}）"
        for k in ("rare", "domain", "hard")]
    run = p_legend.add_run("图例：" + "；".join(legend_parts) + "。")
    _apply_run_fonts(run)
    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out


def _add_formatted_runs(paragraph, text):
    parts = text.split('**')
    for i, part in enumerate(parts):
        run = paragraph.add_run(part)
        _apply_run_fonts(run)
        if i % 2 != 0:
            run.bold = True


def _markdown_table_row(line):
    value = str(line or "").strip()
    if value.startswith("|"):
        value = value[1:]
    if value.endswith("|"):
        value = value[:-1]
    return [cell.strip() for cell in value.split("|")]


def _is_markdown_table_separator(line):
    cells = _markdown_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def _add_markdown_table(doc, rows):
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=1, cols=columns)
    table.style = "Table Grid"
    for row_index, values in enumerate(rows):
        cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
        for col_index in range(columns):
            cell = cells[col_index]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.line_spacing = 1.35
            paragraph.paragraph_format.space_after = 0
            paragraph.paragraph_format.first_line_indent = None
            run = paragraph.add_run(values[col_index] if col_index < len(values) else "")
            _apply_run_fonts(run)
            if row_index == 0:
                run.bold = True
    return table


def markdown_to_word(md_text, theory):
    doc = Document()
    _apply_doc_fonts(doc)
    md_text = re.sub(r'```markdown|```', '', md_text).strip()
    md_text = re.sub(r'<!--.*?-->', '', md_text, flags=re.DOTALL)
    quote_labels = {
        "SYNTHETIC_SOURCE": "真实源文",
        "SIMULATED": "模拟初译",
        "OPTIMIZED": "优化译文",
    }
    quote_labels.update({"SOURCE": "原文", "INITIAL": "初译", "TARGET": "终译"})
    md_text = re.sub(
        r'(?m)^>\s*\[(SYNTHETIC_SOURCE|SIMULATED|OPTIMIZED|SOURCE|INITIAL|TARGET)\s+'
        r'(?:SC-\d{4,}|seg-[A-Za-z0-9_-]+)\]:\s*',
        lambda match: f"> {quote_labels[match.group(1)]}：", md_text)
    lines = md_text.splitlines()
    while lines and not lines[0].strip():
        lines.pop(0)
    if lines and re.match(r"^#\s+翻译实践报告", lines[0], re.IGNORECASE):
        lines.pop(0)
    deduped = []
    for line in lines:
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*$", line.strip())
        if heading:
            title = re.sub(r"\s+#+\s*$", "", heading.group(2)).strip()
            previous = next((value for value in reversed(deduped) if value.strip()), "")
            previous_heading = re.match(r"^#{1,6}\s+(.+?)\s*$", previous.strip())
            if previous_heading:
                previous_title = re.sub(r"\s+#+\s*$", "", previous_heading.group(1)).strip()
                if re.sub(r"^\d+(?:\.\d+)*[.、．]?\s*", "", title).casefold() == \
                        re.sub(r"^\d+(?:\.\d+)*[.、．]?\s*", "", previous_title).casefold():
                    continue
        deduped.append(line)
    title = doc.add_heading(f'翻译实践报告：基于{theory}', 0)
    title.alignment = 1
    index = 0
    while index < len(deduped):
        line = deduped[index].strip()
        if ("|" in line and index + 1 < len(deduped)
                and _is_markdown_table_separator(deduped[index + 1])):
            rows = [_markdown_table_row(line)]
            index += 2
            while index < len(deduped) and "|" in deduped[index] and deduped[index].strip():
                rows.append(_markdown_table_row(deduped[index]))
                index += 1
            _add_markdown_table(doc, rows)
            continue
        line = line.strip()
        if not line:
            index += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
        if heading:
            heading_level = max(1, min(4, len(heading.group(1)) - 1))
            doc.add_heading(heading.group(2).strip(), level=heading_level)
        elif line.startswith(('- ', '* ')):
            p = doc.add_paragraph(style='List Bullet')
            _add_formatted_runs(p, line[2:])
        elif line.startswith('> '):
            p = doc.add_paragraph(style='Intense Quote')
            _add_formatted_runs(p, line[2:])
        else:
            p = doc.add_paragraph()
            _add_formatted_runs(p, line)
        index += 1
    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)
    return doc_io


# ================= 任务持久化（真正的断点续传）=================
def _ensure_output_dir():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def is_onboarded():
    """首次使用引导标记：成功配置并测试通过 AI 引擎后置位。"""
    return (OUTPUT_DIR / ".onboarded").exists()


def mark_onboarded():
    _ensure_output_dir()
    (OUTPUT_DIR / ".onboarded").touch()


def provider_config_path():
    return OUTPUT_DIR / "provider_config.json"


def save_provider_config(provider, model, api_key, base_url=None, reviewer=None,
                         reasoning_effort=None):
    """把 AI 引擎配置落盘（本地单机工具，0600 权限），重启后自动加载。"""
    _ensure_output_dir()
    normalized_base = normalize_openai_base_url(base_url) if base_url else ""
    normalized_reasoning = (str(reasoning_effort or "").strip().lower()
                            if str(reasoning_effort or "").strip().lower()
                            in _REASONING_EFFORT_VALUES else "")
    cfg = {
        "provider": provider or "",
        "model": model or "",
        "api_key": api_key or "",
        "base_url": normalized_base,
        "reasoning_effort": normalized_reasoning,
    }
    if isinstance(reviewer, dict) and reviewer.get("provider") and reviewer.get("model"):
        cfg["reviewer"] = {
            "provider": reviewer.get("provider", ""),
            "model": reviewer.get("model", ""),
            "api_key": reviewer.get("api_key", ""),
            "base_url": normalize_openai_base_url(reviewer.get("base_url"))
            if reviewer.get("base_url") else "",
        }
    path = provider_config_path()
    temp_path = path.with_suffix(path.suffix + ".tmp")
    temp_path.write_text(json.dumps(cfg, ensure_ascii=False, indent=2),
                         encoding="utf-8")
    try:
        temp_path.chmod(0o600)
    except OSError:
        pass
    temp_path.replace(path)
    return path


def load_provider_config():
    """读取已保存的 AI 引擎配置；文件缺失或损坏时返回 None。"""
    try:
        path = provider_config_path()
        if not path.exists():
            return None
        data = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(data, dict) or not data.get("provider"):
            return None
        result = {
            "provider": str(data.get("provider", "")),
            "model": str(data.get("model", "") or ""),
            "api_key": str(data.get("api_key", "") or ""),
            "base_url": normalize_openai_base_url(data.get("base_url"))
            if data.get("base_url") else "",
            "reasoning_effort": str(data.get("reasoning_effort", "") or "")
            .strip().lower(),
        }
        if result["reasoning_effort"] not in _REASONING_EFFORT_VALUES:
            result["reasoning_effort"] = ""
        reviewer = data.get("reviewer")
        if isinstance(reviewer, dict) and reviewer.get("provider"):
            result["reviewer"] = {
                "provider": str(reviewer.get("provider", "")),
                "model": str(reviewer.get("model", "") or ""),
                "api_key": str(reviewer.get("api_key", "") or ""),
                "base_url": normalize_openai_base_url(reviewer.get("base_url"))
                if reviewer.get("base_url") else "",
            }
        return result
    except Exception:  # noqa: BLE001 - 配置文件损坏时按未保存处理
        return None


def job_dir(job_id):
    return OUTPUT_DIR / job_id


def job_state_path(job_id):
    return job_dir(job_id) / "state.json"


RUNTIME_SCHEMA_VERSION = "2"
RUNTIME_STATE_FILE = "runtime_state.json"
RUNTIME_EVENTS_FILE = "runtime_events.jsonl"
RUNTIME_TECHNICAL_LOG = "runtime_technical.log"
RUNTIME_HEARTBEAT_SECONDS = 5
RUNTIME_LEASE_SECONDS = 60
RUNTIME_STALL_SECONDS = 45
RUNTIME_ACTIVE_STATUSES = {
    "resume_requested", "queued", "starting", "running",
    "waiting_external", "cancelling",
}

_RUNTIME_EVENT_META = {
    "resume_requested": ("user", "lifecycle"),
    "pipeline_resumed": ("user", "progress"),
    "llm_request_started": ("user", "external_wait"),
    "llm_response_received": ("user", "progress"),
    "job_completed": ("user", "lifecycle"),
    "job_failed": ("user", "error"),
    "job_cancelled": ("user", "lifecycle"),
    "cancel_requested": ("user", "lifecycle"),
    "interrupted": ("user", "error"),
    "stalled": ("user", "error"),
    "retry_requested": ("user", "lifecycle"),
    "job_queued": ("technical", "orchestration"),
    "worker_started": ("technical", "orchestration"),
    "worker_released": ("technical", "orchestration"),
    "job_checkpointed": ("technical", "checkpoint"),
    "checkpoint_saved": ("technical", "checkpoint"),
    "state_flushed": ("technical", "checkpoint"),
    "lease_acquired": ("technical", "orchestration"),
    "lease_renewed": ("technical", "orchestration"),
    "evidence": ("user", "progress"),
    "literature_evidence": ("user", "progress"),
    "research_model": ("user", "progress"),
    "literature_claims": ("user", "progress"),
    "argument_plan": ("user", "progress"),
    "selected_cases": ("user", "progress"),
    "outline": ("user", "progress"),
    "sections": ("user", "progress"),
    "validation": ("user", "progress"),
    "review": ("user", "progress"),
    "quality_repair": ("user", "progress"),
    "academic_quality": ("user", "progress"),
}


def runtime_state_path(job_id):
    return job_dir(job_id) / RUNTIME_STATE_FILE


def runtime_events_path(job_id):
    return job_dir(job_id) / RUNTIME_EVENTS_FILE


def runtime_technical_log_path(job_id):
    return job_dir(job_id) / RUNTIME_TECHNICAL_LOG


def _utc_now_iso():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _runtime_defaults():
    return {
        "schema_version": RUNTIME_SCHEMA_VERSION,
        "status": "idle",
        "pipeline": "document_pipeline",
        "stage": "",
        "stage_id": "",
        "stage_index": None,
        "stage_total": None,
        "operation": "",
        "operation_id": "",
        "operation_label": "",
        "stage_progress": {},
        "section_id": None,
        "phase": "",
        "phase_label": "",
        "started_at": None,
        "operation_started_at": None,
        "last_heartbeat_at": None,
        "last_progress_at": None,
        "completed_units": 0,
        "total_units": 0,
        "overall_progress": None,
        "last_event": "",
        "events": [],
        "last_technical_event": "",
        "cancel_requested": False,
        "attempt": 0,
        "resume_request_id": None,
        "error": None,
        "worker": {
            "owner_pid": None,
            "worker_id": None,
            "lease_expires_at": None,
        },
    }


def load_runtime_state(job_id):
    """Read the durable worker status without changing the workflow state."""
    if not job_id:
        return _runtime_defaults()
    try:
        raw = json.loads(runtime_state_path(job_id).read_text(encoding="utf-8"))
    except (OSError, TypeError, ValueError):
        return _runtime_defaults()
    if not isinstance(raw, dict):
        return _runtime_defaults()
    result = _runtime_defaults()
    result.update(raw)
    result["schema_version"] = RUNTIME_SCHEMA_VERSION
    if result.get("status") == "cancel_requested":
        result["status"] = "cancelling"
    worker = result.get("worker")
    result["worker"] = {**_runtime_defaults()["worker"], **worker} \
        if isinstance(worker, dict) else dict(_runtime_defaults()["worker"])
    result["events"] = [item for item in result.get("events") or []
                         if isinstance(item, dict)
                         and item.get("visibility") == "user"][-24:]
    return result


def _runtime_event_code(message):
    value = re.sub(r"[^a-z0-9]+", "_", str(message or "").lower()).strip("_")
    return value[:80] or "runtime_event"


def _runtime_event_meta(event_name, visibility=None, category=None):
    default_visibility, default_category = _RUNTIME_EVENT_META.get(
        str(event_name or ""), ("technical", "debug"))
    return visibility or default_visibility, category or default_category


def _append_runtime_event(job_id, event, *, stage=None, operation=None, metadata=None,
                          visibility=None, category=None):
    event_name = event.get("event") if isinstance(event, dict) else _runtime_event_code(event)
    visibility, category = _runtime_event_meta(event_name, visibility, category)
    record = {
        "timestamp": _utc_now_iso(),
        "job_id": job_id,
        "pipeline": event.get("pipeline") if isinstance(event, dict) else None,
        "stage": stage or "",
        "operation": operation or "",
        "event": event_name,
        "message": event.get("message") if isinstance(event, dict) else str(event),
        "visibility": visibility,
        "category": category,
    }
    if record["pipeline"] is None:
        record["pipeline"] = "document_pipeline"
    if metadata:
        record["metadata"] = metadata
    path = runtime_events_path(job_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as stream:
        stream.write(json.dumps(record, ensure_ascii=False, sort_keys=True) + "\n")
        stream.flush()
        os.fsync(stream.fileno())
    return record


def read_runtime_events(job_id, limit=None, *, visibility=None, category=None):
    path = runtime_events_path(job_id)
    if not path.is_file():
        return []
    events = []
    try:
        for line in path.read_text(encoding="utf-8").splitlines():
            try:
                value = json.loads(line)
            except (TypeError, ValueError):
                continue
            if isinstance(value, dict):
                event_visibility, event_category = _runtime_event_meta(
                    value.get("event"), value.get("visibility"), value.get("category"))
                normalized = {**value, "visibility": event_visibility,
                              "category": event_category}
                if visibility and event_visibility != visibility:
                    continue
                if category and event_category != category:
                    continue
                events.append(normalized)
    except OSError:
        return []
    return events[-limit:] if limit else events


def update_runtime_state(job_id, *, event=None, progress=False, heartbeat=False,
                         event_name=None, event_metadata=None,
                         event_visibility=None, event_category=None, **changes):
    """Atomically merge a worker status update into runtime_state.json."""
    if not job_id:
        return _runtime_defaults()
    now = _utc_now_iso()
    with _RUNTIME_LOCK:
        current = load_runtime_state(job_id)
        # ``None`` is normally omitted so optional status fields do not erase
        # an existing value accidentally.  ``error`` is different: callers
        # pass ``error=None`` when a retry or a successful checkpoint clears a
        # stale failure from an earlier attempt.  Keep that explicit clear
        # while retaining the omission behavior for other optional fields.
        cancel_active = current.get("cancel_requested") and changes.get("cancel_requested") is not False
        if cancel_active and changes.get("status") in {"running", "waiting_external"}:
            changes = {**changes, "status": "cancelling", "phase": "cancelling",
                       "phase_label": "正在取消"}
        current.update({key: value for key, value in changes.items()
                        if value is not None or key == "error"})
        if heartbeat:
            current["last_heartbeat_at"] = now
            worker = dict(current.get("worker") or {})
            if worker.get("worker_id"):
                worker["lease_expires_at"] = (
                    datetime.now(timezone.utc) + timedelta(seconds=RUNTIME_LEASE_SECONDS)
                ).isoformat(timespec="seconds")
                current["worker"] = worker
        if progress:
            current["last_progress_at"] = now
        if event:
            message = str(event).strip()
            if message:
                event_name = event_name or _runtime_event_code(message)
                visibility, category = _runtime_event_meta(
                    event_name, event_visibility, event_category)
                inline = {
                    "at": now, "timestamp": now, "event": event_name,
                    "message": message, "visibility": visibility,
                    "category": category,
                }
                if event_metadata:
                    inline["metadata"] = event_metadata
                duplicate = visibility == "user" and bool(current.get("events")) \
                    and current["events"][-1].get("event") == event_name \
                    and current["events"][-1].get("message") == message \
                    and current["events"][-1].get("metadata") == event_metadata
                if visibility == "user":
                    current["last_event"] = message
                    if not duplicate:
                        current["events"] = (current.get("events") or []) + [inline]
                        current["events"] = current["events"][-24:]
                else:
                    current["last_technical_event"] = message
                if not duplicate:
                    _append_runtime_event(
                        job_id,
                        {"event": event_name, "message": message,
                         "pipeline": current.get("pipeline")},
                        stage=current.get("stage_id") or current.get("stage"),
                        operation=current.get("operation_id") or current.get("operation"),
                        metadata=event_metadata, visibility=visibility, category=category)
        directory = job_dir(job_id)
        directory.mkdir(parents=True, exist_ok=True)
        tmp = directory / f"{RUNTIME_STATE_FILE}.tmp"
        tmp.write_text(json.dumps(current, ensure_ascii=False, indent=2),
                       encoding="utf-8")
        tmp.replace(runtime_state_path(job_id))
        return current


def _runtime_pid_alive(pid):
    if not isinstance(pid, int) or pid <= 0:
        return False
    try:
        os.kill(pid, 0)
    except OSError:
        return False
    return True


def _runtime_parse_time(value):
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)
    except (TypeError, ValueError):
        return None


def _runtime_worker_registered(job_id, worker_id=None):
    with _RUNTIME_WORKERS_LOCK:
        worker = _RUNTIME_WORKERS.get(job_id)
        return bool(worker and worker.is_alive() and (
            worker_id is None or getattr(worker, "worker_id", None) == worker_id))


def _annotation_blocks_completion(state):
    """标注是否仍然阻止"业务完成"。

    只有**显式要求**标注（`enable_annotate is True`）且尚未完成时才阻止。
    缺失该字段不能当作"要求标注"：产品自身的默认是关闭
    （`DELIVERY_CONFIG_DEFAULTS["enable_annotate"] = False`），而旧任务里往往
    根本没有这个键。把它当作 True 会让一个已经完成的任务被报成
    `idle_incomplete`，于是概览一边说"可以准备交付"、一边渲染出「继续处理」。
    """
    if not isinstance(state, dict):
        return False
    return state.get("enable_annotate") is True and not state.get("annotations_done")


def _runtime_business_complete(state):
    return bool(state and state.get("p1_done") and state.get("p2_done") and (
        state.get("p3_done") or not state.get("report_enabled", True)) and not
        _annotation_blocks_completion(state))


def _lost_worker_message(runtime, state, *, stalled=False):
    """中断/停滞时给用户一个有信息量的原因。

    以前这里只有一句笼统的"上次运行已中断"，用户无法判断是模型卡住了、还是应用
    被关掉了。现在带上：中断时正在做什么、已经提交了多少段、能否从断点继续。
    """
    label = str(runtime.get("phase_label") or runtime.get("operation_label")
                or runtime.get("phase") or "").strip()
    pairs = len((state or {}).get("pairs") or [])
    paras = len((state or {}).get("paras") or [])
    if stalled:
        parts = ["超过预期时间没有收到运行信号"]
    else:
        parts = ["上次运行的进程已退出（应用被关闭或中断），翻译线程随之终止"]
    if label:
        parts.append(f"中断时正在：{label}")
    if paras:
        parts.append(f"已完成 {pairs}/{paras} 段，可从断点继续")
    return "；".join(parts)


def _runtime_mark_lost(job_id, status, message):
    return update_runtime_state(
        job_id, status=status, phase=status, phase_label=message,
        worker={"owner_pid": None, "worker_id": None, "lease_expires_at": None},
        event=message, event_name=status)


def get_job_runtime_status(job_id, state=None):
    """Return the durable runtime status and reconcile dead local workers."""
    state = state if isinstance(state, dict) else load_job_state(job_id)
    runtime = load_runtime_state(job_id)
    status = runtime.get("status") or "idle"
    if status in RUNTIME_ACTIVE_STATUSES:
        worker = runtime.get("worker") or {}
        pid = worker.get("owner_pid")
        worker_id = worker.get("worker_id")
        lease = _runtime_parse_time(worker.get("lease_expires_at"))
        heartbeat = _runtime_parse_time(runtime.get("last_heartbeat_at"))
        last_progress = _runtime_parse_time(runtime.get("last_progress_at"))
        now = datetime.now(timezone.utc)
        registered = _runtime_worker_registered(job_id, worker_id)
        transition_age = (now - last_progress).total_seconds() if last_progress else None
        if status == "resume_requested" and not pid \
                and transition_age is not None and transition_age <= RUNTIME_STALL_SECONDS:
            return runtime
        if not _runtime_pid_alive(pid):
            return _runtime_mark_lost(job_id, "interrupted",
                                   _lost_worker_message(runtime, state))
        if lease and lease < now or heartbeat and (
                now - heartbeat).total_seconds() > RUNTIME_STALL_SECONDS:
            if registered:
                return _runtime_mark_lost(job_id, "stalled",
                                       _lost_worker_message(
                                           runtime, state, stalled=True))
            return _runtime_mark_lost(job_id, "interrupted",
                                   _lost_worker_message(runtime, state))
        if not registered and pid == os.getpid():
            # Streamlit can re-execute the app module while a worker thread
            # from the same process is still handling a long model request.
            # A module reload recreates the in-memory registry, but the durable
            # lease and heartbeat still prove that this process owns the run.
            # Do not turn that live request into a false "interrupted" state.
            fresh_heartbeat = heartbeat and (
                now - heartbeat).total_seconds() <= RUNTIME_STALL_SECONDS
            lease_valid = not lease or lease >= now
            if fresh_heartbeat and lease_valid:
                return runtime
            if status in {"queued", "starting"} \
                    and transition_age is not None and transition_age <= RUNTIME_STALL_SECONDS:
                return runtime
            return _runtime_mark_lost(job_id, "interrupted",
                                      _lost_worker_message(runtime, state))
        return runtime
    if status == "idle":
        inferred = "completed" if _runtime_business_complete(state) else "idle_incomplete"
        if inferred != status:
            # 只影响本次读取的返回值，**不落盘**：读取不应有写副作用。
            # 此前的实现会在这里 update_runtime_state，于是"看一眼任务"就会给
            # 从未运行过的任务写出 runtime_state.json。真实状态由 worker 在任务
            # 运行时写入；这里的推断每次都能重新得出同样结果，不需要持久化。
            return {**runtime, "status": inferred, "phase": inferred,
                    "phase_label": "已完成" if inferred == "completed" else "尚未完成"}
    return runtime


def build_job_runtime_view(job_id, state=None):
    """Canonical user-facing runtime view shared by every product surface."""
    state = state if isinstance(state, dict) else load_job_state(job_id)
    runtime = get_job_runtime_status(job_id, state)
    status = runtime.get("status") or "idle_incomplete"
    if status == "idle" and not _runtime_business_complete(state):
        status = "idle_incomplete"
    if status == "idle_incomplete" and state and \
            state.get("stage") == "TERMS_PREPARED" and state.get("quality_mode") \
            and state.get("glossary") is not None and not state.get("glossary_frozen") \
            and not state.get("quality_bypass"):
        status = "waiting_manual"
    labels = {
        "resume_requested": "正在恢复任务", "queued": "正在恢复任务",
        "starting": "正在恢复任务", "running": "正在运行",
        "waiting_external": "正在运行",
        "stalled": "暂无运行信号", "interrupted": "上次运行已中断",
        "failed": "当前步骤失败", "cancelling": "正在取消", "cancelled": "任务已取消",
        "completed": "已完成", "idle_incomplete": "未完成",
        "waiting_manual": "待术语确认",
    }
    actions = {
        "resume_requested": ["details"], "queued": ["cancel", "details"],
        "starting": ["cancel", "details"],
        "running": ["cancel", "details"],
        "waiting_external": ["cancel", "details"],
        "stalled": ["retry", "cancel", "details"],
        "interrupted": ["resume", "retry", "details"],
        "failed": ["retry", "resume", "details"],
        "cancelling": ["details"], "cancelled": ["resume", "details"],
        "idle_incomplete": ["resume"], "waiting_manual": ["view"],
        "completed": ["view"],
    }
    events = read_runtime_events(job_id, 5, visibility="user") \
        or runtime.get("events") or []
    completed, total = int(runtime.get("completed_units") or 0), \
        int(runtime.get("total_units") or 0)
    academic_present = bool(state and ((state.get("academic_state") or {}).get("artifacts")
                                       or (state.get("academic_state") or {}).get(
                                           "current_stage") not in {None, "", "not_started"}))
    if state and state.get("p2_done") and (
            state.get("report_enabled", True) or academic_present):
        completed, total = _academic_runtime_progress(job_id, state)
    if status == "completed" and total:
        completed = total
    progress = (completed, total) if total else None
    operation = runtime.get("operation_label") or ""
    is_report = bool(state and state.get("p2_done") and (
        state.get("report_enabled", True) or academic_present))
    surface_label = "实践报告" if is_report else "任务处理"
    if is_report and (not operation or operation in {"准备工作流", "准备任务", "继续处理"}):
        operation = _academic_resume_context(state, runtime)["operation_label"]
    if status in {"resume_requested", "queued", "starting"}:
        detail = "正在读取最近检查点…"
        if not operation or operation in {"准备工作流", "准备任务"}:
            operation = "正在从最近检查点继续学术写作" if is_report \
                else "正在从最近进度继续处理"
    elif status == "waiting_external":
        detail = "正在等待模型响应"
    elif status == "running":
        detail = runtime.get("phase_label") or "已恢复 · 正在执行"
    elif status == "interrupted":
        detail = "当前进度已安全保存，可以继续处理。"
    elif status == "stalled":
        detail = "暂时没有新的运行信号，可以重试当前步骤。"
    elif status == "idle_incomplete":
        detail = "当前进度已保存，可以继续处理。"
    elif status == "failed":
        error = runtime.get("error") or {}
        detail = error.get("message") if isinstance(error, dict) else str(error)
    else:
        detail = runtime.get("phase_label") or ""
    primary_action = {
        "idle_incomplete": "resume", "interrupted": "resume",
        "cancelled": "resume", "failed": "retry", "stalled": "retry",
    }.get(status)
    return {
        "status": status,
        "status_label": labels.get(status, status),
        "runtime_status": status,
        "headline_status": labels.get(status, status),
        "badge": status,
        "surface_label": surface_label,
        "headline": operation or surface_label,
        "detail": detail,
        "current_operation": operation,
        "operation_id": runtime.get("operation_id") or runtime.get("operation") or "",
        "progress": progress,
        "progress_completed": completed,
        "progress_total": total,
        "stage_progress": runtime.get("stage_progress") or {},
        "available_actions": actions.get(status, []),
        "primary_action": primary_action,
        "show_no_worker_warning": status in {
            "idle_incomplete", "interrupted", "cancelled"},
        "last_activity": events[-1].get("message") if events else "",
        "last_activity_at": events[-1].get("timestamp") if events else None,
        "user_events": events,
        "events": events,
        "runtime": runtime,
    }


def _append_runtime_technical_log(job_id, text):
    """向技术日志追加一行。

    进程被直接杀掉时不会有机会写任何东西，所以关键节点必须**主动**留痕：
    worker 启动/释放都记一条，这样"有启动、没有释放"本身就说明了中断方式。
    """
    try:
        path = runtime_technical_log_path(job_id)
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("a", encoding="utf-8") as stream:
            stream.write(f"[{_utc_now_iso()}] {text}\n")
            stream.flush()
    except OSError:
        pass  # 日志失败不得影响任务本身


def _write_runtime_technical_log(job_id, exc):
    path = runtime_technical_log_path(job_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as stream:
        stream.write(f"[{_utc_now_iso()}] {type(exc).__name__}: {exc}\n")
        stream.write(traceback.format_exc())
        stream.write("\n")
        stream.flush()


def _runtime_stage_info(label):
    """Extract the human-facing stage contract from existing status labels."""
    text = str(label or "").strip()

    def nested_progress(value):
        """Parse durable item progress without inventing an ETA."""
        batch_match = re.search(r"批次\s*(\d+)\s*/\s*(\d+)", value)
        page_match = re.search(r"(\d+)\s*/\s*(\d+)\s*页", value)
        sentence_match = re.search(r"(\d+)\s*/\s*(\d+)\s*句", value)
        match = batch_match or page_match or sentence_match
        if not match:
            return {}
        unit = "批次" if batch_match else ("页" if page_match else "句")
        completed, total = int(match.group(1)), int(match.group(2))
        progress = {
            "unit": unit,
            "completed": max(0, completed),
            "total": max(0, total),
            "in_flight": [],
        }
        active_match = re.search(r"处理中\s+((?:#\d+\s*)+)", value)
        if active_match:
            progress["in_flight"] = [
                int(item[1:]) for item in active_match.group(1).split()
                if item.startswith("#") and item[1:].isdigit()
            ]
        return progress

    def operation_without_progress(value):
        operation = re.sub(r"^.*?】\s*", "", value).strip(" .…") or value
        operation = re.sub(
            r"\s*[（(](?=[^）)]*(?:批次|页|句)[^）)]*[）)])[^）)]*[）)]?", "",
            operation,
        )
        # The expression above deliberately handles the normal full-width
        # parenthesis form; keep a small fallback for labels using ASCII
        # parentheses so an individual provider cannot pollute the headline.
        operation = re.sub(r"\s*\([^()]*?(?:批次|页|句)[^()]*\)", "", operation)
        return operation.strip(" .…") or value

    stage_progress = nested_progress(text)
    operation_label = operation_without_progress(text)
    if "扫描 PDF OCR" in text or "OCR" in text and "页" in text:
        return {
            "stage": "document_processing",
            "stage_id": "ocr",
            "operation": "ocr",
            "operation_id": "ocr",
            "operation_label": operation_label,
            "stage_progress": stage_progress,
        }
    if "LLM 原文纠错" in text or "原文纠错与断行" in text:
        return {
            "stage": "document_processing",
            "stage_id": "llm_cleanup",
            "operation": "llm_cleanup",
            "operation_id": "llm_cleanup",
            "operation_label": operation_label,
            "stage_progress": stage_progress,
        }
    if "按句构建" in text or "翻译单元" in text:
        return {
            "stage": "document_processing",
            "stage_id": "segment_build",
            "operation": "segment_build",
            "operation_id": "segment_build",
            "operation_label": operation_label,
            "stage_progress": stage_progress,
        }
    match = re.search(r"学术写作\s+(\d+)\s*/\s*(\d+)", text)
    if match:
        index, total = int(match.group(1)), int(match.group(2))
        section_match = re.search(r"第\s*([\w.-]+)\s*节", text)
        stage_id = "academic_writing"
        for needle, candidate in (
                ("证据库", "evidence"), ("研究问题", "research_model"),
                ("论点", "argument_plan"), ("案例", "selected_cases"),
                ("提纲", "outline"), ("撰写正文", "sections"),
                ("确定性", "validation"), ("审稿", "review"),
                ("修订", "quality_repair"), ("质量", "academic_quality")):
            if needle in text:
                stage_id = candidate
                break
        return {
            "stage": "academic_writing",
            "stage_id": stage_id,
            "stage_index": index,
            "stage_total": total,
            "operation": stage_id,
            "operation_id": "section_rewrite" if section_match else stage_id,
            "operation_label": operation_label,
            "section_id": section_match.group(1) if section_match else None,
            "stage_progress": stage_progress,
        }
    if "文档" in text or "排版" in text:
        stage = "document_processing"
    elif "术语" in text:
        stage = "terminology"
    elif "翻译" in text:
        stage = "translation"
    elif "标注" in text:
        stage = "annotation"
    elif "报告" in text or "学术" in text:
        stage = "academic_writing"
    else:
        stage = "pipeline"
    return {
        "stage": stage,
        "stage_id": stage,
        "operation": stage,
        "operation_id": stage,
        "operation_label": operation_label,
        "stage_progress": stage_progress,
    }


def _runtime_overall_progress(stage_info, state):
    """Return a conservative progress estimate; only completion may be 1.0."""
    if not isinstance(state, dict):
        return 0.0
    if state.get("p1_done") and state.get("p2_done") \
            and (state.get("p3_done") or not state.get("report_enabled", True)) \
            and not _annotation_blocks_completion(state):
        return 1.0
    return None


def _academic_runtime_progress(job_id, state):
    """Count only durable academic artifact checkpoints; no estimated percent."""
    if not isinstance(state, dict):
        return 0, 0
    academic = state.get("academic_state") or {}
    artifacts = academic.get("artifacts") or {}
    names = ("evidence", "research_model", "argument_plan", "selected_cases",
             "outline", "sections", "validation", "review",
             "literature_support_review", "academic_quality", "report")
    completed = sum(1 for name in names if (artifacts.get(name) or {}).get("file"))
    return completed, len(names)


_ACADEMIC_RESUME_META = {
    "evidence": (1, "evidence", "继续构建学术证据"),
    "literature_evidence": (2, "literature_evidence", "继续整理文献证据"),
    "research_model": (3, "research_model", "继续建立研究模型"),
    "literature_claims": (4, "literature_claims", "继续整理文献主张"),
    "argument_plan": (5, "argument_plan", "继续规划论点"),
    "case_analysis": (6, "outline", "继续规划案例分析"),
    "outline": (6, "outline", "继续生成学术提纲"),
    "writing": (7, "sections", "继续撰写报告章节"),
    "validation": (8, "validation", "继续验证报告"),
    "review": (9, "review", "继续执行学术复核"),
    "repair": (10, "quality_repair", "继续修订受影响章节"),
    "academic_quality": (11, "academic_quality", "继续评估学术质量"),
}


def _academic_resume_context(state, runtime=None):
    academic = (state or {}).get("academic_state") or {}
    current_stage = str(academic.get("current_stage") or "")
    index, operation, label = _ACADEMIC_RESUME_META.get(
        current_stage, (1, "academic_writing", "继续学术写作"))
    runtime = runtime or {}
    section_id = runtime.get("section_id")
    if current_stage == "repair" and section_id:
        operation = "section_rewrite"
        label = f"继续重新生成第 {section_id} 节"
    return {
        "pipeline": "academic_writing",
        "stage": "academic_writing",
        "stage_id": operation,
        "stage_index": index,
        "stage_total": 11,
        "operation": operation,
        "operation_id": operation,
        "operation_label": label,
        "stage_progress": {},
        "section_id": section_id or "",
    }


def _runtime_heartbeat_loop(job_id, stop_event):
    while not stop_event.wait(RUNTIME_HEARTBEAT_SECONDS):
        current = load_runtime_state(job_id)
        if current.get("status") not in RUNTIME_ACTIVE_STATUSES:
            return
        update_runtime_state(job_id, heartbeat=True)


def _runtime_status_callback(job_id, state, label):
    stage_info = _runtime_stage_info(label)
    current = load_runtime_state(job_id)
    operation = stage_info.get("operation") or current.get("operation")
    operation_started_at = current.get("operation_started_at")
    if operation != current.get("operation") or not operation_started_at:
        operation_started_at = _utc_now_iso()
    saved_state = load_job_state(job_id)
    state = saved_state or (state if isinstance(state, dict) else {})
    changes = dict(stage_info)
    completed_units, total_units = _academic_runtime_progress(job_id, state)
    changes.update({
        "operation": operation,
        "operation_id": stage_info.get("operation_id") or operation,
        "operation_label": stage_info.get("operation_label"),
        "stage_progress": stage_info.get("stage_progress") or {},
        "section_id": stage_info.get("section_id") or "",
        "operation_started_at": operation_started_at,
        "overall_progress": _runtime_overall_progress(stage_info, state),
        "status": "running",
        "phase": "running",
        "phase_label": "正在处理",
        "completed_units": completed_units,
        "total_units": total_units,
    })
    update_runtime_state(job_id, progress=True, event=label,
                         event_name=stage_info.get("stage_id") or operation,
                         event_visibility="user", event_category="progress",
                         **changes)
    if _runtime_cancel_requested(job_id):
        raise RuntimeError("任务已请求取消")


def _runtime_caption_callback(job_id, text):
    message = str(text or "").strip()
    if not message:
        return
    update_runtime_state(job_id, progress=True, event=message, status="running",
                         event_visibility="user", event_category="progress")
    if _runtime_cancel_requested(job_id):
        raise RuntimeError("任务已请求取消")


def _run_job_worker(job_id, filename, file_bytes, pipeline_kwargs, base_url=None):
    stop_event = threading.Event()
    heartbeat = threading.Thread(target=_runtime_heartbeat_loop,
                                 args=(job_id, stop_event), daemon=True)
    _RUNTIME_CTX.job_id = job_id
    set_llm_base_url(base_url)
    set_llm_reasoning_effort((pipeline_kwargs or {}).get("reasoning_effort"))
    heartbeat.start()
    # 主动留痕：进程被直接杀掉时没有机会写日志，所以"有启动、没有释放"本身就是
    # 最有价值的证据（本次排查正是靠它区分"被关闭"与"抛异常"）。
    _append_runtime_technical_log(
        job_id,
        f"worker started pid={os.getpid()} thread={threading.current_thread().name}"
        f" attempt={load_runtime_state(job_id).get('attempt')}"
        f" resume={bool(load_runtime_state(job_id).get('resume_request_id'))}")
    try:
        kwargs = dict(pipeline_kwargs or {})
        kwargs.pop("on_status", None)
        kwargs.pop("on_caption", None)
        state = load_job_state(job_id) or new_job_state(filename)
        queued = load_runtime_state(job_id)
        update_runtime_state(
            job_id, status="starting", phase="starting",
            phase_label="正在读取最近检查点", heartbeat=True, progress=True,
            event="后台 worker 已接管任务", event_name="worker_started",
            event_visibility="technical", event_category="orchestration",
            pipeline="academic_writing" if kwargs.get("enable_report")
            else "document_pipeline")
        update_runtime_state(
            job_id, status="running", phase="running",
            phase_label="已恢复 · 正在执行" if queued.get("resume_request_id")
            else "正在执行", heartbeat=True)
        result = run_job_pipeline(
            job_id, filename, file_bytes, **kwargs,
            on_status=lambda label: _runtime_status_callback(job_id, state, label),
            on_caption=lambda text: _runtime_caption_callback(job_id, text),
        )
        result_state = load_job_state(job_id) or state
        if _runtime_cancel_requested(job_id):
            update_runtime_state(job_id, status="cancelled", phase="cancelled",
                                 phase_label="已取消", event="任务已取消",
                                 event_name="job_cancelled", progress=True,
                                 event_visibility="user", event_category="lifecycle",
                                 worker={"owner_pid": None, "worker_id": None,
                                         "lease_expires_at": None})
        elif _runtime_business_complete(result_state):
            update_runtime_state(job_id, status="completed", phase="completed",
                                 phase_label="已完成", event="任务已完成",
                                 event_name="job_completed", progress=True,
                                 event_visibility="user", event_category="lifecycle",
                                 heartbeat=True, overall_progress=1.0,
                                 worker={"owner_pid": None, "worker_id": None,
                                         "lease_expires_at": None})
        else:
            update_runtime_state(job_id, status="idle_incomplete", phase="idle_incomplete",
                                 phase_label="等待继续", event="阶段已保存，等待继续",
                                 event_name="job_checkpointed", progress=True,
                                 event_visibility="technical", event_category="checkpoint",
                                 heartbeat=True,
                                 worker={"owner_pid": None, "worker_id": None,
                                         "lease_expires_at": None})
        return result
    # 捕获 BaseException 而不是 Exception：worker 只捕 Exception 时，若线程抛出
    # SystemExit 一类的 BaseException，`finally` 不会执行 → 心跳线程继续续租 →
    # lease 永不过期 → 应用还活着时这个任务会**真的硬卡死**（状态一直停在活跃态，
    # start_job_worker / resume_job 都会拒绝）。这里必须失败闭合。
    except BaseException as exc:  # noqa: BLE001 - worker must publish failure to UI
        cancelled = _runtime_cancel_requested(job_id) or "请求取消" in str(exc)
        if not cancelled:
            _write_runtime_technical_log(job_id, exc)
        current = load_runtime_state(job_id)
        provider_status = provider_error_status(exc)
        user_error_message = (
            provider_error_message(exc, "任务执行失败")
            if provider_status["status"] != "unknown"
            else str(exc)[:500] or "任务执行失败")
        error = None if cancelled else {
            "type": type(exc).__name__,
            "message": user_error_message,
            "stage": current.get("stage_id") or current.get("stage"),
            "operation": current.get("operation_id") or current.get("operation"),
            "timestamp": _utc_now_iso(),
            "technical_log": RUNTIME_TECHNICAL_LOG,
        }
        update_runtime_state(
            job_id, status="cancelled" if cancelled else "failed",
            phase="cancelled" if cancelled else "failed",
            phase_label="已取消" if cancelled else "步骤失败",
            error=error,
            event="任务已取消" if cancelled else f"步骤失败：{user_error_message[:180]}",
            event_name="job_cancelled" if cancelled else "job_failed",
            event_visibility="user",
            event_category="lifecycle" if cancelled else "error",
            progress=True, heartbeat=True,
            worker={"owner_pid": None, "worker_id": None,
                    "lease_expires_at": None})
        return None
    finally:
        _append_runtime_technical_log(
            job_id, f"worker released pid={os.getpid()}"
                    f" thread={threading.current_thread().name}")
        stop_event.set()
        set_llm_base_url(None)
        _RUNTIME_CTX.__dict__.clear()
        update_runtime_state(
            job_id, event="后台运行已释放", event_name="worker_released",
            event_visibility="technical", event_category="orchestration")
        with _RUNTIME_WORKERS_LOCK:
            _RUNTIME_WORKERS.pop(job_id, None)


def start_job_worker(job_id, filename, file_bytes, pipeline_kwargs, base_url=None,
                     resume_request_id=None):
    """Start one resumable pipeline worker; repeated UI reruns are idempotent."""
    if file_bytes is not None:
        try:
            save_source(job_id, file_bytes)
        except Exception:
            pass
    with _RUNTIME_WORKERS_LOCK:
        worker = _RUNTIME_WORKERS.get(job_id)
        if worker and worker.is_alive():
            return False
        runtime = load_runtime_state(job_id)
        status = runtime.get("status") or "idle"
        matching_resume = status == "resume_requested" and resume_request_id \
            and runtime.get("resume_request_id") == resume_request_id
        if status in RUNTIME_ACTIVE_STATUSES and not matching_resume:
            return False
        if not matching_resume:
            runtime = get_job_runtime_status(job_id)
            if runtime.get("status") in RUNTIME_ACTIVE_STATUSES:
                return False
        state = load_job_state(job_id) or new_job_state(filename)
        worker_id = uuid.uuid4().hex
        attempt = int(runtime.get("attempt") or 0) + 1
        started_at = _utc_now_iso()
        lease_expires_at = (datetime.now(timezone.utc) +
                            timedelta(seconds=RUNTIME_LEASE_SECONDS)).isoformat(
                                timespec="seconds")
        completed_units, total_units = (0, 0)
        if state.get("p2_done") and pipeline_kwargs.get("enable_report"):
            completed_units, total_units = _academic_runtime_progress(job_id, state)
        context = _academic_resume_context(state, runtime) \
            if matching_resume and pipeline_kwargs.get("enable_report") else {
                "pipeline": "document_pipeline", "stage": "pipeline",
                "stage_id": "pipeline", "operation": "pipeline",
                "operation_id": "pipeline", "operation_label": "准备任务",
                "stage_progress": {},
                "section_id": "", "stage_index": None, "stage_total": None,
            }
        update_runtime_state(
            job_id, status="queued", phase="starting", phase_label="准备中",
            **context, started_at=runtime.get("started_at") if matching_resume
            else started_at,
            operation_started_at=started_at, last_progress_at=started_at,
            last_heartbeat_at=started_at, cancel_requested=False, error=None,
            attempt=attempt, resume_request_id=resume_request_id or "",
            completed_units=completed_units, total_units=total_units,
            overall_progress=None,
            worker={"owner_pid": os.getpid(), "worker_id": worker_id,
                    "lease_expires_at": lease_expires_at},
            event="已排入后台 worker", event_name="job_queued",
            event_visibility="technical", event_category="orchestration")
        worker = threading.Thread(
            target=_run_job_worker,
            args=(job_id, filename, file_bytes, dict(pipeline_kwargs or {}), base_url),
            name=f"transpraxis-{job_id}", daemon=True)
        worker.worker_id = worker_id
        _RUNTIME_WORKERS[job_id] = worker
        worker.start()
        return True


def resume_job(job_id, filename, pipeline_kwargs, base_url=None,
               resume_request_id=None):
    """Idempotently request resume and publish the transition before worker start."""
    with _RUNTIME_WORKERS_LOCK:
        runtime = get_job_runtime_status(job_id)
        if runtime.get("status") in RUNTIME_ACTIVE_STATUSES:
            return False
        state = load_job_state(job_id) or new_job_state(filename)
        request_id = resume_request_id or uuid.uuid4().hex
        if runtime.get("resume_request_id") == request_id:
            return False
        pipeline_kwargs = dict(pipeline_kwargs or {})
        academic = state.get("academic_state") or {}
        if state.get("p2_done") and (academic.get("artifacts") or academic.get(
                "current_stage") not in {None, "", "not_started"}):
            pipeline_kwargs["enable_report"] = True
        completed_units, total_units = (0, 0)
        context = {
            "pipeline": "document_pipeline", "stage": "pipeline",
            "stage_id": "pipeline", "operation": "pipeline",
            "operation_id": "pipeline", "operation_label": "继续处理",
            "stage_progress": {},
            "section_id": "", "stage_index": None, "stage_total": None,
        }
        if state.get("p2_done") and pipeline_kwargs.get("enable_report"):
            completed_units, total_units = _academic_runtime_progress(job_id, state)
            context = _academic_resume_context(state, runtime)
        now = _utc_now_iso()
        update_runtime_state(
            job_id, status="resume_requested", phase="resume_requested",
            phase_label="正在恢复任务", resume_request_id=request_id,
            started_at=now, operation_started_at=now,
            last_progress_at=now, completed_units=completed_units,
            total_units=total_units, overall_progress=None,
            worker={"owner_pid": None, "worker_id": None,
                    "lease_expires_at": None},
            event="已从断点恢复任务", event_name="resume_requested",
            event_visibility="user", event_category="lifecycle",
            event_metadata={"resume_request_id": request_id}, **context)
        return start_job_worker(
            job_id, filename, None, pipeline_kwargs, base_url=base_url,
            resume_request_id=request_id)


def is_job_worker_alive(job_id):
    if not job_id:
        return False
    with _RUNTIME_WORKERS_LOCK:
        worker = _RUNTIME_WORKERS.get(job_id)
        if worker and worker.is_alive():
            return True
        # If in-memory registry lost the reference (e.g. Streamlit reloaded the module),
        # look for any active worker thread belonging to this job in the current process.
        target_name = f"transpraxis-{job_id}"
        for thread in threading.enumerate():
            if thread.name == target_name and thread.is_alive():
                _RUNTIME_WORKERS[job_id] = thread
                return True
        return False


def request_job_cancel(job_id, *, force=False):
    """Request cancellation between provider calls; an active HTTP call finishes first.

    If no active worker is running, or if force=True, marks the runtime state as cancelled immediately.
    """
    if not job_id:
        return False
    alive = is_job_worker_alive(job_id)
    if alive and not force:
        update_runtime_state(job_id, status="cancelling", cancel_requested=True,
                             phase="cancelling", phase_label="正在取消",
                             event="已请求取消任务", event_name="cancel_requested",
                             progress=True)
        return True

    # If no worker is alive or force cancel was explicitly requested:
    # Immediately transition the job state out of active running status.
    update_runtime_state(
        job_id, status="cancelled", phase="cancelled", phase_label="已取消",
        cancel_requested=False, error=None,
        event="任务已取消", event_name="job_cancelled",
        event_visibility="user", event_category="lifecycle",
        progress=True, heartbeat=True,
        worker={"owner_pid": None, "worker_id": None, "lease_expires_at": None})
    return True


def retry_job_step(job_id):
    """Invalidate only the failed academic operation and its downstream work."""
    runtime = get_job_runtime_status(job_id)
    if runtime.get("status") not in {"failed", "stalled", "interrupted", "cancelled"}:
        return False
    operation = runtime.get("operation_id") or runtime.get("operation") or ""
    scopes = {
        "validation": "validation", "review": "review",
        "literature_support_review": "literature_review",
        "academic_quality": "quality", "quality_repair": "quality",
        "section_rewrite": "section", "sections": "writer", "repair": "writer",
        "outline": "planning",
        "argument_plan": "planning", "selected_cases": "planning",
        "research_model": "planning", "evidence": "all",
    }
    state = load_job_state(job_id)
    if state and (runtime.get("stage") == "academic_writing" or
                  operation in scopes):
        error_message = str((runtime.get("error") or {}).get("message") or "")
        scope = "case_analysis" if operation == "section_rewrite" and \
            "missing case target subsection" in error_message else \
            scopes.get(operation, "writer")
        invalidate_academic_report(job_id, scope, runtime.get("section_id"))
    update_runtime_state(
        job_id, status="idle_incomplete", phase="retry_ready", phase_label="等待重试",
        cancel_requested=False, error=None, event="已准备重试当前步骤",
        event_name="retry_requested", worker={"owner_pid": None, "worker_id": None,
                                               "lease_expires_at": None})
    return True


def _invalidate_final_delivery_state(state):
    """Invalidate only the mutable working approval; snapshot history stays on disk."""
    state["delivery_status"] = "draft"
    state["delivery_approved_by_human"] = False
    state["delivery_approval"] = None
    if state.get("stage") == "FINAL":
        state["stage"] = _state_migration.derive_stage(state)
    return state


def _finalization_artifacts(job_id):
    """Load only the small set of artifacts needed for impact explanations."""
    names = ("selected_cases", "outline", "argument_plan", "sections")
    return {name: load_academic_artifact(job_id, name) for name in names}


def _reset_final_qa(state, reason=""):
    qa = _finalization.normalize_final_qa(state.get("final_qa"))
    qa.update({
        "structural_qa": "NOT_RUN",
        "libreoffice_render": "NOT_RUN",
        "author_visual_review": "NOT_CONFIRMED",
        "word_final_review": "NOT_CONFIRMED",
        "rendered_at": None,
        "page_count": None,
        "updated_at": _finalization.now_iso(),
    })
    if reason:
        notes = dict(qa.get("notes") or {})
        notes["stale_reason"] = reason
        qa["notes"] = notes
    state["final_qa"] = qa
    return qa


def _mark_translation_truth_changed(
    job_id, state, indexes, reason, *, actor="user", action="translation_changed",
    stale_translation_reviews=True,
):
    """Record one canonical CURRENT_TRANSLATION mutation and its impact slice."""
    indexes = sorted({int(index) for index in indexes
                      if isinstance(index, int) or str(index).lstrip("-").isdigit()})
    truth = dict(state.get("translation_truth") or {})
    truth["authority"] = _finalization.CURRENT_TRANSLATION
    truth["version"] = int(truth.get("version") or 0) + 1
    truth["last_changed_at"] = _finalization.now_iso()
    truth["last_change"] = {
        "action": action,
        "actor": actor,
        "reason": reason,
        "segment_indexes": indexes,
        "segment_ids": [
            _finalization.segment_id(
                job_id, index, (state.get("pairs") or [])[index]
                if 0 <= index < len(state.get("pairs") or []) else {})
            for index in indexes
        ],
    }
    state["translation_truth"] = truth
    if stale_translation_reviews:
        _invalidate_translation_reviews(state, indexes, reason)
    changed_segment_ids = list(truth["last_change"]["segment_ids"])
    from transpraxis import academic_writer
    propagated = academic_writer.propagate_artifact_staleness(
        state, input_segment_ids=changed_segment_ids)
    # Make the read-only artifact inputs available to the pure impact helper,
    # then remove them before state is persisted.
    enriched = dict(state)
    enriched["_finalization_artifacts"] = _finalization_artifacts(job_id)
    impact = _finalization.build_dependency_impact(
        enriched, job_id, indexes, reason)
    state["dependency_impact"] = impact
    _finalization.mark_case_reviews_stale(
        state, impact.get("affected_case_ids") or [], reason)
    stale_names = [item.get("id") for item in impact.get("affected") or []
                   if item.get("kind") == "artifact"]
    academic = state.setdefault("academic_state", {})
    # Legacy records have no direct edges and retain their historical
    # invalidation behavior. Canonical records are changed only by the graph
    # propagation above, preserving their own direct inputs.
    if not propagated and stale_names:
        academic_writer._invalidate_names(
            state, [name for name in stale_names
                    if name not in {"delivery_assets", "libreoffice_render"}], reason)
    if propagated:
        if any(name in propagated for name in {"sections", "report", "validation", "review"}):
            state["p3_done"] = False
            academic["status"] = "stale"
        if any(name in propagated for name in {"report", "sections"}):
            state["p3_md"] = ""
            state["p3_sections"] = []
    _reset_final_qa(state, reason)
    _invalidate_final_delivery_state(state)
    state.setdefault("human_actions", []).append({
        "finding_id": f"segment-mutation:{','.join(map(str, indexes)) or 'none'}",
        "action": action,
        "note": reason,
        "timestamp": _finalization.now_iso(),
        "actor": actor,
    })
    return state


def _invalidate_translation_reviews(
    state, indexes, reason, *, review_event_ids=None,
):
    """Revoke review/TM trust while preserving review and decision history."""
    # A whole-book consistency pass is a view over the current translation.
    # Any dependency invalidation makes that pass stale, even when the changed
    # segment did not have a current review event to revoke.
    state["targeted_final_review"] = {
        "status": "not_run", "segment_ids": [],
        "reviewed_segment_ids": [], "failed_segment_ids": [],
    }
    indexes = sorted({int(index) for index in indexes
                      if isinstance(index, int) or str(index).lstrip("-").isdigit()})
    changed = _translation_evidence.mark_runtime_review_stale(
        state, indexes, reason, dependency_change=True,
        review_event_ids=review_event_ids)
    if not any(changed.values()):
        return changed
    tm = load_tm(tm_project_id(state))
    tm_changed = False
    job_lang = state_target_lang(state)
    pairs = state.get("pairs") or []
    for index in indexes:
        if not 0 <= index < len(pairs):
            continue
        pair = pairs[index]
        pair["reviewed"] = False
        pair["review_status"] = "not_reviewed"
        pair["from_tm"] = False
        for key in ("accepted_target", "human_accepted", "accepted_by_human"):
            pair.pop(key, None)
        source = str(pair.get("source") or "")
        if tm_discard(tm, source, job_lang):
            tm_changed = True
    if tm_changed:
        save_tm(tm, tm_project_id(state))
    _recount_reviewed_segments(state)
    return changed


def translation_truth_view(job_id, state=None):
    """Return the user-facing authority/version summary for current targets."""
    state = state if state is not None else load_job_state(job_id) or {}
    truth = dict(state.get("translation_truth") or {})
    pairs = state.get("pairs") or []
    return {
        "authority": truth.get("authority") or _finalization.CURRENT_TRANSLATION,
        "version": int(truth.get("version") or 0),
        "segment_count": len(pairs),
        "last_changed_at": truth.get("last_changed_at"),
        "last_change": dict(truth.get("last_change") or {}),
        "label": "CURRENT_TRANSLATION · 当前工作译文",
    }


def dependency_impact_view(job_id, state=None):
    state = state if state is not None else load_job_state(job_id) or {}
    return _finalization.normalize_dependency_impact(state.get("dependency_impact"))


def compliance_profile_view(job_id, state=None):
    """Evaluate source-backed compliance and explicit project constraints."""
    state = state if state is not None else load_job_state(job_id) or {}
    from transpraxis import compliance
    artifacts = {
        name: load_academic_artifact(job_id, name)
        for name in ("evidence", "report", "validation", "outline",
                     "selected_cases", "literature_sources",
                     "final_docx_validation")
    }
    profile_id = str(state.get("compliance_profile_id") or
                     compliance.DEFAULT_PROFILE_ID)
    profile = compliance.compliance_profile(profile_id)
    result = compliance.evaluate_compliance(
        state, artifacts, profile, state.get("p3_md") or "")
    language = compliance.evaluate_language_constraints(
        state, state.get("p3_md") or "")
    result["language_constraints"] = language
    language_constraints = language.get("constraints") or []
    result["counts"]["pass"] += sum(
        item.get("status") == "pass" for item in language_constraints)
    result["counts"]["fail"] += len(language.get("failures") or [])
    result["counts"]["manual_review"] += sum(
        item.get("status") == "manual_review" for item in language_constraints)
    project = result.setdefault("project_constraints", {})
    project.setdefault("failures", []).extend(
        f"language:{item.get('kind')}:{item.get('value')}"
        for item in language.get("failures") or [])
    if language.get("failures"):
        project["status"] = "fail"
        result["status"] = "fail"
    if language.get("status") == "manual_review" and \
            result.get("status") != "fail":
        result["status"] = "manual_review"
    return result


def current_translation_hash(state=None):
    from transpraxis import academic_evidence
    state = state or {}
    return academic_evidence.stable_hash({
        "pairs": [
            {key: pair.get(key) for key in ("source", "initial_target", "target")}
            for pair in state.get("pairs") or []
        ],
    })


def generate_report_qa(job_id, state=None, *, save_file=False):
    """Build the concise QA report and bind it to current artifact hashes."""
    from transpraxis import academic_evidence, academic_writer
    state = state or load_job_state(job_id) or {}
    report_record = load_academic_artifact(job_id, "report") or {}
    render_record = load_academic_artifact(job_id, "libreoffice_render") or {}
    final_docx = load_academic_artifact(job_id, "final_docx_validation") or {}
    compliance = compliance_profile_view(job_id, state)
    case_review = _finalization.case_review_gate(
        state, load_academic_artifact(job_id, "selected_cases"))
    final_qa = _finalization.normalize_final_qa(state.get("final_qa"))
    final_qa["structural_qa"] = "PASS" if final_docx.get("status") in {
        "pass", "pass_with_warnings"} else "FAIL" if final_docx.get(
        "status") == "fail" else "NOT_RUN"
    translation_hash = current_translation_hash(state)
    report_hash = str(report_record.get("content_hash") or "")
    docx_hash = str(render_record.get("source_docx_hash") or
                    final_docx.get("source_docx_hash") or "")
    markdown = _rendered_qa.render_qa_markdown(
        translation_hash=translation_hash, report_hash=report_hash,
        docx_hash=docx_hash, render_record=render_record,
        pdf_qa=render_record.get("analysis") or {}, compliance=compliance,
        case_review=case_review, final_qa=final_qa,
        placeholders=next((x.get("actual") or [] for x in compliance.get(
            "rules") or [] if x.get("rule_id") == "author_placeholders"), []))
    value = {
        "schema_version": _rendered_qa.VERSION,
        "generated_at": _finalization.now_iso(),
        "translation_truth_hash": translation_hash,
        "report_content_hash": report_hash,
        "source_docx_hash": docx_hash,
        "rendered_pdf_hash": render_record.get("rendered_pdf_hash"),
        "final_qa": final_qa,
        "compliance_status": compliance.get("status"),
        "case_review_status": case_review.get("status"),
        "content_hash": academic_evidence.stable_hash({
            "translation": translation_hash, "report": report_hash,
            "docx": docx_hash, "pdf": render_record.get("rendered_pdf_hash"),
            "compliance": compliance.get("status"),
            "case_review": case_review.get("status"), "final_qa": final_qa,
        }),
        "markdown": markdown,
    }
    if save_file:
        academic_writer._save_artifact(
            state, job_dir(job_id), "report_qa", value, str(value["content_hash"]),
            _rendered_qa.VERSION,
            input_artifact_ids=["report", "final_docx_validation",
                                "libreoffice_render"],
            input_segment_ids=[])
        (job_dir(job_id) / "report-qa.md").write_text(markdown, encoding="utf-8")
        save_job_state(job_id, state)
    return value


def save_compliance_record(job_id, state=None):
    """Persist the current compliance and language-constraint artifacts."""
    from transpraxis import academic_evidence, academic_writer, compliance
    state = state or load_job_state(job_id) or {}
    result = compliance_profile_view(job_id, state)
    academic = state.get("academic_state") or {}
    records = academic.get("artifacts") or {}
    compliance_dependency = academic_evidence.stable_hash({
        "profile_id": result.get("profile_id"),
        "report": (records.get("report") or {}).get("content_hash"),
        "evidence": (records.get("evidence") or {}).get("content_hash"),
        "selected_cases": (records.get("selected_cases") or {}).get("content_hash"),
        "literature_sources": (records.get("literature_sources") or {}).get(
            "content_hash"),
        "outline": (records.get("outline") or {}).get("content_hash"),
    })
    academic_writer._save_artifact(
        state, job_dir(job_id), "compliance", result,
        compliance_dependency, compliance.compliance_profile(
            str(state.get("compliance_profile_id") or compliance.DEFAULT_PROFILE_ID)
        ).get("schema_version"),
        input_artifact_ids=["report", "evidence", "selected_cases",
                            "literature_sources"])
    language = result.get("language_constraints") or {}
    settings = state.get("research_settings") or {}
    language_dependency = academic_evidence.stable_hash({
        "report": (records.get("report") or {}).get("content_hash"),
        "language_constraints": {
            key: settings.get(key) for key in (
                "forbidden_report_phrases", "allowed_theory_labels",
                "required_terminology", "protected_names",
                "protected_work_titles")
        },
    })
    academic_writer._save_artifact(
        state, job_dir(job_id), "language_constraints", language,
        language_dependency, compliance.VERSION,
        input_artifact_ids=["report"])
    state["compliance_record"] = result
    state["language_constraint_record"] = language
    save_job_state(job_id, state)
    return result


def _case_artifact_case(job_id, case_id):
    selected = load_academic_artifact(job_id, "selected_cases") or {}
    case = next((item for item in selected.get("cases") or []
                 if str(item.get("case_id")) == str(case_id)), None)
    return selected, case


def _mark_case_downstream_stale(job_id, state, case_id, reason, *, actor="user",
                                action="case_changed", root_stale=True):
    from transpraxis import academic_writer
    root_id = f"case:{case_id}"
    before_root = dict((state.get("academic_state") or {}).get(
        "artifacts", {}).get(root_id) or {})
    propagated = academic_writer.propagate_artifact_staleness(
        state, input_artifact_ids=[f"case:{case_id}"])
    if not root_stale and before_root:
        academic = state.setdefault("academic_state", {})
        academic.setdefault("artifacts", {})[root_id] = before_root
        academic_writer._write_status_mirror(academic, root_id, before_root)
    enriched = dict(state)
    enriched["_finalization_artifacts"] = _finalization_artifacts(job_id)
    impact = _finalization.build_dependency_impact(
        enriched, job_id, [], reason, changed_case_ids=[str(case_id)])
    state["dependency_impact"] = impact
    stale_names = [item.get("id") for item in impact.get("affected") or []
                   if item.get("kind") == "artifact"]
    academic = state.setdefault("academic_state", {})
    # The old fallback is retained only for pre-Stage-2 records. Canonical
    # records are invalidated by their case direct edge and reverse traversal.
    if not propagated and stale_names:
        academic_writer._invalidate_names(state, stale_names, reason)
    if propagated and any(name in propagated for name in {"sections", "report"}):
        state["p3_done"] = False
        academic["status"] = "stale"
    _reset_final_qa(state, reason)
    _invalidate_final_delivery_state(state)
    state.setdefault("human_actions", []).append({
        "finding_id": f"case:{case_id}", "action": action, "note": reason,
        "timestamp": _finalization.now_iso(), "actor": actor,
    })
    return state


def review_academic_case(job_id, case_id, status, note="", actor="user"):
    """Record one author decision without changing provenance or content."""
    state = load_job_state(job_id)
    if state is None:
        return None, False, "任务不存在"
    selected, case = _case_artifact_case(job_id, case_id)
    if case is None:
        return state, False, "找不到案例"
    from transpraxis import case_provenance
    normalized = case_provenance.with_provenance(case)
    review_status = str(status or "").strip().lower()
    if review_status not in case_provenance.REVIEW_STATUSES:
        return state, False, "案例审校状态无效"
    state.setdefault("case_reviews", {})[str(case_id)] = {
        "review_status": review_status,
        "case_origin": normalized.get("case_origin"),
        "text_role": dict(normalized.get("text_role") or {}),
        "review_reason": str(note or "")[:700],
        "note": str(note or "")[:700],
        "reviewed_at": _finalization.now_iso(),
        "updated_at": _finalization.now_iso(),
        "actor": actor,
        "translation_truth_version": int(
            (state.get("translation_truth") or {}).get("version") or 0),
        "content_stale": False,
        "stale_reason": None,
        "stale_at": None,
    }
    state.setdefault("human_actions", []).append({
        "finding_id": f"case:{case_id}",
        "action": "case_approved" if review_status == "approved"
        else "case_excluded",
        "note": "批准案例纳入学术分析" if review_status == "approved"
        else f"排除案例：{str(note or '人工排除')[:180]}",
        "timestamp": _finalization.now_iso(),
        "actor": actor,
    })
    _mark_case_downstream_stale(
        job_id, state, str(case_id),
        "作者更新案例审核状态，需重组案例相关写作下游",
        actor=actor, action="case_review_changed", root_stale=False)
    save_job_state(job_id, state)
    return state, True, "已保存案例审校状态"


def replace_rejected_case(job_id, case_id, *, actor="user"):
    """Replace one author-rejected case from the already validated pool."""
    from transpraxis import academic_writer, academic_quality, case_provenance
    state = load_job_state(job_id)
    if state is None:
        return None, False, ["任务不存在"]
    selected, old_case = _case_artifact_case(job_id, case_id)
    if selected is None or old_case is None:
        return state, False, ["找不到案例"]
    review = (state.get("case_reviews") or {}).get(str(case_id)) or {}
    if str(review.get("review_status")) != "rejected":
        return state, False, ["只能替换已被作者排除的案例"]
    selected_ids = {str(x.get("case_id")) for x in selected.get("cases") or []}
    reviews = state.get("case_reviews") or {}
    overrides = state.get("case_review_overrides") or {}
    candidates = []
    if case_provenance.is_synthetic(old_case):
        synthetic = load_academic_artifact(job_id, "synthetic_validation") or {}
        for candidate in synthetic.get("items") or []:
            cid = str(candidate.get("case_id") or "")
            validation = candidate.get("validation") or {}
            review_record = reviews.get(cid) if isinstance(reviews, dict) else None
            override = overrides.get(cid) if isinstance(overrides, dict) else None
            if not cid or cid in selected_ids or candidate is old_case:
                continue
            if review_record and review_record.get("review_status") == "rejected":
                continue
            if override and override.get("baseline_status") == "rejected":
                continue
            if not validation.get("academic_case_eligible"):
                continue
            difficulty = candidate.get("difficulty") or {}
            evidence = candidate.get("synthetic_evidence") or {}
            old_group = str(old_case.get("difficulty_group") or "")
            target_match = bool(old_group and str(difficulty.get("group") or "") == old_group)
            candidates.append((not target_match,
                               difficulty.get("academic_value") != "high",
                               difficulty.get("confidence") != "high",
                               evidence.get("material_difference") != "pass",
                               candidate.get("segment_index", 0), candidate))
    else:
        evidence_artifact = load_academic_artifact(job_id, "evidence") or {}
        argument_plan = load_academic_artifact(job_id, "argument_plan") or {}
        candidate = academic_quality.select_replacement_case(
            str(case_id), list(old_case.get("supports_claims") or []),
            selected, argument_plan, evidence_artifact)
        if candidate:
            candidates.append((False, False, False, False,
                               candidate.get("segment_index", 0), candidate))
    candidates.sort(key=lambda item: item[:-1])
    if not candidates:
        return state, False, ["没有可用的已验证替换候选；请保持该案例排除并调整案例数量"]
    candidate = case_provenance.with_provenance(dict(candidates[0][-1]))
    candidate.update({
        "supports_claims": sorted(set(old_case.get("supports_claims") or [])),
        "research_questions": sorted(set(old_case.get("research_questions") or [])),
        "argument_role": old_case.get("argument_role", "supporting"),
        "difficulty_group": old_case.get("difficulty_group") or
        candidate.get("difficulty_group"),
        "difficulty_subsection": old_case.get("difficulty_subsection"),
        "strategy_subsection": old_case.get("strategy_subsection"),
        "target_subsection": old_case.get("target_subsection"),
        "review_status": "unreviewed",
        "replacement_of": str(case_id),
        "selection_rationale": f"replacement of rejected {case_id}: existing validated pool",
    })
    selected["cases"] = [
        candidate if str(x.get("case_id")) == str(case_id)
        else x for x in selected.get("cases") or []
    ]
    record = academic_writer.artifact_record(state, "selected_cases")
    academic_writer._save_artifact(
        state, job_dir(job_id), "selected_cases", selected,
        str(record.get("dependency_hash") or ""),
        str(record.get("version") or "case-review-v1"))
    _mark_case_downstream_stale(
        job_id, state, str(candidate.get("case_id")),
        "作者排除案例后从已验证候选池替换，需重组其写作下游",
        actor=actor, action="case_replaced", root_stale=False)
    state.setdefault("case_reviews", {})[str(candidate.get("case_id"))] = {
        "review_status": "unreviewed",
        "case_origin": candidate.get("case_origin"),
        "text_role": dict(candidate.get("text_role") or {}),
        "review_reason": "替换被排除案例；需作者重新终审",
        "reviewed_at": _finalization.now_iso(),
        "updated_at": _finalization.now_iso(),
        "actor": actor,
        "translation_truth_version": int(
            (state.get("translation_truth") or {}).get("version") or 0),
        "content_stale": False,
    }
    state.setdefault("human_actions", []).append({
        "finding_id": f"case:{candidate.get('case_id')}",
        "action": "case_replaced",
        "note": f"从已验证候选池替换 {case_id}",
        "timestamp": _finalization.now_iso(),
        "actor": actor,
    })
    save_job_state(job_id, state)
    return state, True, [str(candidate.get("case_id"))]


def update_synthetic_baseline(job_id, case_id, text, *, status="modified", note="", actor="user"):
    """Modify or reject a synthetic baseline; never mutate translation truth."""
    state = load_job_state(job_id)
    if state is None:
        return None, False, "任务不存在"
    _selected, case = _case_artifact_case(job_id, case_id)
    if case is None:
        return state, False, "找不到案例"
    from transpraxis import case_provenance
    if not case_provenance.is_synthetic(case):
        return state, False, "只有合成对照案例可以修改或拒绝模拟初译"
    if status not in {"modified", "rejected", "approved"}:
        return state, False, "模拟初译状态无效"
    record = dict((state.setdefault("case_review_overrides", {}).get(str(case_id)) or {}))
    if status == "modified":
        text = str(text or "").strip()
        if not text:
            return state, False, "模拟初译不能为空"
        record["synthetic_baseline_text"] = text
    record.update({
        "baseline_status": status,
        "note": str(note or "")[:700],
        "updated_at": _finalization.now_iso(),
        "actor": actor,
    })
    state["case_review_overrides"][str(case_id)] = record
    reason = ("修改模拟初译，需重跑该案例下游"
              if status == "modified" else "拒绝模拟初译，需替换或重新确认该案例")
    _finalization.mark_case_reviews_stale(state, [str(case_id)], reason)
    _mark_case_downstream_stale(
        job_id, state, str(case_id), reason, actor=actor,
        action="synthetic_baseline_modified" if status == "modified"
        else "synthetic_baseline_rejected")
    save_job_state(job_id, state)
    return state, True, "已保存模拟初译决定"


def record_final_qa(job_id, field, status, note="", actor="user"):
    """Persist one of the four independent final-QA facts."""
    if field not in _finalization.QA_FIELDS:
        raise ValueError(f"未知最终 QA 项：{field}")
    state = load_job_state(job_id)
    if state is None:
        return None
    allowed = {"PASS", "FAIL", "NOT_RUN"} if field in {
        "structural_qa", "libreoffice_render"} else {"CONFIRMED", "NOT_CONFIRMED"}
    if status not in allowed:
        raise ValueError(f"{field} 状态无效：{status}")
    qa = _finalization.normalize_final_qa(state.get("final_qa"))
    qa[field] = status
    qa["translation_truth_version"] = int(
        (state.get("translation_truth") or {}).get("version") or 0)
    notes = dict(qa.get("notes") or {})
    if note:
        notes[field] = str(note)[:700]
    qa["notes"] = notes
    qa["updated_at"] = _finalization.now_iso()
    state["final_qa"] = qa
    save_job_state(job_id, state)
    return state


def run_libreoffice_render_qa(job_id, state=None):
    """Render through LibreOffice, then run separate deterministic PDF QA."""
    state = state or load_job_state(job_id)
    if state is None:
        raise ValueError(f"找不到任务 {job_id}")
    docx = report_docx_bytes(job_id, state)
    document_kind = "report"
    if docx is None:
        report_record = (state.get("academic_state") or {}).get(
            "artifacts", {}).get("report") or {}
        if state.get("report_enabled") and state.get("p3_done") and \
                report_record.get("status") in {"stale", "missing", "failed"}:
            raise RuntimeError("当前实践报告 artifact 已 stale，请先按影响范围重建报告")
        try:
            docx = build_delivery_assets(job_id, state).get("translation.docx")
            document_kind = "translation"
        except Exception as exc:
            raise RuntimeError(f"当前 DOCX 不可生成：{str(exc)[:180]}") from exc
    if not docx:
        raise RuntimeError("当前没有可渲染的 DOCX")
    previous_render = load_academic_artifact(job_id, "libreoffice_render") or {}
    current_qa = _finalization.normalize_final_qa(state.get("final_qa"))
    if (previous_render.get("rendered_pdf_hash") or
            previous_render.get("qa_status") == "PASS" or
            current_qa.get("author_visual_review") == "CONFIRMED" or
            current_qa.get("word_final_review") == "CONFIRMED"):
        _reset_final_qa(state, "LibreOffice render rerun; author and Word reviews reset")
        if current_qa.get("structural_qa") in {"PASS", "FAIL"}:
            state["final_qa"]["structural_qa"] = current_qa["structural_qa"]
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        qa = _finalization.normalize_final_qa(state.get("final_qa"))
        qa.update({"libreoffice_render": "NOT_RUN", "rendered_at": None,
                   "page_count": None, "source_docx_hash": _rendered_qa.sha256(docx),
                   "rendered_pdf_hash": None,
                   "updated_at": _finalization.now_iso(),
                   "translation_truth_version": int(
                       (state.get("translation_truth") or {}).get("version") or 0)})
        qa.setdefault("notes", {})["libreoffice_render"] = \
            "LibreOffice not installed; render NOT_RUN"
        state["final_qa"] = qa
        from transpraxis import academic_writer
        academic_writer._save_artifact(
            state, job_dir(job_id), "libreoffice_render", {
                "schema_version": _rendered_qa.VERSION,
                "status": "not_run", "qa_status": "NOT_RUN",
                "render_engine": "libreoffice", "render_engine_version": None,
                "source_docx_hash": _rendered_qa.sha256(docx),
                "rendered_pdf_hash": None, "rendered_at": None, "page_count": None,
                "stale_reason": "LibreOffice unavailable; render was not run",
                "analysis": {"warnings": [], "manual_reviews": [{
                    "type": "libreoffice_unavailable", "severity": "manual_review"}]},
            }, "no-engine", _rendered_qa.VERSION,
            input_artifact_ids=["final_docx_validation"], status="missing")
        generate_report_qa(job_id, state, save_file=True)
        save_job_state(job_id, state)
        return state, qa
    from transpraxis import academic_evidence, academic_writer
    version_result = subprocess.run([soffice, "--version"], capture_output=True,
                                    text=True, timeout=10, check=False)
    engine_version = (version_result.stdout or version_result.stderr or "").strip()
    with tempfile.TemporaryDirectory(prefix=f"transpraxis-lo-{job_id}-") as tmp:
        tmp_path = Path(tmp)
        source_path = tmp_path / "current.docx"
        source_path.write_bytes(docx)
        profile = tmp_path / "profile"
        profile.mkdir()
        command = [soffice, "--headless", "-env:UserInstallation=file://"
                   + str(profile), "--convert-to", "pdf", "--outdir", tmp,
                   str(source_path)]
        result = subprocess.run(command, capture_output=True, text=True,
                                timeout=120, check=False)
        pdf_path = tmp_path / "current.pdf"
        if result.returncode != 0 or not pdf_path.is_file():
            detail = (result.stderr or result.stdout or "未知 LibreOffice 错误").strip()
            qa = _finalization.normalize_final_qa(state.get("final_qa"))
            qa.update({"libreoffice_render": "FAIL",
                       "rendered_at": _finalization.now_iso(),
                       "translation_truth_version": int(
                           (state.get("translation_truth") or {}).get("version") or 0),
                       "updated_at": _finalization.now_iso()})
            qa.setdefault("notes", {})["libreoffice_render"] = detail[:700]
            state["final_qa"] = qa
            render_value = {
                "status": "fail", "qa_status": "FAIL", "document_kind": document_kind,
                "detail": detail[:700],
                "render_engine": "libreoffice", "render_engine_version": engine_version,
                "source_docx_hash": _rendered_qa.sha256(docx),
                "rendered_pdf_hash": None,
                "rendered_at": _finalization.now_iso(), "page_count": None,
                "stale_reason": {"code": "render_failed",
                                 "source_type": "artifact",
                                 "source_id": "final_docx_validation"},
            }
            academic_writer._save_artifact(
                state, job_dir(job_id), "libreoffice_render", render_value,
                academic_evidence.stable_hash({
                    "final_docx_validation": (state.get("academic_state") or {}).get(
                        "artifacts", {}).get("final_docx_validation", {}).get(
                            "content_hash"),
                    "version": _finalization.VERSION,
                }), _finalization.VERSION,
                input_artifact_ids=["final_docx_validation"], status="failed",
                stale_reason={"code": "render_failed", "source_type": "artifact",
                              "source_id": "final_docx_validation"})
            generate_report_qa(job_id, state, save_file=True)
            save_job_state(job_id, state)
            raise RuntimeError(f"LibreOffice 渲染失败：{detail[:180]}")
        pdf_bytes = pdf_path.read_bytes()
    output = job_dir(job_id) / "libreoffice-render.pdf"
    output.write_bytes(pdf_bytes)
    analysis = _rendered_qa.analyze_pdf(pdf_bytes)
    page_count = int(analysis.get("page_count") or 0)
    page_metrics = analysis.get("pages") or []
    render_status = "PASS" if page_count and not analysis.get(
        "definite_failures") else "FAIL"
    qa = _finalization.normalize_final_qa(state.get("final_qa"))
    qa.update({
        "libreoffice_render": render_status,
        "rendered_at": _finalization.now_iso(),
        "page_count": page_count,
        "source_docx_hash": _rendered_qa.sha256(docx),
        "rendered_pdf_hash": _rendered_qa.sha256(pdf_bytes),
        "translation_truth_version": int(
            (state.get("translation_truth") or {}).get("version") or 0),
        "updated_at": _finalization.now_iso(),
    })
    qa.setdefault("notes", {})["document_kind"] = document_kind
    qa["page_metrics"] = page_metrics
    state["final_qa"] = qa
    render_value = _rendered_qa.build_render_record(
        document_kind=document_kind, source_docx=docx, rendered_pdf=pdf_bytes,
        engine="libreoffice", engine_version=engine_version, analysis=analysis)
    render_value["qa_status"] = render_status
    render_value["status"] = "pass" if render_status == "PASS" else "fail"
    academic_writer._save_artifact(
        state, job_dir(job_id), "libreoffice_render", render_value,
        academic_evidence.stable_hash({
            "final_docx_validation": (state.get("academic_state") or {}).get(
                "artifacts", {}).get("final_docx_validation", {}).get("content_hash"),
            "version": _finalization.VERSION,
        }), _finalization.VERSION,
        input_artifact_ids=["final_docx_validation"],
        status="valid" if qa["libreoffice_render"] == "PASS" else "failed")
    generate_report_qa(job_id, state, save_file=True)
    save_job_state(job_id, state)
    return state, qa


def _reconcile_final_delivery_snapshot(job_id, state):
    latest = _snapshots.latest_snapshot(job_dir(job_id))
    if latest and state.get("delivery_status") == "final" \
            and _snapshots.state_identity(state) != latest.get("translation_state_identity"):
        return _invalidate_final_delivery_state(state)
    return state


def new_job_state(filename):
    state = {
        "filename": filename,
        "p1_done": False,
        "p2_done": False,
        "p3_done": False,
        "report_enabled": False,
        "paras": [],
        "pairs": [],
        "auto_terms": {},
        "findings": [],
        "review_stats": {
            "reviewed_segments": 0, "batches_reviewed": 0,
            "blocking": 0, "actionable": 0, "informational": 0, "review_failed": 0,
        },
        "tm_used_count": 0,
        "has_blocking": False,
        "p3_md": "",
        "p3_sections": [],
        "theory": "",
        "warnings": [],
        "annotations": {},
        "annotations_done": False,
        "annotations_done_offset": 0,
    }
    # 术语治理 / 交付门禁新增字段（默认值集中在 state_migration，保持单一来源）
    state.update(_state_migration._default_new_fields())
    return state


def load_job_state(job_id):
    """加载任务状态。使用文件签名缓存（mtime/size）消除重复反序列化，并返回安全深拷贝。"""
    return _job_repo.load_job_state_cached(job_id)


def save_job_state(job_id, state):
    """原子写入（先写临时文件再替换），避免中断写坏 state.json，并同步派生 summary.json。"""
    _job_repo.save_job_state_atomic(job_id, state)


# ================= 未保存译文草稿的持久化 =================
# 草稿**不是**文档状态的一部分：它不进 state.json、不改 reviewed/review_status、
# 不进入交付资产。它只回答一个问题——"用户敲进去但还没点保存的内容"，在会话被
# 重建（刷新、关标签页、进程重启）之后还在不在。此前它只活在 st.session_state 里，
# 刷新即丢；会话会重置这件事我们改不了，但"草稿只存在会话里"是可以改的。
TRANSLATION_DRAFTS_VERSION = 1
TRANSLATION_DRAFTS_FILE = "translation_drafts.json"


def translation_drafts_path(job_id):
    return job_dir(job_id) / TRANSLATION_DRAFTS_FILE


def load_translation_drafts(job_id):
    """读回未保存草稿，返回 `{segment_id: {"text", "baseline", "updated_at"}}`。

    读不到、格式不对、文件损坏一律返回空字典：草稿是"尽力而为"的兜底，
    它的读取失败绝不能让工作台打不开。
    """
    path = translation_drafts_path(job_id)
    if not path.is_file():
        return {}
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
    except Exception:  # noqa: BLE001 - 草稿损坏按"没有草稿"处理
        return {}
    if not isinstance(raw, dict):
        return {}
    items = raw.get("drafts")
    if not isinstance(items, list):
        return {}
    drafts = {}
    for item in items:
        if not isinstance(item, dict):
            continue
        segment_id = str(item.get("segment_id") or "")
        if not segment_id:
            continue
        drafts[segment_id] = {
            "segment_id": segment_id,
            "text": str(item.get("text") or ""),
            "baseline": str(item.get("baseline") or ""),
            "updated_at": str(item.get("updated_at") or ""),
        }
    return drafts


def save_translation_drafts(job_id, drafts):
    """原子写入草稿文件。没有草稿时**删除文件**，不留空壳。"""
    d = job_dir(job_id)
    d.mkdir(parents=True, exist_ok=True)
    path = d / TRANSLATION_DRAFTS_FILE
    items = []
    for segment_id, record in (drafts or {}).items():
        if not isinstance(record, dict):
            continue
        items.append({
            "segment_id": str(segment_id),
            "text": str(record.get("text") or ""),
            "baseline": str(record.get("baseline") or ""),
            "updated_at": str(record.get("updated_at")
                              or datetime.now(timezone.utc).isoformat()),
        })
    if not items:
        path.unlink(missing_ok=True)
        return
    items.sort(key=lambda item: item["segment_id"])
    payload = {
        "version": TRANSLATION_DRAFTS_VERSION,
        "job_id": str(job_id),
        "updated_at": datetime.now(timezone.utc).isoformat(),
        "drafts": items,
    }
    tmp = d / (TRANSLATION_DRAFTS_FILE + ".tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(path)


def clear_translation_drafts(job_id):
    """丢弃本任务全部未保存草稿（文件层面）。"""
    translation_drafts_path(job_id).unlink(missing_ok=True)


def _recount_reviewed_segments(state):
    """Keep the segment review count in sync with the canonical pair records."""
    state.setdefault("review_stats", {})["reviewed_segments"] = sum(
        bool(pair.get("reviewed")) for pair in state.get("pairs") or [])


# ================= 段落身份 / 排除 / 结构快照 =================
# `pairs` 与 `paras` 是贯穿流水线的位置契约：审校发现、检查点、交付资产全部按
# index 引用同一条记录。拆分、合并、插入、删除会改变**每一个后续 index 的含义**，
# 而界面需要"同一段在结构操作之后仍然是同一段"（编辑框、选中态、Agent 候选都不
# 能跟着旧索引漂移）。`segment_uid` 只为界面身份服务，数据契约保持按索引不变。
SEGMENT_UID_FIELD = "segment_uid"
_UID_SEQ_FIELD = "segment_uid_seq"

SEGMENT_STRUCTURE_OPERATION_LABELS = {
    "source_edit": "修改原文",
    "split": "拆分段落",
    "merge": "合并段落",
    "insert": "插入空段",
    "delete": "删除空段",
    "exclude": "排除段落",
    "include": "恢复段落",
}


def segment_uid(pair):
    """一段的稳定身份；没有时返回空串（调用方负责退化到索引）。"""
    return str((pair or {}).get(SEGMENT_UID_FIELD) or "")


def is_excluded(pair):
    """这一段是否被用户显式排除在翻译与交付之外。"""
    return bool((pair or {}).get("excluded"))


def active_pair_indexes(state):
    """参与翻译、审校与交付的段落索引（排除项不属于任何一项）。"""
    return [index for index, pair in enumerate((state or {}).get("pairs") or [])
            if isinstance(pair, dict) and not is_excluded(pair)]


def next_segment_uid(state, job_id=""):
    """分配一个新的稳定身份。单调递增，撤销不回退（已用过的身份不再复用）。"""
    seq = int(state.get(_UID_SEQ_FIELD) or 0)
    state[_UID_SEQ_FIELD] = seq + 1
    return f"seg-{job_id}-u{seq:06d}" if job_id else f"u{seq:06d}"


def assign_missing_segment_uids(state, *, job_id=""):
    """给缺少稳定身份的段落补齐身份，返回新分配的数量。

    旧任务（v0.4 之前）没有这个字段，补齐动作发生在**第一次结构编辑或人工保存**
    上，因此普通渲染不会因为"读了一次状态"就改写磁盘。

    初始身份刻意沿用导出资产既有的格式 `seg-<job_id>-<index>`：升级瞬间不会发生
    身份漂移，浏览器里已经存在的选中态、编辑框与未保存草稿都继续有效。结构操作
    产生的新段落用 `seg-<job_id>-u<n>`，与索引格式不可能碰撞，且永不复用。
    """
    pairs = (state or {}).get("pairs")
    if not isinstance(pairs, list):
        return 0
    seq = int(state.get(_UID_SEQ_FIELD) or 0)
    seen = set()
    assigned = 0
    for index, pair in enumerate(pairs):
        if not isinstance(pair, dict):
            continue
        uid = segment_uid(pair)
        if uid and uid not in seen:
            seen.add(uid)
            continue
        if job_id:
            pair[SEGMENT_UID_FIELD] = f"seg-{job_id}-{index:04d}"
        else:
            pair[SEGMENT_UID_FIELD] = f"u{seq:06d}"
            seq += 1
        seen.add(pair[SEGMENT_UID_FIELD])
        assigned += 1
    state[_UID_SEQ_FIELD] = max(seq, len(pairs), len(seen))
    return assigned


def excluded_segment_records(state):
    """被排除段落的清单——导出、交付说明与「恢复」入口的唯一来源。

    排除采用"移出 pairs/paras + 在此留档"的模型：只要不在工作列表里，它就自然
    不进入翻译、不计入待完成数量、不参与审校门禁、不进导出；而原文与当时的译文
    在这里完整保留，所以恢复是精确还原而不是重新猜。
    """
    records = []
    for position, item in enumerate((state or {}).get("excluded_segments") or []):
        if not isinstance(item, dict):
            continue
        records.append({
            "excluded_id": str(item.get("excluded_id") or f"x{position:06d}"),
            "segment_index": item.get("index"),
            "source": str(item.get("source") or ""),
            "target": str(item.get("target") or ""),
            "reason": str(item.get("reason") or ""),
            "excluded_at": str(item.get("excluded_at") or ""),
            "excluded_by": str(item.get("excluded_by") or ""),
            "restorable": bool(item.get("pair") is not None
                               or str(item.get("source") or "").strip()),
        })
    # 兼容极早期状态：pair 上直接带 excluded 标记。
    for index, pair in enumerate((state or {}).get("pairs") or []):
        if not isinstance(pair, dict) or not is_excluded(pair):
            continue
        records.append({
            "excluded_id": segment_uid(pair) or f"index-{index}",
            "segment_index": index,
            "source": str(pair.get("source") or ""),
            "target": str(pair.get("target") or ""),
            "reason": str(pair.get("excluded_reason") or ""),
            "excluded_at": str(pair.get("excluded_at") or ""),
            "excluded_by": str(pair.get("excluded_by") or ""),
            "restorable": False,
        })
    return records


def excluded_segments_summary(state):
    """排除范围的只读摘要（交付页/导入报告共用，避免各处自己数一遍）。"""
    records = excluded_segment_records(state)
    working = len((state or {}).get("pairs") or []) \
        or len((state or {}).get("paras") or [])
    return {
        "count": len(records),
        "working_segments": working,
        "total_segments": working + len(records),
        "records": records,
    }


def excluded_segments_manifest(state):
    """交付包里如实说明"哪些内容没有进入译文文档"。"""
    records = excluded_segment_records(state)
    return {
        "version": 1,
        "rule": "被排除段落不进入译文的生成式输出，也不占用待完成与审校配额；"
                "原文与排除时的译文在此完整保留，可在工作台恢复。",
        "count": len(records),
        "records": records,
    }


def _excluded_segments_md(manifest):
    lines = [
        "# 已排除段落清单",
        "",
        f"- 排除数量：{manifest.get('count', 0)}",
        f"- 规则：{manifest.get('rule', '')}",
        "",
        "| 原位置 | 排除原因 | 原文 |",
        "| --- | --- | --- |",
    ]
    for record in manifest.get("records") or []:
        index = record.get("segment_index")
        position = f"第 {int(index) + 1} 段" if isinstance(index, int) else "—"
        source = " ".join(str(record.get("source") or "").split())
        if len(source) > 160:
            source = source[:159] + "…"
        reason = str(record.get("reason") or "未填写")
        lines.append(f"| {position} | {reason} | {source or '（空）'} |")
    lines.append("")
    return "\n".join(lines)


def _segment_source_list(state):
    """当前的段落原文序列（有译文时以 pairs 为准，否则用 paras）。

    排除/恢复需要用原文做锚点：索引会随结构操作漂移，原文文本不会。
    """
    pairs = (state or {}).get("pairs") or []
    if pairs:
        return [str((pair or {}).get("source") or "") for pair in pairs]
    return [str(item or "") for item in (state or {}).get("paras") or []]


def _exclusion_restore_index(state, record):
    """被排除段落应当插回的位置：优先后锚点，其次前锚点，最后退回原索引。"""
    sources = _segment_source_list(state)
    after = str((record or {}).get("after_source") or "")
    before = str((record or {}).get("before_source") or "")
    if after:
        for position, text in enumerate(sources):
            if text == after:
                return position
    if before:
        for position, text in enumerate(sources):
            if text == before:
                return position + 1
    return max(0, min(int((record or {}).get("index") or 0), len(sources)))


def can_undo_translation_segment_structure(state):
    """最近一次结构操作是否可撤销（供界面决定是否显示入口）。"""
    history = [record for record in (state or {}).get("segment_structure_history") or []
               if isinstance(record, dict)]
    if not history:
        return {"available": False, "label": "", "operation": "", "index": None}
    record = history[-1]
    operation = str(record.get("operation") or "")
    return {
        "available": isinstance(record.get("restore_pairs"), list),
        "label": SEGMENT_STRUCTURE_OPERATION_LABELS.get(operation, "段落结构已更新"),
        "operation": operation,
        "index": record.get("index"),
    }


def translation_terms_for_pair(state, pair):
    """Resolve only glossary entries explicitly attached to one pair."""
    ids = {str(item) for item in pair.get("glossary_entry_ids") or []}
    entries = state.get("glossary") or []
    if not entries:
        frozen = state.get("glossary_frozen") or {}
        entries = frozen.get("entries") or [] if isinstance(frozen, dict) else []
    if not entries:
        versions = state.get("glossary_versions") or []
        if versions and isinstance(versions[-1], dict):
            entries = versions[-1].get("entries") or []
    if not ids or not isinstance(entries, list):
        return []
    terms = []
    for entry in entries:
        if not isinstance(entry, dict) or str(entry.get("id") or "") not in ids:
            continue
        source = str(entry.get("source") or "")
        if source:
            terms.append((source, str(entry.get("preferred") or entry.get("target") or "—"),
                          "项目术语"))
    return terms[:8]


def translation_visible_indexes(state, search="", status_filter="全部",
                               filter_terms=False, filter_edited=False,
                               filter_issues=False, filter_tm=False,
                               issue_indexes=None):
    """Return visible pair indexes without mutating the loaded state."""
    pairs = state.get("pairs") or []
    query = str(search or "").strip().casefold()
    issue_indexes = set(issue_indexes or [])
    visible = []
    for index, pair in enumerate(pairs):
        source = str(pair.get("source") or "")
        target = str(pair.get("target") or "")
        if query:
            paragraph_query = query.lstrip("#").strip()
            if paragraph_query.isdigit():
                if int(paragraph_query) != index + 1:
                    continue
            elif query not in f"{source}\n{target}".casefold():
                continue
        if status_filter == "待审" and pair.get("reviewed"):
            continue
        if status_filter == "已审校" and not pair.get("reviewed"):
            continue
        if filter_terms and not translation_terms_for_pair(state, pair):
            continue
        if filter_edited and not (pair.get("human_edited") or
                                  pair.get("source_human_edited")):
            continue
        if filter_tm and not pair.get("from_tm"):
            continue
        if filter_issues and index not in issue_indexes:
            continue
        visible.append(index)
    return visible


def save_translation_edit(job_id, index, target, actor="user"):
    """Save one human translation edit through the state/business layer.

    A manual edit is a new working translation: it preserves the paragraph
    identity, clears segment review, removes TM reuse provenance, and invalidates
    mutable final approval while leaving historical frozen snapshots untouched.
    """
    from transpraxis import delivery as _delivery

    state = load_job_state(job_id)
    if state is None:
        raise ValueError(f"找不到任务 {job_id}")
    pairs = state.get("pairs") or []
    if not (0 <= index < len(pairs)):
        raise IndexError(f"段落索引超出范围：{index}")
    pair = pairs[index]
    if is_excluded(pair):
        raise ValueError("该段落已被排除，不会进入翻译与交付；请先恢复本段再编辑译文。")
    assign_missing_segment_uids(state, job_id=job_id)
    new_target = str(target or "").strip()
    if not pair.get("human_edited"):
        pair["_translation_edit_restore"] = {
            "target": pair.get("target") or "",
            "reviewed": bool(pair.get("reviewed")),
            "review_status": pair.get("review_status", "not_reviewed"),
            "target_provenance": pair.get("target_provenance", "generated"),
            "from_tm": bool(pair.get("from_tm")),
        }
    pair["target"] = new_target
    pair["human_edited"] = True
    pair["reviewed"] = False
    pair["review_status"] = "not_reviewed"
    pair["target_provenance"] = "human_edit"
    pair["from_tm"] = False
    _recount_reviewed_segments(state)

    _mark_translation_truth_changed(
        job_id, state, [index], "人工修改 CURRENT_TRANSLATION；相关案例与学术下游需要重建",
        actor=actor, action="translation_edit")
    _recheck_delivery_invariants_for_segments(state, [index], actor=actor)
    save_job_state(job_id, state)
    return state


def _recheck_delivery_invariants_for_segments(state, indexes, *, actor="user"):
    """Refresh deterministic target blockers after a working-text edit.

    A target invariant describes the text that existed when it was recorded.
    When a user edits that target, a resolved old invariant must not continue
    to block delivery; if the same problem remains, the existing finding stays
    open and the current validation report remains authoritative.
    """
    from transpraxis import delivery as _delivery

    indexes = {int(index) for index in indexes}
    report = validate_delivery_translation_state(state)
    active = {
        (issue.get("segment_index"), issue.get("code"))
        for issue in report.get("issues") or []
    }
    for finding in state.setdefault("findings", []):
        if not isinstance(finding, dict) \
                or finding.get("type") != "delivery_invariant" \
                or finding.get("segment_index") not in indexes \
                or finding.get("resolved"):
            continue
        key = (finding.get("segment_index"), finding.get("invariant_code"))
        if key in active:
            continue
        finding["resolved"] = True
        finding["resolution"] = {
            "action": "target_rechecked",
            "note": "CURRENT_TRANSLATION 编辑后重新通过目标文本门禁",
            "timestamp": _finalization.now_iso(),
            "actor": actor,
        }
        _delivery.add_human_action(
            state, _delivery.finding_id(finding), "target_rechecked",
            "CURRENT_TRANSLATION 编辑后重新通过目标文本门禁", actor)
    _record_delivery_validation_findings(state, report)
    state["delivery_status"] = _delivery.compute_delivery_status(state)
    return report


def restore_translation_edit(job_id, index, actor="user"):
    """Restore the translation state that existed before the first edit."""
    from transpraxis import delivery as _delivery

    state = load_job_state(job_id)
    if state is None:
        raise ValueError(f"找不到任务 {job_id}")
    pairs = state.get("pairs") or []
    if not (0 <= index < len(pairs)):
        raise IndexError(f"段落索引超出范围：{index}")
    pair = pairs[index]
    restore = pair.pop("_translation_edit_restore", None)
    if restore is None and not pair.get("human_edited"):
        return state
    restore = restore or {}
    pair["target"] = restore.get("target") or pair.get("initial_target") or ""
    pair["reviewed"] = bool(restore.get("reviewed"))
    pair["review_status"] = restore.get("review_status", "not_reviewed")
    pair["target_provenance"] = restore.get("target_provenance", "generated")
    pair["from_tm"] = bool(restore.get("from_tm"))
    pair.pop("human_edited", None)
    _recount_reviewed_segments(state)

    _mark_translation_truth_changed(
        job_id, state, [index], "恢复前版本也改变了 CURRENT_TRANSLATION；相关下游需要重建",
        actor=actor, action="translation_restore")
    save_job_state(job_id, state)
    return state


def mutate_translation_segments(
    job_id, index, operation, *, source_text=None, source_a=None,
    source_b=None, target_a=None, target_b=None, exclude_reason=None,
    excluded_id=None, actor="user",
):
    """Apply a user initiated CAT segment-structure edit.

    ``pairs`` and ``paras`` are a positional contract throughout the pipeline:
    review findings, checkpoints and delivery assets all refer to the same
    segment index.  This is therefore deliberately one state-layer operation
    instead of a UI-only splice.  Structural edits are rejected while a worker
    is active, stale every existing review dependency, and persist a compact
    history record for auditability (including a content snapshot, so
    :func:`undo_translation_segment_structure` can restore the previous text).

    Supported operations:

    ``source_edit``
        Correct the source text in place while retaining the current target.
    ``split``
        Replace one segment with two explicitly supplied source/target pairs.
    ``merge``
        Join the selected segment with the following segment using a newline.
    ``insert``
        Add an empty manual segment immediately after ``index``.
    ``delete``
        Remove an empty manual segment.  Existing content cannot be deleted by
        this path, which prevents an accidental paragraph loss.
    ``exclude``
        Take a segment out of translation, review and delivery without touching
        its text.  The source text is always retained, so ``include`` restores
        it exactly.  This is the general entry point for OCR junk, page numbers
        and repeated running heads that no character-level repair can fix.
    ``include``
        Undo an exclusion.

    The returned value is the persisted state, matching the existing edit APIs.
    """
    state = load_job_state(job_id)
    if state is None:
        raise ValueError(f"找不到任务 {job_id}")

    runtime = get_job_runtime_status(job_id, state)
    runtime_status = str(runtime.get("status") or "")
    if runtime_status in RUNTIME_ACTIVE_STATUSES:
        raise RuntimeError("任务正在运行，段落结构编辑将在任务完成或暂停后可用。")

    operation = str(operation or "").strip().lower()
    aliases = {
        "source": "source_edit", "edit_source": "source_edit",
        "split_segment": "split", "merge_segment": "merge",
        "insert_segment": "insert", "delete_segment": "delete",
        "exclude_segment": "exclude", "include_segment": "include",
        "restore_segment": "include",
    }
    operation = aliases.get(operation, operation)
    if operation not in {"source_edit", "split", "merge", "insert", "delete",
                         "exclude", "include"}:
        raise ValueError(f"不支持的段落结构操作：{operation or '—'}")

    # 结构操作是界面身份最容易漂移的地方：先给所有段落补齐稳定身份，再动列表。
    assign_missing_segment_uids(state, job_id=job_id)

    pairs = list(state.get("pairs") or [])
    paras = list(state.get("paras") or [])
    if pairs and len(paras) != len(pairs):
        # A completed translation must never be silently repaired here: a
        # mismatch means another invariant has already been violated and the
        # user needs the existing delivery gate to surface it.
        raise ValueError("当前任务的源文与双语段落数量不一致，无法安全编辑段落结构。")
    if not pairs and not paras:
        raise ValueError("当前任务还没有可编辑的段落。")
    if not pairs and operation not in {"exclude", "include"}:
        raise ValueError("当前任务还没有可编辑的译文段落。")
    if not all(isinstance(pair, dict) for pair in pairs):
        raise ValueError("当前任务包含无效的段落记录，无法安全编辑。")

    # 「恢复」的目标不在 pairs/paras 里（它正是被移出去的那一段），所以它的位置
    # 由原文锚点反推，而不是当作索引来校验。
    include_record = None
    if operation == "include":
        wanted = str(excluded_id or "")
        include_record = next(
            (item for item in state.get("excluded_segments") or []
             if isinstance(item, dict)
             and str(item.get("excluded_id") or "") == wanted), None)
        if include_record is None:
            raise ValueError("找不到要恢复的已排除段落。")
        index = _exclusion_restore_index(state, include_record)

    limit = len(pairs) if pairs else len(paras)
    if isinstance(index, bool) or not isinstance(index, int):
        raise ValueError("段落索引无效。")
    if operation == "include":
        if not 0 <= index <= limit:
            raise IndexError(f"段落索引超出范围：{index}")
    elif not 0 <= index < limit:
        raise IndexError(f"段落索引超出范围：{index}")

    # 操作前的内容快照：撤销需要的是**原文本身**，只记"操作名 + 位置"恢复不了
    # 被拆分/合并掉的内容。capture 是操作前被替换的切片；undo 描述撤销时要用
    # 多少次替换回去（insert 是插入、delete 是删除，形状与 capture 不同）。
    _SHAPES = {
        "source_edit": (index, 1, 1),
        "split": (index, 1, 2),
        "merge": (index, 2, 1),
        "insert": (index + 1, 0, 1),
        "delete": (index, 1, 0),
        # 排除把段落移出 pairs/paras；撤销要把同一条记录放回原位。
        "exclude": (index, 1, 0),
        # 恢复把已排除段落插回；撤销要再把它移出去，并还原排除记录。
        "include": (index, 0, 1),
    }
    undo_index, capture_span, undo_span = _SHAPES[operation]

    def normalized(value):
        return str(value or "").replace("\r\n", "\n").strip()

    def reset_trust(pair, source, target, *, initial_target=None,
                    source_edited=False, manual_segment=False):
        """Copy a pair while revoking trust tied to the old segmentation."""
        result = dict(pair or {})
        result["source"] = normalized(source)
        result["target"] = normalized(target)
        if initial_target is not None:
            result["initial_target"] = normalized(initial_target)
        else:
            result["initial_target"] = normalized(
                result.get("initial_target") or result.get("target"))
        result["reviewed"] = False
        result["review_status"] = "not_reviewed"
        result["from_tm"] = False
        result["target_provenance"] = "human_edit"
        result.pop("_translation_edit_restore", None)
        if source_edited:
            result["source_human_edited"] = True
            result["source_provenance"] = "human_edit"
        if manual_segment:
            result["manual_segment"] = True
            result["segment_origin"] = "manual"
        return result

    old_indexes = list(range(len(pairs)))
    # 撤销要恢复的是**内容**：只记"操作名 + 位置"恢复不了被拆分/合并掉的原句。
    # 快照只存操作前被替换的那一小段（0–2 条记录），50 条历史也不会影响体积。
    _restore_pairs = ([json.loads(json.dumps(item))
                       for item in pairs[undo_index:undo_index + capture_span]]
                      if capture_span else [])
    _restore_paras = ([str(item or "")
                       for item in paras[undo_index:undo_index + capture_span]]
                      if capture_span else [])
    reason_by_operation = {
        "source_edit": "人工修订了源文段落，原有译文审校需要重新确认",
        "split": "人工拆分了段落，原有译文审校需要重新确认",
        "merge": "人工合并了段落，原有译文审校需要重新确认",
        "insert": "人工插入了段落，原有译文审校需要重新确认",
        "delete": "人工删除了空段落，原有译文审校需要重新确认",
        "exclude": "人工排除了段落，交付范围发生变化，原有审校需要重新确认",
        "include": "人工恢复了被排除的段落，交付范围发生变化，原有审校需要重新确认",
    }
    reason = reason_by_operation[operation]

    # Stale the old ordinal dependencies before changing the list.  The normal
    # truth mutation below uses the new indexes; this pre-pass also covers a
    # deleted tail segment whose old index no longer exists afterward.  A source
    # text correction keeps list positions stable, so it only touches the
    # edited segment; every other operation changes the ordinal mapping.
    #
    # `exclude` / `include` 也属于"改变序号含义"的一类：移出或插回一段会让**它
    # 之后每一段**的索引都位移一位，而审校发现、检查点与交付资产全部按索引引用
    # 同一条记录。只把 `index` 判过期，会让后面那些段的旧发现静默指到别的段落上
    # ——这正是"结构操作后问题定位对不上"的根因。
    pre_stale_indexes = [index] if operation == "source_edit" else old_indexes
    _translation_evidence.mark_runtime_review_stale(
        state, pre_stale_indexes, reason, dependency_change=True)

    # 排除/恢复对导出清单的增量；撤销时按这两份增量反向操作。
    excluded_ids_added = []
    excluded_records_restored = []

    if operation == "source_edit":
        new_source = normalized(source_text)
        if not new_source:
            raise ValueError("原文不能为空；如需留白，请插入空段。")
        current = pairs[index]
        updated_pair = reset_trust(
            current, new_source, current.get("target") or "",
            initial_target=current.get("initial_target"), source_edited=True)
        pairs[index] = updated_pair
        paras[index] = new_source
        affected_hint = [index]
    elif operation == "split":
        left_source, right_source = normalized(source_a), normalized(source_b)
        if not left_source or not right_source:
            raise ValueError("拆分后的两段原文都不能为空。")
        current = pairs[index]
        left_target, right_target = normalized(target_a), normalized(target_b)
        original_ids = list(current.get("glossary_entry_ids") or [])
        left = reset_trust(
            current, left_source, left_target, initial_target=left_target,
            source_edited=True)
        right = reset_trust(
            current, right_source, right_target, initial_target=right_target,
            source_edited=True)
        # 拆出来的两段是**新**工作单元：必须拿到新身份，否则右半段会继承原段
        # 的身份，界面上的编辑框/选中态就会认错段落。
        left[SEGMENT_UID_FIELD] = next_segment_uid(state, job_id)
        right[SEGMENT_UID_FIELD] = next_segment_uid(state, job_id)
        if original_ids:
            left["glossary_entry_ids"] = list(original_ids)
            right["glossary_entry_ids"] = list(original_ids)
        pairs[index:index + 1] = [left, right]
        paras[index:index + 1] = [left_source, right_source]
        affected_hint = [index, index + 1]
    elif operation == "merge":
        if index + 1 >= len(pairs):
            raise ValueError("最后一段没有下一段可合并。")
        current, following = pairs[index], pairs[index + 1]
        merged_source = "\n".join(
            item for item in (normalized(current.get("source")),
                              normalized(following.get("source"))) if item)
        merged_target = "\n".join(
            item for item in (normalized(current.get("target")),
                              normalized(following.get("target"))) if item)
        merged_initial = "\n".join(
            item for item in (normalized(current.get("initial_target")),
                              normalized(following.get("initial_target"))) if item)
        merged = reset_trust(
            current, merged_source, merged_target,
            initial_target=merged_initial, source_edited=True)
        # 合并保留左段的身份（用户是在左段上发起操作的），右段身份随之消失。
        merged[SEGMENT_UID_FIELD] = segment_uid(current) or next_segment_uid(state, job_id)
        merged_ids = list(dict.fromkeys(
            list(current.get("glossary_entry_ids") or []) +
            list(following.get("glossary_entry_ids") or [])))
        if merged_ids:
            merged["glossary_entry_ids"] = merged_ids
        merged["merged_segment_count"] = int(
            current.get("merged_segment_count") or 1) + int(
                following.get("merged_segment_count") or 1)
        pairs[index:index + 2] = [merged]
        paras[index:index + 2] = [merged_source]
        affected_hint = [index]
    elif operation == "insert":
        inserted = reset_trust(
            {}, "", "", initial_target="", manual_segment=True)
        inserted[SEGMENT_UID_FIELD] = next_segment_uid(state, job_id)
        pairs.insert(index + 1, inserted)
        paras.insert(index + 1, "")
        affected_hint = [index + 1]
    elif operation == "exclude":
        # 「排除」把这一段整体移出 pairs/paras：不进入翻译、不计入待完成数量、
        # 不参与审校门禁、不进导出。原文（连同当时的译文）完整存进
        # `state["excluded_segments"]`，因此恢复是精确的、不是重新猜。
        sources = _segment_source_list(state)
        source_value = str(paras[index] or "")
        target_value = str(pairs[index].get("target") or "") if pairs else ""
        if not normalized(source_value) and not normalized(target_value):
            raise ValueError("该段落没有内容，无需排除；如需清理空段请使用「删除空段」。")
        excluded_id_value = f"x{int(state.get('exclusion_seq') or 0):06d}"
        state["exclusion_seq"] = int(state.get("exclusion_seq") or 0) + 1
        excluded_record = {
            "excluded_id": excluded_id_value,
            "index": index,
            "source": source_value,
            "target": target_value,
            "pair": json.loads(json.dumps(pairs[index])) if pairs else None,
            "in_pairs": bool(pairs),
            "before_source": sources[index - 1] if index > 0 else "",
            "after_source": sources[index + 1] if index + 1 < len(sources) else "",
            "reason": str(exclude_reason or "").strip(),
            "excluded_at": _finalization.now_iso(),
            "excluded_by": str(actor or "user"),
        }
        state["excluded_segments"] = list(
            state.get("excluded_segments") or []) + [excluded_record]
        excluded_ids_added = [excluded_id_value]
        if pairs:
            pairs.pop(index)
        paras.pop(index)
        affected_hint = [max(0, min(index, len(pairs) - 1))] if pairs else []
    elif operation == "include":
        excluded_record = dict(include_record)
        state["excluded_segments"] = [
            item for item in state.get("excluded_segments") or []
            if not (isinstance(item, dict)
                    and str(item.get("excluded_id") or "")
                    == str(excluded_record.get("excluded_id") or ""))
        ]
        # 排除发生在翻译之前时没有 pair 记录，此时只需要把原文插回 paras。
        restored_pair = excluded_record.get("pair")
        if pairs and isinstance(restored_pair, dict):
            # 内容原样回来，但**不**连带把旧的"已审校"标记也一起点亮：它的审校
            # 依赖在排除那一刻已经被判过期，重新显示"已审校"会是假绿（撤销路径
            # 出于同样的理由在 `undo_translation_segment_structure` 里清标记）。
            restored_pair = json.loads(json.dumps(restored_pair))
            restored_pair["reviewed"] = False
            restored_pair["review_status"] = "not_reviewed"
            restored_pair.pop("accepted_target", None)
            restored_pair.pop("human_accepted", None)
            restored_pair.pop("accepted_by_human", None)
            pairs.insert(index, restored_pair)
        paras.insert(index, str(excluded_record.get("source") or ""))
        affected_hint = [index]
        excluded_records_restored = [excluded_record]
    else:  # delete
        candidate = pairs[index]
        if not candidate.get("manual_segment"):
            raise ValueError("只能删除手动插入且仍为空的段落；"
                             "导入内容请使用「排除本段」。")
        if normalized(candidate.get("source")) or normalized(candidate.get("target")):
            raise ValueError("该段落已有内容，请先清空原文和译文后再删除。")
        pairs.pop(index)
        paras.pop(index)
        affected_hint = [max(0, min(index, len(pairs) - 1))] if pairs else []

    state["pairs"] = pairs
    state["paras"] = paras
    # Splitting, merging, inserting or deleting changes the meaning of every
    # following ordinal.  Revoke their trust together so no old review badge or
    # TM reuse survives under a different source/target pairing.  A plain
    # source correction only resets the pair already rebuilt above, and an
    # exclusion changes the delivery scope without touching other segments.
    if operation not in {"source_edit", "exclude", "include"}:
        for pair in pairs:
            pair["reviewed"] = False
            pair["review_status"] = "not_reviewed"
            pair["from_tm"] = False
            for key in ("accepted_target", "human_accepted", "accepted_by_human"):
                pair.pop(key, None)
    state["annotations"] = {}
    state["annotations_done"] = False
    state["annotations_done_offset"] = 0
    state["segment_structure_version"] = int(
        state.get("segment_structure_version") or 0) + 1
    history = list(state.get("segment_structure_history") or [])
    history.append({
        "operation": operation,
        "index": index,
        "affected_indexes": affected_hint,
        "timestamp": _finalization.now_iso(),
        "actor": actor,
        # 撤销要恢复的是**内容**：只记"操作名 + 位置"恢复不了被拆分/合并掉的原句。
        # 快照只存操作前被替换的那一小段（0–2 条记录），50 条历史也不影响体积。
        "undo_index": undo_index,
        "undo_span": undo_span,
        "restore_pairs": _restore_pairs,
        "restore_paras": _restore_paras,
        "excluded_ids_added": list(excluded_ids_added),
        "excluded_records_restored": [
            json.loads(json.dumps(item)) for item in excluded_records_restored],
    })
    state["segment_structure_history"] = history[-50:]

    # 结构操作会成片清掉 `reviewed`（或把它留给恢复的那一段），计数必须跟着重算，
    # 否则报告模板里的 "已审校段落" 与段落本身脱节。
    _recount_reviewed_segments(state)

    all_new_indexes = list(range(len(pairs)))
    # source_edit 原地改字，影响面就是那一段；其余操作改的是序号映射，整张表都要
    # 按新索引重新对账（排除/恢复也一样）。
    trust_indexes = (list(affected_hint) or [index]
                     if operation == "source_edit"
                     else all_new_indexes)
    _mark_translation_truth_changed(
        job_id, state, trust_indexes, reason, actor=actor,
        action=f"segment_{operation}")
    _recheck_delivery_invariants_for_segments(
        state, all_new_indexes, actor=actor)
    save_job_state(job_id, state)
    return state


def undo_translation_segment_structure(job_id, *, actor="user"):
    """撤销最近一次段落结构操作，恢复到操作前的原文/译文。

    只撤销**一步**：历史记录被弹出并留档到 `segment_structure_undo_log`，因此
    连续点两次撤销不会"撤销掉撤销"。被恢复的段落会失去审校通过标记——它们的
    内容虽然是操作前的原文，但下游依赖已经在操作时被判定为过期，重新点亮
    "已审校"会是假绿。
    """
    state = load_job_state(job_id)
    if state is None:
        raise ValueError(f"找不到任务 {job_id}")

    runtime = get_job_runtime_status(job_id, state)
    if str(runtime.get("status") or "") in RUNTIME_ACTIVE_STATUSES:
        raise RuntimeError("任务正在运行，段落结构编辑将在任务完成或暂停后可用。")

    history = [record for record in state.get("segment_structure_history") or []
               if isinstance(record, dict)]
    if not history:
        raise ValueError("没有可撤销的段落结构操作。")
    record = history[-1]
    restore_pairs = record.get("restore_pairs")
    if not isinstance(restore_pairs, list):
        raise ValueError("最近一次结构操作没有保存内容快照，无法撤销。")

    pairs = list(state.get("pairs") or [])
    paras = list(state.get("paras") or [])
    if len(paras) < len(pairs):
        raise ValueError("当前任务的源文与双语段落数量不一致，无法安全撤销。")

    undo_index = int(record.get("undo_index", record.get("index") or 0) or 0)
    undo_span = int(record.get("undo_span") or 0)
    undo_index = max(0, min(undo_index, max(len(pairs), len(paras))))
    pairs_end = max(undo_index, min(len(pairs), undo_index + undo_span))
    paras_end = max(undo_index, min(len(paras), undo_index + undo_span))

    restored = [dict(item) for item in restore_pairs if isinstance(item, dict)]
    for pair in restored:
        pair["reviewed"] = False
        pair["review_status"] = "not_reviewed"
    restored_paras = [str(item or "") for item in record.get("restore_paras") or []]
    if len(restored_paras) != len(restored):
        restored_paras = [str(pair.get("source") or "") for pair in restored]

    # 排除发生在翻译之前时 pairs 本来就是空的，撤销只动 paras。
    if pairs:
        pairs[undo_index:pairs_end] = restored
    paras[undo_index:paras_end] = restored_paras

    # 排除/恢复对清单的增量按反方向还原，保证撤销之后"已排除清单"和段落列表一致。
    added_ids = {str(item) for item in record.get("excluded_ids_added") or []}
    if added_ids:
        state["excluded_segments"] = [
            item for item in state.get("excluded_segments") or []
            if not (isinstance(item, dict)
                    and str(item.get("excluded_id") or "") in added_ids)]
    put_back = [dict(item) for item in record.get("excluded_records_restored") or []
                if isinstance(item, dict)]
    if put_back:
        state["excluded_segments"] = list(
            state.get("excluded_segments") or []) + put_back
    state["pairs"] = pairs
    state["paras"] = paras
    state["annotations"] = {}
    state["annotations_done"] = False
    state["annotations_done_offset"] = 0
    state["segment_structure_version"] = int(
        state.get("segment_structure_version") or 0) + 1
    state["segment_structure_history"] = history[:-1]
    undo_log = list(state.get("segment_structure_undo_log") or [])
    undo_log.append({**record, "undone_at": _finalization.now_iso(),
                     "undone_by": actor})
    state["segment_structure_undo_log"] = undo_log[-20:]
    # 恢复的段落失去了 `reviewed`，计数必须跟着降下来。
    _recount_reviewed_segments(state)

    all_indexes = list(range(len(pairs)))
    _mark_translation_truth_changed(
        job_id, state, all_indexes,
        "撤销了最近一次段落结构操作；恢复的译文需要重新审校",
        actor=actor, action="segment_undo")
    _recheck_delivery_invariants_for_segments(state, all_indexes, actor=actor)
    save_job_state(job_id, state)
    return state


def list_jobs():
    """列出本地任务。**纯读取**：目录不存在时返回空列表，不创建它。"""
    jobs = []
    if not OUTPUT_DIR.is_dir():
        return jobs
    for d in sorted(OUTPUT_DIR.iterdir()):
        if not d.is_dir():
            continue
        s = _job_repo.load_job_state_cached(d.name)
        if s is not None:
            jobs.append({"job_id": d.name, "state": s})
    return jobs


def list_job_summaries():
    """列出本地任务的轻量级摘要（派生读取模型）。"""
    return _job_repo.list_job_summaries()


def read_runtime_view(job_id):
    """读取任务的轻量级运行时视图，避免每次轮询读取并解析完整 state.json。"""
    return _job_repo.read_runtime_view(job_id)


def delete_job(job_id, *, allow_active=False):
    """永久删除一个翻译任务（含其输出目录）。

    正在运行的任务默认拒绝删除：worker 还在往这个目录写状态，删掉会留下一个
    半写状态并让运行时报错。返回是否真的删除。
    """
    if not str(job_id or "").strip():
        return False
    d = job_dir(job_id)
    if not d.exists():
        return False
    if not allow_active and job_is_active(job_id):
        raise ValueError("任务正在运行，无法删除；请先取消或等它结束。")
    shutil.rmtree(d)
    _job_repo.invalidate_cache(job_id)
    return True


def file_job_id(file_bytes):
    """**文档身份**：文件内容哈希。同一份文档重传得到同一个值。

    注意：这**不是任务身份**。同一份文档在不同项目 / 不同目标语言下是彼此
    独立的本地化任务，内容哈希无法区分它们——用内容哈希当任务 ID 会让
    "同一个文件、另一个项目、另一种目标语言"静默打开旧任务。任务身份见
    `task_job_id` / `resolve_task_id`。
    """
    return hashlib.sha256(file_bytes).hexdigest()[:16]


# 任务身份 = 文档身份 + 本地化上下文（项目 + 目标语言 + 源语言）。
# 文档内容相同不等于本地化任务相同：项目决定注入哪一套项目记忆与术语，
# 目标语言决定译文本身。把它们丢掉，就等于让内容哈希替用户决定
# "这两个语义不同的活是同一个活"。
_TASK_ID_FIELDS = ("project_id", "target_lang", "source_lang")


def task_context(project_id=None, target_lang=None, source_lang=None):
    """本地化上下文的规范化三元组（用于任务身份与"是否同一个任务"的比较）。"""
    return {
        "project_id": resolved_project_id({"project_id": project_id}),
        "target_lang": normalize_language(target_lang),
        "source_lang": normalize_language(source_lang),
    }


def task_context_of(state):
    """一个已存在任务记录的本地化上下文（缺字段按"未指定"处理）。"""
    state = state or {}
    return task_context(state.get("project_id"), state.get("target_lang"),
                        state.get("source_lang"))


def task_job_id(file_bytes, *, project_id=None, target_lang=None,
                source_lang=None):
    """**任务身份**：文档身份 + 本地化上下文，确定性推导，不读磁盘。"""
    context = task_context(project_id, target_lang, source_lang)
    digest = "\x1f".join([file_job_id(file_bytes)]
                         + [context[field] for field in _TASK_ID_FIELDS])
    return hashlib.sha256(digest.encode("utf-8")).hexdigest()[:16]


def resolve_task_id(file_bytes, *, project_id=None, target_lang=None,
                    source_lang=None):
    """这份文档在**这个本地化上下文**下应当打开的任务 ID。

    三条规则，顺序即优先级：

    1. 按上下文推导出的任务已存在 -> 续做它（"同一个文件、同一个项目、
       同一种目标语言"就是同一个任务，重传即续传）；
    2. 否则，若历史版本留下的、以**内容哈希**命名的旧任务存在，且它记录的
       项目 / 目标语言与请求的上下文**可证明一致** -> 沿用它（旧任务不因
       这次修复而失联）；
    3. 否则 -> 用推导出的新 ID 建一个独立任务。

    规则 2 里的"可证明一致"是关键：旧任务没有目标语言记录时**不会**被认领，
    宁可让用户重建一个任务，也不把另一种语言的旧任务当成这个活。
    """
    context = task_context(project_id, target_lang, source_lang)
    candidate = task_job_id(file_bytes, project_id=project_id,
                            target_lang=target_lang, source_lang=source_lang)
    if load_job_state(candidate) is not None:
        return candidate
    legacy_id = file_job_id(file_bytes)
    if legacy_id != candidate:
        legacy_state = load_job_state(legacy_id)
        if legacy_state is not None and task_context_of(legacy_state) == context:
            return legacy_id
    return candidate


def save_source(job_id, file_bytes):
    """留存源文件，刷新页面后即使不重新上传也能继续（如重做阶段一）。"""
    d = job_dir(job_id)
    d.mkdir(parents=True, exist_ok=True)
    (d / "source.bin").write_bytes(file_bytes)


def load_source(job_id):
    p = job_dir(job_id) / "source.bin"
    return p.read_bytes() if p.is_file() else None


def save_report_template(job_id, filename, template_bytes):
    """Parse and persist the uploaded DOCX template before academic stages run."""
    from transpraxis import report_template

    raw = _bytes(template_bytes)
    contract = report_template.parse_docx_template(filename, raw)
    directory = job_dir(job_id)
    directory.mkdir(parents=True, exist_ok=True)
    (directory / "report-template.docx").write_bytes(raw)
    (directory / "template-contract.json").write_text(
        json.dumps(contract, ensure_ascii=False, indent=2), encoding="utf-8")
    state = load_job_state(job_id)
    if state is not None:
        identity = contract.get("template_identity") or {}
        state["report_template"] = {
            "filename": identity.get("filename"),
            "template_id": identity.get("template_id"),
            "template_hash": identity.get("sha256"),
            "schema_version": contract.get("schema_version"),
            "status": "parsed",
        }
        state["report_template_contract"] = contract
        settings = dict(state.get("research_settings") or {})
        settings["report_template_contract"] = contract
        state["research_settings"] = settings
        save_job_state(job_id, state)
    return contract


def clear_report_template(job_id):
    """Remove the saved report template and its contract from a task."""
    directory = job_dir(job_id)
    for name in ("report-template.docx", "template-contract.json"):
        path = directory / name
        if path.is_file():
            path.unlink()
    state = load_job_state(job_id)
    if state is not None:
        state["report_template"] = None
        state["report_template_contract"] = None
        settings = dict(state.get("research_settings") or {})
        settings.pop("report_template_contract", None)
        settings.pop("template_contract", None)
        state["research_settings"] = settings
        save_job_state(job_id, state)
    return state


def load_report_template(job_id):
    """Load the persisted DOCX bytes and canonical contract, if configured."""
    from transpraxis import report_template

    directory = job_dir(job_id)
    docx_path = directory / "report-template.docx"
    contract_path = directory / "template-contract.json"
    state = load_job_state(job_id)
    contract = None
    if contract_path.is_file():
        try:
            value = json.loads(contract_path.read_text(encoding="utf-8"))
            contract = value if isinstance(value, dict) else None
        except (OSError, ValueError):
            contract = None
    if contract is None and state:
        contract = state.get("report_template_contract") or \
            (state.get("research_settings") or {}).get("report_template_contract")
    if not contract or not docx_path.is_file():
        return None
    raw = docx_path.read_bytes()
    return {
        "bytes": raw,
        "contract": contract,
        "metadata": state.get("report_template") if state else None,
        "summary": report_template.contract_summary(contract),
    }


def _bytes(value):
    return value.getvalue() if hasattr(value, "getvalue") else bytes(value)


def report_docx_bytes(job_id, state=None, frozen_assets=None):
    """Render the structured report with its template, or use the legacy fallback."""
    if frozen_assets and frozen_assets.get("stage3_report.docx"):
        return frozen_assets["stage3_report.docx"]
    state = state or load_job_state(job_id) or {}
    report = state.get("p3_md")
    if not report:
        return None
    template = load_report_template(job_id)
    if template:
        from transpraxis import academic_evidence, academic_writer, compliance
        from transpraxis import final_docx, report_template
        artifact = load_academic_artifact(job_id, "report")
        report_record = academic_writer.artifact_record(state, "report")
        if isinstance((state.get("academic_state") or {}).get("artifacts", {}).get(
                "report"), dict) and report_record.get("status") != "valid":
            return None
        if not artifact:
            raise report_template.TemplateParseError(
                "模板化报告缺少结构化 report artifact，请重新生成报告。")
        if artifact.get("report_status") not in {
                "generated", "review_required", "literature_required"} or artifact.get(
                "template_compliance") not in {"pass", "pass_with_warnings"}:
            return None
        rendered = _bytes(report_template.render_report_docx(
            artifact, template["bytes"], template["contract"]))
        source_docx_hash = _rendered_qa.sha256(rendered)
        previous_docx = academic_writer.artifact_record(
            state, "final_docx_validation")
        previous_qa = _finalization.normalize_final_qa(state.get("final_qa"))
        previous_docx_hash = str(
            previous_docx.get("source_docx_hash") or
            previous_qa.get("source_docx_hash") or "")
        if (previous_docx_hash and previous_docx_hash != source_docx_hash) or (
                isinstance(previous_docx, dict) and previous_docx and (
                    previous_qa.get("author_visual_review") == "CONFIRMED" or
                    previous_qa.get("word_final_review") == "CONFIRMED")):
            academic_writer.propagate_artifact_staleness(
                state, input_artifact_ids=["final_docx_validation"])
            _reset_final_qa(state, "DOCX bytes changed; render and human reviews reset")
        final_validation = final_docx.validate_final_docx(rendered, artifact)
        final_validation.update({
            "source_docx_hash": source_docx_hash,
            "report_content_hash": (state.get("academic_state") or {}).get(
                "artifacts", {}).get("report", {}).get("content_hash") or
                artifact.get("content_hash"),
            "layout_facts": compliance.inspect_docx_layout(rendered),
        })
        final_validation["content_hash"] = academic_evidence.stable_hash({
            key: value for key, value in final_validation.items()
            if key != "content_hash"
        })
        final_validation_dep = academic_evidence.stable_hash({
            "report": (state.get("academic_state") or {}).get("artifacts", {}).get(
                "report", {}).get("content_hash") or artifact.get("content_hash"),
            "template": template.get("contract", {}).get("template_identity") or
            template.get("contract", {}).get("template_hash"),
            "version": final_docx.SCHEMA_VERSION,
        })
        academic_writer._save_artifact(
            state, job_dir(job_id), "final_docx_validation", final_validation,
            final_validation_dep, final_docx.SCHEMA_VERSION,
            input_artifact_ids=["report"],
            status="valid" if final_validation.get("status") != "fail" else "failed")
        qa = _finalization.normalize_final_qa(state.get("final_qa"))
        qa.update({
            "structural_qa": "FAIL" if final_validation.get("status") == "fail"
            else "PASS",
            "source_docx_hash": source_docx_hash,
            "updated_at": _finalization.now_iso(),
        })
        state["final_qa"] = qa
        save_compliance_record(job_id, state)
        generate_report_qa(job_id, state, save_file=True)
        save_job_state(job_id, state)
        if final_validation.get("status") == "fail":
            return None
        return rendered
    if state.get("report_status") in {
            "incomplete", "failed_template_validation", "review_required"}:
        return None
    from transpraxis import academic_writer
    if isinstance((state.get("academic_state") or {}).get("artifacts", {}).get(
            "report"), dict) and academic_writer.artifact_record(
                state, "report").get("status") != "valid":
        return None
    return _bytes(markdown_to_word(report, state.get("theory") or ""))


def validate_translation_target(source, target, *, segment_index=None,
                                allow_json=False):
    """Public target invariant entry point used by runtime and delivery."""
    return _translation_target.validate_translation_target(
        source, target, segment_index=segment_index, allow_json=allow_json)


def validate_translation_pairs(pairs, glossary=None, target_lang=""):
    """Validate pairs independently of the model response parser."""
    pairs = list(pairs or [])
    target_report = _translation_target.validate_translation_pairs(pairs)
    qa_findings = check_translation_batch(
        [str(pair.get("source") or "") for pair in pairs if isinstance(pair, dict)],
        [str(pair.get("target") or "") for pair in pairs if isinstance(pair, dict)],
        glossary or [], target_lang or "简体中文")
    return {
        **target_report,
        "qa_findings": qa_findings,
        "blocking_findings": [
            finding for finding in qa_findings
            if finding.get("severity") == "blocking"
        ],
    }


def validate_delivery_translation_state(state):
    """Run the final, state-level translation gate before any delivery asset."""
    state = state if isinstance(state, dict) else {}
    pairs = state.get("pairs") or []
    source_quality_gate = _audit_source_quality_for_state(state)
    source_quality_findings = [
        _source_quality_finding(flag)
        for flag in source_quality_gate.get("flags") or []
        if isinstance(flag, dict) and isinstance(flag.get("segment_index"), int)
    ]
    report = validate_translation_pairs(
        pairs, state.get("glossary") or [], state.get("target_lang") or "简体中文")
    report["source_quality_gate"] = source_quality_gate
    report["source_quality_findings"] = source_quality_findings
    report["blocking_findings"] = list(report.get("blocking_findings") or [])
    report["blocking_findings"].extend(source_quality_findings)
    issues = list(report.get("issues") or [])
    if state.get("p2_done") and len(pairs) != len(state.get("paras") or []):
        issues.append({
            "code": "pair_count_mismatch",
            "message": "双语 pairs 数量与源文段落数量不一致，不能生成完整交付物",
            "severity": "blocking",
        })
    entity_findings = _entity_registry.EntityRegistry(
        state.get("entity_registry") or []
    ).consistency_findings()
    return {
        **report,
        "issues": issues,
        "entity_findings": entity_findings,
        "blocking": bool(issues or report.get("blocking_findings")),
        "status": "fail" if issues or report.get("blocking_findings") else (
            "review_required" if entity_findings else "pass"),
    }


def _record_delivery_validation_findings(state, report):
    """Persist only new invariant/entity findings for the review queue."""
    state["source_quality_gate"] = dict(report.get("source_quality_gate") or {})
    state["delivery_validation"] = {
        "status": report.get("status"),
        "blocking": bool(report.get("blocking")),
        "checked_pairs": report.get("checked_pairs", 0),
        "issues": list(report.get("issues") or []),
        "entity_conflicts": [
            {
                "source": item.get("source"),
                "preferred_target": item.get("preferred_target"),
                "observed_targets": item.get("observed_targets"),
            }
            for item in report.get("entity_findings") or []
        ],
    }
    existing = {
        (item.get("type"), item.get("segment_index"), item.get("invariant_code"),
         item.get("reason"))
        for item in state.setdefault("findings", [])
        if isinstance(item, dict) and not item.get("resolved")
    }
    findings = _translation_target.target_invariant_findings(report)
    findings.extend(
        dict(item) for item in report.get("source_quality_findings") or []
        if isinstance(item, dict)
    )
    for item in report.get("entity_findings") or []:
        findings.append({
            **item,
            "segment_index": item.get("segment_index"),
            "segment_id": item.get("segment_id"),
            "summary": item.get("reason"),
            "explanation": item.get("reason"),
            "recommendation": "核对全文同一实体的译名，并在人工作区确认一个一致形式。",
            "confidence": None,
            "diagnostic_version": 1,
        })
    for finding in findings:
        key = (finding.get("type"), finding.get("segment_index"),
               finding.get("invariant_code"), finding.get("reason"))
        if key not in existing:
            state["findings"].append(finding)
            existing.add(key)
    stats = state.setdefault("review_stats", {})
    stats["blocking"] = sum(
        1 for item in state["findings"]
        if item.get("severity") == "blocking" and not item.get("resolved")
    )
    stats["actionable"] = sum(
        1 for item in state["findings"]
        if item.get("severity") == "actionable" and not item.get("resolved")
    )
    stats["informational"] = sum(
        1 for item in state["findings"]
        if item.get("severity") == "informational" and not item.get("resolved")
    )
    state["has_blocking"] = stats["blocking"] > 0
    if report.get("blocking") and state.get("delivery_status") in ("approved", "final"):
        # A state-level mutation (for example a manually injected target) must
        # invalidate the mutable approval even when no snapshot exists yet.
        _invalidate_final_delivery_state(state)
    return state


def _academic_workspace_archive(job_id):
    from transpraxis import academic_writer

    names = ("research_model", "argument_plan", "selected_cases", "outline",
             "case_analysis_plans", "human_evidence_questions")
    output = io.BytesIO()
    written = 0
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as bundle:
        for name in names:
            filename = academic_writer.ARTIFACT_FILES[name]
            path = job_dir(job_id) / filename
            if path.is_file():
                bundle.writestr(filename, path.read_bytes())
                written += 1
    return output.getvalue() if written else None


def _delivery_asset_bundle(job_id, state, target_lang, provider, model):
    """Build the configured delivery set used by both preview and snapshots."""
    from transpraxis import assets as _assets
    from transpraxis import report_evidence as _report_evidence

    validation = validate_delivery_translation_state(state)
    _record_delivery_validation_findings(state, validation)
    if validation["blocking"]:
        save_job_state(job_id, state)
        reasons = "；".join(
            issue.get("message", "目标文本未通过交付检查")
            for issue in validation.get("issues") or []
        )
        raise RuntimeError(
            "交付被 Translation Target Invariant 阻止"
            + (f"：{reasons}" if reasons else "")
        )

    source = load_source(job_id)
    snapshot_state = dict(state)
    if source is not None:
        snapshot_state["_source_bin"] = source
    filename = state.get("filename", "")
    configured = state.get("delivery_config")
    if isinstance(configured, dict) and configured:
        config = normalize_delivery_config(
            configured, enable_report=state.get("report_enabled", False),
            enable_annotate=state.get("enable_annotate", False))
        bundle = {}
        pairs = state.get("pairs") or []
        glossary = state.get("glossary") or []
        if state.get("source_cleanup_done"):
            bundle["source_cleanup.json"] = (
                json.dumps(state.get("source_cleanup") or {},
                           ensure_ascii=False, indent=2) + "\n"
            ).encode("utf-8")
        if state.get("p2_done") and pairs:
            if config["deliver_plain_docx"]:
                bundle["translation.docx"] = _bytes(translations_to_word(pairs))
            if config["deliver_bilingual_docx"]:
                bundle["bilingual.docx"] = _bytes(pairs_to_word(pairs))
            if config["deliver_pdf"]:
                bundle["translation.pdf"] = translations_to_pdf(pairs)
            if config["enable_annotate"]:
                bundle["annotated_bilingual.docx"] = _bytes(pairs_to_word(
                    pairs, annotations=state.get("annotations"),
                    colors=ANNOTATION_COLORS))
            if config["deliver_terms_xlsx"]:
                bundle["terms.xlsx"] = _bytes(glossary_to_excel(
                    glossary, state.get("auto_terms")))
            if config["deliver_tbx"]:
                bundle["terms.tbx"] = _assets.build_tbx(glossary)
            if config["deliver_tmx"]:
                bundle["memory.tmx"] = _assets.build_tmx(state, job_id=job_id)
            if config["deliver_jsonl"]:
                bundle["bilingual.jsonl"] = _assets.build_jsonl(
                    state, job_id=job_id).encode("utf-8")
            if config["deliver_evidence"]:
                bundle["segment_evidence.jsonl"] = \
                    _report_evidence.export_segment_evidence_jsonl(
                        state, job_id).encode("utf-8")
            if config["deliver_review_report"]:
                bundle["review_report.md"] = findings_report_md(state).encode("utf-8")
        # 被排除的内容不进译文文档，但必须在交付包里留下痕迹：否则"我排除了
        # 12 段"会变成一次静默删除，用户在交付物上再也看不到它。
        excluded_manifest = excluded_segments_manifest(state)
        if excluded_manifest["count"]:
            bundle["excluded_segments.md"] = _excluded_segments_md(
                excluded_manifest).encode("utf-8")
            bundle["excluded_segments.json"] = (
                json.dumps(excluded_manifest, ensure_ascii=False, indent=2)
                + "\n").encode("utf-8")
        if config["deliver_cases"]:
            selected_cases = load_academic_artifact(job_id, "selected_cases")
            if selected_cases:
                bundle["selected_cases.json"] = (
                    json.dumps(selected_cases, ensure_ascii=False, indent=2) + "\n"
                ).encode("utf-8")
        if config["deliver_academic_workspace"]:
            workspace = _academic_workspace_archive(job_id)
            if workspace:
                bundle["academic_workspace.zip"] = workspace
        if config["enable_report"] and state.get("p3_md"):
            report_docx = report_docx_bytes(job_id, state)
            if report_docx is not None:
                bundle["report.docx"] = report_docx
            bundle["report.md"] = state["p3_md"].encode("utf-8")
        generated_assets = sorted([*bundle, "delivery_manifest.json"])
        manifest = _assets.build_delivery_manifest(
            snapshot_state, job_id, target_lang, provider, model,
            generated_assets=generated_assets, source_filename=filename,
            translator_config=state.get("translator_config"),
            reviewer_config=state.get("reviewer_config"))
        bundle["delivery_manifest.json"] = (
            json.dumps(manifest, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
        return bundle

    # Existing jobs without a saved selection retain their historical bundle.
    bundle = {}
    if state.get("p1_done") and state.get("paras"):
        # ``paras`` is the sentence-sized CAT contract; the stage-1 source
        # artifact should retain the recovered paragraph layout when available.
        cleaned_source = state.get("source_paragraphs") or state.get("paras")
        bundle["stage1_cleaned.docx"] = _bytes(paragraphs_to_word(cleaned_source))
    if state.get("source_cleanup_done"):
        bundle["source_cleanup.json"] = (
            json.dumps(state.get("source_cleanup") or {},
                       ensure_ascii=False, indent=2) + "\n"
        ).encode("utf-8")
    if state.get("auto_terms"):
        bundle["auto_terms.xlsx"] = _bytes(dict_to_excel(state["auto_terms"]))
    if state.get("p2_done") and state.get("pairs"):
        bundle["stage2_bilingual.docx"] = _bytes(pairs_to_word(
            state["pairs"], annotations=state.get("annotations"), colors=ANNOTATION_COLORS))
    if state.get("p3_md"):
        report_docx = report_docx_bytes(job_id, state)
        if report_docx is not None:
            bundle["stage3_report.docx"] = report_docx
    bundle.update(_assets.export_all(
        snapshot_state, job_id, target_lang, provider, model,
        source_filename=filename, source_bin=source))
    # 旧任务没有保存过 delivery_config，但排除段仍然必须在交付包中
    # 留下可追溯清单；否则历史任务会把排除误表现成静默删除。
    excluded_manifest = excluded_segments_manifest(state)
    if excluded_manifest["count"]:
        bundle["excluded_segments.md"] = _excluded_segments_md(
            excluded_manifest).encode("utf-8")
        bundle["excluded_segments.json"] = (
            json.dumps(excluded_manifest, ensure_ascii=False, indent=2)
            + "\n").encode("utf-8")
    if state.get("p2_done"):
        bundle["segment_evidence.jsonl"] = _report_evidence.export_segment_evidence_jsonl(
            state, job_id).encode("utf-8")
        if state.get("findings"):
            bundle["review_report.md"] = findings_report_md(state).encode("utf-8")
    manifest = _assets.build_delivery_manifest(
        snapshot_state, job_id, target_lang, provider, model,
        generated_assets=sorted(bundle), source_filename=filename,
        translator_config=state.get("translator_config"),
        reviewer_config=state.get("reviewer_config"))
    bundle["delivery_manifest.json"] = (
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    return bundle


def build_delivery_assets(job_id, state=None, target_lang="", provider="", model=""):
    state = state or load_job_state(job_id) or {}
    return _delivery_asset_bundle(
        job_id, state, target_lang or state.get("target_lang") or "",
        provider or state.get("provider") or "",
        model or state.get("model") or "")


def _reset_source_dependent_state(state):
    """Invalidate artifacts derived from source paragraphs after cleanup."""
    state.update({
        "profile_done": False,
        "document_profile": None,
        "profile_warnings": [],
        "semantic_units": [],
        "section_digests": [],
        "document_synopsis": None,
        "understanding_done": False,
        "understanding_warnings": [],
        "context_packet_log": [],
        "llm_usage": _usage.empty_usage(),
        "knowledge_feedback_policy": {},
        "targeted_final_review": {
            "status": "not_run", "segment_ids": [],
            "reviewed_segment_ids": [], "failed_segment_ids": [],
        },
        "knowledge_candidates": [],
        "translation_continuity": [],
        "knowledge_events": [],
        "confirmed_style_rules": [],
        "knowledge_feedback_failures": 0,
        "entity_registry": [],
        "translation_failures": [],
        "review_evidence": [],
        "repair_overlays": [],
        "auto_term_entries": [],
        "auto_terms": {},
        "glossary": [],
        "glossary_draft": [],
        "glossary_frozen": None,
        "glossary_versions": [],
        "glossary_injection_log": [],
        "pairs": [],
        "findings": [],
        "has_blocking": False,
        "review_stats": {
            "reviewed_segments": 0, "batches_reviewed": 0,
            "blocking": 0, "actionable": 0,
            "informational": 0, "review_failed": 0,
        },
        "p2_done": False,
        "p3_done": False,
        "p3_md": "",
        "p3_sections": [],
        "annotations": {},
        "annotations_done": False,
        "annotations_done_offset": 0,
        "batch_plan": {},
        "delivery_status": "draft",
        "delivery_approved_by_human": False,
        "delivery_approval": None,
        "delivery_validation": {},
        "delivery_snapshots": [],
        "latest_delivery_snapshot_version": None,
        "exported_assets": [],
        "translation_truth": {
            "authority": "CURRENT_TRANSLATION",
            "version": 0,
            "last_changed_at": None,
            "last_change": None,
        },
        "source_quality_gate": {},
    })


def rescan_pdf_source(job_id, *, filename=None, provider, api_key, model,
                      base_url=None, on_status=None, max_batch_chars=None,
                      max_batch_items=None, parallelism=None, ocr_workers=None,
                      ocr_queue_size=None, confidence_threshold=None,
                      max_retries=None, request_interval_seconds=None,
                      segmentation_mode=None,
                      _allow_active=False):
    """Re-extract and LLM-clean a saved PDF without running later stages.

    This is the explicit "重新从 PDF 扫一遍" operation.  It deliberately
    invalidates every artifact derived from the previous source text, including
    partial translation pairs, so the next resume cannot pair old translations
    with newly merged or corrected source paragraphs.
    """
    if not str(job_id or "").strip():
        raise ValueError("缺少任务 ID")
    if not _allow_active and (is_job_worker_alive(job_id) or job_is_active(job_id)):
        raise RuntimeError("任务正在运行，无法重新扫描原文")
    state = load_job_state(job_id)
    if state is None:
        raise ValueError(f"找不到任务 {job_id}")
    resolved_filename = str(filename or state.get("filename") or "")
    if not resolved_filename.lower().endswith(".pdf"):
        raise ValueError("原文重新扫描只支持 PDF 任务")
    source = load_source(job_id)
    if source is None:
        raise ValueError("任务没有已保存的 PDF 源文件，请重新上传后再试")

    previous_pairs = len(state.get("pairs") or [])
    previous_paragraphs = len(state.get("paras") or [])
    old_base_url = getattr(_LLM_CTX, "base_url", None)
    effective_cleanup_base_url = base_url if base_url is not None else old_base_url
    saved_config = state.get("pipeline_config") or {}
    effective_parallelism = _runtime_int_option(
        parallelism, saved_config.get("source_cleanup_concurrency"),
        "FOLIOTHREAD_LLM_CLEANUP_CONCURRENCY",
        _source_cleanup.DEFAULT_PARALLELISM, maximum=16)
    effective_ocr_workers = _runtime_int_option(
        ocr_workers, saved_config.get("ocr_workers"), "FOLIOTHREAD_OCR_WORKERS",
        OCR_WORKERS_DEFAULT, maximum=8)
    effective_ocr_queue_size = _runtime_int_option(
        ocr_queue_size, saved_config.get("ocr_queue_size"),
        "FOLIOTHREAD_OCR_QUEUE_SIZE", effective_ocr_workers * 2, maximum=32)
    effective_confidence_threshold = _runtime_float_option(
        confidence_threshold, saved_config.get("source_cleanup_confidence_threshold"),
        "FOLIOTHREAD_OCR_CONFIDENCE_THRESHOLD",
        OCR_CONFIDENCE_THRESHOLD_DEFAULT, maximum=100.0)
    effective_max_retries = _runtime_int_option(
        max_retries, saved_config.get("source_cleanup_max_retries"),
        "FOLIOTHREAD_LLM_CLEANUP_MAX_RETRIES",
        _source_cleanup.DEFAULT_MAX_RETRIES, minimum=0, maximum=6)
    effective_request_interval = _runtime_float_option(
        request_interval_seconds, saved_config.get("source_cleanup_request_interval"),
        "FOLIOTHREAD_LLM_CLEANUP_REQUEST_INTERVAL", 0.0, maximum=60.0)
    profiler = PipelineProfiler(
        job_dir(job_id) / "performance.json", run_id=uuid.uuid4().hex
    )

    def report(label):
        if on_status:
            on_status(label)
        # A direct invocation has no worker thread, so mirror the same durable
        # progress surface that the normal worker provides.
        _runtime_status_callback(job_id, state, label)

    update_runtime_state(
        job_id, status="running", phase="source_cleanup",
        stage="document_processing", stage_id="document_processing",
        operation="source_cleanup", operation_id="source_cleanup",
        operation_label="重新扫描 PDF 原文", phase_label="正在扫描并纠错原文",
        cancel_requested=False, error=None, progress=True,
        event="开始重新扫描 PDF 原文", event_name="source_rescan_started",
        event_visibility="user", event_category="progress")
    set_llm_base_url(effective_cleanup_base_url)
    try:
        report("【阶段一】重新从 PDF 提取原文…")
        raw_paragraphs, extraction_warnings, extraction_report = \
            extract_document_paragraphs_with_report(
                resolved_filename, source, ocr_max_pages=None, on_progress=report,
                profiler=profiler, checkpoint_dir=job_dir(job_id),
                ocr_workers=effective_ocr_workers,
                ocr_queue_size=effective_ocr_queue_size,
                cancel_check=lambda: _runtime_cancel_requested(job_id))
        if not raw_paragraphs:
            detail = extraction_warnings[-1] if extraction_warnings else "未知提取错误"
            raise ValueError(f"未提取到有效文本：{detail}")

        report("【阶段一】使用最新纠错功能整理 OCR 错字与断行…")
        inherited_job_id = getattr(_RUNTIME_CTX, "job_id", None)
        inherited_reasoning = getattr(_LLM_CTX, "reasoning_effort", None)

        def cleanup_call(provider_name, api_key_value, model_name,
                         system_prompt, user_prompt, **kwargs):
            if inherited_job_id:
                _RUNTIME_CTX.job_id = inherited_job_id
            inherited_kwargs = {}
            if effective_cleanup_base_url:
                inherited_kwargs["base_url"] = effective_cleanup_base_url
            if inherited_reasoning:
                inherited_kwargs["reasoning_effort"] = inherited_reasoning
            return call_llm(
                provider_name, api_key_value, model_name, system_prompt,
                user_prompt, **inherited_kwargs, **kwargs)
        cleaned, cleanup_meta, cleanup_warnings = cleanup_source_paragraphs(
            raw_paragraphs, provider, api_key, model, call_llm_fn=cleanup_call,
            on_progress=report, max_batch_chars=max_batch_chars,
            max_batch_items=max_batch_items, parallelism=effective_parallelism,
            confidence_by_index=((extraction_report.get("ocr") or {}).get(
                "paragraph_confidences") if isinstance(extraction_report, dict) else None),
            confidence_threshold=effective_confidence_threshold,
            checkpoint_dir=job_dir(job_id),
            cancel_check=lambda: _runtime_cancel_requested(job_id),
            max_retries=effective_max_retries,
            request_interval_seconds=effective_request_interval,
            profiler=profiler)
        cleanup_meta = dict(cleanup_meta or {})
        segments, segmentation_meta = _apply_source_segmentation(
            state, cleaned, mode=segmentation_mode)
        cleanup_meta.update({
            "forced_rescan": True,
            "previous_pair_count": previous_pairs,
            "previous_paragraph_count": previous_paragraphs,
            "extraction_warning_count": len(extraction_warnings),
            "parallelism": effective_parallelism,
            "confidence_threshold": effective_confidence_threshold,
            "ocr_workers": effective_ocr_workers,
            "ocr_queue_size": effective_ocr_queue_size,
            "max_batch_chars": max_batch_chars,
            "max_batch_items": max_batch_items,
            "paragraph_count": len(cleaned),
            "segment_count": len(segments),
        })

        warnings = state.setdefault("warnings", [])
        for warning in extraction_warnings:
            if warning not in warnings:
                warnings.append(warning)
            _append_runtime_technical_log(job_id, f"document extraction: {warning}")
        for warning in cleanup_warnings:
            if warning not in warnings:
                warnings.append(warning)
            _append_runtime_technical_log(job_id, f"source cleanup: {warning}")

        _reset_source_dependent_state(state)
        state.update({
            "filename": resolved_filename,
            "source_paras": list(raw_paragraphs),
            "paras": list(segments),
            "source_cleanup": cleanup_meta,
            "source_cleanup_done": True,
            "segmentation": segmentation_meta,
            "extraction_report": extraction_report,
            "source_quality_gate": audit_source_quality(segments),
            "p1_done": True,
            "source_page_count": 0,
        })
        try:
            with fitz.open(stream=source, filetype="pdf") as source_pdf:
                state["source_page_count"] = source_pdf.page_count
        except Exception as exc:  # noqa: BLE001 - extraction already succeeded
            _append_runtime_technical_log(job_id, f"source page count: {exc}")
        pipeline_config = state.get("pipeline_config")
        if isinstance(pipeline_config, dict):
            pipeline_config["enable_source_cleanup"] = True
        save_source(job_id, source)
        state["performance"] = profiler.persist()
        save_job_state(job_id, state)
        update_runtime_state(
            job_id, status="idle_incomplete", phase="idle_incomplete",
            phase_label="原文纠错已完成，等待继续翻译", operation="source_cleanup",
            operation_id="source_cleanup", operation_label="原文纠错已完成",
            error=None, progress=True, heartbeat=True,
            event=(f"原文重新扫描完成：{len(raw_paragraphs)} 项提取为 "
                   f"{len(cleaned)} 段、{len(segments)} 个翻译单元，修复 "
                   f"{len(cleanup_meta.get('changed_segments') or [])} 项"),
            event_name="source_rescan_completed", event_visibility="user",
            event_category="progress",
            worker={"owner_pid": None, "worker_id": None,
                    "lease_expires_at": None})
        return state
    except BaseException as exc:
        _write_runtime_technical_log(job_id, exc)
        provider_status = provider_error_status(exc)
        user_error_message = (
            provider_error_message(exc, "原文重新扫描失败")
            if provider_status["status"] != "unknown"
            else str(exc)[:500] or "原文重新扫描失败")
        update_runtime_state(
            job_id, status="failed", phase="failed", phase_label="原文重新扫描失败",
            error={"type": type(exc).__name__, "message": user_error_message,
                   "stage": "document_processing", "operation": "source_cleanup",
                   "timestamp": _utc_now_iso(),
                   "technical_log": RUNTIME_TECHNICAL_LOG},
            event=f"原文重新扫描失败：{user_error_message[:180]}",
            event_name="source_rescan_failed", event_visibility="user",
            event_category="error", progress=True, heartbeat=True,
            worker={"owner_pid": None, "worker_id": None,
                    "lease_expires_at": None})
        raise
    finally:
        profiler.persist()
        set_llm_base_url(old_base_url)


def create_delivery_snapshot(job_id, state, target_lang="", provider="", model=""):
    validation = validate_delivery_translation_state(state)
    _record_delivery_validation_findings(state, validation)
    if validation["blocking"]:
        save_job_state(job_id, state)
        raise RuntimeError("最终交付未通过 Translation Target Invariant")
    source = load_source(job_id)
    frozen = state.get("glossary_frozen") or {}
    source_hash = hashlib.sha256(source).hexdigest() if source else \
        str(frozen.get("source_hash") or "")
    manifest = _snapshots.create_snapshot(
        job_dir(job_id), job_id, state,
        _delivery_asset_bundle(job_id, state, target_lang, provider, model),
        source_identity={
            "filename": state.get("filename", ""),
            "source_hash": source_hash,
            "source_page_count": state.get("source_page_count"),
        },
        active_terminology_version={
            "version": frozen.get("version"),
            "glossary_hash": frozen.get("glossary_hash"),
            "status": "已冻结" if frozen else "未冻结",
        })
    return manifest


def list_delivery_snapshots(job_id):
    return _snapshots.list_snapshots(job_dir(job_id))


def delivery_snapshot_status(job_id, state=None):
    state = state if state is not None else load_job_state(job_id)
    latest = _snapshots.latest_snapshot(job_dir(job_id))
    if latest is None:
        return {"latest": None, "current": False, "diverged": False, "integrity": False}
    current_identity = _snapshots.state_identity(state or {})
    matches = current_identity == latest.get("translation_state_identity")
    assets = delivery_snapshot_assets(job_id, latest["snapshot_version"])
    integrity = bool(assets) and all(data is not None for data in assets.values())
    return {
        "latest": latest,
        "current": bool(matches and integrity and state
                        and state.get("delivery_status") == "final"),
        "diverged": not matches,
        "integrity": integrity,
    }


def delivery_snapshot_assets(job_id, version):
    manifest = next((item for item in list_delivery_snapshots(job_id)
                     if item.get("snapshot_version") == int(version)), None)
    if manifest is None:
        return {}
    return {
        item["name"]: _snapshots.load_asset(job_dir(job_id), version, item["name"])
        for item in manifest.get("assets") or []
    }


def delivery_snapshot_archive(job_id, version):
    return _snapshots.archive(job_dir(job_id), version)


def progress_label(state):
    academic = state.get("academic_state") or {}
    if academic.get("status") == "failed":
        return "翻译完成 · 学术写作失败"
    if academic.get("status") in ("in_progress", "stale"):
        return "翻译完成 · 学术写作中"
    if academic.get("quality_status") == "review_required":
        return "翻译完成 · 报告待学术复核"
    if academic.get("quality_status") == "fail":
        return "翻译完成 · 报告验证失败"
    if state.get("p1_done") and state.get("p2_done") and \
            (state.get("p3_done") or not state.get("report_enabled", True)):
        return "已完成"
    if state.get("p2_done"):
        return "报告生成中"
    if state.get("p1_done"):
        return "翻译中"
    return "待处理"


def recovery_summary(job_id, state=None):
    """Read-only durable progress summary used by History and the workspace."""
    state = state if state is not None else load_job_state(job_id)
    return _checkpoint.recovery_summary(job_dir(job_id), state or {})


def load_context_artifacts(job_id, state=None):
    """Load persisted context artifacts, tolerating older or partial jobs."""
    state = state if state is not None else load_job_state(job_id)
    state = state or {}

    def read_json(name, default):
        path = job_dir(job_id) / name
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except (OSError, TypeError, ValueError):
            return default

    units = state.get("semantic_units") or read_json("semantic_units.json", [])
    digests = state.get("section_digests") or read_json("section_digests.json", [])
    synopsis = state.get("document_synopsis")
    stored_synopsis = read_json("document_synopsis.json", None)
    if not isinstance(synopsis, dict):
        synopsis = {}
    if not synopsis or (synopsis.get("status") == "pending"
                        and isinstance(stored_synopsis, dict)
                        and stored_synopsis.get("status") != "pending"):
        synopsis = stored_synopsis or {}
    return {
        "semantic_units": units if isinstance(units, list) else [],
        "section_digests": digests if isinstance(digests, list) else [],
        "document_synopsis": synopsis if isinstance(synopsis, dict) else {},
        "warnings": list(state.get("understanding_warnings") or []),
    }


def task_status_label(state, job_id=""):
    """User-facing task status derived from persisted workflow state."""
    from transpraxis import delivery as _delivery

    if job_id:
        snapshot = delivery_snapshot_status(job_id, state)
        latest = snapshot.get("latest") or {}
        version = latest.get("snapshot_version")
        if snapshot.get("current"):
            return (f"已冻结交付 v{version}" if version is not None else
                    "已冻结交付")
        if snapshot.get("diverged"):
            return (f"工作版本已偏离冻结交付 v{version}" if version is not None else
                    "工作版本已偏离冻结交付")

    if state.get("delivery_status") == "final":
        return "已冻结交付"
    if state.get("p2_done") and _delivery.unresolved_blocking(state):
        return "待审校"
    academic = state.get("academic_state") or {}
    if state.get("p2_done") and state.get("p3_done") and (
            academic.get("quality_status") in ("review_required", "fail", "failed")
            or academic.get("status") == "failed"):
        return "待学术复核"
    if state.get("p2_done") and (
            state.get("p3_done") or not state.get("report_enabled", True)):
        if state.get("report_enabled", True):
            final_qa = state.get("final_qa") or {}
            report_status = state.get("report_status") or academic.get("report_status")
            if (report_status != "generated" or
                    final_qa.get("author_visual_review") != "CONFIRMED" or
                    final_qa.get("word_final_review") != "CONFIRMED"):
                return "暂不满足交付条件"
        return "可以冻结交付"
    if state.get("p1_done") and not state.get("p2_done"):
        if (state.get("stage") == "TERMS_PREPARED" and state.get("quality_mode")
                and state.get("glossary") is not None
                and not state.get("glossary_frozen")
                and not state.get("quality_bypass")):
            return "待术语确认"
        return "处理中断"
    if state.get("p2_done") and not state.get("p3_done"):
        return "处理中断"
    return "待处理"


# ================= 术语审核状态（草稿 / 锁定 / 拒绝 / 冻结）=================
def save_glossary_draft(job_id, entries):
    """保存术语审核草稿（不冻结）。刷新/重启后从 TERMS_PREPARED 恢复。"""
    state = load_job_state(job_id)
    if state is None:
        return None
    norm = normalize_glossary(entries)
    state["glossary_draft"] = norm
    state["glossary"] = norm
    state["stage"] = "TERMS_PREPARED"
    save_job_state(job_id, state)
    return state


def _glossary_mutation_note(action, detail, actor):
    return {
        "action": action,
        "finding_id": detail.get("finding_id") or "glossary",
        "note": detail.get("note") or "",
        "timestamp": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "actor": str(actor or "user"),
        "actor_type": "human",
        "previous_status": detail.get("previous_status"),
        "new_status": detail.get("new_status"),
    }


def update_glossary_entry(job_id, entry_id, *, preferred=None, domain=None,
                          status=None, note=None, actor="user"):
    """原地修改一条术语条目并生成新的术语版本。返回 (state, ok, message)。

    与 `save_glossary_draft` 的区别：后者是「翻译前准备术语」，会把任务阶段改回
    TERMS_PREPARED。语言资产管理中心对**已完成任务**做一次术语修订，不应该把
    任务打回准备阶段，因此这里直接改条目、记录人工动作，然后走既有的
    `freeze_glossary` —— 术语决策变化必须产生新版本并失效受影响段落，这条
    不变量不能被绕过。
    """
    state = load_job_state(job_id)
    if state is None:
        return None, False, "任务不存在"
    entries = normalize_glossary(state.get("glossary") or [])
    target = next((item for item in entries
                   if str(item.get("id")) == str(entry_id)), None)
    if target is None:
        return state, False, "找不到该术语"
    before = {key: target.get(key)
              for key in ("preferred", "target", "domain", "status", "note")}
    if preferred is not None:
        text = str(preferred or "").strip()
        if not text:
            return state, False, "推荐译法不能为空"
        target["preferred"] = text
        target["target"] = text
        target["proposed_target"] = text
    if domain is not None:
        target["domain"] = str(domain or "").strip()
    if status is not None:
        if status not in _models.STATUSES:
            return state, False, f"非法术语状态：{status}"
        target["status"] = status
    if note is not None:
        target["note"] = str(note or "").strip()
    after = {key: target.get(key) for key in before}
    if before == after:
        return state, True, "没有需要保存的改动"
    entries = normalize_glossary(entries)
    state["glossary"] = entries
    state["glossary_draft"] = entries
    state.setdefault("human_actions", []).append(_glossary_mutation_note(
        "glossary_entry_updated",
        {"finding_id": f"glossary:{entry_id}",
         "note": "在语言资产管理中心修改术语条目",
         "previous_status": before.get("status"),
         "new_status": after.get("status")}, actor))
    save_job_state(job_id, state)
    frozen = freeze_glossary(job_id, entries=entries, frozen_by=actor)
    version = (frozen or {}).get("glossary_frozen", {}).get("version") \
        if isinstance(frozen, dict) else None
    suffix = f"，术语版本已更新为 v{version}" if version else ""
    return frozen or state, True, f"已更新术语并生成新的术语版本{suffix}"


def delete_glossary_entry(job_id, entry_id, actor="user"):
    """删除一条术语条目并生成新的术语版本。返回 (state, ok, message)。"""
    state = load_job_state(job_id)
    if state is None:
        return None, False, "任务不存在"
    entries = normalize_glossary(state.get("glossary") or [])
    kept = [item for item in entries if str(item.get("id")) != str(entry_id)]
    if len(kept) == len(entries):
        return state, False, "找不到该术语"
    state["glossary"] = kept
    state["glossary_draft"] = kept
    state.setdefault("human_actions", []).append(_glossary_mutation_note(
        "glossary_entry_deleted",
        {"finding_id": f"glossary:{entry_id}",
         "note": "在语言资产管理中心删除术语条目"}, actor))
    save_job_state(job_id, state)
    frozen = freeze_glossary(job_id, entries=kept, frozen_by=actor)
    return frozen or state, True, "已删除该术语并生成新的术语版本"


def add_glossary_entry(job_id, source, target, *, domain="", scope="document",
                       status="locked", actor="user"):
    """在指定任务中新增一条术语并生成新的术语版本。

    返回 (state, ok, message, entry_id)。长期保存位置只有「项目术语」——
    术语写在任务的术语表里并通过冻结版本生效；当前没有独立的全局术语库。
    """
    source = str(source or "").strip()
    target = str(target or "").strip()
    if not source:
        return None, False, "原术语不能为空", ""
    if not target:
        return None, False, "推荐译法不能为空", ""
    if status not in _models.STATUSES:
        return None, False, f"非法术语状态：{status}", ""
    state = load_job_state(job_id)
    if state is None:
        return None, False, "任务不存在", ""
    entries = normalize_glossary(state.get("glossary") or [])
    existing = next((item for item in entries
                     if str(item.get("source") or "").casefold()
                     == source.casefold()), None)
    if existing is not None:
        current = str(existing.get("preferred") or existing.get("target") or "").strip()
        if current.casefold() != target.casefold():
            return state, False, f"该术语已存在，推荐译法为「{current}」；请先编辑原条目", \
                str(existing.get("id") or "")
        return state, False, "该术语已存在", str(existing.get("id") or "")
    entry = _models.normalize_glossary_entry({
        "source": source, "target": target, "preferred": target,
        "behavior": "translate", "status": status, "scope": scope,
        "domain": str(domain or "").strip(),
        "note": "由人工在语言资产管理中心新增",
        "evidence": [{
            "evidence_type": "user", "source_name": "语言资产管理中心",
            "note": "人工新增术语", "quote": "", "url": "",
        }],
    })
    if entry is None:
        return state, False, "术语内容无效", ""
    entries.append(entry)
    entries = normalize_glossary(entries)
    state["glossary"] = entries
    state["glossary_draft"] = entries
    state.setdefault("human_actions", []).append(_glossary_mutation_note(
        "glossary_entry_added",
        {"finding_id": f"glossary:{entry.get('id')}",
         "note": f"在语言资产管理中心新增术语「{source}」"}, actor))
    save_job_state(job_id, state)
    frozen = freeze_glossary(job_id, entries=entries, frozen_by=actor)
    return frozen or state, True, "已新增术语并生成新的术语版本", str(entry.get("id") or "")


def _apply_glossary_staleness(state, job_id=None):
    """把受冻结术语表变更影响的段落标记 stale，并清除其 TM 信任。

    权威集合：每次调用都重新计算（先清旧标记/旧 stale finding，
    再标记当前受影响段）。stale 段：
    - stale_due_to_glossary = True，reviewed = False，from_tm = False；
    - 追加 blocking finding（type=glossary_stale），交付回到 review_required；
    - 从 translation_memory.json 中清除，防止后续任务精确命中旧译文。
    """
    from transpraxis import delivery as _delivery
    from transpraxis.terminology import stale_segments_for_glossary

    stale = stale_segments_for_glossary(state)
    pairs = state.get("pairs") or []
    current_glossary = normalize_glossary(
        (state.get("glossary_frozen") or {}).get("entries")
        or state.get("glossary") or [])
    current_glossary_hash = _models.glossary_hash(current_glossary)
    stale_event_ids = []
    for event in state.get("review_evidence") or []:
        if not isinstance(event, dict) or not (
                event.get("review_scope") == "current_translation"
                or event.get("phase") == "formal_review"):
            continue
        event_hash = str((event.get("translation_core") or {}).get(
            "glossary_hash") or event.get("glossary_hash") or "")
        if event_hash and event_hash != current_glossary_hash \
                and event.get("review_event_id"):
            stale_event_ids.append(str(event["review_event_id"]))
    if stale or stale_event_ids:
        _invalidate_translation_reviews(
            state, stale,
            "canonical glossary changed; independent review must be refreshed",
            review_event_ids=stale_event_ids)
    for p in pairs:
        p.pop("stale_due_to_glossary", None)
    state["findings"] = [f for f in state.get("findings") or []
                         if f.get("type") != "glossary_stale"]
    if not stale:
        return state, []

    state["knowledge_candidates"] = _knowledge.discard_candidates_for_segments(
        state.get("knowledge_candidates") or [], stale)
    state["translation_continuity"] = list(state["knowledge_candidates"])

    fg = state.get("glossary_frozen") or {}
    for i in stale:
        p = pairs[i]
        p["stale_due_to_glossary"] = True
        p["reviewed"] = False
        p["from_tm"] = False
        p["review_status"] = "not_reviewed"
        for key in ("accepted_target", "human_accepted", "accepted_by_human",
                    "target_provenance"):
            p.pop(key, None)
        state["findings"].append({
            "segment_id": i, "segment_index": i,
            "severity": "blocking",
            "type": "glossary_stale",
            "entry_id": None,
            "category": "terminology_consistency",
            "summary": "冻结术语表已变更，本段需要重新确认",
            "source_span": None, "target_span": None,
            "explanation": f"冻结术语表已从当前版本变更为 v{fg.get('version')}，本段原有译文可能不再符合最新术语规则。",
            "recommendation": "检查本段涉及的术语；确认译法后标记已处理，或使用最新术语重新翻译。",
            "confidence": None, "detector": "Terminology QA",
            "diagnostic_version": 1,
            "reason": f"冻结术语表已变更（当前 v{fg.get('version')}），本段需复核或重译",
        })

    # 受影响段不得继续作为可信翻译记忆
    tm = load_tm(tm_project_id(state))
    dirty = False
    job_lang = state_target_lang(state)
    for i in stale:
        src = pairs[i]["source"]
        if tm_discard(tm, src, job_lang):
            dirty = True
    if dirty:
        save_tm(tm, tm_project_id(state))
    state["delivery_approved_by_human"] = False
    state["delivery_approval"] = None

    stats = state.setdefault("review_stats", {})
    stats["blocking"] = sum(1 for f in state["findings"]
                            if f["severity"] == "blocking")
    stats["actionable"] = sum(1 for f in state["findings"]
                              if f["severity"] == "actionable")
    stats["informational"] = sum(1 for f in state["findings"]
                                 if f["severity"] == "informational")
    state["has_blocking"] = stats["blocking"] > 0
    state["delivery_status"] = _delivery.compute_delivery_status(state)
    if job_id:
        from transpraxis import academic_writer
        segment_ids = [
            _finalization.segment_id(job_id, index, pairs[index])
            for index in stale if 0 <= index < len(pairs)
        ]
        propagated = academic_writer.propagate_artifact_staleness(
            state, input_segment_ids=segment_ids)
        enriched = dict(state)
        enriched["_finalization_artifacts"] = _finalization_artifacts(job_id)
        state["dependency_impact"] = _finalization.build_dependency_impact(
            enriched, job_id, stale, "冻结术语表变化")
        affected_cases = [name.split(":", 1)[1] for name in propagated
                          if str(name).startswith("case:")]
        _finalization.mark_case_reviews_stale(
            state, affected_cases, "冻结术语表变化影响了本案例绑定段落")
        if propagated:
            academic = state.setdefault("academic_state", {})
            if any(name in propagated for name in {"sections", "report", "validation", "review"}):
                state["p3_done"] = False
                state["p3_md"] = ""
                state["p3_sections"] = []
                academic["status"] = "stale"
            _reset_final_qa(state, "冻结术语表变化；相关案例与学术下游需重新检查")
            _invalidate_final_delivery_state(state)
    return state, stale


def set_glossary_entry_status(job_id, entry_ids, status):
    """批量修改术语状态（candidate/provisional/locked/rejected）。"""
    if status not in ("candidate", "provisional", "locked", "rejected"):
        raise ValueError(f"非法状态：{status}")
    state = load_job_state(job_id)
    if state is None:
        return None
    ids = set(entry_ids or [])
    for e in state.get("glossary") or []:
        if e.get("id") in ids:
            e["status"] = status
    save_job_state(job_id, state)
    return state


def review_knowledge_candidate(job_id, candidate_id, decision, actor="user"):
    """Apply an explicit human decision to one persisted knowledge candidate.

    This preserves the existing task-level mutation contract: ``project_term``
    locks the candidate in the task glossary and freezes its version. Callers
    that need cross-task reuse must then explicitly pass that confirmed state
    through ``promote_job_to_project`` (the Memory gate). There is intentionally
    no global glossary store.
    """
    allowed = {"project_term", "task_only", "rejected"}
    if decision not in allowed:
        raise ValueError(f"非法知识候选决策：{decision}")
    state = load_job_state(job_id)
    if state is None:
        return None, False, "任务不存在"
    actor = str(actor or "").strip()
    if not actor:
        return state, False, "人工确认必须记录 actor"
    candidate = next((item for item in state.get("knowledge_candidates") or []
                      if _knowledge.candidate_id(item) == str(candidate_id)), None)
    if candidate is None:
        return state, False, "找不到待确认词条"
    if candidate.get("decision"):
        return state, False, "该词条已经处理过"

    context = _knowledge.candidate_context(candidate, state)
    timestamp = datetime.now(timezone.utc).isoformat(timespec="seconds")
    previous_status = str(candidate.get("status") or "emergent_candidate")

    def record(action, note, event):
        candidate["decision"] = decision
        candidate["decision_at"] = timestamp
        candidate["decision_by"] = actor
        candidate["decision_note"] = note
        candidate["status"] = {
            "project_term": "promoted_project_term",
            "task_only": "accepted_task",
            "rejected": "rejected",
        }[decision]
        state["translation_continuity"] = [
            dict(item) for item in state.get("knowledge_candidates") or []
            if isinstance(item, dict)
        ]
        state.setdefault("knowledge_events", []).append({
            "type": "human_candidate_decision",
            "candidate_id": context["candidate_id"],
            "source": context["source"],
            "observed_target": context["proposed_target"],
            "decision": decision,
            "scope": "document" if decision == "project_term" else "task",
            "timestamp": timestamp,
            "actor": actor,
            "actor_type": "human",
            "note": note,
            "previous_status": previous_status,
            "new_status": candidate["status"],
            **event,
        })
        state.setdefault("human_actions", []).append({
            "action": action,
            "finding_id": f"knowledge:{context['candidate_id']}",
            "note": note,
            "timestamp": timestamp,
            "actor": actor,
            "actor_type": "human",
            "previous_status": previous_status,
            "new_status": candidate["status"],
        })

    if decision == "project_term":
        if context["conflicts"]:
            return state, False, "与现有项目术语冲突，未覆盖现有术语"
        entries = normalize_glossary(
            state.get("glossary") or (state.get("glossary_frozen") or {}).get("entries") or [])
        matching = [entry for entry in entries
                    if entry.get("source", "").casefold() == context["source"].casefold()]
        target = context["proposed_target"]
        changed = False
        if matching:
            entry = matching[0]
            current = str(entry.get("preferred") or entry.get("target") or "").strip()
            if current.casefold() != target.casefold():
                return state, False, "与现有术语译名冲突，未覆盖现有术语"
            if entry.get("status") != "locked":
                entry["status"] = "locked"
                entry["preferred"] = target
                entry["target"] = target
                changed = True
            entry_id = entry.get("id")
        else:
            entry = _models.normalize_glossary_entry({
                "source": context["source"],
                "proposed_target": target,
                "target": target,
                "preferred": target,
                "behavior": "translate",
                "status": "locked",
                "scope": "document",
                "occurrences": context["occurrences"],
                "note": "由人工从翻译流知识候选提升为项目术语",
                "evidence": [{
                    "evidence_type": "user",
                    "source_name": "译页知识候选",
                    "note": f"来自第 {(context['first_observed_segment'] + 1) if context['first_observed_segment'] is not None else '?'} 段的人工确认",
                    "quote": context["source_context"],
                    "url": "",
                }],
            })
            if entry is None:
                return state, False, "词条内容无效，未加入项目术语"
            entries.append(entry)
            entry_id = entry.get("id")
            changed = True
        candidate["promotion_entry_id"] = entry_id
        record("knowledge_project_term", "人工确认并加入项目术语；术语版本将更新", {
            "entry_id": entry_id,
        })
        state["glossary"] = entries
        save_job_state(job_id, state)
        if changed or not state.get("glossary_frozen"):
            state = freeze_glossary(job_id, entries=entries, frozen_by=actor)
        return state, True, "已加入项目术语，并通过术语版本冻结流程保存"

    if decision == "task_only":
        record("knowledge_task_only", "人工确认仅在本任务采用，不加入项目术语", {})
        save_job_state(job_id, state)
        return state, True, "已记录为仅本任务采用"

    record("knowledge_rejected", "人工拒绝该知识候选", {})
    save_job_state(job_id, state)
    return state, True, "已拒绝该知识候选"


def confirm_translation_style_rule(
    job_id, rule, actor, *, actor_type, note="", source_finding_id="",
):
    """Explicitly promote one human-confirmed style rule into Project Memory."""
    actor = str(actor or "").strip()
    text = str(rule or "").strip()
    if actor_type != "human" or not actor:
        raise ValueError("only an identified human may confirm style knowledge")
    if not text:
        raise ValueError("confirmed style rule must not be empty")
    state = load_job_state(job_id)
    if state is None:
        raise ValueError(f"找不到任务 {job_id}")
    rule_id = _models.stable_id(text.casefold(), prefix="s")
    existing = next((item for item in state.get("confirmed_style_rules") or []
                     if isinstance(item, dict) and item.get("rule_id") == rule_id), None)
    if existing is not None:
        return state, existing
    timestamp = datetime.now(timezone.utc).isoformat(timespec="seconds")
    record = {
        "rule_id": rule_id,
        "rule": text,
        "status": "confirmed",
        "confirmed_by": actor,
        "confirmed_at": timestamp,
        "note": str(note or "").strip(),
        "source_finding_id": str(source_finding_id or "").strip(),
    }
    state.setdefault("confirmed_style_rules", []).append(record)
    state.setdefault("human_actions", []).append({
        "finding_id": record["source_finding_id"] or f"style:{rule_id}",
        "action": "confirm_style_rule",
        "actor": actor,
        "actor_type": "human",
        "timestamp": timestamp,
        "note": record["note"] or text,
        "previous_status": "candidate",
        "new_status": "confirmed",
        "rule_id": rule_id,
    })
    excluded = _translation_review_finding(state, record["source_finding_id"])
    excluded_segment = excluded.get("segment_id") if excluded else None
    indexes = [index for index in range(len(state.get("pairs") or []))
               if index != excluded_segment]
    _invalidate_translation_reviews(
        state, indexes, "confirmed style knowledge changed review dependencies")
    if indexes:
        _invalidate_final_delivery_state(state)
    save_job_state(job_id, state)
    return state, record


def freeze_glossary(job_id, entries=None, frozen_by="user"):
    """冻结术语表：生成新版本 + 确定性 glossary_hash。

    修改后再次冻结 -> 新版本追加到 glossary_versions，不悄悄覆盖旧冻结状态。
    相同 canonical 内容再次冻结 -> 不创建新版本（幂等）。
    冻结新版本后立即对已翻译段落执行术语依赖失效（stale 标记 + TM 清除）。
    """
    state = load_job_state(job_id)
    if state is None:
        return None
    was_final = state.get("delivery_status") == "final"
    norm = normalize_glossary(entries if entries is not None
                              else state.get("glossary") or [])
    versions = state.get("glossary_versions") or []
    if versions and isinstance(versions[-1], dict) \
            and _models.entries_equal(versions[-1].get("entries") or [], norm):
        # 相同内容：不创建新版本（决策 A），保持原 frozen 快照
        state["glossary"] = norm
        state["glossary_frozen"] = versions[-1]
        state["stage"] = "GLOSSARY_FROZEN"
        save_job_state(job_id, state)
        return state
    source_hash = ""
    src = job_dir(job_id) / "source.bin"
    if src.is_file():
        source_hash = hashlib.sha256(src.read_bytes()).hexdigest()
    version = len(versions) + 1
    frozen = {
        "version": version,
        "source_hash": source_hash,
        "entries": norm,
        "frozen_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "glossary_hash": _models.glossary_hash(norm),
        "frozen_by": frozen_by,
    }
    state["glossary"] = norm
    state["glossary_frozen"] = frozen
    versions.append(frozen)
    state["glossary_versions"] = versions
    state["stage"] = "GLOSSARY_FROZEN"
    if state.get("delivery_status") not in ("approved", "final"):
        state["delivery_status"] = "draft"
    # 术语决策变化 -> 立即失效受影响段落
    state, stale = _apply_glossary_staleness(state, job_id)
    if was_final and not stale:
        _invalidate_final_delivery_state(state)
    save_job_state(job_id, state)
    return state


def unfreeze_glossary(job_id):
    """返回修改：解除冻结（翻译开始前），回到 TERMS_PREPARED；旧冻结版本保留。"""
    state = load_job_state(job_id)
    if state is None:
        return None
    if state.get("p2_done"):
        # 翻译已开始：不允许解除冻结，只能生成新版本（见 freeze_glossary）
        return state
    state["glossary_frozen"] = None
    state["stage"] = "TERMS_PREPARED"
    if state.get("delivery_status") not in ("approved", "final"):
        state["delivery_status"] = "draft"
    save_job_state(job_id, state)
    return state


def save_document_profile(job_id, profile):
    """人工填写/修改文档画像后保存。"""
    state = load_job_state(job_id)
    if state is None:
        return None
    normalized = _models.normalize_document_profile(profile)
    changed = normalized != _models.normalize_document_profile(
        state.get("document_profile"))
    state["document_profile"] = normalized
    state["profile_done"] = True
    if changed and state.get("p2_done"):
        _invalidate_translation_reviews(
            state, list(range(len(state.get("pairs") or []))),
            "document profile changed; review context must be refreshed")
        _invalidate_final_delivery_state(state)
    save_job_state(job_id, state)
    return state


def set_entity_translation(job_id, source_form, target, *, entity_type="other_proper_noun",
                           locked=True, actor="user", note=""):
    """Persist a human entity choice; it outranks generated continuity hints."""
    state = load_job_state(job_id)
    if state is None:
        return None
    registry = _entity_registry.EntityRegistry(state.get("entity_registry") or [])
    record = (registry.lock if locked else registry.accept)(
        source_form, target, entity_type=entity_type, note=note)
    state["entity_registry"] = registry.to_list()
    state.setdefault("human_actions", []).append({
        "finding_id": f"entity:{record.get('id') or source_form}",
        "action": "entity_locked" if locked else "entity_accepted",
        "note": note or f"人工确认实体译名：{source_form} -> {target}",
        "timestamp": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "actor": actor,
    })
    if state.get("p2_done"):
        state["targeted_final_review"] = {
            "status": "not_run", "segment_ids": [],
            "reviewed_segment_ids": [], "failed_segment_ids": [],
        }
    if (state.get("delivery_status") in ("approved", "final")
            or state.get("delivery_approved_by_human")):
        _invalidate_final_delivery_state(state)
    save_job_state(job_id, state)
    return state


def write_profile_artifacts(job_id, document_profile, style_profile):
    """把 Step 01 的画像产物落盘为版本化 artifact：
    document_profile.json / style_profile.json（含 style_profile_id 哈希）。
    返回写入的 style_profile_id，失败返回 None。
    """
    try:
        from transpraxis.style_profile import style_profile_id
        job_root = job_dir(job_id)
        job_root.mkdir(parents=True, exist_ok=True)
        (job_root / "document_profile.json").write_text(
            json.dumps(_models.normalize_document_profile(document_profile),
                       ensure_ascii=False, indent=2), encoding="utf-8")
        profile_id = style_profile_id(style_profile or {})
        artifact = dict(style_profile or {})
        artifact["style_profile_id"] = profile_id
        (job_root / "style_profile.json").write_text(
            json.dumps(artifact, ensure_ascii=False, indent=2), encoding="utf-8")
        return profile_id
    except Exception:  # noqa: BLE001 - artifact 落盘失败不影响主流程
        return None


def bypass_freeze(job_id, frozen_by="user"):
    """快速模式跳过人工冻结：允许以 provisional 术语直接翻译（记录审计标记）。"""
    state = load_job_state(job_id)
    if state is None:
        return None
    state["quality_bypass"] = True
    state.setdefault("human_actions", []).append({
        "action": "bypass_freeze",
        "note": "快速模式：跳过人工术语冻结，以 provisional 术语直接翻译",
        "timestamp": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "actor": frozen_by,
    })
    save_job_state(job_id, state)
    return state


# ================= 交付状态与人工处理记录 =================
def _translation_review_finding(state, finding_id):
    """Resolve either the Core finding ID or the existing review-card ID."""
    from transpraxis import delivery as _delivery
    selector = str(finding_id or "")
    candidates = [item for item in state.get("findings") or []
                  if isinstance(item, dict) and (
                      str(item.get("finding_id") or "") == selector
                      or _delivery.finding_id(item) == selector)]
    return next((item for item in reversed(candidates)
                 if isinstance(item.get("segment_id"), int)
                 and (_translation_evidence.current_review_event(
                     state, item["segment_id"]) or {}).get(
                         "review_event_id") == item.get("review_event_id")),
                candidates[-1] if candidates else None)


def _promote_human_reviewed_segment(state, segment_id, actor):
    """Promote one currently clean human-confirmed pair through existing TM truth."""
    from transpraxis import delivery as _delivery
    pairs = state.get("pairs") or []
    if not isinstance(segment_id, int) or not 0 <= segment_id < len(pairs):
        return
    if any(_delivery._segment_index(item) == segment_id
           for item in _delivery.unresolved_findings(state)):
        return
    pair = pairs[segment_id]
    deterministic = check_translation_batch(
        [str(pair.get("source") or "")], [str(pair.get("target") or "")],
        state.get("glossary") or [], state.get("target_lang") or "简体中文")
    if any(item.get("severity") in {"blocking", "actionable"}
           for item in deterministic):
        return
    pair["reviewed"] = True
    pair["review_status"] = "reviewed_human"
    pair["accepted_target"] = pair.get("target") or ""
    pair["target_provenance"] = "human_accepted"
    pair["human_accepted"] = True
    pair["accepted_by_human"] = actor
    if state.get("use_tm", True) and _tm_eligible(pair.get("source"), pair.get("target")):
        tm = load_tm(tm_project_id(state))
        tm_put(tm, str(pair.get("source") or ""), str(pair.get("target") or ""),
               state_target_lang(state))
        save_tm(tm, tm_project_id(state))


def decide_translation_review_finding(
    job_id, finding_id, decision, actor, *, actor_type, note="",
):
    """Apply one fail-closed Translation Core HumanDecision to persisted state."""
    from transpraxis import delivery as _delivery
    state = load_job_state(job_id)
    if state is None:
        raise ValueError(f"找不到任务 {job_id}")
    finding = _translation_review_finding(state, finding_id)
    if finding is None or not finding.get("finding_id"):
        raise ValueError("finding has no Translation Core identity")
    state, updated, record = _translation_evidence.record_runtime_human_decision(
        state, finding["finding_id"], decision, actor, actor_type=actor_type,
        note=note)
    segment_id = updated.get("segment_id")
    if decision in {"accept_resolution", "dismiss"}:
        _promote_human_reviewed_segment(state, segment_id, actor)
    _recount_reviewed_segments(state)
    queue = _delivery.review_queue_findings(state)
    stats = state.setdefault("review_stats", {})
    for severity in ("blocking", "actionable", "informational"):
        stats[severity] = sum(item.get("severity") == severity for item in queue)
    state["has_blocking"] = bool(_delivery.unresolved_blocking(state))
    state["delivery_status"] = _delivery.compute_delivery_status(state)
    save_job_state(job_id, state)
    return state, updated, record


def _targeted_final_review_ids(state, limit=8):
    """Select a small, stable set of segments for the whole-book recheck.

    Batch review already covers ordinary semantic issues.  The final pass is
    reserved for risks that only become visible after all pairs are present:
    entity drift, locked-term conflicts, and low-authority knowledge conflicts.
    Resolved findings are deliberately excluded so a human decision is not
    reopened by a later resume.
    """
    pairs = state.get("pairs") or []
    if not pairs:
        return []
    severity_rank = {"blocking": 0, "actionable": 1, "informational": 2}
    selected = {}
    for finding in state.get("findings") or []:
        if not isinstance(finding, dict) or finding.get("resolved"):
            continue
        finding_type = str(finding.get("type") or "").strip().lower()
        category = str(finding.get("category") or "").strip().lower()
        is_global_consistency_risk = (
            finding_type in {"knowledge_conflict", "entity_conflict"}
            or (finding_type == "glossary" and finding.get("conflict"))
            or (finding.get("conflict") and category == "terminology_consistency")
        )
        if not is_global_consistency_risk:
            continue
        raw_index = finding.get("segment_id", finding.get("segment_index"))
        try:
            segment_id = int(raw_index)
        except (TypeError, ValueError):
            continue
        if isinstance(raw_index, bool) or not 0 <= segment_id < len(pairs):
            continue
        score = severity_rank.get(str(finding.get("severity") or ""), 3)
        previous = selected.get(segment_id)
        if previous is None or score < previous:
            selected[segment_id] = score
    bounded_limit = max(0, int(limit or 0))
    return [segment_id for segment_id, _score in sorted(
        selected.items(), key=lambda item: (item[1], item[0]))[:bounded_limit]]


def review_translation_segments(
    job_id, indexes, provider, api_key, model, target_lang, *, style_rules="",
    call_llm_fn=None, base_url=None, review_focus="", focus_findings=None,
):
    """Run the existing independent reviewer on current persisted segments.

    ``review_focus`` and ``focus_findings`` are optional context for the
    post-translation whole-book pass.  They do not change the normal manual
    refresh path and keep the global conflict visible to the reviewer.
    """
    from transpraxis import delivery as _delivery
    from transpraxis.terminology import (
        glossary_block as _glossary_block,
        select_glossary_for_segments as _select_glossary,
    )
    state = load_job_state(job_id)
    if state is None:
        raise ValueError(f"找不到任务 {job_id}")
    pairs = state.get("pairs") or []
    segment_ids = sorted({int(index) for index in indexes
                          if str(index).lstrip("-").isdigit()
                          and 0 <= int(index) < len(pairs)})
    if not segment_ids:
        return state, {"reviewed_segment_ids": [], "failed_segment_ids": []}
    review_call = call_llm_fn or _model_roles.make_role_call(
        call_llm, {"provider": provider, "api_key": api_key,
                   "model": model, "base_url": base_url})
    glossary = normalize_glossary(
        state.get("glossary") or (state.get("glossary_frozen") or {}).get(
            "entries") or [])
    reviewed, failed_ids = [], []
    tm = load_tm(tm_project_id(state)) if state.get("use_tm", True) else {}
    tm_changed = False
    for segment_id in segment_ids:
        pair = pairs[segment_id]
        source = str(pair.get("source") or "")
        target = str(pair.get("target") or "")
        section_profile = _batch_section_profile(
            state.get("document_profile"), segment_id, 1)
        selected, _ = _select_glossary(
            [source], glossary + _knowledge.provisional_hints(
                state.get("knowledge_candidates") or [],
                authoritative_entries=glossary),
            state.get("document_profile"), section_profile)
        glossary_text = _glossary_block(selected)
        deterministic = _globalize_batch_findings(check_translation_batch(
            [source], [target], glossary, target_lang,
            section_profile=section_profile), segment_id)
        review_context = _runtime_review_context(
            state, segment_id, 1, glossary_text, style_rules, target_lang)
        if review_focus:
            risk_items = [
                {
                    key: item.get(key)
                    for key in (
                        "segment_id", "type", "category", "severity", "reason",
                        "source", "preferred_target", "observed_target",
                        "observed_targets",
                    ) if item.get(key) is not None
                } for item in (focus_findings or [])
                if isinstance(item, dict)
                and item.get("segment_id", item.get("segment_index")) == segment_id
            ][:4]
            review_context["review_focus"] = str(review_focus)
            review_context["whole_book_consistency_risks"] = risk_items
        packet = _translation_evidence.build_runtime_review_packet(
            state, [pair], [segment_id], glossary,
            deterministic_checks=deterministic, review_context=review_context)
        evidence_index = _translation_evidence.TranslationEvidenceIndex(
            state.get("paras") or [item.get("source") or "" for item in pairs],
            pairs, glossary, state.get("document_profile"),
            state.get("document_synopsis"), state.get("section_digests"),
            state.get("findings"))
        findings, failed, trace = \
            _translation_evidence.review_translation_batch_with_evidence(
                [source], [target], glossary_text, style_rules, target_lang,
                provider, api_key, model, evidence_index,
                call_llm=review_call, segment_ids=[segment_id],
                translation_core_packet=packet)
        event_id = (
            f"translation-review-{job_id}-refresh-{segment_id}-"
            f"{len(state.get('review_evidence') or [])}")
        records = [_review_finding_record(item, event_id) for item in findings]
        state.setdefault("findings", []).extend(records)
        _translation_evidence.register_runtime_review_event(
            state, trace, records, event_id, [segment_id])
        focus_risk_present = bool(review_focus and any(
            isinstance(item, dict)
            and item.get("segment_id", item.get("segment_index")) == segment_id
            and not item.get("resolved")
            for item in (focus_findings or [])))
        if failed:
            failed_ids.append(segment_id)
            pair["reviewed"] = False
            pair["review_status"] = "review_failed"
        elif not focus_risk_present and not any(
                item.get("severity") in {"blocking", "actionable"}
                for item in [*deterministic, *records]):
            reviewed.append(segment_id)
            pair["reviewed"] = True
            pair["review_status"] = "reviewed_clean"
            pair["accepted_target"] = target
            pair["target_provenance"] = "reviewed"
            if state.get("use_tm", True) and _tm_eligible(source, target):
                if tm_put(tm, source, target, state_target_lang(state)):
                    tm_changed = True
        else:
            pair["reviewed"] = False
            pair["review_status"] = "reviewed_with_findings"
        _checkpoint.append_event(job_dir(job_id), {
            "phase": "semantic_review_done", "segment_ids": [segment_id],
            "refresh": True, "failed": failed,
            **({"review_focus": str(review_focus)} if review_focus else {}),
        })
    if tm_changed:
        save_tm(tm, tm_project_id(state))
    _recount_reviewed_segments(state)
    queue = _delivery.review_queue_findings(state)
    stats = state.setdefault("review_stats", {})
    for severity in ("blocking", "actionable", "informational"):
        stats[severity] = sum(item.get("severity") == severity for item in queue)
    stats["review_failed"] = sum(
        item.get("decision") == "failed" and _translation_evidence._current_translation_review(
            item) for item in state.get("review_evidence") or [] if isinstance(item, dict))
    state["has_blocking"] = bool(_delivery.unresolved_blocking(state))
    state["delivery_status"] = _delivery.compute_delivery_status(state)
    save_job_state(job_id, state)
    return state, {
        "reviewed_segment_ids": reviewed,
        "failed_segment_ids": failed_ids,
    }


def _run_targeted_final_review(
    job_id, state, reviewer_config, target_lang, style_rules, *,
    enable_review=True, on_status=None,
):
    """Run the bounded whole-book consistency pass once per translation.

    The existing segment reviewer is reused so evidence, review events and
    delivery gates remain on one path.  The usage ledger lives on the caller's
    state object; the review helper reloads state from disk, so it is merged
    back explicitly after the call to avoid losing the newly recorded usage.
    """
    record = state.get("targeted_final_review") or {}
    if record.get("status") in {"completed", "skipped"}:
        return state
    if not enable_review:
        state["targeted_final_review"] = {
            "status": "skipped", "reason": "review_disabled",
            "segment_ids": [], "reviewed_segment_ids": [],
            "failed_segment_ids": [],
        }
        save_job_state(job_id, state)
        return state
    segment_ids = _targeted_final_review_ids(state)
    if not segment_ids:
        state["targeted_final_review"] = {
            "status": "skipped", "reason": "no_cross_book_risk",
            "segment_ids": [], "reviewed_segment_ids": [],
            "failed_segment_ids": [],
        }
        save_job_state(job_id, state)
        return state
    if on_status:
        on_status(f"【阶段二终检】全书一致性定向复核（{len(segment_ids)} 段）...")
    focus_findings = []
    selected = set(segment_ids)
    for finding in state.get("findings") or []:
        if not isinstance(finding, dict) or finding.get("resolved"):
            continue
        raw_index = finding.get("segment_id", finding.get("segment_index"))
        try:
            segment_id = int(raw_index)
        except (TypeError, ValueError):
            continue
        if isinstance(raw_index, bool) or segment_id not in selected:
            continue
        item = {
            key: finding.get(key)
            for key in (
                "segment_id", "type", "category", "severity", "reason",
                "source", "preferred_target", "observed_target", "observed_targets",
            ) if finding.get(key) is not None
        }
        item["segment_id"] = segment_id
        focus_findings.append(item)
    state.setdefault("llm_usage", _usage.empty_usage())
    final_review_call = _model_roles.make_role_call(
        _usage.tracked_call(
            call_llm, state["llm_usage"], role="targeted_final_review"),
        reviewer_config)
    try:
        reviewed_state, result = review_translation_segments(
            job_id, segment_ids, reviewer_config["provider"],
            reviewer_config["api_key"], reviewer_config["model"], target_lang,
            style_rules=style_rules, call_llm_fn=final_review_call,
            base_url=reviewer_config.get("base_url"),
            review_focus="whole_book_consistency",
            focus_findings=focus_findings,
        )
        # review_translation_segments reloads and saves its own state.  Carry
        # over the in-memory ledger mutated by the tracked call wrapper.
        reviewed_state["llm_usage"] = state["llm_usage"]
        reviewed_ids = list(result.get("reviewed_segment_ids") or [])
        failed_ids = list(result.get("failed_segment_ids") or [])
        final_status = "failed" if failed_ids else "completed"
        reviewed_state["targeted_final_review"] = {
            "status": final_status,
            "segment_ids": list(segment_ids),
            "reviewed_segment_ids": reviewed_ids,
            "failed_segment_ids": failed_ids,
        }
        if failed_ids:
            warning = f"全书一致性定向复核有 {len(failed_ids)} 段未完成"
            warnings = reviewed_state.setdefault("warnings", [])
            if warning not in warnings:
                warnings.append(warning)
        save_job_state(job_id, reviewed_state)
        return reviewed_state
    except Exception as exc:
        state["targeted_final_review"] = {
            "status": "failed", "segment_ids": list(segment_ids),
            "reviewed_segment_ids": [], "failed_segment_ids": list(segment_ids),
            "error": str(exc)[:240],
        }
        warning = f"全书一致性定向复核失败：{str(exc)[:200]}"
        warnings = state.setdefault("warnings", [])
        if warning not in warnings:
            warnings.append(warning)
        save_job_state(job_id, state)
        return state


def mark_findings_resolved(job_id, finding_ids, action, note="", actor="user"):
    """Legacy action surface; Core findings route through HumanDecision."""
    from transpraxis import delivery as _delivery
    state = load_job_state(job_id)
    if state is None:
        return None
    legacy_ids = []
    mapped = {"human_fixed": "accept_resolution", "preserved": "dismiss"}
    for selector in finding_ids or []:
        finding = _translation_review_finding(state, selector)
        if finding is not None and finding.get("finding_id") \
                and finding.get("input_fingerprint"):
            if action not in mapped:
                raise ValueError(f"Translation Core finding 不支持旧 action：{action}")
            state, updated, _record = \
                _translation_evidence.record_runtime_human_decision(
                    state, finding["finding_id"], mapped[action], actor,
                    actor_type="human", note=note)
            if mapped[action] in {"accept_resolution", "dismiss"}:
                _promote_human_reviewed_segment(
                    state, updated.get("segment_id"), actor)
        else:
            legacy_ids.append(selector)
    if legacy_ids:
        state, _marked = _delivery.mark_findings(
            state, legacy_ids, action, note, actor)
    _recount_reviewed_segments(state)
    state["has_blocking"] = bool(_delivery.unresolved_blocking(state))
    state["delivery_status"] = _delivery.compute_delivery_status(state)
    save_job_state(job_id, state)
    return state


def approve_delivery(job_id, note="", accept_blocking=False, actor="user",
                     target_lang="", provider="", model=""):
    """人工交付确认 -> final + immutable snapshot."""
    from transpraxis import delivery as _delivery
    state = load_job_state(job_id)
    if state is None:
        return None, False, ["任务不存在"]
    if state.get("report_enabled") and state.get("report_status") in {
            "incomplete", "failed_template_validation", "review_required"}:
        return state, False, ["实践报告尚未完成，不能冻结最终交付"]
    academic_records = (state.get("academic_state") or {}).get("artifacts") or {}
    # A pre-v0.4 job may have only p3_md and the old delivery formats.  It can
    # still be read and delivered through its historical path; strict report
    # compliance/QA begins once a structured report artifact is present.
    strict_report_gate = state.get("report_enabled") and (
        isinstance(academic_records.get("report"), dict) or
        any(name in academic_records for name in (
            "compliance", "final_docx_validation", "libreoffice_render",
            "report_qa")))
    case_gate = _finalization.case_review_gate(
        state, load_academic_artifact(job_id, "selected_cases"),
        require_artifact_status=bool(strict_report_gate))
    if case_gate.get("status") == "blocked":
        labels = ", ".join(case_gate.get("blocked_case_ids") or [])
        return state, False, [
            f"案例人工终审未通过：{labels}。作者拒绝或过期案例不能进入最终交付。"
        ]
    if strict_report_gate:
        compliance = compliance_profile_view(job_id, state)
        profile_rules = compliance.get("profile_compliance") or {}
        project = compliance.get("project_constraints") or {}
        if profile_rules.get("status") == "fail":
            return state, False, [
                "Default profile compliance failed: " +
                ", ".join(profile_rules.get("blocking_failures") or [])]
        if project.get("status") == "fail":
            return state, False, [
                "Project compliance constraint failed: " +
                ", ".join(project.get("failures") or [])]
        from transpraxis import academic_writer
        report_record = academic_writer.artifact_record(state, "report")
        docx_record = academic_writer.artifact_record(
            state, "final_docx_validation")
        render_record_state = academic_writer.artifact_record(
            state, "libreoffice_render")
        stale_artifacts = [name for name, record in (
            ("report", report_record),
            ("DOCX", docx_record),
            ("LibreOffice render", render_record_state),
        ) if record.get("status") in {"stale", "missing", "failed"}]
        if stale_artifacts:
            return state, False, [
                "Finalization artifacts are stale or unavailable: " +
                ", ".join(stale_artifacts)]
        final_docx = load_academic_artifact(job_id, "final_docx_validation") or {}
        render_record = load_academic_artifact(job_id, "libreoffice_render") or {}
        qa = _finalization.normalize_final_qa(state.get("final_qa"))
        structural = "PASS" if final_docx.get("status") in {
            "pass", "pass_with_warnings"} else "FAIL" if final_docx.get(
            "status") == "fail" else "NOT_RUN"
        qa_reasons = []
        if structural != "PASS":
            qa_reasons.append(f"Structural QA={structural}")
        if render_record.get("qa_status") != "PASS":
            qa_reasons.append(f"LibreOffice Render={render_record.get('qa_status', 'NOT_RUN')}")
        if qa.get("author_visual_review") != "CONFIRMED":
            qa_reasons.append("Author Visual Review=NOT_CONFIRMED")
        if qa.get("word_final_review") != "CONFIRMED":
            qa_reasons.append("Word Final Review=NOT_CONFIRMED")
        if qa_reasons:
            return state, False, ["Final QA gate blocked: " + "; ".join(qa_reasons)]
    validation = validate_delivery_translation_state(state)
    if validation["blocking"]:
        _record_delivery_validation_findings(state, validation)
        save_job_state(job_id, state)
        reasons = [issue.get("message", "译文未通过交付检查")
                   for issue in validation.get("issues") or []]
        # `blocking` can also come from deterministic QA blocking findings
        # (`core.validate_translation_pairs` -> `blocking_findings`), which carry a
        # segment and a summary.  Without them the user is told only
        # "译文未通过最终交付检查" and cannot tell what to fix.
        for finding in validation.get("blocking_findings") or []:
            index = finding.get("segment_index")
            location = "" if index is None else f"第 {int(index) + 1} 段 "
            category = str(finding.get("category") or "check")
            summary = (finding.get("summary") or finding.get("reason")
                       or "译文未通过交付检查")
            reasons.append(f"{location}[{category}] {summary}")
        return state, False, reasons or ["译文未通过最终交付检查"]
    state, ok, errors = _delivery.approve_delivery(state, note, actor, accept_blocking)
    if not ok:
        save_job_state(job_id, state)
        return state, ok, errors
    try:
        manifest = create_delivery_snapshot(
            job_id, state,
            target_lang=state.get("target_lang") or target_lang or "",
            provider=state.get("provider") or provider or "",
            model=state.get("model") or model or "")
    except Exception as exc:  # fail closed: final state is not saved without bytes
        persisted = load_job_state(job_id) or state
        return persisted, False, [f"无法冻结最终交付版本：{str(exc)[:200]}"]
    state.setdefault("delivery_snapshots", []).append({
        "version": manifest["snapshot_version"],
        "created_at": manifest["created_at"],
        "approval": dict(manifest.get("approval") or {}),
        "asset_count": len(manifest.get("assets") or []),
        "translation_state_identity": manifest.get("translation_state_identity"),
    })
    state["latest_delivery_snapshot_version"] = manifest["snapshot_version"]
    save_job_state(job_id, state)
    return state, True, []


def retranslate_segments(job_id, indexes, provider, api_key, model, target_lang,
                         style_rules="", glossary=None, on_status=None,
                         on_caption=None, actor="user", reviewer_provider=None,
                         reviewer_api_key=None, reviewer_model=None,
                         reviewer_base_url=None):
    """定点重译（抽取自 scripts/fix_segments.py 的能力）。"""
    from transpraxis import delivery as _delivery
    state, fixed = _delivery.retranslate_segments(
        job_id, indexes, provider, api_key, model, target_lang,
        style_rules, glossary, on_status, on_caption, actor)
    if fixed and state.get("translation_core_review_required"):
        review_provider = provider if reviewer_provider is None else reviewer_provider
        review_api_key = api_key if reviewer_api_key is None else reviewer_api_key
        review_model = model if reviewer_model is None else reviewer_model
        state, _ = review_translation_segments(
            job_id, fixed, review_provider, review_api_key, review_model,
            target_lang, style_rules=style_rules, base_url=reviewer_base_url)
    return state, fixed


def delivery_status_label(state):
    labels = {"draft": "草稿（draft）", "review_required": "待审（review_required）",
              "approved": "已批准（approved）", "final": "最终交付（final）"}
    return labels.get(state.get("delivery_status"), str(state.get("delivery_status")))


def academic_status_label(state):
    status = state.get("report_status") or \
        (state.get("academic_state") or {}).get("report_status") or \
        (state.get("academic_state") or {}).get("quality_status") or \
        (state.get("academic_state") or {}).get("status") or "not_started"
    labels = {
        "not_started": "尚未开始", "stale": "需要更新（按影响范围）",
        "in_progress": "学术写作中", "failed": "学术写作失败",
        "pass": "验证通过", "pass_with_warnings": "通过（有警告）",
        "review_required": "需要人工学术复核", "fail": "验证失败",
        "generated": "报告已生成", "incomplete": "当前证据不足，报告不完整",
        "failed_template_validation": "报告生成未通过模板校验",
    }
    return labels.get(status, status)


def invalidate_academic_report(job_id, scope="all", section_id=None):
    """Invalidate academic artifacts only; translation stages remain intact."""
    from transpraxis import academic_writer
    state = load_job_state(job_id)
    if state is None:
        raise ValueError(f"找不到任务 {job_id}")
    academic_writer.invalidate_academic_state(state, scope, section_id)
    save_job_state(job_id, state)
    return state


def load_academic_artifact(job_id, name):
    """Read a canonical academic JSON artifact for UI/CLI inspection."""
    from transpraxis import academic_writer
    ARTIFACT_FILES = academic_writer.ARTIFACT_FILES
    if name not in ARTIFACT_FILES:
        raise ValueError(f"未知学术 artifact：{name}")
    path = job_dir(job_id) / ARTIFACT_FILES[name]
    if not path.is_file():
        return None
    try:
        return academic_writer._read_artifact(path)
    except Exception:
        return None


def record_human_evidence(job_id, question_id, answer, interface="academic_workspace"):
    """Record a human author answer for an open evidence question.

    The answer is stored verbatim with provenance; the question is marked
    answered; the case-analysis/section staleness is propagated through the
    existing dependency-hash architecture (only affected sections rewrite).
    """
    from transpraxis import academic_writer, human_evidence
    state = load_job_state(job_id)
    if state is None:
        raise ValueError(f"找不到任务 {job_id}")
    questions = academic_writer._read_artifact(
        job_dir(job_id) / academic_writer.ARTIFACT_FILES["human_evidence_questions"])
    if not questions:
        raise ValueError("当前没有待回答的人类证据问题；请先生成学术报告。")
    evidence = academic_writer._read_artifact(
        job_dir(job_id) / academic_writer.ARTIFACT_FILES["evidence"])
    entry, updated_questions = human_evidence.record_human_answer(
        questions, question_id, answer, evidence or {}, interface,
        existing=state.get("human_evidence") or [])
    entries = list(state.get("human_evidence") or [])
    entries = [x for x in entries if x.get("question_id") != question_id]
    entries.append(entry)
    state["human_evidence"] = entries
    academic_writer._write_artifact(
        job_dir(job_id) / academic_writer.ARTIFACT_FILES["human_evidence_questions"],
        updated_questions)
    record = state.get("academic_state", {}).get("artifacts", {}).get(
        "human_evidence_questions")
    if record:
        record["content_hash"] = updated_questions["content_hash"]
        record["updated_at"] = academic_writer._now()
    save_job_state(job_id, state)
    return entry


# ================= 阶段三：证据约束型学术写作 =================
def generate_mti_report(bilingual_pairs, termbase_dict, theory, provider, api_key,
                        model, state, job_id, on_status=None,
                        research_settings=None, literature_sources=None):
    """Compatibility wrapper for the evidence-grounded academic pipeline.

    ``bilingual_pairs`` and ``termbase_dict`` remain in the signature for old
    callers; the canonical inputs are now the saved translation state and the
    durable academic artifacts beside ``state.json``.
    """
    from transpraxis import academic_writer
    report_md = academic_writer.run_academic_pipeline(
        state, job_id, theory, provider, api_key, model,
        artifact_dir=job_dir(job_id), call_llm=call_llm,
        save_state=lambda current: save_job_state(job_id, current),
        research_settings=research_settings, literature_sources=literature_sources,
        on_status=on_status,
    )
    save_compliance_record(job_id, load_job_state(job_id) or state)
    return report_md


# ================= 主流程：单文档完整流水线 =================
def run_job_pipeline(job_id, filename, file_bytes, *, provider, api_key, model,
                     target_lang, auto_term, enable_report, translation_theory,
                     user_glossary=None, style_rules="", enable_review=True,
                     enable_annotate=True, use_tm=True,
                     strict_terminology_governance=False, mode=None,
                     research_settings=None, literature_sources=None,
                     delivery_config=None, batch_size=None, max_batch_chars=None,
                     enable_understanding=None, reviewer_provider=None,
                     reviewer_api_key=None, reviewer_model=None,
                     reviewer_base_url=None, translator_base_url=None,
                     auxiliary_provider=None, auxiliary_api_key=None,
                     auxiliary_model=None, auxiliary_base_url=None,
                     knowledge_feedback_interval=None,
                     enable_source_cleanup=None,
                     reasoning_effort=None, source_cleanup_concurrency=None,
                     ocr_workers=None, ocr_queue_size=None,
                     source_cleanup_confidence_threshold=None,
                     source_cleanup_max_retries=None,
                     source_cleanup_request_interval=None,
                     segmentation_mode=None,
                     translation_concurrency=None,
                     on_status=None, on_caption=None):
    """执行单个文档的完整流程；每个里程碑实时落盘，刷新/重启后均可继续。

    strict_terminology_governance=True：翻译前建立文档画像，并要求自动候选
    术语完成审核/冻结。导入的锁定术语视为已固定，不会阻塞翻译。
    ``mode`` 仅保留给旧调用方；quality 映射到严格术语治理，quick 映射到关闭。
    ``auxiliary_*`` 可把画像、摘要、术语抽取与连续性观察路由到较低成本模型；
    不提供时沿用正文模型。``knowledge_feedback_interval`` 仅控制普通模式的
    批后连续性观察频率，严格审校模式仍逐批观察。
    """
    _ensure_output_dir()
    if translator_base_url is not None:
        set_llm_base_url(translator_base_url)
    set_llm_reasoning_effort(reasoning_effort)
    base = new_job_state(filename)
    state = load_job_state(job_id) or base
    state = {**base, **state}  # 兼容旧版本状态缺字段
    state = _state_migration.migrate_state(state)
    effective_segmentation_mode = _source_segmentation_mode(
        state, segmentation_mode)
    previous_performance = state.get("performance")
    saved_runtime_config = state.get("pipeline_config") or {}
    effective_cleanup_concurrency = _runtime_int_option(
        source_cleanup_concurrency,
        saved_runtime_config.get("source_cleanup_concurrency"),
        "FOLIOTHREAD_LLM_CLEANUP_CONCURRENCY",
        _source_cleanup.DEFAULT_PARALLELISM, maximum=16)
    effective_ocr_workers = _runtime_int_option(
        ocr_workers, saved_runtime_config.get("ocr_workers"),
        "FOLIOTHREAD_OCR_WORKERS", OCR_WORKERS_DEFAULT, maximum=8)
    effective_ocr_queue_size = _runtime_int_option(
        ocr_queue_size, saved_runtime_config.get("ocr_queue_size"),
        "FOLIOTHREAD_OCR_QUEUE_SIZE", effective_ocr_workers * 2, maximum=32)
    effective_cleanup_confidence_threshold = _runtime_float_option(
        source_cleanup_confidence_threshold,
        saved_runtime_config.get("source_cleanup_confidence_threshold"),
        "FOLIOTHREAD_OCR_CONFIDENCE_THRESHOLD",
        OCR_CONFIDENCE_THRESHOLD_DEFAULT, maximum=100.0)
    effective_cleanup_max_retries = _runtime_int_option(
        source_cleanup_max_retries,
        saved_runtime_config.get("source_cleanup_max_retries"),
        "FOLIOTHREAD_LLM_CLEANUP_MAX_RETRIES",
        _source_cleanup.DEFAULT_MAX_RETRIES, minimum=0, maximum=6)
    effective_cleanup_request_interval = _runtime_float_option(
        source_cleanup_request_interval,
        saved_runtime_config.get("source_cleanup_request_interval"),
        "FOLIOTHREAD_LLM_CLEANUP_REQUEST_INTERVAL", 0.0, maximum=60.0)
    profiler = PipelineProfiler(
        job_dir(job_id) / "performance.json", run_id=uuid.uuid4().hex
    )

    def persist_performance():
        # Resuming a task after ingestion must not replace a useful prior
        # timing artifact with an empty run merely because later translation
        # stages do not touch the PDF profiler.
        if (isinstance(previous_performance, dict)
                and previous_performance.get("stages")
                and not profiler.snapshot().get("stages")):
            return previous_performance
        return profiler.persist()

    # An explicit UI-triggered source rescan is a checkpoint-only operation:
    # it must not fall through into profiling, terminology, or translation in
    # the same worker turn.  The request is persisted before the worker starts
    # so a browser refresh cannot lose the intent.
    if state.get("source_rescan_requested"):
        options = state.pop("source_rescan_options", {})
        state.pop("source_rescan_requested", None)
        save_job_state(job_id, state)
        return rescan_pdf_source(
            job_id, filename=filename, provider=provider, api_key=api_key,
            model=model, base_url=translator_base_url, on_status=on_status,
            max_batch_chars=(options.get("max_batch_chars")
                             if isinstance(options, dict) else None),
            max_batch_items=(options.get("max_batch_items")
                             if isinstance(options, dict) else None),
            parallelism=(options.get("parallelism", effective_cleanup_concurrency)
                         if isinstance(options, dict) else effective_cleanup_concurrency),
            ocr_workers=(options.get("ocr_workers") if isinstance(options, dict) else None),
            ocr_queue_size=(options.get("ocr_queue_size") if isinstance(options, dict) else None),
            confidence_threshold=(options.get("confidence_threshold")
                                  if isinstance(options, dict) else None),
            max_retries=(options.get("max_retries") if isinstance(options, dict) else None),
            request_interval_seconds=(options.get("request_interval_seconds")
                                      if isinstance(options, dict) else None),
            segmentation_mode=(options.get("segmentation_mode", segmentation_mode)
                               if isinstance(options, dict) else segmentation_mode),
            _allow_active=True)

    # ---- 项目记忆注入（蓝图 §3.2 的跨任务复用）----
    # 必须在 state/pipeline_config 定型之前完成，否则 state["style_rules"] 与
    # 运行时实际使用的风格会不一致（存储值与实际值分叉）。
    # 只对尚未开始工作的任务注入：进行中的任务不能因为项目记忆变化而改术语。
    project_entries: list = []
    if not state.get("glossary") and not str(state.get("glossary_frozen") or ""):
        project_entries, project_style, project_meta = _project_memory_injection(state)
        if project_entries or project_style:
            known = {str(e.get("source") or "").casefold()
                     for e in normalize_glossary(list(user_glossary or []))}
            added = [e for e in project_entries
                     if str(e.get("source") or "").casefold() not in known]
            user_glossary = list(user_glossary or []) + added
            if project_style:
                style_rules = "；".join(
                    part for part in (project_style, str(style_rules or "").strip())
                    if part)
            state["project_memory"] = {
                **project_meta,
                "injected_entry_ids": [e["id"] for e in added],
                "injected_at": _project_now_iso(),
            }

    previous_target_lang = str(state.get("target_lang") or "")
    previous_style_rules = str(state.get("style_rules") or "")
    review_required = bool(state.get("translation_core_review_required")) \
        if state.get("p2_done") else bool(enable_review)
    saved_understanding = (state.get("pipeline_config") or {}).get(
        "enable_understanding")
    saved_source_cleanup = (state.get("pipeline_config") or {}).get(
        "enable_source_cleanup")
    is_pdf_for_config = str(filename or state.get("filename") or "").lower().endswith(".pdf")
    if mode is not None:
        strict_terminology_governance = mode == "quality"
    translator_config = _model_roles.normalize_role_config(
        None, fallback_provider=provider, fallback_model=model,
        fallback_api_key=api_key, fallback_base_url=translator_base_url)
    reviewer_config = _model_roles.normalize_role_config(
        {
            "provider": reviewer_provider,
            "model": reviewer_model,
            "api_key": reviewer_api_key,
            "base_url": reviewer_base_url,
        },
        fallback_provider=provider, fallback_model=model,
        fallback_api_key=api_key)
    saved_auxiliary = state.get("auxiliary_config") or {}
    auxiliary_config = _model_roles.normalize_role_config(
        {
            "provider": auxiliary_provider or saved_auxiliary.get("provider"),
            "model": auxiliary_model or saved_auxiliary.get("model"),
            "api_key": auxiliary_api_key,
            "base_url": auxiliary_base_url or saved_auxiliary.get("base_url"),
        },
        fallback_provider=provider, fallback_model=model,
        fallback_api_key=api_key)
    pipeline_config = {
        "target_lang": target_lang,
        "auto_term": bool(auto_term),
        "enable_report": bool(enable_report),
        "translation_theory": translation_theory,
        "style_rules": style_rules,
        "enable_review": bool(enable_review),
        "enable_annotate": bool(enable_annotate),
        "use_tm": bool(use_tm),
        "strict_terminology_governance": bool(strict_terminology_governance),
        # 批次参数随任务保存：恢复任务必须沿用原配置，否则同一任务的批次划分
        # 会在中途变化。
        "batch_size": int(batch_size or BATCH_SIZE),
        "max_batch_chars": int(max_batch_chars or TRANSLATION_MAX_BATCH_CHARS),
        "enable_understanding": enable_understanding,
        "reasoning_effort": (str(reasoning_effort or "").strip().lower()
                              if str(reasoning_effort or "").strip().lower()
                              in _REASONING_EFFORT_VALUES else ""),
        "translator": _model_roles.public_role_config(translator_config),
        "reviewer": _model_roles.public_role_config(reviewer_config),
    }
    # Auxiliary work (profile, synopsis, glossary observation) can use a
    # cheaper model without changing the authoritative translator/reviewer.
    # Keep this outside the historical pipeline_config shape unless explicitly
    # configured, so old task snapshots remain byte-for-byte compatible.
    auxiliary_explicit = any(
        value is not None and str(value).strip()
        for value in (auxiliary_provider, auxiliary_model, auxiliary_base_url)
    )
    if auxiliary_explicit or saved_auxiliary:
        state["auxiliary_config"] = _model_roles.public_role_config(auxiliary_config)
    if knowledge_feedback_interval is not None:
        try:
            normalized_feedback_interval = max(1, int(knowledge_feedback_interval))
        except (TypeError, ValueError):
            normalized_feedback_interval = 1
        pipeline_config["knowledge_feedback_interval"] = normalized_feedback_interval
    if translation_concurrency is not None or "translation_concurrency" in saved_runtime_config:
        pipeline_config["translation_concurrency"] = _runtime_int_option(
            translation_concurrency,
            saved_runtime_config.get("translation_concurrency"),
            "FOLIOTHREAD_TRANSLATION_CONCURRENCY",
            DEFAULT_TRANSLATION_CONCURRENCY,
            minimum=1,
            maximum=16,
        )
    if segmentation_mode is not None or "segmentation_mode" in saved_runtime_config:
        pipeline_config["segmentation_mode"] = effective_segmentation_mode
    state["pipeline_config"] = pipeline_config
    state.setdefault("llm_usage", _usage.empty_usage())
    auxiliary_call = _model_roles.make_role_call(
        _usage.tracked_call(call_llm, state["llm_usage"], role="auxiliary"),
        auxiliary_config)
    state.update(
        target_lang=target_lang, auto_term_enabled=bool(auto_term),
        report_enabled=bool(enable_report), theory=translation_theory,
        style_rules=style_rules, enable_review=bool(enable_review),
        enable_annotate=bool(enable_annotate), use_tm=bool(use_tm),
        provider=provider, model=model,
        translation_core_review_required=review_required,
        translator_config=_model_roles.public_role_config(translator_config),
        reviewer_config=_model_roles.public_role_config(reviewer_config),
    )
    if state.get("p2_done") and review_required and (
            previous_target_lang != str(target_lang or "")
            or previous_style_rules != str(style_rules or "")):
        _invalidate_translation_reviews(
            state, list(range(len(state.get("pairs") or []))),
            "target language or style review context changed")
        _invalidate_final_delivery_state(state)
    if delivery_config is not None:
        state["delivery_config"] = normalize_delivery_config(
            delivery_config, enable_report=enable_report,
            enable_annotate=enable_annotate)
    elif state.get("delivery_config"):
        state["delivery_config"] = normalize_delivery_config(
            state["delivery_config"], enable_report=enable_report,
            enable_annotate=enable_annotate)
    if enable_understanding is None:
        if isinstance(saved_understanding, bool):
            # A resumed task keeps the decision it was created with.  This is
            # important for old Quick jobs whose legacy state has no new flag.
            enable_understanding = saved_understanding
        elif state.get("p1_done") or state.get("p2_done"):
            # Legacy in-progress jobs had no understanding pass outside strict
            # terminology governance; do not unexpectedly add API cost on resume.
            enable_understanding = bool(
                state.get("profile_done") or state.get("understanding_done")
                or state.get("quality_mode"))
        else:
            # New tasks: Quick skips it; Standard and Academic keep it enabled.
            enable_understanding = mode != "quick"
    if enable_source_cleanup is None:
        if isinstance(saved_source_cleanup, bool):
            enable_source_cleanup = saved_source_cleanup
        else:
            enable_source_cleanup = str(filename or "").lower().endswith(".pdf")
    pipeline_config["enable_understanding"] = bool(enable_understanding)
    # Source cleanup is automatically enabled for PDF jobs.  Persist the key
    # only for PDF tasks (or an explicit opt-in) so legacy DOCX pipeline
    # configurations retain their stable shape.
    if is_pdf_for_config or enable_source_cleanup:
        pipeline_config["enable_source_cleanup"] = bool(enable_source_cleanup)
        pipeline_config.update({
            "source_cleanup_concurrency": effective_cleanup_concurrency,
            "ocr_workers": effective_ocr_workers,
            "ocr_queue_size": effective_ocr_queue_size,
            "source_cleanup_confidence_threshold": effective_cleanup_confidence_threshold,
            "source_cleanup_max_retries": effective_cleanup_max_retries,
            "source_cleanup_request_interval": effective_cleanup_request_interval,
        })
    else:
        pipeline_config.pop("enable_source_cleanup", None)
    state["pipeline_config"] = pipeline_config
    state["quality_mode"] = bool(strict_terminology_governance)
    # Persist the selected mode before extraction starts.  If the worker is
    # cancelled while rasterizing/OCRing, the resumed call keeps the same CAT
    # contract even when the process environment or UI defaults changed.
    if (not state.get("p1_done")
            or (not state.get("source_cleanup_done")
                and not (state.get("segmentation") or {}).get("mode"))):
        state.setdefault("segmentation", {})["mode"] = effective_segmentation_mode
    warnings = state.setdefault("warnings", [])

    if enable_report:
        from transpraxis import academic_writer
        academic_writer.prepare_academic_inputs(
            state, translation_theory, research_settings, literature_sources)
        academic_writer.sync_versions(state)
    save_job_state(job_id, state)

    # 术语依赖失效：必须在“全部完成”早退之前执行，
    # 否则冻结术语表变更后的旧译文会继续以 reviewed/final/TM 状态存在。
    state, stale_segs = _apply_glossary_staleness(state, job_id)
    if stale_segs:
        save_job_state(job_id, state)

    # Legacy PDF tasks created before the source-cleanup checkpoint may already
    # have partial translation pairs.  Re-run the source stage once before
    # continuing, then invalidate every artifact derived from the old text so
    # translations cannot be silently paired with corrected source segments.
    is_pdf = str(filename or state.get("filename") or "").lower().endswith(".pdf")
    if (state.get("p1_done") and not state.get("source_cleanup_done")
            and is_pdf and enable_source_cleanup and not state.get("p2_done")):
        legacy_source = list(state.get("source_paras") or state.get("paras") or [])
        legacy_segments = list(state.get("paras") or [])
        if legacy_source:
            if on_status:
                on_status("【阶段一】恢复 LLM 原文纠错与断行整理…")
            cleaned, cleanup_meta, cleanup_warnings = cleanup_source_paragraphs(
                legacy_source, provider, api_key, model,
                on_progress=on_status,
                parallelism=effective_cleanup_concurrency,
                confidence_threshold=effective_cleanup_confidence_threshold,
                checkpoint_dir=job_dir(job_id),
                cancel_check=lambda: _runtime_cancel_requested(job_id),
                max_retries=effective_cleanup_max_retries,
                request_interval_seconds=effective_cleanup_request_interval,
                profiler=profiler)
            segments, segmentation_meta = _apply_source_segmentation(
                state, cleaned, mode=effective_segmentation_mode)
            state["source_paras"] = legacy_source
            state["paras"] = segments
            state["segmentation"] = segmentation_meta
            state["source_cleanup"] = cleanup_meta
            state["source_cleanup_done"] = True
            state["source_quality_gate"] = audit_source_quality(segments)
            cleanup_meta.update({
                "paragraph_count": len(cleaned),
                "segment_count": len(segments),
            })
            for warning in cleanup_warnings:
                if warning not in warnings:
                    warnings.append(warning)
                _append_runtime_technical_log(job_id, f"source cleanup: {warning}")
            # A sentence split changes the positional source/pair contract even
            # when the OCR cleanup text itself is byte-for-byte unchanged.
            if cleaned != legacy_source or segments != legacy_segments:
                cleanup_meta["previous_pair_count"] = len(state.get("pairs") or [])
                _reset_source_dependent_state(state)
                state["paras"] = segments
                state["source_paras"] = legacy_source
                state["source_paragraphs"] = list(cleaned)
                state["segmentation"] = segmentation_meta
                state["source_cleanup"] = cleanup_meta
                state["source_cleanup_done"] = True
                state["source_quality_gate"] = audit_source_quality(segments)
            state["p1_done"] = True
            save_job_state(job_id, state)

    # Whole-book consistency review is a post-translation checkpoint.  Run it
    # before the completed-task fast path as well, so older jobs that already
    # have p2_done can receive the upgrade once without repeating it later.
    if state.get("p2_done"):
        state = _run_targeted_final_review(
            job_id, state, reviewer_config, target_lang, style_rules,
            enable_review=enable_review, on_status=on_status)

    # 全部完成 -> 直接返回
    if state["p1_done"] and state["p2_done"] and (not enable_report or state["p3_done"]) \
            and (not enable_annotate or state.get("annotations_done")):
        state["stage"] = _state_migration.derive_stage(state)
        state["performance"] = persist_performance()
        save_job_state(job_id, state)
        return state

    # ---------------- 阶段一：排版清洗 ----------------
    if not state["p1_done"]:
        if file_bytes is None:
            file_bytes = load_source(job_id)
        if file_bytes is not None:
            try:
                save_source(job_id, file_bytes)
            except Exception:
                pass

        if state.get("source_paras"):
            raw_paragraphs = list(state.get("source_paras") or [])
            paragraphs = list(raw_paragraphs)
            extraction_warnings = []
            extraction_report = state.get("extraction_report") or {}
        else:
            if file_bytes is None:
                raise ValueError("缺少源文件，请重新上传后再继续")
            if on_status:
                on_status("【阶段一】排版解析与段落重建（确定性提取）...")

            # One extraction path for every format: the same call that produces the
            # paragraphs also produces the import-scope report, so a DOCX table can
            # never be dropped silently just because the pipeline took a shortcut.
            paragraphs, extraction_warnings, extraction_report = \
                extract_document_paragraphs_with_report(
                    filename, file_bytes, ocr_max_pages=None, on_progress=on_status,
                    profiler=profiler, checkpoint_dir=job_dir(job_id),
                    ocr_workers=effective_ocr_workers,
                    ocr_queue_size=effective_ocr_queue_size,
                    cancel_check=lambda: _runtime_cancel_requested(job_id))

            raw_paragraphs = list(paragraphs)
            # Persist the extraction boundary before remote cleanup starts.  If the
            # process is cancelled while waiting for a model, OCR pages and cleanup
            # batches can resume from their own checkpoints without losing the raw
            # source audit trail.
            state["source_paras"] = raw_paragraphs
            state["extraction_report"] = extraction_report
            save_job_state(job_id, state)
        cleanup_meta = {
            "status": "skipped",
            "input_count": len(raw_paragraphs),
            "output_count": len(raw_paragraphs),
        }
        cleanup_warnings = []
        if filename.lower().endswith(".pdf") and enable_source_cleanup and paragraphs:
            if on_status:
                on_status("【阶段一】LLM 原文纠错与断行整理…")
            paragraphs, cleanup_meta, cleanup_warnings = cleanup_source_paragraphs(
                raw_paragraphs, provider, api_key, model,
                on_progress=on_status,
                parallelism=effective_cleanup_concurrency,
                confidence_by_index=((extraction_report.get("ocr") or {}).get(
                    "paragraph_confidences") if isinstance(extraction_report, dict) else None),
                confidence_threshold=effective_cleanup_confidence_threshold,
                checkpoint_dir=job_dir(job_id),
                cancel_check=lambda: _runtime_cancel_requested(job_id),
                max_retries=effective_cleanup_max_retries,
                request_interval_seconds=effective_cleanup_request_interval,
                profiler=profiler)
        elif profiler is not None:
            profiler.skipped("deterministic_cleanup", metadata={"reason": "source_cleanup_disabled"})
            profiler.skipped("llm_cleanup", metadata={"reason": "source_cleanup_disabled"})

        for warning in extraction_warnings:
            if warning not in warnings:
                warnings.append(warning)
            _append_runtime_technical_log(job_id, f"document extraction: {warning}")
        for warning in cleanup_warnings:
            if warning not in warnings:
                warnings.append(warning)
            _append_runtime_technical_log(job_id, f"source cleanup: {warning}")
        if not paragraphs:
            save_job_state(job_id, state)
            detail = extraction_warnings[-1] if extraction_warnings else "未知提取错误"
            raise ValueError(f"未提取到有效文本：{detail}")
        cleaned_paragraphs = list(paragraphs)
        segments, segmentation_meta = _apply_source_segmentation(
            state, cleaned_paragraphs, mode=effective_segmentation_mode)
        state["source_cleanup"] = cleanup_meta
        state["source_cleanup_done"] = True
        state["source_cleanup"].update({
            "paragraph_count": len(cleaned_paragraphs),
            "segment_count": len(segments),
        })
        if on_status:
            on_status(f"【阶段一】按句构建翻译单元（0/{len(segments)} 句）…")
        segment_stage = profiler.start_stage(
            "segment_build", concurrency=1, item_count=len(segments),
            metadata={
                "source": "cleaned_paragraphs",
                "mode": segmentation_meta.get("mode"),
                "paragraph_count": len(cleaned_paragraphs),
            }
        ) if profiler is not None else None
        state["paras"] = segments
        state["source_quality_gate"] = _audit_source_quality_for_state(state)
        if segment_stage is not None:
            for index in range(len(segments)):
                segment_stage.item(index, segment_stage.started_monotonic,
                                   metadata={"kind": "segment"})
            segment_stage.finish(status="completed")
        if on_status:
            on_status(f"【阶段一】按句构建翻译单元（{len(segments)}/{len(segments)} 句）…")
        if filename.lower().endswith(".pdf"):
            if file_bytes is not None:
                try:
                    with fitz.open(stream=file_bytes, filetype="pdf") as source_pdf:
                        state["source_page_count"] = source_pdf.page_count
                except Exception:
                    pass
            elif not state.get("source_page_count"):
                pages = ((extraction_report or {}).get("extracted") or {}).get("pages")
                if pages:
                    state["source_page_count"] = pages
                elif "total_pages" in ((extraction_report or {}).get("ocr") or {}):
                    state["source_page_count"] = extraction_report["ocr"]["total_pages"]
        state["p1_done"] = True
        if file_bytes is not None:
            try:
                save_source(job_id, file_bytes)  # 留存源文件，刷新后无需重新上传
            except Exception:
                pass
        for required_stage in ("pdf_classify", "layout_recovery", "rasterize", "ocr",
                               "deterministic_cleanup", "llm_cleanup", "segment_build"):
            if profiler.get(required_stage) is None:
                profiler.skipped(required_stage, metadata={"reason": "not_applicable"})
        state["performance"] = profiler.persist()
        save_job_state(job_id, state)

    if _runtime_cancel_requested(job_id):
        raise RuntimeError("任务已请求取消")

    # ---------------- 阶段 1.2：文档画像（长文理解；失败仅警告，不阻断） ----------------
    if enable_understanding and not state.get("profile_done"):
        if _runtime_cancel_requested(job_id):
            raise RuntimeError("任务已请求取消")
        if on_status:
            on_status("【阶段1.2】文档画像（分布式采样 + 结构化校验）...")
        from transpraxis.document_profile import profile_document
        profile, profile_warnings = profile_document(
            state["paras"], auxiliary_config["provider"],
            auxiliary_config["api_key"], auxiliary_config["model"], target_lang,
            call_llm=auxiliary_call)
        state["document_profile"] = profile
        state["profile_done"] = True
        for w in profile_warnings:
            if w not in warnings:
                warnings.append(w)
        if on_caption and profile:
            on_caption(f"✅ 文档画像完成：领域「{profile.get('domain') or '未知'}」"
                       f"· 文本类型「{profile.get('genre') or '未知'}」")
        elif on_caption:
            on_caption("⚠️ 文档画像失败，已跳过（可在 UI 中人工填写）。")
        save_job_state(job_id, state)

    if _runtime_cancel_requested(job_id):
        raise RuntimeError("任务已请求取消")

    # ---------------- 阶段 1.3：全文语义理解 ----------------
    if enable_understanding and not state.get("understanding_done"):
        if _runtime_cancel_requested(job_id):
            raise RuntimeError("任务已请求取消")
        if on_status:
            on_status("【阶段1.3】全文语义理解（语义单元摘要 + 全文概要）...")
        if on_caption:
            on_caption("🧭 正在建立全文概要与当前单元摘要...")
        understanding_call = auxiliary_call
        units, digests, synopsis, understanding_warnings = \
            _context.build_document_understanding(
                state["paras"], state.get("document_profile"),
                auxiliary_config["provider"], auxiliary_config["api_key"],
                auxiliary_config["model"], target_lang, call_llm=understanding_call,
                checkpoint_dir=job_dir(job_id))
        state["semantic_units"] = units
        state["section_digests"] = digests
        state["document_synopsis"] = synopsis
        state["understanding_warnings"] = list(understanding_warnings)
        state["understanding_done"] = True
        for warning in understanding_warnings:
            if warning not in warnings:
                warnings.append(warning)
        _context.write_understanding_artifacts(
            job_dir(job_id), units, digests, synopsis)
        if on_caption:
            status = synopsis.get("status") or "unavailable"
            on_caption(f"✅ 全文理解完成：{len(digests)} 个语义单元 · 概要状态 {status}")
        save_job_state(job_id, state)

    if _runtime_cancel_requested(job_id):
        raise RuntimeError("任务已请求取消")

    # ---------------- 阶段 1.5：智能抽取术语 ----------------
    if auto_term and not state["auto_terms"]:
        if _runtime_cancel_requested(job_id):
            raise RuntimeError("任务已请求取消")
        if on_status:
            on_status("【阶段1.5】正在 AI 智能抽取全文核心术语...")
        if on_caption:
            on_caption("🤖 正在从全文分布式样本中提取专业术语...")
        from transpraxis.terminology import extract_auto_terms_v2
        entries, extract_warnings = extract_auto_terms_v2(
            state["paras"], target_lang, auxiliary_config["provider"],
            auxiliary_config["api_key"], auxiliary_config["model"],
            document_profile=state.get("document_profile"),
            call_llm=auxiliary_call,
            cancel_check=lambda: _runtime_cancel_requested(job_id))
        state["auto_term_entries"] = entries
        state["auto_terms"] = {e["source"]: e["target"] for e in entries}
        if entries:
            if on_caption:
                on_caption(f"✅ 成功提取 {len(entries)} 个候选术语（全部出现位置已记录）")
        else:
            msg = "术语抽取失败（限流或返回格式异常），已跳过该步骤；可稍后点击“继续处理”重试。"
            if msg not in warnings and msg not in extract_warnings:
                warnings.append(msg)
        for w in extract_warnings:
            if w not in warnings:
                warnings.append(w)
        save_job_state(job_id, state)

    if _runtime_cancel_requested(job_id):
        raise RuntimeError("任务已请求取消")
    legacy_auto = [{"source": k, "target": v, "behavior": "translate",
                    "status": "provisional"}
                   for k, v in (state["auto_terms"] or {}).items()]
    auto_entries = normalize_glossary(state.get("auto_term_entries") or legacy_auto)
    user_entries = normalize_glossary(list(user_glossary or []))
    if not strict_terminology_governance:
        for e in auto_entries:
            if e["status"] == "candidate":
                e["status"] = "provisional"
    working = normalize_glossary(state.get("glossary") or [])
    if not working:
        # 项目记忆已在本函数开头注入到 user_glossary/style_rules，这里只做合并。
        working = normalize_glossary(user_entries + auto_entries)
    else:
        # 新上传/新抽取的术语若不在已保存审核表中，追加（不覆盖人工审核结果）
        known = {e["source"].casefold() for e in working}
        for e in user_entries + auto_entries:
            if e["source"].casefold() not in known:
                working.append(e)
    if not strict_terminology_governance:
        for e in working:
            if e["status"] == "candidate":
                e["status"] = "provisional"
    state["glossary"] = working
    glossary = working
    final_termbase = glossary_to_terms(glossary)

    # ---------------- 严格术语治理门禁：候选术语需人工审核/冻结后才能翻译 ----------------
    if strict_terminology_governance:
        pending = [e for e in working
                   if (e.get("status") or "").lower() == "candidate"]
        if not state.get("glossary_frozen") and not state.get("quality_bypass") \
                and pending:
            if on_status:
                on_status(f"⏸ 严格术语治理：{len(pending)} 条候选术语等待人工审核与冻结...")
            msg = (f"严格术语治理：术语表尚未冻结（{len(pending)} 条自动抽取的"
                   "候选术语待审核），翻译未开始。请在“术语准备与审核”面板"
                   "完成审核并冻结后继续；导入术语库中的锁定术语已视为固定，"
                   "无需再次审核。")
            if msg not in warnings:
                warnings.append(msg)
            state["stage"] = _state_migration.derive_stage(state)
            save_job_state(job_id, state)
            return state

    # ---------------- 阶段二：双语翻译（批次 + 确定性检查 + 独立审校 + 翻译记忆）----------------
    if not state["p2_done"]:
        if _runtime_cancel_requested(job_id):
            raise RuntimeError("任务已请求取消")
        if on_status:
            on_status("【阶段二】双语翻译与术语严格注入（批次翻译 + 确定性检查 + 独立审校）...")
        _batch_cfg = state.get("pipeline_config") or {}
        translate_stage(state, job_id, glossary, provider, api_key, model, target_lang,
                        style_rules, enable_review, use_tm=use_tm,
                        document_profile=state.get("document_profile"),
                        translator_config=translator_config,
                        reviewer_config=reviewer_config,
                        auxiliary_config=auxiliary_config,
                        on_status=on_status, on_caption=on_caption,
                        batch_size=_batch_cfg.get("batch_size", batch_size),
                        max_batch_chars=_batch_cfg.get("max_batch_chars", max_batch_chars),
                        knowledge_feedback_interval=_batch_cfg.get(
                            "knowledge_feedback_interval", knowledge_feedback_interval),
                        translation_concurrency=_batch_cfg.get(
                            "translation_concurrency", translation_concurrency))
        state["p2_done"] = True
        from transpraxis import delivery as _delivery
        state["delivery_status"] = _delivery.compute_delivery_status(state)
        save_job_state(job_id, state)

    if _runtime_cancel_requested(job_id):
        raise RuntimeError("任务已请求取消")

    if state.get("p2_done"):
        state = _run_targeted_final_review(
            job_id, state, reviewer_config, target_lang, style_rules,
            enable_review=enable_review, on_status=on_status)

    if _runtime_cancel_requested(job_id):
        raise RuntimeError("任务已请求取消")

    # ---------------- 阶段 2.5：三色自动标注 ----------------
    if enable_annotate and state["p2_done"] and not state.get("annotations_done"):
        if _runtime_cancel_requested(job_id):
            raise RuntimeError("任务已请求取消")
        if on_status:
            on_status("【阶段 2.5】自动标注学习重点（红=生僻词 / 黄=专业名词 / 青绿=难点句）...")
        annotate_stage(state, job_id, glossary, provider, api_key, model, target_lang,
                       on_caption=on_caption)

    if _runtime_cancel_requested(job_id):
        raise RuntimeError("任务已请求取消")

    # ---------------- 阶段三：报告生成 ----------------
    if enable_report and not state["p3_done"]:
        if _runtime_cancel_requested(job_id):
            raise RuntimeError("任务已请求取消")
        if on_status:
            on_status(f"【阶段三】基于《{translation_theory}》生成报告...")
        report_md = generate_mti_report(state["pairs"], final_termbase, translation_theory,
                                        provider, api_key, model, state, job_id,
                                        on_status=on_status,
                                        research_settings=research_settings,
                                        literature_sources=literature_sources)
        if not report_md.strip():
            if state.get("report_status") == "blocked_final_case_policy":
                state["stage"] = _state_migration.derive_stage(state)
                save_job_state(job_id, state)
                return state
            raise RuntimeError("报告内容为空，请点击“继续处理”重试")
        state["p3_md"] = report_md
        state["theory"] = translation_theory
        state["p3_done"] = True
        save_job_state(job_id, state)

    state["stage"] = _state_migration.derive_stage(state)
    state["performance"] = persist_performance()
    save_job_state(job_id, state)
    return state
